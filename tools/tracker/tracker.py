#!/usr/bin/env python3
"""SmartDoc master implementation tracker: SmartDoc_Master_Implementation_Tracker.xlsx.

The workbook in the repository root is the source of truth for implementation status.
Every command loads it, applies one change and rewrites the sheets this tool manages.
MASTER, CHANGE_LOG and the input tables (gates, findings, benchmarks, test runs, plans,
backlog, session state) are data; everything else is Excel formulas over MASTER, so a
status typed straight into MASTER shows up on every sheet at the next recalculation.

    python tools/tracker/tracker.py init
    python tools/tracker/tracker.py set EDIT-001 --status IN_PROGRESS
    python tools/tracker/tracker.py set EDIT-001 --status VERIFIED --evidence "..." --tests "..." --log "..."
    python tools/tracker/tracker.py add EDIT-001A --parent EDIT-001 --title "..."
    python tools/tracker/tracker.py log EDIT-001 "Fixed ... Added regression test X. E2E passed." --result PASS
    python tools/tracker/tracker.py testrun backend "pytest -q" --result PASS --passed 643 --failed 0
    python tools/tracker/tracker.py bench BENCH-001 --current 4.2 --threshold 10
    python tools/tracker/tracker.py gate GATE-013 --status PASS --evidence "..."
    python tools/tracker/tracker.py state          # SESSION_STATE <- docs/session-state.json
    python tools/tracker/tracker.py show [--phase 1] [--status BLOCKED] [--open]
    python tools/tracker/tracker.py render         # rewrite the managed sheets, data unchanged

openpyxl stores formulas without results; Excel computes them on open. excel_recalc.py
(Windows + Excel + pywin32) stores the values and checks every formula for errors.
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import os
import re
import subprocess
import sys
import time
import warnings
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import seed  # noqa: E402

warnings.filterwarnings("ignore", module="openpyxl")

ROOT = HERE.parents[1]
DEFAULT_WORKBOOK = ROOT / "SmartDoc_Master_Implementation_Tracker.xlsx"
DEFAULT_STATE_JSON = ROOT / "docs" / "session-state.json"
RESUME = "Прочети docs/AI-CONTINUATION.md и продължи от next task."
AUDIT_REPORT = "https://claude.ai/artifact/WKskiyexTDLAHe1BfjvWR1"
MAX_ROW = 2000  # formulas cover MASTER rows 2..MAX_ROW

STATUSES = [name for name, _ in seed.STATUS_MEANINGS]
ACTIVE = ("IN_PROGRESS", "IMPLEMENTED", "TESTING")
NEEDS_REASON = ("BLOCKED", "FAILED", "DEFERRED")
PRIORITIES = ("P0", "P1", "P2", "P3")
RISKS = ("Low", "Medium", "High", "Critical")
RESULTS = ("PASS", "FAIL", "PARTIAL", "N/A")
GATE_STATUSES = ("NOT_EVALUATED", "PASS", "FAIL")
BUCKETS = ("NOW", "NEXT", "LATER", "ARCHITECTURE_ONLY")
CATEGORIES = list(seed.CATEGORY_SHEETS)
ID_RE = re.compile(r"^[A-Z][A-Z0-9]*-\d{3}[A-Z]?$")


class TrackerError(Exception):
    pass


# -- sheets -----------------------------------------------------------------------------

VIEW_SHEETS = [
    "DOCUMENT_FIDELITY", "DOCX_OOXML", "PDF", "EDITOR", "FORMATTING", "AI", "SECURITY",
    "PERFORMANCE", "TESTING", "PRODUCT_FEATURES", "TRANSLATION", "PDF_TO_EDITABLE",
    "BILLING_PLANS", "INFRA",
]
SHEET_ORDER = ["DASHBOARD", "MASTER", "CRITICAL_FIXES", *VIEW_SHEETS, "CHANGE_LOG", "SESSION_STATE", "RELEASE_GATES"]
SHEET_INFO = {
    "DASHBOARD": ("Overview", "Totals, progress by phase and category, latest test runs, next action."),
    "MASTER": ("Every task", "Status, evidence, tests and commit per task. The one place task data is edited."),
    "CRITICAL_FIXES": ("Audit findings and P0 tasks", "The 20 findings of the 2026-09-26 audit (current known risks) with fix progress, then every P0 task."),
    "DOCUMENT_FIDELITY": ("Content preservation", "Canonical model, fidelity reports, capability matrix, change/review model."),
    "DOCX_OOXML": ("DOCX / OOXML preservation", "Import and export fidelity for Word documents."),
    "PDF": ("PDF engine", "PDF export and PDF parsing."),
    "EDITOR": ("Editor integrity", "Tiptap mapping and save integrity."),
    "FORMATTING": ("Formatting and Document Health", "Formatting engine, Format by Example, Document Health."),
    "AI": ("AI fidelity", "Fidelity checks, proposals and review, budgets."),
    "SECURITY": ("Security and accounts", "Security hardening and account essentials."),
    "PERFORMANCE": ("Performance", "Measured benchmarks (baseline vs current) and performance tasks."),
    "TESTING": ("Testing", "Test runs and testing tasks."),
    "PRODUCT_FEATURES": ("Product features", "Feature backlog (NOW / NEXT / LATER / ARCHITECTURE_ONLY) and product tasks."),
    "TRANSLATION": ("Translation and fonts", "Translation MVP and FontResolver."),
    "PDF_TO_EDITABLE": ("PDF to editable", "PDF to editable document conversion."),
    "BILLING_PLANS": ("Plans and billing", "Plan hypotheses, usage units, entitlements, Stripe."),
    "INFRA": ("Infrastructure", "Tracker, jobs, storage, observability, deployment, documentation."),
    "CHANGE_LOG": ("Change log", "Every recorded change with files, tests, result and commit."),
    "SESSION_STATE": ("Where the work stands", "Current phase and task, last verification, next action. Mirrors docs/session-state.json."),
    "RELEASE_GATES": ("Release gates", "GATE-001..015. A gate is set to PASS only with evidence from its verification run."),
}
SHEET_CATEGORIES: dict[str, list[str]] = {}
for _cat, _sheet in seed.CATEGORY_SHEETS.items():
    SHEET_CATEGORIES.setdefault(_sheet, []).append(_cat)

# -- columns ----------------------------------------------------------------------------

# key, header, width, wrap
MASTER_COLS = [
    ("id", "ID", 12, False),
    ("phase", "Phase", 7, False),
    ("category", "Category", 14, False),
    ("title", "Feature/Task", 42, True),
    ("description", "Description", 60, True),
    ("priority", "Priority", 9, False),
    ("status", "Status", 14, False),
    ("evidence", "Evidence", 45, True),
    ("tests", "Tests", 40, True),
    ("owner", "Owner/System", 24, True),
    ("deps", "Dependencies", 22, True),
    ("risk", "Risk", 9, False),
    ("started", "Started", 17, False),
    ("completed", "Completed", 17, False),
    ("last_verified", "Last Verified", 17, False),
    ("commit", "Commit", 10, False),
    ("notes", "Notes", 45, True),
    ("parent", "Parent ID", 11, False),
    ("audit", "Audit Ref", 16, False),
]
MCOL = {key: get_column_letter(i) for i, (key, *_rest) in enumerate(MASTER_COLS, 1)}
DATE_KEYS = ("started", "completed", "last_verified")

CHANGELOG_COLS = [
    ("timestamp", "Timestamp", 17, False),
    ("task", "Task ID", 16, False),
    ("phase", "Phase", 7, False),
    ("change", "Change", 70, True),
    ("files", "Files", 45, True),
    ("tests", "Tests", 40, True),
    ("result", "Result", 9, False),
    ("commit", "Commit", 10, False),
    ("notes", "Notes", 40, True),
]

# header, MASTER key, width, kind
VIEW_COLS = [
    ("ID", "id", 12, "id"),
    ("Category", "category", 14, "text"),
    ("Phase", "phase", 7, "num"),
    ("Feature/Task", "title", 44, "text"),
    ("Priority", "priority", 9, "text"),
    ("Status", "status", 14, "status"),
    ("Evidence", "evidence", 40, "text"),
    ("Tests", "tests", 34, "text"),
    ("Last Verified", "last_verified", 16, "date"),
    ("Commit", "commit", 10, "text"),
    ("Notes", "notes", 40, "text"),
    ("Parent ID", "parent", 11, "text"),
]

SESSION_FIELDS = [
    "Current Phase", "Current Task", "Last Completed", "Last Verified", "Current Branch",
    "Current Commit", "Tests Passed", "Tests Failed", "Open Blockers", "Next Task",
    "Next Command", "Next Test", "Checkpoint Time", "Resume Instructions",
]
SESSION_EXTRA = ["Production Code Changed", "DB Migrations Applied"]
SESSION_FIRST_ROW = 5

FINDING_COLS = [("id", "Finding ID"), ("finding", "Finding"), ("tasks", "Fix Tasks"), ("notes", "Notes")]
GATE_COLS = [("id", "Gate ID"), ("gate", "Gate"), ("linked", "Linked Tasks"), ("status", "Status"),
             ("evidence", "Evidence"), ("last_verified", "Last Verified"), ("notes", "Notes")]
BENCH_COLS = [("id", "Bench ID"), ("scenario", "Scenario"), ("metric", "Metric"), ("unit", "Unit"),
              ("baseline", "Baseline"), ("source", "Baseline Source"), ("current", "Current"),
              ("measured", "Measured On"), ("threshold", "Threshold"), ("notes", "Notes")]
RUN_COLS = [("id", "Run ID"), ("suite", "Suite"), ("result", "Result"), ("command", "Command"), ("passed", "Passed"),
            ("failed", "Failed"), ("notes", "Notes"), ("skipped", "Skipped"), ("date", "Date"),
            ("duration", "Duration (s)"), ("commit", "Commit")]
PLAN_COLS = [("unit", "Usage unit / feature"), ("free", "FREE"), ("pro", "PRO"), ("business", "BUSINESS"), ("notes", "Notes")]
BACKLOG_COLS = [("feature", "Feature"), ("bucket", "Bucket"), ("tasks", "Linked Tasks"), ("notes", "Notes")]

# -- styles -----------------------------------------------------------------------------

NAVY = "1F3864"


def _font(**kw) -> Font:
    kw.setdefault("name", "Arial")
    kw.setdefault("size", 10)
    return Font(**kw)


F_BODY = _font()
F_BOLD = _font(bold=True)
F_HEAD = _font(bold=True, color="FFFFFF")
F_TITLE = _font(size=16, bold=True, color=NAVY)
F_SUB = _font(italic=True, color="595959")
F_SECTION = _font(size=12, bold=True, color=NAVY)
F_LABEL = _font(bold=True, color="404040")
F_KPI = _font(size=11, bold=True)
F_INPUT = _font(color="0000FF")
F_LINK = _font(color="0563C1", underline="single")
FILL_HEAD = PatternFill("solid", fgColor=NAVY)
FILL_LABEL = PatternFill("solid", fgColor="D9E1F2")
FILL_KPI = PatternFill("solid", fgColor="F2F2F2")
_THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
A_WRAP = Alignment(wrap_text=True, vertical="top")
A_TOP = Alignment(vertical="top")
A_CENTER = Alignment(horizontal="center", vertical="top")
A_HEAD = Alignment(wrap_text=True, vertical="center")
DATE_FMT = "yyyy-mm-dd hh:mm"
PCT_FMT = "0.0%"

# status -> (fill, font colour, bold, italic)
STATUS_STYLES = {
    "DONE": ("C6EFCE", "006100", True, False),
    "VERIFIED": ("E2EFDA", "375623", False, False),
    "IN_PROGRESS": ("FFEB9C", "9C5700", True, False),
    "IMPLEMENTED": ("DDEBF7", "1F4E78", False, False),
    "TESTING": ("DDEBF7", "1F4E78", False, True),
    "BLOCKED": ("F8CBAD", "833C0B", True, False),
    "FAILED": ("FFC7CE", "9C0006", True, False),
    "DEFERRED": (None, "808080", False, True),
    "UNKNOWN": ("E4DFEC", "5B2C6F", False, False),
    # gates, findings, benchmarks, results, buckets
    "PASS": ("C6EFCE", "006100", True, False),
    "FAIL": ("FFC7CE", "9C0006", True, False),
    "PARTIAL": ("FFEB9C", "9C5700", True, False),
    "RESOLVED": ("C6EFCE", "006100", True, False),
    "OPEN": ("FCE4D6", "833C0B", False, False),
    "WITHIN": ("C6EFCE", "006100", True, False),
    "OVER": ("FFC7CE", "9C0006", True, False),
    "NOW": ("DDEBF7", "1F4E78", True, False),
    "NOT IN MASTER": ("FFC7CE", "9C0006", True, True),
}


def _dxf_fill(color: str | None) -> PatternFill | None:
    return PatternFill(start_color=color, end_color=color, fill_type="solid") if color else None


def add_status_cf(ws, rng: str, names=None) -> None:
    for name in names or STATUS_STYLES:
        if name not in STATUS_STYLES:
            continue
        fill, color, bold, italic = STATUS_STYLES[name]
        ws.conditional_formatting.add(rng, CellIsRule(
            operator="equal", formula=[f'"{name}"'], fill=_dxf_fill(fill),
            font=Font(color=color, bold=bold, italic=italic)))


def add_priority_cf(ws, rng: str) -> None:
    ws.conditional_formatting.add(rng, CellIsRule(
        operator="equal", formula=['"P0"'], fill=_dxf_fill("FFC7CE"), font=Font(color="9C0006", bold=True)))
    ws.conditional_formatting.add(rng, CellIsRule(
        operator="equal", formula=['"P1"'], font=Font(color="C65911", bold=True)))


def add_list_dv(ws, rng: str, values, title: str) -> None:
    dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True,
                        showErrorMessage=True, errorStyle="stop", errorTitle=f"Invalid {title}",
                        error=f"Pick a {title} from the list.")
    dv.add(rng)
    ws.add_data_validation(dv)


# -- helpers ----------------------------------------------------------------------------

def now() -> dt.datetime:
    return dt.datetime.now().replace(microsecond=0)


def mref(key: str) -> str:
    col = MCOL[key]
    return f"MASTER!${col}$2:${col}${MAX_ROW}"


def split_ids(value) -> list[str]:
    return [x.strip() for x in str(value or "").split(",") if x.strip()]


def norm(value):
    if isinstance(value, str):
        return value.strip()
    return "" if value is None else value


def put(ws, row: int, col: int, value, font=F_BODY, align=A_TOP, fmt: str | None = None,
        fill: PatternFill | None = None, border: bool = True):
    cell = ws.cell(row=row, column=col)
    cell.value = None if value == "" else value
    cell.font = font
    cell.alignment = align
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = fill
    if border:
        cell.border = BOX
    return cell


def header(ws, row: int, headers, col: int = 1) -> None:
    for i, text in enumerate(headers):
        put(ws, row, col + i, text, font=F_HEAD, align=A_HEAD, fill=FILL_HEAD)


def need(widths: dict, col: int, width: float) -> None:
    widths[col] = max(widths.get(col, 0), width)


def apply_widths(ws, widths: dict) -> None:
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width


def page_setup(ws) -> None:
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def linked_count(list_cell: str, condition: str | None = None) -> str:
    """SUMPRODUCT counting MASTER rows whose ID is in the comma list in `list_cell`."""
    match = f'ISNUMBER(SEARCH(","&{mref("id")}&",",","&SUBSTITUTE({list_cell}," ","")&","))'
    return f"SUMPRODUCT({match}*({condition}))" if condition else f"SUMPRODUCT({match}*1)"


def status_is(*names: str) -> str:
    return "+".join(f'({mref("status")}="{n}")' for n in names)


def list_total(list_cell: str) -> str:
    compact = f'SUBSTITUTE({list_cell}," ","")'
    return f'=IF(LEN({compact})=0,0,LEN({compact})-LEN(SUBSTITUTE({compact},",",""))+1)'


# -- data -------------------------------------------------------------------------------

class Data:
    def __init__(self):
        self.tasks: list[dict] = []
        self.changelog: list[dict] = []
        self.session: dict[str, str] = {}
        self.findings: list[dict] = []
        self.gates: list[dict] = []
        self.benches: list[dict] = []
        self.runs: list[dict] = []
        self.plans: list[dict] = []
        self.backlog: list[dict] = []

    def task(self, tid: str) -> dict:
        for t in self.tasks:
            if t["id"] == tid:
                return t
        raise TrackerError(f"unknown task {tid}")

    def has(self, tid: str) -> bool:
        return any(t["id"] == tid for t in self.tasks)

    def validate(self) -> None:
        seen = set()
        for t in self.tasks:
            tid = t["id"]
            if tid in seen:
                raise TrackerError(f"duplicate task ID {tid} in MASTER")
            seen.add(tid)
            if t["status"] not in STATUSES:
                raise TrackerError(f"{tid}: status {t['status']!r} is not one of {', '.join(STATUSES)}")
            if t["priority"] not in PRIORITIES:
                raise TrackerError(f"{tid}: priority {t['priority']!r} is not one of {', '.join(PRIORITIES)}")
            if t["category"] not in CATEGORIES:
                raise TrackerError(f"{tid}: category {t['category']!r} is not one of {', '.join(CATEGORIES)}")
            if not isinstance(t["phase"], int) or not 0 <= t["phase"] <= 18:
                raise TrackerError(f"{tid}: phase {t['phase']!r} must be a whole number 0-18")
        for t in self.tasks:
            if t["parent"] and t["parent"] not in seen:
                raise TrackerError(f"{t['id']}: parent {t['parent']} does not exist")

    def counts(self) -> dict:
        by = {s: 0 for s in STATUSES}
        for t in self.tasks:
            by[t["status"]] += 1
        total = len(self.tasks)
        done = by["DONE"]
        verified = by["VERIFIED"] + done
        return {
            "total": total, "done": done, "verified": verified,
            "in_progress": sum(by[s] for s in ACTIVE), "blocked": by["BLOCKED"], "failed": by["FAILED"],
            "deferred": by["DEFERRED"], "remaining": total - done,
            "completion_pct": round(100 * done / total, 1) if total else 0.0,
            "verified_pct": round(100 * verified / total, 1) if total else 0.0,
            "p0_remaining": sum(1 for t in self.tasks if t["priority"] == "P0" and t["status"] != "DONE"),
            "p1_remaining": sum(1 for t in self.tasks if t["priority"] == "P1" and t["status"] != "DONE"),
            "by_status": by,
        }


def from_seed() -> Data:
    d = Data()
    for tid, parent, phase, category, title, description, priority, owner, deps, risk, audit in seed.TASKS:
        d.tasks.append(dict(
            id=tid, phase=phase, category=category, title=title, description=description, priority=priority,
            status="NOT_STARTED", evidence="", tests="", owner=owner, deps=deps, risk=risk, started=None,
            completed=None, last_verified=None, commit="", notes="", parent=parent, audit=audit))
    d.findings = [dict(id=i, finding=f, tasks=t, notes="") for i, f, t in seed.AUDIT_FINDINGS]
    d.gates = [dict(id=i, gate=g, linked=t, status="NOT_EVALUATED", evidence="", last_verified=None, notes="")
               for i, g, t in seed.RELEASE_GATES]
    d.benches = [dict(id=i, scenario=s, metric=m, unit=u, baseline=b, source=src, current=None, measured=None,
                      threshold=None, notes="") for i, s, m, b, u, src in seed.BENCHMARKS]
    plan_notes = {
        "Documents / month": "BUSINESS: per workspace",
        "Exports / month": "BUSINESS: 'high export limits', no number given",
    }
    d.plans = [dict(unit=u, free=f, pro=p, business=b, notes=plan_notes.get(u, ""))
               for u, f, p, b in seed.PLAN_DEFAULTS]
    d.backlog = [dict(feature=f, bucket=b, tasks=t, notes=n) for f, b, t, n in seed.PRODUCT_BACKLOG]
    d.session = {name: "" for name in SESSION_FIELDS + SESSION_EXTRA}
    d.session["Resume Instructions"] = RESUME
    return d


# -- reading ----------------------------------------------------------------------------

def read_table(ws, first_header: str, columns) -> list[dict]:
    header_row = None
    for (cell,) in ws.iter_rows(min_col=1, max_col=1):
        if isinstance(cell.value, str) and cell.value.strip() == first_header:
            header_row = cell.row
            break
    if header_row is None:
        raise TrackerError(f"{ws.title}: table starting with '{first_header}' not found")
    positions = {}
    for cell in ws[header_row]:
        if isinstance(cell.value, str):
            positions[cell.value.strip()] = cell.column
    missing = [h for _, h in columns if h not in positions]
    if missing:
        raise TrackerError(f"{ws.title}: columns missing: {', '.join(missing)}")
    rows = []
    row = header_row + 1
    key_col = positions[columns[0][1]]
    while norm(ws.cell(row, key_col).value) != "":
        rows.append({key: norm(ws.cell(row, positions[h]).value) for key, h in columns})
        row += 1
    return rows


def sheet(wb, name: str):
    if name not in wb.sheetnames:
        raise TrackerError(f"sheet {name} is missing from the workbook")
    return wb[name]


def load(path: Path):
    if not path.exists():
        raise TrackerError(f"{path.name} does not exist; run `tracker.py init` first")
    wb = load_workbook(path)
    d = Data()
    d.tasks = read_table(sheet(wb, "MASTER"), "ID", [(k, h) for k, h, *_ in MASTER_COLS])
    for t in d.tasks:
        t["status"] = str(t["status"]).upper()
        t["priority"] = str(t["priority"]).upper()
        try:
            t["phase"] = int(t["phase"])
        except (TypeError, ValueError):
            raise TrackerError(f"{t['id']}: phase {t['phase']!r} is not a number") from None
        for key in DATE_KEYS:
            if t[key] == "":
                t[key] = None
    d.changelog = read_table(sheet(wb, "CHANGE_LOG"), "Timestamp", [(k, h) for k, h, *_ in CHANGELOG_COLS])
    known = set(SESSION_FIELDS + SESSION_EXTRA)
    d.session = {name: "" for name in known}
    for row in read_table(sheet(wb, "SESSION_STATE"), "Field", [("field", "Field"), ("value", "Value")]):
        if row["field"] in known:
            d.session[row["field"]] = row["value"]
    d.findings = read_table(sheet(wb, "CRITICAL_FIXES"), "Finding ID", FINDING_COLS)
    d.gates = read_table(sheet(wb, "RELEASE_GATES"), "Gate ID", GATE_COLS)
    d.benches = read_table(sheet(wb, "PERFORMANCE"), "Bench ID", BENCH_COLS)
    d.runs = read_table(sheet(wb, "TESTING"), "Run ID", RUN_COLS)
    d.plans = read_table(sheet(wb, "BILLING_PLANS"), "Usage unit / feature", PLAN_COLS)
    d.backlog = read_table(sheet(wb, "PRODUCT_FEATURES"), "Feature", BACKLOG_COLS)
    d.validate()
    return wb, d


# -- writing ----------------------------------------------------------------------------

def title_block(ws, name: str) -> None:
    title, description = SHEET_INFO[name]
    ws.cell(1, 1, f"{name} — {title}").font = F_TITLE
    ws.cell(2, 1, description).font = F_SUB


def view_formula(key: str, kind: str, row: int) -> str:
    rng = mref(key)
    match = f"MATCH($A{row},{mref('id')},0)"
    if kind == "status":
        return f'=IFERROR(INDEX({rng},{match})&"","NOT IN MASTER")'
    if kind == "num":
        return f'=IFERROR(INDEX({rng},{match}),"")'
    if kind == "date":
        return f'=IFERROR(IF(INDEX({rng},{match})="","",INDEX({rng},{match})),"")'
    return f'=IFERROR(INDEX({rng},{match})&"","")'


def task_view(ws, top: int, ids: list[str], caption: str, widths: dict) -> tuple[int, int]:
    """Formula view of MASTER rows; returns (header row, last row)."""
    ws.cell(top, 1, caption).font = F_SECTION
    header_row = top + 1
    header(ws, header_row, [h for h, *_ in VIEW_COLS])
    for col, (_, _, width, _) in enumerate(VIEW_COLS, 1):
        need(widths, col, width)
    row = header_row
    for tid in ids:
        row += 1
        for col, (_, key, _, kind) in enumerate(VIEW_COLS, 1):
            value = tid if kind == "id" else view_formula(key, kind, row)
            put(ws, row, col, value,
                font=F_BOLD if kind == "id" else F_BODY,
                align=A_CENTER if kind == "num" else A_TOP,
                fmt=DATE_FMT if kind == "date" else None)
    if not ids:
        row += 1
        put(ws, row, 1, "No tasks in this view yet.", font=F_SUB, border=False)
        return header_row, row
    add_status_cf(ws, f"F{header_row + 1}:F{row}")
    add_priority_cf(ws, f"E{header_row + 1}:E{row}")
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(VIEW_COLS))}{row}"
    return header_row, row


def category_summary(ws, categories: list[str]) -> None:
    def total(statuses=None):
        parts = []
        for c in categories:
            if statuses is None:
                parts.append(f'COUNTIF({mref("category")},"{c}")')
            else:
                parts += [f'COUNTIFS({mref("category")},"{c}",{mref("status")},"{s}")' for s in statuses]
        return "=" + "+".join(parts)

    labels = ["Tasks", "Done", "Verified", "In progress", "Blocked", "Failed", "Done %", "Verified %"]
    values = [total(), total(["DONE"]), total(["VERIFIED"]) + "+B4", total(ACTIVE), total(["BLOCKED"]),
              total(["FAILED"]), "=IF(A4=0,0,B4/A4)", "=IF(A4=0,0,C4/A4)"]
    for col, (label, value) in enumerate(zip(labels, values), 1):
        put(ws, 3, col, label, font=F_LABEL, fill=FILL_LABEL)
        put(ws, 4, col, value, font=F_KPI, fill=FILL_KPI, align=A_CENTER,
            fmt=PCT_FMT if label.endswith("%") else "0")


def view_ids(d: Data, sheet_name: str) -> list[str]:
    cats = SHEET_CATEGORIES.get(sheet_name, [])
    return [t["id"] for t in d.tasks if t["category"] in cats]


def write_master(ws, d: Data) -> None:
    header(ws, 1, [h for _, h, *_ in MASTER_COLS])
    for row, t in enumerate(d.tasks, 2):
        for col, (key, _, _, wrap) in enumerate(MASTER_COLS, 1):
            value = t.get(key)
            put(ws, row, col, "" if value is None else value,
                font=F_BOLD if key == "id" else F_BODY,
                align=A_CENTER if key == "phase" else (A_WRAP if wrap else A_TOP),
                fmt=DATE_FMT if key in DATE_KEYS else None)
    last = max(2, len(d.tasks) + 1)
    widths = {}
    for col, (_, _, width, _) in enumerate(MASTER_COLS, 1):
        need(widths, col, width)
    apply_widths(ws, widths)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(MASTER_COLS))}{last}"
    col = MCOL
    add_list_dv(ws, f"{col['status']}2:{col['status']}{MAX_ROW}", STATUSES, "status")
    add_list_dv(ws, f"{col['priority']}2:{col['priority']}{MAX_ROW}", PRIORITIES, "priority")
    add_list_dv(ws, f"{col['risk']}2:{col['risk']}{MAX_ROW}", RISKS, "risk")
    add_list_dv(ws, f"{col['category']}2:{col['category']}{MAX_ROW}", CATEGORIES, "category")
    phase_dv = DataValidation(type="whole", operator="between", formula1="0", formula2="18", allow_blank=True,
                              showErrorMessage=True, errorTitle="Invalid phase", error="Phase is a whole number 0-18.")
    phase_dv.add(f"{col['phase']}2:{col['phase']}{MAX_ROW}")
    ws.add_data_validation(phase_dv)
    add_status_cf(ws, f"{col['status']}2:{col['status']}{MAX_ROW}", names=STATUSES)
    add_priority_cf(ws, f"{col['priority']}2:{col['priority']}{MAX_ROW}")
    page_setup(ws)


def write_changelog(ws, d: Data) -> None:
    header(ws, 1, [h for _, h, *_ in CHANGELOG_COLS])
    for row, entry in enumerate(d.changelog, 2):
        for col, (key, _, _, wrap) in enumerate(CHANGELOG_COLS, 1):
            put(ws, row, col, entry.get(key, ""), align=A_WRAP if wrap else A_TOP,
                fmt=DATE_FMT if key == "timestamp" else None)
    widths = {}
    for col, (_, _, width, _) in enumerate(CHANGELOG_COLS, 1):
        need(widths, col, width)
    apply_widths(ws, widths)
    ws.freeze_panes = "A2"
    last = max(2, len(d.changelog) + 1)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(CHANGELOG_COLS))}{last}"
    add_list_dv(ws, "G2:G5000", RESULTS, "result")
    add_status_cf(ws, "G2:G5000", names=["PASS", "FAIL", "PARTIAL"])
    page_setup(ws)


def write_session(ws, d: Data) -> None:
    title_block(ws, "SESSION_STATE")
    put(ws, 3, 1, f"Resume with: {RESUME}", font=F_BOLD, border=False)
    header(ws, SESSION_FIRST_ROW - 1, ["Field", "Value"])
    row = SESSION_FIRST_ROW
    for name in SESSION_FIELDS:
        put(ws, row, 1, name, font=F_LABEL, fill=FILL_LABEL)
        put(ws, row, 2, d.session.get(name, ""), align=A_WRAP)
        row += 1
    kpi = DASHBOARD_KPI_ROWS
    excel_status = ('="Tasks "&DASHBOARD!B{t}&" · done "&DASHBOARD!B{d}&" · verified "&DASHBOARD!B{v}'
                    '&" · in progress "&DASHBOARD!B{p}&" · blocked "&DASHBOARD!B{b}&" · failed "&DASHBOARD!B{f}'
                    '&" · P0 open "&DASHBOARD!B{p0}&" · P1 open "&DASHBOARD!B{p1}').format(
        t=kpi["total"], d=kpi["done"], v=kpi["verified"], p=kpi["in_progress"], b=kpi["blocked"],
        f=kpi["failed"], p0=kpi["p0"], p1=kpi["p1"])
    put(ws, row, 1, "Excel Status", font=F_LABEL, fill=FILL_LABEL)
    put(ws, row, 2, excel_status, align=A_WRAP)
    row += 1
    for name in SESSION_EXTRA:
        put(ws, row, 1, name, font=F_LABEL, fill=FILL_LABEL)
        put(ws, row, 2, d.session.get(name, ""), align=A_WRAP)
        row += 1
    apply_widths(ws, {1: 26, 2: 120})
    ws.freeze_panes = f"A{SESSION_FIRST_ROW}"
    page_setup(ws)


def write_findings_sheet(ws, d: Data) -> None:
    title_block(ws, "CRITICAL_FIXES")
    ws.cell(2, 1).value = (f"{SHEET_INFO['CRITICAL_FIXES'][1]} Report: {AUDIT_REPORT} — "
                           "a finding is RESOLVED when every linked fix task is DONE.")
    widths = {}
    top = 7
    first, last = top + 1, top + len(d.findings)
    labels = ["Findings", "Resolved", "In progress", "Blocked", "Open", "P0 tasks", "P0 done", "P0 remaining"]
    rng = f"B{first}:B{last}"
    values = [f"=COUNTA(A{first}:A{last})", f'=COUNTIF({rng},"RESOLVED")', f'=COUNTIF({rng},"IN_PROGRESS")',
              f'=COUNTIF({rng},"BLOCKED")', f'=COUNTIF({rng},"OPEN")', f'=COUNTIF({mref("priority")},"P0")',
              f'=COUNTIFS({mref("priority")},"P0",{mref("status")},"DONE")', "=F4-G4"]
    for col, (label, value) in enumerate(zip(labels, values), 1):
        put(ws, 3, col, label, font=F_LABEL, fill=FILL_LABEL)
        put(ws, 4, col, value, font=F_KPI, fill=FILL_KPI, align=A_CENTER, fmt="0")
    ws.cell(top - 1, 1, "Audit findings (current known risks)").font = F_SECTION
    cols = [("Finding ID", 11), ("Status", 14), ("Progress", 10), ("Finding", 60), ("Tasks", 8), ("Done", 8),
            ("Fix Tasks", 40), ("Active", 8), ("Blocked/Failed", 10), ("Notes", 40)]
    header(ws, top, [c for c, _ in cols])
    for col, (_, width) in enumerate(cols, 1):
        need(widths, col, width)
    for row, f in zip(range(first, last + 1), d.findings):
        g = f"$G{row}"
        put(ws, row, 1, f["id"], font=F_BOLD)
        put(ws, row, 2, f'=IF(E{row}=0,"OPEN",IF(F{row}=E{row},"RESOLVED",IF(I{row}>0,"BLOCKED",'
                        f'IF(H{row}+F{row}>0,"IN_PROGRESS","OPEN"))))')
        put(ws, row, 3, f"=IF(E{row}=0,0,F{row}/E{row})", fmt=PCT_FMT, align=A_CENTER)
        put(ws, row, 4, f["finding"], align=A_WRAP)
        put(ws, row, 5, list_total(g), align=A_CENTER)
        put(ws, row, 6, "=" + linked_count(g, status_is("DONE")), align=A_CENTER)
        put(ws, row, 7, f["tasks"], align=A_WRAP)
        put(ws, row, 8, "=" + linked_count(g, status_is("IN_PROGRESS", "IMPLEMENTED", "TESTING", "VERIFIED")),
            align=A_CENTER)
        put(ws, row, 9, "=" + linked_count(g, status_is("BLOCKED", "FAILED")), align=A_CENTER)
        put(ws, row, 10, f["notes"], align=A_WRAP)
    add_status_cf(ws, rng, names=["RESOLVED", "IN_PROGRESS", "BLOCKED", "OPEN"])
    ws.conditional_formatting.add(f"C{first}:C{last}", DataBarRule(
        start_type="num", start_value=0, end_type="num", end_value=1, color="63BE7B"))
    p0 = [t["id"] for t in d.tasks if t["priority"] == "P0"]
    task_view(ws, last + 3, p0, "P0 tasks (live from MASTER)", widths)
    apply_widths(ws, widths)
    ws.freeze_panes = "A5"
    page_setup(ws)


def write_gates(ws, d: Data) -> None:
    title_block(ws, "RELEASE_GATES")
    top = 6
    first, last = top + 1, top + len(d.gates)
    rng = f"C{first}:C{last}"
    labels = ["Gates", "PASS", "FAIL", "NOT_EVALUATED"]
    values = [f"=COUNTA(A{first}:A{last})", f'=COUNTIF({rng},"PASS")', f'=COUNTIF({rng},"FAIL")',
              f'=COUNTIF({rng},"NOT_EVALUATED")']
    for col, (label, value) in enumerate(zip(labels, values), 1):
        put(ws, 3, col, label, font=F_LABEL, fill=FILL_LABEL)
        put(ws, 4, col, value, font=F_KPI, fill=FILL_KPI, align=A_CENTER, fmt="0")
    cols = [("Gate ID", 11), ("Gate", 40), ("Status", 15), ("Linked Tasks", 34), ("Linked Done", 12),
            ("Linked Total", 12), ("Evidence", 60), ("Last Verified", 17), ("Notes", 40)]
    header(ws, top, [c for c, _ in cols])
    for row, g in zip(range(first, last + 1), d.gates):
        put(ws, row, 1, g["id"], font=F_BOLD)
        put(ws, row, 2, g["gate"], align=A_WRAP)
        put(ws, row, 3, g["status"] or "NOT_EVALUATED")
        put(ws, row, 4, g["linked"], align=A_WRAP)
        put(ws, row, 5, "=" + linked_count(f"$D{row}", status_is("DONE")), align=A_CENTER)
        put(ws, row, 6, list_total(f"$D{row}"), align=A_CENTER)
        put(ws, row, 7, g["evidence"], align=A_WRAP)
        put(ws, row, 8, g["last_verified"] or "", fmt=DATE_FMT)
        put(ws, row, 9, g["notes"], align=A_WRAP)
    add_list_dv(ws, rng, GATE_STATUSES, "gate status")
    add_status_cf(ws, rng, names=["PASS", "FAIL"])
    apply_widths(ws, {i: w for i, (_, w) in enumerate(cols, 1)})
    ws.freeze_panes = "A5"
    page_setup(ws)


def write_testing(ws, d: Data) -> None:
    title_block(ws, "TESTING")
    category_summary(ws, SHEET_CATEGORIES["TESTING"])
    widths = {}
    ws.cell(6, 1, "Test runs (oldest first)").font = F_SECTION
    cols = [("Run ID", 10), ("Suite", 16), ("Result", 9), ("Command", 50), ("Passed", 9), ("Failed", 9),
            ("Notes", 40), ("Skipped", 9), ("Date", 16), ("Duration (s)", 11), ("Commit", 10)]
    header(ws, 7, [c for c, _ in cols])
    for col, (_, width) in enumerate(cols, 1):
        need(widths, col, width)
    row = 7
    for run in d.runs:
        row += 1
        for col, (key, _) in enumerate(RUN_COLS, 1):
            put(ws, row, col, run.get(key, ""), font=F_BOLD if key == "id" else F_BODY,
                align=A_WRAP if key in ("command", "notes") else A_TOP,
                fmt=DATE_FMT if key == "date" else None)
    if d.runs:
        add_list_dv(ws, f"C8:C{row}", RESULTS, "result")
        add_status_cf(ws, f"C8:C{row}", names=["PASS", "FAIL", "PARTIAL"])
    else:
        row += 1
        put(ws, row, 2, "No test runs recorded yet.", font=F_SUB, border=False)
    task_view(ws, row + 3, view_ids(d, "TESTING"), "Testing tasks (live from MASTER)", widths)
    apply_widths(ws, widths)
    ws.freeze_panes = "A5"
    page_setup(ws)


def write_performance(ws, d: Data) -> None:
    title_block(ws, "PERFORMANCE")
    category_summary(ws, SHEET_CATEGORIES["PERFORMANCE"])
    widths = {}
    ws.cell(6, 1, "Measured results — brief §80-81: measure, then set regression thresholds relative to "
                  "the baseline (no invented SLOs)").font = F_SECTION
    cols = [("Bench ID", 11), ("Metric", 14), ("Unit", 7), ("Scenario", 50), ("Baseline", 10), ("Status", 16),
            ("Baseline Source", 44), ("Current", 10), ("Measured On", 16), ("Threshold", 10), ("Notes", 36)]
    header(ws, 7, [c for c, _ in cols])
    for col, (_, width) in enumerate(cols, 1):
        need(widths, col, width)
    row = 7
    keys = ["id", "metric", "unit", "scenario", "baseline", None, "source", "current", "measured", "threshold", "notes"]
    for bench in d.benches:
        row += 1
        for col, key in enumerate(keys, 1):
            if key is None:
                value = (f'=IF(H{row}="","NOT_REMEASURED",IF(J{row}="","NO_THRESHOLD",'
                         f'IF(H{row}<=J{row},"WITHIN","OVER")))')
            else:
                value = bench.get(key)
                value = "" if value is None else value
            put(ws, row, col, value, font=F_BOLD if key == "id" else F_BODY,
                align=A_WRAP if key in ("scenario", "source", "notes") else A_TOP,
                fmt=DATE_FMT if key == "measured" else None)
    if d.benches:
        add_status_cf(ws, f"F8:F{row}", names=["WITHIN", "OVER"])
    task_view(ws, row + 3, view_ids(d, "PERFORMANCE"), "Performance tasks (live from MASTER)", widths)
    apply_widths(ws, widths)
    ws.freeze_panes = "A5"
    page_setup(ws)


def write_billing(ws, d: Data) -> None:
    title_block(ws, "BILLING_PLANS")
    category_summary(ws, SHEET_CATEGORIES["BILLING_PLANS"])
    widths = {}
    ws.cell(6, 1, "Initial plan hypothesis — brief §64: configurable defaults, not final prices; confirm "
                  "against measured AI and infrastructure cost. Blue = input.").font = F_SECTION
    cols = [("Usage unit / feature", 34), ("FREE", 16), ("PRO", 20), ("BUSINESS", 16), ("Notes", 40)]
    header(ws, 7, [c for c, _ in cols])
    for col, (_, width) in enumerate(cols, 1):
        need(widths, col, width)
    row = 7
    for plan in d.plans:
        row += 1
        put(ws, row, 1, plan["unit"], font=F_BOLD)
        for col, key in ((2, "free"), (3, "pro"), (4, "business")):
            value = plan[key]
            put(ws, row, col, value, font=F_INPUT, align=A_CENTER,
                fmt="#,##0" if isinstance(value, (int, float)) else None)
        put(ws, row, 5, plan["notes"], align=A_WRAP)
    task_view(ws, row + 3, view_ids(d, "BILLING_PLANS"), "Billing and plan tasks (live from MASTER)", widths)
    apply_widths(ws, widths)
    ws.freeze_panes = "A5"
    page_setup(ws)


def write_product(ws, d: Data) -> None:
    title_block(ws, "PRODUCT_FEATURES")
    category_summary(ws, SHEET_CATEGORIES["PRODUCT_FEATURES"])
    widths = {}
    ws.cell(6, 1, "Feature backlog — brief §61-62").font = F_SECTION
    cols = [("Feature", 44), ("Bucket", 19), ("Linked Tasks", 34), ("Linked Done", 12), ("Linked Total", 12),
            ("Notes", 50)]
    header(ws, 7, [c for c, _ in cols])
    for col, (_, width) in enumerate(cols, 1):
        need(widths, col, width)
    row = 7
    for item in d.backlog:
        row += 1
        put(ws, row, 1, item["feature"], font=F_BOLD)
        put(ws, row, 2, item["bucket"])
        put(ws, row, 3, item["tasks"], align=A_WRAP)
        put(ws, row, 4, "=" + linked_count(f"$C{row}", status_is("DONE")), align=A_CENTER)
        put(ws, row, 5, list_total(f"$C{row}"), align=A_CENTER)
        put(ws, row, 6, item["notes"], align=A_WRAP)
    if d.backlog:
        add_list_dv(ws, f"B8:B{row}", BUCKETS, "bucket")
        add_status_cf(ws, f"B8:B{row}", names=["NOW"])
    task_view(ws, row + 3, view_ids(d, "PRODUCT_FEATURES"), "Product tasks (live from MASTER)", widths)
    apply_widths(ws, widths)
    ws.freeze_panes = "A5"
    page_setup(ws)


def write_category_sheet(ws, d: Data, name: str) -> None:
    title_block(ws, name)
    cats = SHEET_CATEGORIES[name]
    ws.cell(2, 1).value = (f"{SHEET_INFO[name][1]} Categories: {', '.join(cats)}. Rows are formulas over MASTER — "
                           "change task data in MASTER or with tools/tracker/tracker.py.")
    category_summary(ws, cats)
    widths = {}
    task_view(ws, 6, view_ids(d, name), "Tasks (live from MASTER)", widths)
    apply_widths(ws, widths)
    ws.freeze_panes = "A5"
    page_setup(ws)


# DASHBOARD layout: fixed rows so SESSION_STATE can reference the KPI cells.
DASHBOARD_KPIS = [
    ("total", "Total tasks", "=COUNTIF({id},\"?*\")", "Rows in MASTER"),
    ("done", "Completed (DONE)", "=COUNTIF({st},\"DONE\")", "Implemented + tested + verified, commit recorded"),
    ("verified", "Verified (VERIFIED + DONE)", "=COUNTIF({st},\"VERIFIED\")+B{done}", "Tests pass and evidence is recorded"),
    ("in_progress", "In progress", "=COUNTIF({st},\"IN_PROGRESS\")+COUNTIF({st},\"IMPLEMENTED\")+COUNTIF({st},\"TESTING\")",
     "IN_PROGRESS + IMPLEMENTED + TESTING"),
    ("blocked", "Blocked", "=COUNTIF({st},\"BLOCKED\")", "Reason in MASTER Notes"),
    ("failed", "Failed", "=COUNTIF({st},\"FAILED\")", "See Notes and CHANGE_LOG"),
    ("deferred", "Deferred", "=COUNTIF({st},\"DEFERRED\")", "Postponed on purpose; still counted as remaining"),
    ("not_started", "Not started", "=COUNTIF({st},\"NOT_STARTED\")", ""),
    ("remaining", "Remaining (not DONE)", "=B{total}-B{done}", "Everything that is not DONE"),
    ("completion", "Completion %", "=IF(B{total}=0,0,B{done}/B{total})", "DONE / total"),
    ("verified_pct", "Verified %", "=IF(B{total}=0,0,B{verified}/B{total})", "(VERIFIED + DONE) / total"),
    ("p0", "P0 remaining", "=COUNTIFS({pr},\"P0\",{st},\"<>DONE\")", "P0 tasks not DONE"),
    ("p1", "P1 remaining", "=COUNTIFS({pr},\"P1\",{st},\"<>DONE\")", "P1 tasks not DONE"),
    ("findings", "Audit findings resolved", "=CRITICAL_FIXES!B4", "Of 20; see CRITICAL_FIXES"),
    ("gates_pass", "Release gates passed", "=RELEASE_GATES!B4", "Of 15; see RELEASE_GATES"),
    ("gates_fail", "Release gates failed", "=RELEASE_GATES!C4", ""),
]
DASHBOARD_KPI_ROWS = {key: 5 + i for i, (key, *_rest) in enumerate(DASHBOARD_KPIS)}
DASHBOARD_NEXT = ["Current Phase", "Current Task", "Last Completed", "Last Verified", "Next Task", "Next Command",
                  "Next Test", "Tests Passed", "Tests Failed", "Open Blockers", "Checkpoint Time"]


def write_dashboard(ws, d: Data, written_at: dt.datetime) -> None:
    ws.cell(1, 1, "SmartDoc — Master Implementation Tracker").font = F_TITLE
    ws.cell(2, 1, "Production-hardening brief. MASTER holds the task data; every number here is a formula. "
                  f"Updated with tools/tracker/tracker.py. Last written {written_at:%Y-%m-%d %H:%M} (local time)."
             ).font = F_SUB
    refs = {"id": mref("id"), "st": mref("status"), "pr": mref("priority"), **DASHBOARD_KPI_ROWS}

    # Overall (A-C)
    ws.cell(4, 1, "Overall").font = F_SECTION
    for key, label, formula, note in DASHBOARD_KPIS:
        row = DASHBOARD_KPI_ROWS[key]
        put(ws, row, 1, label, font=F_LABEL, fill=FILL_LABEL)
        put(ws, row, 2, formula.format(**refs), font=F_KPI, fill=FILL_KPI, align=A_CENTER,
            fmt=PCT_FMT if label.endswith("%") else "0")
        put(ws, row, 3, note, font=F_SUB)
    for key in ("blocked", "failed", "p0", "gates_fail"):
        cell = f"B{DASHBOARD_KPI_ROWS[key]}"
        ws.conditional_formatting.add(cell, CellIsRule(operator="greaterThan", formula=["0"],
                                                       fill=_dxf_fill("FFC7CE"), font=Font(color="9C0006", bold=True)))
    row = 5 + len(DASHBOARD_KPIS) + 1

    # Tasks by status (A-C)
    ws.cell(row, 1, "Tasks by status").font = F_SECTION
    header(ws, row + 1, ["Status", "Tasks", "Meaning"])
    first = row + 2
    for i, (status, meaning) in enumerate(seed.STATUS_MEANINGS):
        r = first + i
        put(ws, r, 1, status, font=F_BOLD)
        put(ws, r, 2, f'=COUNTIF({mref("status")},A{r})', align=A_CENTER, fmt="0")
        put(ws, r, 3, meaning)
    add_status_cf(ws, f"A{first}:A{first + len(STATUSES) - 1}", names=STATUSES)
    row = first + len(STATUSES) + 1

    # Latest test runs (A-C)
    ws.cell(row, 1, "Latest test runs (TESTING)").font = F_SECTION
    header(ws, row + 1, ["Suite", "Result", "Passed · failed · skipped"])
    row += 2
    latest = d.runs[-6:]
    run_first_row = 8  # first data row of the TESTING runs table
    if not latest:
        put(ws, row, 1, "No test runs recorded yet.", font=F_SUB, border=False)
        row += 1
    for run in reversed(latest):
        r_testing = run_first_row + d.runs.index(run)
        put(ws, row, 1, f'=TESTING!B{r_testing}&" — "&TESTING!A{r_testing}', font=F_BOLD)
        put(ws, row, 2, f'=TESTING!C{r_testing}&""', align=A_CENTER)
        put(ws, row, 3, f'=TESTING!E{r_testing}&" passed · "&TESTING!F{r_testing}&" failed · "'
                        f'&TESTING!H{r_testing}&" skipped"')
        row += 1
    if latest:
        add_status_cf(ws, f"B{row - len(latest)}:B{row - 1}", names=["PASS", "FAIL", "PARTIAL"])
    row += 1

    # Sheets (A-C)
    ws.cell(row, 1, "Sheets").font = F_SECTION
    row += 1
    for name in SHEET_ORDER:
        cell = put(ws, row, 1, name, font=F_LINK)
        cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{name}'!A1", display=name)
        put(ws, row, 2, SHEET_INFO[name][1], border=False)
        row += 1
    row += 1

    # How to read and update (A)
    ws.cell(row, 1, "How to read and update").font = F_SECTION
    notes = [
        "Status is always text; colours only highlight it. Completed = DONE. Verified = VERIFIED + DONE.",
        "In progress = IN_PROGRESS + IMPLEMENTED + TESTING. Remaining = every task that is not DONE (DEFERRED included).",
        "DONE needs evidence, tests and a commit; VERIFIED needs evidence and tests (never fake progress).",
        "Change task data in MASTER, or run: python tools/tracker/tracker.py set <ID> --status <STATUS> ... (tools/tracker/README.md).",
        "Other sheets are rebuilt by the tool on every write. Input tables that are kept: CRITICAL_FIXES fix tasks and notes, "
        "RELEASE_GATES, PERFORMANCE measurements, TESTING runs, BILLING_PLANS values, PRODUCT_FEATURES backlog, "
        "CHANGE_LOG, SESSION_STATE values.",
        f"Resume a session with: {RESUME}",
    ]
    for text in notes:
        row += 1
        put(ws, row, 1, text, border=False)

    # Next action (E-F)
    ws.cell(4, 5, "Next action (SESSION_STATE)").font = F_SECTION
    for i, name in enumerate(DASHBOARD_NEXT):
        r = 5 + i
        src = SESSION_FIRST_ROW + SESSION_FIELDS.index(name)
        put(ws, r, 5, name, font=F_LABEL, fill=FILL_LABEL)
        put(ws, r, 6, f'=IF(SESSION_STATE!$B${src}="","",SESSION_STATE!$B${src})', border=False)

    # Progress by phase (E-N)
    phase_top = 5 + len(DASHBOARD_NEXT) + 1
    ws.cell(phase_top, 5, "Progress by phase").font = F_SECTION
    cols = ["Phase", "Name", "Tasks", "Done", "Verified", "In progress", "Blocked", "Failed", "Done %", "Verified %"]
    header(ws, phase_top + 1, cols, col=5)
    first = phase_top + 2
    for i, (number, name) in enumerate(sorted(seed.PHASES.items())):
        r = first + i
        _progress_row(ws, r, number, name, "phase")
    last = first + len(seed.PHASES) - 1
    _bars(ws, f"M{first}:N{last}")

    # Progress by category (E-N)
    cat_top = last + 2
    ws.cell(cat_top, 5, "Progress by category").font = F_SECTION
    header(ws, cat_top + 1, ["Category", "Sheet", *cols[2:]], col=5)
    first = cat_top + 2
    for i, cat in enumerate(CATEGORIES):
        r = first + i
        _progress_row(ws, r, cat, seed.CATEGORY_SHEETS[cat], "category", link=True)
    _bars(ws, f"M{first}:N{first + len(CATEGORIES) - 1}")

    apply_widths(ws, {1: 34, 2: 13, 3: 50, 4: 3, 5: 16, 6: 36, 7: 9, 8: 9, 9: 10, 10: 11, 11: 9, 12: 9, 13: 10, 14: 11})
    ws.freeze_panes = "A4"
    page_setup(ws)


def _progress_row(ws, r: int, key, label: str, field: str, link: bool = False) -> None:
    rng, st = mref(field), mref("status")
    put(ws, r, 5, key, font=F_BOLD, align=A_CENTER if field == "phase" else A_TOP)
    cell = put(ws, r, 6, label, font=F_LINK if link else F_BODY)
    if link:
        cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{label}'!A1", display=label)
    put(ws, r, 7, f"=COUNTIF({rng},$E{r})", align=A_CENTER, fmt="0")
    put(ws, r, 8, f'=COUNTIFS({rng},$E{r},{st},"DONE")', align=A_CENTER, fmt="0")
    put(ws, r, 9, f'=COUNTIFS({rng},$E{r},{st},"VERIFIED")+H{r}', align=A_CENTER, fmt="0")
    put(ws, r, 10, "=" + "+".join(f'COUNTIFS({rng},$E{r},{st},"{s}")' for s in ACTIVE), align=A_CENTER, fmt="0")
    put(ws, r, 11, f'=COUNTIFS({rng},$E{r},{st},"BLOCKED")', align=A_CENTER, fmt="0")
    put(ws, r, 12, f'=COUNTIFS({rng},$E{r},{st},"FAILED")', align=A_CENTER, fmt="0")
    put(ws, r, 13, f"=IF(G{r}=0,0,H{r}/G{r})", align=A_CENTER, fmt=PCT_FMT)
    put(ws, r, 14, f"=IF(G{r}=0,0,I{r}/G{r})", align=A_CENTER, fmt=PCT_FMT)


def _bars(ws, rng: str) -> None:
    ws.conditional_formatting.add(rng, DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1,
                                                   color="63BE7B"))


WRITERS = {
    "MASTER": write_master,
    "CHANGE_LOG": write_changelog,
    "SESSION_STATE": write_session,
    "CRITICAL_FIXES": write_findings_sheet,
    "RELEASE_GATES": write_gates,
    "TESTING": write_testing,
    "PERFORMANCE": write_performance,
    "BILLING_PLANS": write_billing,
    "PRODUCT_FEATURES": write_product,
}
TAB_COLORS = {"DASHBOARD": NAVY, "MASTER": NAVY, "CRITICAL_FIXES": "C00000", "CHANGE_LOG": "7F7F7F",
              "SESSION_STATE": "7F7F7F", "RELEASE_GATES": "548235"}


def render(wb, d: Data) -> None:
    d.validate()
    written_at = now()
    for index, name in enumerate(SHEET_ORDER):
        if name in wb.sheetnames:
            wb.remove(wb[name])
        ws = wb.create_sheet(name, index)
        if name in TAB_COLORS:
            ws.sheet_properties.tabColor = TAB_COLORS[name]
        if name == "DASHBOARD":
            write_dashboard(ws, d, written_at)
        elif name in WRITERS:
            WRITERS[name](ws, d)
        else:
            write_category_sheet(ws, d, name)
    wb.active = 0
    wb.calculation.fullCalcOnLoad = True
    strip_sensitivity_labels(wb)


def strip_sensitivity_labels(wb) -> None:
    """Excel on this machine stamps Microsoft sensitivity-label properties (MSIP_Label_*, with the
    organisation's tenant ID) into docProps/custom.xml; the repository is public, so they never ship."""
    props = wb.custom_doc_props
    props.props = [p for p in props.props if not p.name.startswith("MSIP_Label_")]


def lock_file(path: Path) -> Path | None:
    for candidate in (path.with_name("~$" + path.name), path.with_name("~$" + path.name[2:])):
        if candidate.exists():
            return candidate
    return None


def save(wb, path: Path) -> None:
    lock = lock_file(path)
    if lock:
        raise TrackerError(f"{path.name} is open in Excel ({lock.name} exists); close it and run the command again")
    tmp = path.with_name(f".{path.stem}.tmp-{os.getpid()}.xlsx")
    wb.save(tmp)
    # Windows briefly locks a file that was just written (indexing, antivirus): retry.
    for attempt in range(8):
        try:
            os.replace(tmp, path)
            return
        except PermissionError:
            time.sleep(0.25 * (attempt + 1))
    tmp.unlink(missing_ok=True)
    raise TrackerError(f"cannot replace {path.name}; is it open in Excel?")


# -- commands ---------------------------------------------------------------------------

def git(*args: str) -> str:
    try:
        return subprocess.run(["git", "-C", str(ROOT), *args], capture_output=True, text=True, check=True,
                              encoding="utf-8").stdout.strip()
    except (OSError, subprocess.CalledProcessError):
        return ""


def append_log(d: Data, tasks: str, change: str, files="", tests="", result="", commit="", notes="", phase=None):
    ids = split_ids(tasks)
    for tid in ids:
        if tid != "-" and not d.has(tid):
            raise TrackerError(f"CHANGE_LOG: unknown task {tid}")
    if phase is None:
        phase = d.task(ids[0])["phase"] if ids and d.has(ids[0]) else ""
    if result and result not in RESULTS:
        raise TrackerError(f"result must be one of {', '.join(RESULTS)}")
    if len(change.strip()) < 25:
        raise TrackerError("describe the change properly (what changed, which test proves it), not just 'fixed bug'")
    d.changelog.append(dict(timestamp=now(), task=", ".join(ids), phase=phase, change=change.strip(), files=files,
                            tests=tests, result=result or "N/A", commit=commit, notes=notes))


def cmd_init(args) -> None:
    path = args.workbook
    if path.exists():
        raise TrackerError(f"{path.name} already exists; use `render` to rebuild its sheets")
    d = from_seed()
    wb = Workbook()
    wb.remove(wb.active)
    append_log(d, "INFRA-001", f"Created {path.name} from the 2026-09-26 audit findings and the brief's phases: "
                               f"{len(d.tasks)} tasks, {len(d.findings)} findings, {len(d.gates)} release gates.",
               files="tools/tracker/tracker.py, tools/tracker/seed.py", result="N/A")
    render(wb, d)
    save(wb, path)
    print(f"created {path} with {len(d.tasks)} tasks")


def cmd_render(args) -> None:
    wb, d = load(args.workbook)
    render(wb, d)
    save(wb, args.workbook)
    print(f"rewrote {args.workbook.name}")


def cmd_set(args) -> None:
    wb, d = load(args.workbook)
    t = d.task(args.id)
    stamp = now()
    for key, value in (("evidence", args.evidence), ("tests", args.tests), ("commit", args.commit),
                       ("notes", args.notes), ("owner", args.owner), ("deps", args.deps), ("title", args.title),
                       ("description", args.description)):
        if value is not None:
            t[key] = value
    for key, value in (("evidence", args.add_evidence), ("tests", args.add_tests), ("notes", args.add_notes)):
        if value:
            t[key] = f"{t[key]}; {value}" if t[key] else value
    if args.priority:
        t["priority"] = args.priority
    if args.risk:
        t["risk"] = args.risk
    if args.status:
        status = args.status
        if status in ("VERIFIED", "DONE") and not (t["evidence"] and t["tests"]):
            raise TrackerError(f"{t['id']}: {status} needs --evidence and --tests (never fake progress)")
        if status == "DONE" and not t["commit"]:
            raise TrackerError(f"{t['id']}: DONE needs --commit")
        if status in NEEDS_REASON and not t["notes"]:
            raise TrackerError(f"{t['id']}: {status} needs a reason in --notes")
        old = t["status"]
        t["status"] = status
        if status not in ("NOT_STARTED", "DEFERRED", "UNKNOWN") and not t["started"]:
            t["started"] = stamp
        if status == "DONE":
            if old != "DONE" or not t["completed"]:
                t["completed"] = stamp
        else:
            t["completed"] = None
        if status in ("VERIFIED", "DONE"):
            t["last_verified"] = stamp
    if args.verified:
        if not t["tests"]:
            raise TrackerError(f"{t['id']}: --verified needs the tests that were run (--tests)")
        t["last_verified"] = stamp
    if args.log:
        derived = {"VERIFIED": "PASS", "DONE": "PASS", "FAILED": "FAIL"}.get(t["status"], "N/A")
        append_log(d, t["id"], args.log, files=args.files or "", tests=args.tests or t["tests"],
                   result=args.result or derived, commit=t["commit"])
    render(wb, d)
    save(wb, args.workbook)
    print(f"{t['id']}: {t['status']}")


def _descends_from(d: Data, t: dict, ancestor: str) -> bool:
    seen = set()
    while t["parent"] and t["parent"] not in seen:
        if t["parent"] == ancestor:
            return True
        seen.add(t["parent"])
        t = d.task(t["parent"])
    return False


def cmd_add(args) -> None:
    wb, d = load(args.workbook)
    tid = args.id
    if not ID_RE.match(tid):
        raise TrackerError(f"{tid}: IDs look like CORE-001 or DOCX-101A")
    if d.has(tid):
        raise TrackerError(f"{tid} already exists (IDs are never reused)")
    parent = d.task(args.parent) if args.parent else None
    phase = args.phase if args.phase is not None else (parent["phase"] if parent else None)
    category = args.category or (parent["category"] if parent else None)
    if phase is None or category is None:
        raise TrackerError("--phase and --category are required without --parent")
    t = dict(id=tid, phase=phase, category=category, title=args.title, description=args.description or "",
             priority=args.priority or (parent["priority"] if parent else "P2"), status="NOT_STARTED", evidence="",
             tests="", owner=args.owner or (parent["owner"] if parent else ""), deps=args.deps or "",
             risk=args.risk or (parent["risk"] if parent else "Medium"), started=None, completed=None,
             last_verified=None, commit="", notes=args.notes or "", parent=args.parent or "",
             audit=args.audit or (parent["audit"] if parent else ""))
    if parent:
        index = max(i for i, x in enumerate(d.tasks) if x["id"] == parent["id"] or _descends_from(d, x, parent["id"])) + 1
    else:
        same = [i for i, x in enumerate(d.tasks) if x["phase"] == phase]
        later = [i for i, x in enumerate(d.tasks) if x["phase"] > phase]
        index = same[-1] + 1 if same else (later[0] if later else len(d.tasks))
    d.tasks.insert(index, t)
    render(wb, d)
    save(wb, args.workbook)
    print(f"added {tid} (phase {phase}, {category})")


def cmd_log(args) -> None:
    wb, d = load(args.workbook)
    append_log(d, args.tasks, args.change, files=args.files or "", tests=args.tests or "", result=args.result or "",
               commit=args.commit or "", notes=args.notes or "", phase=args.phase)
    render(wb, d)
    save(wb, args.workbook)
    print(f"logged change #{len(d.changelog)}")


def cmd_testrun(args) -> None:
    wb, d = load(args.workbook)
    if args.result not in RESULTS:
        raise TrackerError(f"result must be one of {', '.join(RESULTS)}")
    run_id = f"RUN-{len(d.runs) + 1:04d}"
    d.runs.append(dict(id=run_id, date=now(), suite=args.suite, command=args.command, result=args.result,
                       passed=args.passed, failed=args.failed, skipped=args.skipped,
                       duration=args.duration if args.duration is not None else "",
                       commit=args.commit or git("rev-parse", "--short", "HEAD"), notes=args.notes or ""))
    render(wb, d)
    save(wb, args.workbook)
    print(f"{run_id}: {args.suite} {args.result} ({args.passed} passed, {args.failed} failed)")


def cmd_bench(args) -> None:
    wb, d = load(args.workbook)
    bench = next((b for b in d.benches if b["id"] == args.id), None)
    if bench is None:
        if not (args.scenario and args.metric and args.unit):
            raise TrackerError(f"{args.id} is new: give --scenario, --metric, --unit (and --baseline, --source)")
        bench = dict(id=args.id, scenario=args.scenario, metric=args.metric, unit=args.unit, baseline=args.baseline,
                     source=args.source or "", current=None, measured=None, threshold=None, notes="")
        d.benches.append(bench)
    if args.current is not None:
        bench["current"] = args.current
        bench["measured"] = now()
    if args.threshold is not None:
        bench["threshold"] = args.threshold
    if args.notes is not None:
        bench["notes"] = args.notes
    render(wb, d)
    save(wb, args.workbook)
    print(f"{args.id}: current={bench['current']} threshold={bench['threshold']}")


def cmd_gate(args) -> None:
    wb, d = load(args.workbook)
    gate = next((g for g in d.gates if g["id"] == args.id), None)
    if gate is None:
        raise TrackerError(f"unknown gate {args.id}")
    if args.status in ("PASS", "FAIL") and not (args.evidence or gate["evidence"]):
        raise TrackerError(f"{args.id}: {args.status} needs --evidence")
    gate["status"] = args.status
    if args.evidence:
        gate["evidence"] = args.evidence
    if args.notes is not None:
        gate["notes"] = args.notes
    gate["last_verified"] = now()
    render(wb, d)
    save(wb, args.workbook)
    print(f"{args.id}: {args.status}")


def cmd_finding(args) -> None:
    wb, d = load(args.workbook)
    finding = next((f for f in d.findings if f["id"] == args.id), None)
    if finding is None:
        raise TrackerError(f"unknown finding {args.id}")
    if args.tasks is not None:
        for tid in split_ids(args.tasks):
            d.task(tid)
        finding["tasks"] = args.tasks
    if args.notes is not None:
        finding["notes"] = args.notes
    render(wb, d)
    save(wb, args.workbook)
    print(f"{args.id}: {finding['tasks']}")


def _join(*parts) -> str:
    return " — ".join(str(p) for p in parts if p not in (None, "", [], {}))


def session_from_json(state: dict) -> dict:
    phase = state.get("current_phase") or {}
    task = state.get("current_task") or {}
    last = state.get("last_completed_task") or {}
    verified = state.get("last_verified") or {}
    results = state.get("test_results") or {}
    nxt = state.get("next_exact_task") or {}
    commit = state.get("current_commit", "")
    if state.get("working_tree_clean") is False:
        commit = f"{commit} (+ uncommitted changes)"
    verified_text = _join(verified.get("task"), verified.get("what"), verified.get("result"), verified.get("at"))
    if verified.get("command"):
        verified_text += f" · command: {verified['command']}"
    return {
        "Current Phase": _join(phase.get("number"), phase.get("name"), phase.get("status")),
        "Current Task": _join(task.get("id"), task.get("details")),
        "Last Completed": _join(last.get("id"), last.get("summary")),
        "Last Verified": verified_text,
        "Current Branch": state.get("current_branch", ""),
        "Current Commit": commit,
        "Tests Passed": "; ".join(results.get("passed") or []) or "—",
        "Tests Failed": "; ".join(results.get("failed") or []) or "none",
        "Open Blockers": "; ".join(state.get("known_blockers") or []) or "none",
        "Next Task": _join(nxt.get("id"), nxt.get("description")),
        "Next Command": nxt.get("command", ""),
        "Next Test": nxt.get("test", ""),
        "Checkpoint Time": state.get("last_checkpoint", ""),
        "Resume Instructions": state.get("resume_instructions") or RESUME,
        "Production Code Changed": "yes" if state.get("production_code_changed") else "no",
        "DB Migrations Applied": "yes" if state.get("database_migrations_applied") else "no",
    }


def cmd_state(args) -> None:
    path = args.json
    if not path.exists():
        raise TrackerError(f"{path} does not exist")
    state = json.loads(path.read_text(encoding="utf-8"))
    wb, d = load(args.workbook)
    if not args.no_git:
        branch, commit = git("rev-parse", "--abbrev-ref", "HEAD"), git("rev-parse", "--short", "HEAD")
        if branch:
            state["current_branch"] = branch
            state["current_commit"] = commit
            state["working_tree_clean"] = git("status", "--porcelain") == ""
    counts = d.counts()
    state["current_excel_status"] = {
        "workbook": args.workbook.name, **{k: v for k, v in counts.items() if k != "by_status"},
        "by_status": counts["by_status"],
        "phases": {str(p): f"{sum(t['status'] == 'DONE' for t in d.tasks if t['phase'] == p)}/"
                           f"{sum(t['phase'] == p for t in d.tasks)} done" for p in sorted(seed.PHASES)},
        "findings_resolved": sum(1 for f in d.findings
                                 if split_ids(f["tasks"]) and all(d.has(x) and d.task(x)["status"] == "DONE"
                                                                   for x in split_ids(f["tasks"]))),
        "gates_passed": sum(g["status"] == "PASS" for g in d.gates),
    }
    if args.touch:
        state["last_checkpoint"] = dt.datetime.now().astimezone().replace(microsecond=0).isoformat()
    path.write_text(json.dumps(state, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    d.session.update(session_from_json(state))
    render(wb, d)
    save(wb, args.workbook)
    print(f"SESSION_STATE updated from {path.name}; excel status: {counts['done']}/{counts['total']} done, "
          f"{counts['verified']} verified, P0 open {counts['p0_remaining']}")


def cmd_show(args) -> None:
    _, d = load(args.workbook)
    c = d.counts()
    print(f"Tasks {c['total']} · DONE {c['done']} · verified {c['verified']} · in progress {c['in_progress']} · "
          f"blocked {c['blocked']} · failed {c['failed']} · remaining {c['remaining']} · "
          f"P0 open {c['p0_remaining']} · P1 open {c['p1_remaining']}")
    for p, name in sorted(seed.PHASES.items()):
        tasks = [t for t in d.tasks if t["phase"] == p]
        done = sum(t["status"] == "DONE" for t in tasks)
        active = sum(t["status"] in (*ACTIVE, "VERIFIED") for t in tasks)
        print(f"  phase {p:>2} {name:<45} {done:>3}/{len(tasks):<3} done  {active:>3} active")
    rows = d.tasks
    if args.phase is not None:
        rows = [t for t in rows if t["phase"] == args.phase]
    if args.status:
        rows = [t for t in rows if t["status"] == args.status]
    if args.priority:
        rows = [t for t in rows if t["priority"] == args.priority]
    if args.open:
        rows = [t for t in rows if t["status"] not in ("DONE", "DEFERRED")]
    if any((args.phase is not None, args.status, args.priority, args.open)):
        for t in rows:
            print(f"  {t['id']:<11} {t['priority']} {t['status']:<12} {t['title']}")


def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(description="SmartDoc master implementation tracker")
    p.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    sub = p.add_subparsers(dest="cmd", required=True)

    sub.add_parser("init", help="create the workbook from seed.py")
    sub.add_parser("render", help="rewrite the managed sheets without changing data")

    s = sub.add_parser("set", help="update a task")
    s.add_argument("id")
    s.add_argument("--status", choices=STATUSES)
    s.add_argument("--evidence")
    s.add_argument("--add-evidence")
    s.add_argument("--tests")
    s.add_argument("--add-tests")
    s.add_argument("--commit")
    s.add_argument("--notes")
    s.add_argument("--add-notes")
    s.add_argument("--owner")
    s.add_argument("--deps")
    s.add_argument("--title")
    s.add_argument("--description")
    s.add_argument("--priority", choices=PRIORITIES)
    s.add_argument("--risk", choices=RISKS)
    s.add_argument("--verified", action="store_true", help="record a re-verification now")
    s.add_argument("--log", help="also append this CHANGE_LOG entry")
    s.add_argument("--files", help="files for the CHANGE_LOG entry")
    s.add_argument("--result", choices=RESULTS, help="result for the CHANGE_LOG entry")

    a = sub.add_parser("add", help="add a task (IDs are never reused)")
    a.add_argument("id")
    a.add_argument("--title", required=True)
    a.add_argument("--parent")
    a.add_argument("--phase", type=int, choices=range(0, 19))
    a.add_argument("--category", choices=CATEGORIES)
    a.add_argument("--description")
    a.add_argument("--priority", choices=PRIORITIES)
    a.add_argument("--owner")
    a.add_argument("--deps")
    a.add_argument("--risk", choices=RISKS)
    a.add_argument("--audit")
    a.add_argument("--notes")

    lg = sub.add_parser("log", help="append a CHANGE_LOG entry")
    lg.add_argument("tasks", help="task ID(s), comma-separated, or - for none")
    lg.add_argument("change")
    lg.add_argument("--files")
    lg.add_argument("--tests")
    lg.add_argument("--result", choices=RESULTS)
    lg.add_argument("--commit")
    lg.add_argument("--notes")
    lg.add_argument("--phase", type=int)

    t = sub.add_parser("testrun", help="record a test run on TESTING")
    t.add_argument("suite")
    t.add_argument("command")
    t.add_argument("--result", required=True, choices=RESULTS)
    t.add_argument("--passed", type=int, required=True)
    t.add_argument("--failed", type=int, required=True)
    t.add_argument("--skipped", type=int, default=0)
    t.add_argument("--duration", type=float)
    t.add_argument("--commit")
    t.add_argument("--notes")

    b = sub.add_parser("bench", help="record a benchmark measurement on PERFORMANCE")
    b.add_argument("id")
    b.add_argument("--current", type=float)
    b.add_argument("--threshold", type=float)
    b.add_argument("--notes")
    b.add_argument("--scenario")
    b.add_argument("--metric")
    b.add_argument("--unit")
    b.add_argument("--baseline", type=float)
    b.add_argument("--source")

    g = sub.add_parser("gate", help="set a release gate")
    g.add_argument("id")
    g.add_argument("--status", required=True, choices=GATE_STATUSES)
    g.add_argument("--evidence")
    g.add_argument("--notes")

    f = sub.add_parser("finding", help="edit an audit finding's fix tasks or notes")
    f.add_argument("id")
    f.add_argument("--tasks")
    f.add_argument("--notes")

    st = sub.add_parser("state", help="update SESSION_STATE from docs/session-state.json")
    st.add_argument("--json", type=Path, default=DEFAULT_STATE_JSON)
    st.add_argument("--no-git", action="store_true")
    st.add_argument("--touch", action="store_true", help="set last_checkpoint to now")

    sh = sub.add_parser("show", help="print progress")
    sh.add_argument("--phase", type=int)
    sh.add_argument("--status", choices=STATUSES)
    sh.add_argument("--priority", choices=PRIORITIES)
    sh.add_argument("--open", action="store_true")
    return p


COMMANDS = {"init": cmd_init, "render": cmd_render, "set": cmd_set, "add": cmd_add, "log": cmd_log,
            "testrun": cmd_testrun, "bench": cmd_bench, "gate": cmd_gate, "finding": cmd_finding,
            "state": cmd_state, "show": cmd_show}


def main(argv=None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    args = build_parser().parse_args(argv)
    try:
        COMMANDS[args.cmd](args)
    except TrackerError as exc:
        print(f"tracker: {exc}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
