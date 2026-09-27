#!/usr/bin/env python3
"""Self-test for tracker.py: python tools/tracker/selftest.py (exit code 0 = pass).

Works on a throwaway copy in a temporary directory; never touches the real workbook.
"""
from __future__ import annotations

import contextlib
import io
import json
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
import seed  # noqa: E402
import tracker  # noqa: E402

FAILURES: list[str] = []


def check(condition: bool, message: str) -> None:
    if not condition:
        FAILURES.append(message)
        print(f"FAIL {message}")


def run(wb: Path, *args: str) -> int:
    out, err = io.StringIO(), io.StringIO()
    try:
        with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
            return tracker.main(["--workbook", str(wb), *args])
    except SystemExit as exc:  # argparse rejects invalid choices
        return int(exc.code or 0)


def master(wb: Path) -> dict:
    _, d = tracker.load(wb)
    return {t["id"]: t for t in d.tasks}


def main() -> int:
    with tempfile.TemporaryDirectory() as tmp:
        wb = Path(tmp) / "tracker.xlsx"
        check(run(wb, "init") == 0, "init succeeds")
        check(run(wb, "init") == 1, "init refuses to overwrite")
        book = load_workbook(wb)
        check(book.sheetnames == tracker.SHEET_ORDER, f"sheet order {book.sheetnames}")
        required = ["MASTER", "CRITICAL_FIXES", "DOCUMENT_FIDELITY", "DOCX_OOXML", "PDF", "EDITOR", "FORMATTING", "AI",
                    "SECURITY", "PERFORMANCE", "TESTING", "PRODUCT_FEATURES", "TRANSLATION", "PDF_TO_EDITABLE",
                    "BILLING_PLANS", "INFRA", "CHANGE_LOG", "SESSION_STATE", "RELEASE_GATES"]
        check(all(name in book.sheetnames for name in required), "all 19 required sheets exist")
        ms = book["MASTER"]
        headers = [c.value for c in ms[1]]
        spec = ["ID", "Phase", "Category", "Feature/Task", "Description", "Priority", "Status", "Evidence", "Tests",
                "Owner/System", "Dependencies", "Risk", "Started", "Completed", "Last Verified", "Commit", "Notes"]
        check(headers[:17] == spec, f"MASTER columns follow the brief: {headers[:17]}")
        check("Parent ID" in headers, "MASTER has Parent ID")
        check(ms.max_row == len(seed.TASKS) + 1, f"MASTER rows {ms.max_row}")
        dv_ranges = {str(dv.sqref): dv.formula1 for dv in ms.data_validations.dataValidation}
        check(any("NOT_STARTED" in (f or "") and "DONE" in (f or "") for f in dv_ranges.values()),
              "MASTER status validation list")
        check(any('"P0,P1,P2,P3"' == f for f in dv_ranges.values()), "MASTER priority validation list")
        check(len(ms.conditional_formatting) > 0, "MASTER conditional formatting")
        dash = book["DASHBOARD"]
        check(dash["B5"].value == '=COUNTIF(MASTER!$A$2:$A$2000,"?*")', f"dashboard total formula {dash['B5'].value}")
        check(str(book["EDITOR"]["F8"].value).startswith("=IFERROR(INDEX(MASTER!$G$2:$G$2000"), "view status formula")
        check(book["CHANGE_LOG"]["B2"].value == "INFRA-001", "init logged to CHANGE_LOG")

        # status rules
        check(run(wb, "set", "EDIT-001", "--status", "WHATEVER") == 2, "arbitrary status rejected")
        check(run(wb, "set", "EDIT-001", "--status", "VERIFIED") == 1, "VERIFIED without evidence rejected")
        check(run(wb, "set", "EDIT-001", "--status", "BLOCKED") == 1, "BLOCKED without a reason rejected")
        check(run(wb, "set", "EDIT-001", "--status", "IN_PROGRESS") == 0, "IN_PROGRESS accepted")
        t = master(wb)["EDIT-001"]
        check(t["started"] is not None and t["completed"] is None, "started stamped")
        check(run(wb, "set", "EDIT-001", "--status", "VERIFIED", "--evidence", "frontend/editor/x.ts",
                  "--tests", "x.test.ts", "--log", "Mapped lists inside table cells; x.test.ts passes.") == 0,
              "VERIFIED with evidence")
        check(master(wb)["EDIT-001"]["last_verified"] is not None, "last verified stamped")
        check(run(wb, "set", "EDIT-001", "--status", "DONE") == 1, "DONE without commit rejected")
        check(run(wb, "set", "EDIT-001", "--status", "DONE", "--commit", "abc1234") == 0, "DONE with commit")
        t = master(wb)["EDIT-001"]
        check(t["completed"] is not None and t["commit"] == "abc1234", "completed stamped")
        check(run(wb, "set", "EDIT-001", "--status", "FAILED", "--notes", "regressed in e2e") == 0, "FAILED")
        check(master(wb)["EDIT-001"]["completed"] is None, "completed cleared when no longer DONE")

        # add
        check(run(wb, "add", "EDIT-001A", "--parent", "EDIT-001", "--title", "Lists in cells") == 0, "add subtask")
        _, d = tracker.load(wb)
        ids = [x["id"] for x in d.tasks]
        check(ids.index("EDIT-001A") == ids.index("EDIT-001") + 1, "subtask inserted after its parent")
        sub = d.task("EDIT-001A")
        check(sub["phase"] == 1 and sub["category"] == "EDITOR" and sub["parent"] == "EDIT-001", "subtask inherits")
        check(run(wb, "add", "EDIT-001A", "--parent", "EDIT-001", "--title", "again") == 1, "duplicate ID rejected")
        check(run(wb, "add", "edit-9", "--phase", "1", "--category", "EDITOR", "--title", "bad") == 1, "bad ID rejected")

        # log, runs, benches, gates, findings
        check(run(wb, "log", "EDIT-001", "fixed bug") == 1, "vague change log rejected")
        check(run(wb, "log", "EDIT-001,EDIT-002", "Fixed serialization of list nodes inside table cells; "
                  "regression test added.", "--result", "PASS") == 0, "log accepted")
        check(run(wb, "testrun", "backend", "pytest -q", "--result", "PASS", "--passed", "643", "--failed", "0") == 0,
              "testrun")
        check(run(wb, "testrun", "backend", "pytest -q", "--result", "PASS", "--passed", "644", "--failed", "0") == 0,
              "second testrun")
        check(run(wb, "bench", "BENCH-001", "--current", "4.2", "--threshold", "10") == 0, "bench update")
        check(run(wb, "bench", "BENCH-099", "--current", "1") == 1, "new bench needs a description")
        check(run(wb, "gate", "GATE-013", "--status", "PASS") == 1, "gate PASS without evidence rejected")
        check(run(wb, "gate", "GATE-013", "--status", "PASS", "--evidence", "alembic check clean") == 0, "gate PASS")
        check(run(wb, "finding", "AUD-10", "--notes", "table export") == 0, "finding notes")
        _, d = tracker.load(wb)
        check([r["id"] for r in d.runs] == ["RUN-0001", "RUN-0002"], "runs kept in order")
        check(d.benches[0]["current"] == 4.2 and d.benches[0]["threshold"] == 10, "bench values kept")
        check(next(g for g in d.gates if g["id"] == "GATE-013")["status"] == "PASS", "gate kept")
        check(next(f for f in d.findings if f["id"] == "AUD-10")["notes"] == "table export", "finding kept")
        check(len(d.changelog) == 3, f"changelog entries {len(d.changelog)}")
        check(d.plans[0]["free"] == 5 and d.backlog[0]["bucket"] == "NOW", "plans and backlog kept")

        # state
        state = Path(tmp) / "session-state.json"
        state.write_text(json.dumps({
            "current_phase": {"number": 1, "name": "Editor integrity", "status": "IN_PROGRESS"},
            "current_task": {"id": "EDIT-002", "details": "blocks in list items"},
            "next_exact_task": {"id": "EDIT-002", "description": "map code blocks", "command": "npm test"},
            "test_results": {"passed": ["backend 643/643"], "failed": []},
            "known_blockers": [], "production_code_changed": True, "database_migrations_applied": False,
        }), encoding="utf-8")
        check(run(wb, "state", "--json", str(state), "--no-git", "--touch") == 0, "state sync")
        saved = json.loads(state.read_text(encoding="utf-8"))
        check(saved["current_excel_status"]["total"] == len(seed.TASKS) + 1, "excel status written to JSON")
        check("last_checkpoint" in saved, "checkpoint stamped")
        _, d = tracker.load(wb)
        check(d.session["Current Task"] == "EDIT-002 — blocks in list items", f"session {d.session['Current Task']}")
        check(d.session["Production Code Changed"] == "yes", "extra session field")

        # render keeps data and foreign sheets, drops sensitivity labels; lock files block writes
        before = master(wb)
        book = load_workbook(wb)
        book.create_sheet("MY_NOTES")["A1"] = "keep me"
        from openpyxl.packaging.custom import StringProperty
        book.custom_doc_props.append(StringProperty(name="MSIP_Label_0000_SiteId", value="tenant"))
        book.custom_doc_props.append(StringProperty(name="Keep", value="yes"))
        book.save(wb)
        check(run(wb, "render") == 0, "render")
        check(master(wb) == before, "render round trip keeps MASTER data")
        book = load_workbook(wb)
        check("MY_NOTES" in book.sheetnames and book["MY_NOTES"]["A1"].value == "keep me", "foreign sheet kept")
        names = [p.name for p in book.custom_doc_props.props]
        check(names == ["Keep"], f"sensitivity labels dropped, other properties kept: {names}")
        lock = wb.with_name("~$" + wb.name)
        lock.write_text("")
        check(run(wb, "set", "EDIT-002", "--status", "IN_PROGRESS") == 1, "open workbook (lock file) blocks writes")
        lock.unlink()
        check(not list(Path(tmp).glob(".*.tmp-*.xlsx")), "no temp files left behind")

    print("selftest:", "PASS" if not FAILURES else f"{len(FAILURES)} FAILED")
    return 0 if not FAILURES else 1


if __name__ == "__main__":
    sys.exit(main())
