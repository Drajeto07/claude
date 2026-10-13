"""Builds docs/SMARTDOC_2_0_MASTER_TRACKER.xlsx from docs/SMARTDOC_2_0_TRACKER.json (the source of
truth, also read by the next session). Every run rewrites the workbook: edit the JSON, not the xlsx.

    backend/venv/Scripts/python.exe tools/smartdoc2/build_tracker.py

Sheets: Summary (formulas over the others), Roadmap, Tasks, Bugs, Tests, Decisions, Release Readiness.
Arial, frozen headers, filters, status colours by conditional formatting. Recalculate after building
(tools/tracker/excel_recalc.py) so the summary formulas carry values for readers that don't compute."""

import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "docs" / "SMARTDOC_2_0_TRACKER.json"
TARGET = ROOT / "docs" / "SMARTDOC_2_0_MASTER_TRACKER.xlsx"

STATUSES = ["BACKLOG", "READY", "IN_PROGRESS", "BLOCKED", "TESTING", "VERIFIED", "DONE"]
STATUS_FILL = {
    "BACKLOG": "E5E7EB", "READY": "DBEAFE", "IN_PROGRESS": "EDE9FE", "BLOCKED": "FEE2E2",
    "TESTING": "FEF3C7", "VERIFIED": "D1FAE5", "DONE": "BBF7D0",
}
HEADER_FILL = PatternFill("solid", fgColor="4338CA")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=10)
BODY_FONT = Font(name="Arial", size=10)
TITLE_FONT = Font(name="Arial", bold=True, size=14, color="312E81")
THIN = Side(style="thin", color="E5E7EB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

SHEETS = {
    "Roadmap": ("roadmap", [("Phase", "phase", 8), ("Milestone", "milestone", 34), ("Objective", "objective", 50), ("Status", "status", 14),
                            ("Dependencies", "dependencies", 14), ("Target outcome", "outcome", 50)]),
    "Tasks": ("tasks", [("Task ID", "id", 9), ("Phase", "phase", 7), ("Category", "category", 12), ("Title", "title", 38), ("Description", "description", 60),
                        ("Priority", "priority", 8), ("Status", "status", 13), ("Dependencies", "dependencies", 14), ("Acceptance criteria", "acceptance", 50),
                        ("Owner", "owner", 9), ("Branch", None, 20), ("Commit", "commit", 10), ("Preview deployment", "preview", 24),
                        ("Test result", "test_result", 24), ("Last updated", None, 12), ("Notes", "notes", 50)]),
    "Bugs": ("bugs", [("Bug ID", "id", 8), ("Severity", "severity", 9), ("Reproduction steps", "steps", 50), ("Expected result", "expected", 30),
                      ("Actual result", "actual", 50), ("Status", "status", 13), ("Fix commit", "fix", 18), ("Verification", "verification", 50)]),
    "Tests": ("tests", [("Test ID", "id", 8), ("Area", "area", 12), ("Scenario", "scenario", 50), ("Environment", "environment", 18),
                        ("Result", "result", 22), ("Evidence", "evidence", 40), ("Date", "date", 12)]),
    "Decisions": ("decisions", [("Decision ID", "id", 11), ("Topic", "topic", 22), ("Approved decision", "decision", 70), ("Date", "date", 12),
                                ("Impact", "impact", 18), ("Related tasks", "tasks", 18)]),
    "Release Readiness": ("release", [("Requirement", "requirement", 46), ("Preview status", "preview", 28), ("Production status", "production", 30),
                                      ("Blocking issue", "blocking", 34), ("Approval required", "approval", 20)]),
}


def _validate(data: dict) -> None:
    for row in data["roadmap"] + data["tasks"] + data["bugs"]:
        status = row.get("status")
        if row in data["bugs"]:
            continue
        if status not in STATUSES:
            sys.exit(f"Unknown status {status!r} in {row.get('id') or row.get('phase')}")
    ids = [task["id"] for task in data["tasks"]]
    if len(ids) != len(set(ids)):
        sys.exit("Duplicate task ids")


def _sheet(wb: Workbook, title: str, rows: list[dict], columns: list, data: dict) -> None:
    ws = wb.create_sheet(title)
    for index, (header, _, width) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=index, value=header)
        cell.font, cell.fill, cell.border = HEADER_FONT, HEADER_FILL, BORDER
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(index)].width = width
    for r, row in enumerate(rows, start=2):
        for c, (header, key, _) in enumerate(columns, start=1):
            if key is None:
                value = data["branch"] if header == "Branch" else data["updated"]
            else:
                value = row.get(key, "")
            cell = ws.cell(row=r, column=c, value=value)
            cell.font, cell.border = BODY_FONT, BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "B2" if title == "Tasks" else "A2"
    last = get_column_letter(len(columns))
    ws.auto_filter.ref = f"A1:{last}{max(len(rows) + 1, 2)}"
    status_col = next((i for i, (h, _, _) in enumerate(columns, start=1) if h == "Status"), None)
    if status_col:
        letter = get_column_letter(status_col)
        span = f"{letter}2:{letter}{max(len(rows) + 1, 2)}"
        for status, colour in STATUS_FILL.items():
            ws.conditional_formatting.add(span, FormulaRule(formula=[f'${letter}2="{status}"'], fill=PatternFill("solid", fgColor=colour)))


def _summary(wb: Workbook, data: dict) -> None:
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = f"{data['project']} -- master tracker"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Branch {data['branch']} · source docs/SMARTDOC_2_0_TRACKER.json · updated {data['updated']}"
    ws["A2"].font = Font(name="Arial", size=9, italic=True, color="6B7280")
    headers = ["Status", "Tasks", "Roadmap milestones"]
    for c, header in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=c, value=header)
        cell.font, cell.fill, cell.border = HEADER_FONT, HEADER_FILL, BORDER
    task_rows = len(data["tasks"]) + 1
    road_rows = len(data["roadmap"]) + 1
    for r, status in enumerate(STATUSES, start=5):
        ws.cell(row=r, column=1, value=status)
        ws.cell(row=r, column=2, value=f"=COUNTIF(Tasks!$G$2:$G${task_rows},A{r})")
        ws.cell(row=r, column=3, value=f"=COUNTIF(Roadmap!$D$2:$D${road_rows},A{r})")
        for c in range(1, 4):
            ws.cell(row=r, column=c).font, ws.cell(row=r, column=c).border = BODY_FONT, BORDER
        ws.cell(row=r, column=1).fill = PatternFill("solid", fgColor=STATUS_FILL[status])
    total = 5 + len(STATUSES)
    ws.cell(row=total, column=1, value="Total").font = Font(name="Arial", bold=True, size=10)
    ws.cell(row=total, column=2, value=f"=SUM(B5:B{total - 1})").font = Font(name="Arial", bold=True, size=10)
    ws.cell(row=total, column=3, value=f"=SUM(C5:C{total - 1})").font = Font(name="Arial", bold=True, size=10)
    ws.cell(row=total + 1, column=1, value="Tasks done or verified").font = BODY_FONT
    ws.cell(row=total + 1, column=2, value=f'=IFERROR((B{5 + STATUSES.index("DONE")}+B{5 + STATUSES.index("VERIFIED")})/B{total},0)')
    ws.cell(row=total + 1, column=2).number_format = "0%"
    ws.cell(row=total + 2, column=1, value="Open P0 tasks").font = BODY_FONT
    ws.cell(row=total + 2, column=2, value=f'=COUNTIFS(Tasks!$F$2:$F${task_rows},"P0",Tasks!$G$2:$G${task_rows},"<>DONE",Tasks!$G$2:$G${task_rows},"<>VERIFIED")')
    ws.cell(row=total + 3, column=1, value="Open bugs").font = BODY_FONT
    ws.cell(row=total + 3, column=2, value=f'=COUNTA(Bugs!$A$2:$A${len(data["bugs"]) + 1})-COUNTIF(Bugs!$F$2:$F${len(data["bugs"]) + 1},"DONE")-COUNTIF(Bugs!$F$2:$F${len(data["bugs"]) + 1},"VERIFIED")')
    ws.cell(row=total + 5, column=1, value="Statuses: BACKLOG, READY, IN_PROGRESS, BLOCKED, TESTING, VERIFIED, DONE. Edit the JSON source and rebuild.").font = Font(name="Arial", size=9, color="6B7280")
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 20


def main() -> None:
    data = json.loads(SOURCE.read_text(encoding="utf-8"))
    _validate(data)
    wb = Workbook()
    _summary(wb, data)
    for title, (key, columns) in SHEETS.items():
        _sheet(wb, title, data[key], columns, data)
    wb.save(TARGET)
    print(f"Wrote {TARGET.relative_to(ROOT)}: {len(data['tasks'])} tasks, {len(data['roadmap'])} milestones")


if __name__ == "__main__":
    main()
