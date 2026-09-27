"""What a Word import changes without keeping it is named in the import report
(tracker FID-002), and the file's own properties survive an export (DOCX-012)."""

import io
import zipfile
from datetime import datetime, timezone

from docx import Document as DocxDocument
from docx.enum.text import WD_UNDERLINE
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from PIL import Image as PILImage

from app.export.docx_export import build_docx
from app.models.document import Document, DocumentMetadata, Element, ElementType, InlineRun
from app.services.ingestion_service import build_document_from_docx


def _png() -> bytes:
    buffer = io.BytesIO()
    PILImage.new("RGB", (20, 10), "blue").save(buffer, format="PNG")
    return buffer.getvalue()


def _save(document) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _report(data: bytes):
    return build_document_from_docx(data, "detect.docx", None).importReport


def _items(report):
    return {item.feature: item for item in report.items}


def _run(paragraph, text: str, properties: str):
    paragraph._p.append(parse_xml(f"<w:r {nsdecls('w')}><w:rPr>{properties}</w:rPr><w:t xml:space=\"preserve\">{text}</w:t></w:r>"))


def test_each_unkept_feature_is_named_with_an_example():
    document = DocxDocument()
    _run(document.add_paragraph(), "the answer key", "<w:vanish/>")
    _run(document.add_paragraph(), "shouting", "<w:caps/>")  # kept since DOCX-013: not named
    _run(document.add_paragraph(), "dash-dot underlined", '<w:u w:val="dotDash"/>')  # shown as dashed: named
    document.add_paragraph("Write to someone@example.com or see https://example.com/page.")
    document.add_paragraph()._p.append(
        parse_xml(f"<w:sdt {nsdecls('w')}><w:sdtPr><w:alias w:val=\"Status\"/></w:sdtPr><w:sdtContent><w:r><w:t>Draft</w:t></w:r></w:sdtContent></w:sdt>")
    )
    picture = document.add_paragraph().add_run().add_picture(io.BytesIO(_png()))
    blip_fill = picture._inline.graphic.graphicData.pic.blipFill
    blip_fill.append(parse_xml(f'<a:srcRect {nsdecls("a")} l="10000" r="5000"/>'))
    picture._inline.graphic.graphicData.pic.spPr.find(qn("a:xfrm")).set("rot", "5400000")
    table = document.add_table(rows=1, cols=2)
    table.style = "Light Grid Accent 1"
    cell = table.cell(0, 0).paragraphs[0]
    cell.text = "a bullet in a cell"
    cell._p.get_or_add_pPr().append(parse_xml(f'<w:numPr {nsdecls("w")}><w:ilvl w:val="0"/><w:numId w:val="1"/></w:numPr>'))
    rtl = document.add_paragraph("שלום עולם")
    rtl._p.get_or_add_pPr().append(parse_xml(f"<w:bidi {nsdecls('w')}/>"))  # the paragraph's direction: kept (DOCX-014)
    rtl.runs[0]._r.get_or_add_rPr().append(parse_xml(f"<w:rtl {nsdecls('w')}/>"))  # Word's marking on the run: named

    report = _report(_save(document))
    items = _items(report)

    hidden = items["docx.hidden_text"]  # kept hidden (DOCX-025): named, not a content change
    assert hidden.policy == "detected_preserved" and not hidden.contentChanged and hidden.sourceState == "e.g. “the answer key”"
    assert "docx.caps" not in items
    assert items["docx.underline_variant"].sourceState == "e.g. “dash-dot underlined”"
    assert items["docx.autolink"].count == 2
    assert items["docx.content_control"].sourceState == "e.g. “Draft”"
    assert "docx.image.crop" in items and "docx.image.rotation" in items
    assert "docx.table.geometry" in items
    assert items["docx.table.cell_list"].sourceState == "e.g. “a bullet in a cell”"
    assert "docx.rtl" in items
    # None of these loses a word: the content check still verifies the text.
    assert report.contentStatus == "verified", report.content.samples


def test_styles_count_too():
    document = DocxDocument()
    document.styles["Heading 1"].font.underline = WD_UNDERLINE.DOT_DASH
    document.add_heading("Chapter one", level=1)

    items = _items(_report(_save(document)))

    assert items["docx.underline_variant"].sourceState == "e.g. “Chapter one”"


def test_charts_and_shapes_are_named_for_what_they_are():
    document = DocxDocument()

    def drawing(uri: str, content: str) -> str:
        return (
            f'<w:r {nsdecls("w", "wp", "a")}><w:drawing><wp:inline><wp:extent cx="100" cy="100"/><wp:docPr id="7" name="x"/>'
            f'<a:graphic><a:graphicData uri="{uri}">{content}</a:graphicData></a:graphic></wp:inline></w:drawing></w:r>'
        )

    document.add_paragraph("A chart:")._p.append(parse_xml(drawing("http://schemas.openxmlformats.org/drawingml/2006/chart", "")))
    document.add_paragraph("A shape:")._p.append(
        parse_xml(drawing("http://schemas.microsoft.com/office/word/2010/wordprocessingShape", "<a:ln/>"))
    )

    items = _items(_report(_save(document)))

    assert items["docx.chart"].contentChanged and items["docx.shape"].contentChanged
    assert "docx.image.linked" not in items  # what these are, not a misleading "linked picture"


def _with_custom_properties(data: bytes) -> bytes:
    properties = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties" '
        'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
        '<property fmtid="{D5CDD505-2E9C-101B-9397-08002B2CF9AE}" pid="2" name="MSIP_Label_00000000-0000-0000-0000-000000000000_SiteId">'
        "<vt:lpwstr>11111111-2222-3333-4444-555555555555</vt:lpwstr></property>"
        '<property fmtid="{D5CDD505-2E9C-101B-9397-08002B2CF9AE}" pid="3" name="Client"><vt:lpwstr>Secret Client Ltd</vt:lpwstr></property>'
        '<property fmtid="{D5CDD505-2E9C-101B-9397-08002B2CF9AE}" pid="4" name="Matter"><vt:lpwstr>M-42</vt:lpwstr></property>'
        "</Properties>"
    )
    source, target = io.BytesIO(data), io.BytesIO()
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as after:
        for item in before.infolist():
            after.writestr(item, before.read(item.filename))
        after.writestr("docProps/custom.xml", properties)
    return target.getvalue()


def test_metadata_that_is_not_kept_is_reported_without_its_values():
    document = DocxDocument()
    document.add_paragraph("Body.")

    report = _report(_with_custom_properties(_save(document)))
    items = _items(report)

    assert items["docx.metadata.custom_properties"].count == 2
    assert "docx.metadata.sensitivity_label" in items
    written = repr(report.model_dump())
    for value in ("Secret Client", "M-42", "11111111", "Client", "Matter"):
        assert value not in written


def test_the_files_own_properties_go_back_into_word_not_the_templates():
    document = DocxDocument()
    properties = document.core_properties
    properties.author, properties.subject, properties.keywords = "Ana Petrova", "Quarterly report", "finance, q3"
    properties.created = datetime(2020, 5, 17, 9, 30, tzinfo=timezone.utc)
    document.add_paragraph("Body.")

    imported = build_document_from_docx(_save(document), "report.docx", None)
    exported = DocxDocument(io.BytesIO(build_docx(imported))).core_properties

    assert (exported.author, exported.subject, exported.keywords) == ("Ana Petrova", "Quarterly report", "finance, q3")
    assert exported.created == datetime(2020, 5, 17, 9, 30, tzinfo=timezone.utc)
    assert exported.title == imported.metadata.title

    pasted = Document(
        metadata=DocumentMetadata(title="Notes", createdAt=datetime(2026, 1, 2, tzinfo=timezone.utc)),
        elements=[Element(type=ElementType.PARAGRAPH, content="x", inline=[InlineRun(text="x")], order=0)],
    )
    fresh = DocxDocument(io.BytesIO(build_docx(pasted))).core_properties
    assert (fresh.author, fresh.comments, fresh.title) == ("", "", "Notes")  # never "python-docx"
    assert fresh.created == datetime(2026, 1, 2, tzinfo=timezone.utc)


def test_unsafe_links_and_spacing_paragraphs_are_reported():
    document = DocxDocument()
    paragraph = document.add_paragraph("Click ")
    rel_id = paragraph.part.relate_to("javascript:alert(1)", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    paragraph._p.append(parse_xml(f'<w:hyperlink {nsdecls("w", "r")} r:id="{rel_id}"><w:r><w:t>here</w:t></w:r></w:hyperlink>'))
    document.add_paragraph("")
    document.add_paragraph("")
    document.add_paragraph("After the gap.")

    imported = build_document_from_docx(_save(document), "links.docx", None)
    items = _items(imported.importReport)

    assert "docx.link.unsafe" in items
    assert items["docx.empty_paragraph"].count == 2
    assert imported.elements[0].content == "Click here"  # the words stay
    assert not any(mark.type == "link" for run in imported.elements[0].inline for mark in run.marks)


def _sections(*starts: WD_SECTION) -> DocxDocument:
    document = DocxDocument()
    document.add_paragraph("Section 1.")
    for number, start in enumerate(starts, start=2):
        document.add_section(start)
        document.add_paragraph(f"Section {number}.")
    return document


def _kinds(document: DocxDocument) -> list[str]:
    """Each element's type; a section break's with how the section after it starts."""
    return [
        f"section_break:{element.sectionBreak.start}" if element.sectionBreak else element.type.value
        for element in build_document_from_docx(_save(document), "sections.docx", None).elements
    ]


def test_a_section_break_says_how_the_next_section_starts_as_word_does():
    # A section's own type says how it starts, so it's the section after the break that counts (DOCX-015).
    assert _kinds(_sections(WD_SECTION.CONTINUOUS)) == ["paragraph", "section_break:continuous", "paragraph"]
    assert _kinds(_sections(WD_SECTION.NEW_PAGE)) == ["paragraph", "section_break:nextPage", "paragraph"]
    assert _kinds(_sections(WD_SECTION.CONTINUOUS, WD_SECTION.NEW_PAGE)) == [
        "paragraph", "section_break:continuous", "paragraph", "section_break:nextPage", "paragraph",
    ]
    assert _kinds(_sections(WD_SECTION.NEW_PAGE, WD_SECTION.ODD_PAGE)) == [
        "paragraph", "section_break:nextPage", "paragraph", "section_break:oddPage", "paragraph",
    ]


def test_what_sections_change_is_reported():
    document = _sections(WD_SECTION.ODD_PAGE)
    wide = document.sections[0]
    wide.orientation, wide.page_width, wide.page_height = WD_ORIENT.LANDSCAPE, wide.page_height, wide.page_width
    last = document.sections[-1]._sectPr
    last.append(parse_xml(f'<w:pgNumType {nsdecls("w")} w:fmt="lowerRoman"/>'))
    last.append(parse_xml(f'<w:lnNumType {nsdecls("w")} w:countBy="1"/>'))
    last.append(parse_xml(f'<w:vAlign {nsdecls("w")} w:val="center"/>'))
    last.append(parse_xml(
        f'<w:pgBorders {nsdecls("w")}><w:top w:val="single" w:sz="4" w:space="24" w:color="auto"/></w:pgBorders>'
    ))

    imported = build_document_from_docx(_save(document), "sections.docx", None)
    items = _items(imported.importReport)

    assert items["docx.sections.page_setup"].count == 1  # the landscape first section: kept as its section break
    for key in ("docx.sections.page_setup", "docx.sections.break_type", "docx.sections.page_numbering"):
        assert items[key].policy == "detected_not_editable", key  # kept for a Word export (DOCX-015)
    for key in ("docx.sections.line_numbers", "docx.sections.vertical_alignment", "docx.sections.page_borders"):
        assert items[key].policy == "lossy", key
    assert [element.type.value for element in imported.elements] == ["paragraph", "section_break", "paragraph"]
    settings = imported.elements[1].sectionBreak
    assert (settings.start, settings.orientation) == ("oddPage", "landscape")

    plain = _items(_report(_save(_sections(WD_SECTION.NEW_PAGE))))
    assert not any(key.startswith("docx.sections.") for key in plain)  # same setup, ordinary break: nothing to say
