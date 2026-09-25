import io
import json
from pathlib import Path

import pytest
from docx import Document as DocxDocument
from fastapi import UploadFile
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session as OrmSession

from app.ai.factory import get_ai_provider
from app.ai.schemas import (
    AIBlock,
    AIBlockType,
    AIFormattingRule,
    AIInstructionExtractionResponse,
    AIStructureResponse,
    AIStyleAnalysisResponse,
    AIStyleFlag,
)
from app.config import get_settings
from app.db.models import Document as DocumentRow
from app.main import app
from tests.fakes import FakeAIProvider

# https: the session cookie is Secure, and a cookie jar never returns it over plain http.
client = TestClient(app, base_url="https://testserver")
FIXTURES_DIR = Path(__file__).parent / "fixtures"


@pytest.fixture(autouse=True)
def signed_in(api_db):
    """Every test runs as a freshly registered user on a fresh database."""
    client.cookies.clear()
    response = client.post("/api/auth/register", json={"email": "owner@example.com", "password": "long enough password"})
    assert response.status_code == 201
    yield
    client.cookies.clear()


def test_health_check():
    response = client.get("/api/health")
    assert response.status_code == 200
    assert response.json() == {"status": "ok"}


def test_create_and_fetch_document_round_trip():
    # Markdown-looking input takes the deterministic parser path, so this
    # needs no AI provider override to be predictable.
    create_response = client.post("/api/documents", json={"text": "# Hello\n\nWorld is a longer paragraph here."})
    assert create_response.status_code == 201
    created = create_response.json()
    assert len(created["elements"]) == 2
    assert created["schemaVersion"] == 1

    get_response = client.get(f"/api/documents/{created['id']}")
    assert get_response.status_code == 200
    assert get_response.json() == created


def test_get_unknown_document_returns_404():
    response = client.get("/api/documents/does-not-exist")
    assert response.status_code == 404


def test_delete_document_returns_204_and_document_is_actually_gone():
    document_id = _create_document()

    delete_response = client.delete(f"/api/documents/{document_id}")

    assert delete_response.status_code == 204
    assert client.get(f"/api/documents/{document_id}").status_code == 404


def test_delete_unknown_document_returns_404():
    response = client.delete("/api/documents/does-not-exist")
    assert response.status_code == 404


def test_delete_document_also_drops_its_undo_history():
    document_id = _create_document()
    client.post(f"/api/documents/{document_id}/format", data={"templateId": "academic-default"})

    client.delete(f"/api/documents/{document_id}")

    # A deleted id must never resurrect through an unrelated undo/redo call.
    assert client.post(f"/api/documents/{document_id}/undo").status_code == 404


def test_create_document_rejects_blank_text():
    response = client.post("/api/documents", json={"text": "   "})
    assert response.status_code == 422


def test_plain_prose_routes_to_ai_provider():
    fake_response = AIStructureResponse(
        document_type="essay",
        document_type_confidence=0.7,
        blocks=[AIBlock(type=AIBlockType.PARAGRAPH, text="Ordinary unstructured prose text for the API test.", confidence=0.8)],
    )
    app.dependency_overrides[get_ai_provider] = lambda: FakeAIProvider([fake_response])
    try:
        response = client.post("/api/documents", json={"text": "Ordinary unstructured prose text for the API test."})
    finally:
        app.dependency_overrides.pop(get_ai_provider, None)

    assert response.status_code == 201
    body = response.json()
    assert body["documentType"] == "essay"
    assert body["elements"][0]["confidence"] == 0.8


def test_upload_txt_file():
    response = client.post(
        "/api/documents/upload",
        files={"file": ("notes.txt", b"# A Heading\n\nSome body text.", "text/plain")},
    )
    assert response.status_code == 201
    body = response.json()
    assert body["metadata"]["sourceType"] == "uploaded_txt"
    assert body["metadata"]["originalFilename"] == "notes.txt"


def test_upload_docx_file():
    doc = DocxDocument()
    doc.add_heading("Uploaded Title", level=1)
    doc.add_paragraph("Body paragraph.")
    buf = io.BytesIO()
    doc.save(buf)

    response = client.post(
        "/api/documents/upload",
        files={
            "file": (
                "report.docx",
                buf.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert response.status_code == 201
    body = response.json()
    assert body["elements"][0]["content"] == "Uploaded Title"
    assert body["metadata"]["sourceType"] == "uploaded_docx"


def test_upload_pdf_file():
    # The fixture's extracted text is plain prose (no Markdown cues), so this
    # exercises the upload -> PDF-text-extraction -> AI-routing path -- needs
    # a fake provider, same as test_plain_prose_routes_to_ai_provider.
    pdf_bytes = (FIXTURES_DIR / "sample.pdf").read_bytes()
    fake_response = AIStructureResponse(
        document_type="general",
        document_type_confidence=0.6,
        blocks=[AIBlock(type=AIBlockType.PARAGRAPH, text="Hello World Sample PDF Text", confidence=0.75)],
    )
    app.dependency_overrides[get_ai_provider] = lambda: FakeAIProvider([fake_response])
    try:
        response = client.post(
            "/api/documents/upload",
            files={"file": ("sample.pdf", pdf_bytes, "application/pdf")},
        )
    finally:
        app.dependency_overrides.pop(get_ai_provider, None)

    assert response.status_code == 201
    body = response.json()
    assert "Hello World Sample PDF Text" in body["elements"][0]["content"]


def test_upload_rejects_unsupported_extension():
    response = client.post(
        "/api/documents/upload",
        files={"file": ("image.png", b"not text", "image/png")},
    )
    assert response.status_code == 400


def test_upload_rejects_corrupt_docx():
    response = client.post(
        "/api/documents/upload",
        files={
            "file": (
                "bad.docx",
                b"this is not a real docx",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert response.status_code == 400


def test_upload_rejects_oversized_file():
    settings = get_settings()
    original_max = settings.max_upload_size_mb
    settings.max_upload_size_mb = 0
    try:
        response = client.post(
            "/api/documents/upload",
            files={"file": ("small.txt", b"just a few bytes", "text/plain")},
        )
        assert response.status_code == 413
    finally:
        settings.max_upload_size_mb = original_max


def test_upload_rejects_oversized_file_even_when_client_reports_no_size(monkeypatch):
    """Regression: the limit must hold even when file.size is None (some
    multipart encoders never send a per-part Content-Length) -- the fix
    checks the actual bytes read instead of trusting client-reported size."""
    # UploadFile.size is a plain instance attribute (set in __init__), not a
    # class-level property -- overriding it with a getter/no-op-setter
    # property forces every read to see None while leaving __init__'s own
    # `self.size = size` assignment harmless, faithfully simulating a
    # request where Starlette could never determine the part's size.
    monkeypatch.setattr(UploadFile, "size", property(lambda self: None, lambda self, value: None), raising=False)
    settings = get_settings()
    original_max = settings.max_upload_size_mb
    settings.max_upload_size_mb = 0
    try:
        response = client.post(
            "/api/documents/upload",
            files={"file": ("small.txt", b"just a few bytes", "text/plain")},
        )
        assert response.status_code == 413
    finally:
        settings.max_upload_size_mb = original_max


def _create_document() -> str:
    response = client.post("/api/documents", json={"text": "# Hello\n\nWorld is a longer paragraph here."})
    return response.json()["id"]


def test_format_document_with_builtin_template():
    document_id = _create_document()

    response = client.post(f"/api/documents/{document_id}/format", data={"templateId": "academic-default"})

    assert response.status_code == 200
    payload = response.json()
    assert payload["aiUnavailable"] is False
    body = payload["document"]
    assert body["templateId"] == "academic-default"
    assert body["resolvedStyles"]["Paragraph"]["font-family"] == "Times New Roman"
    assert body["elements"][0]["styleRef"] == "Heading 1"


def test_format_document_with_instructions_overrides_template():
    document_id = _create_document()
    instruction_response = AIInstructionExtractionResponse(
        rules=[AIFormattingRule(target="Heading 1", property="color", value="red")]
    )
    app.dependency_overrides[get_ai_provider] = lambda: FakeAIProvider([instruction_response])
    try:
        response = client.post(
            f"/api/documents/{document_id}/format",
            data={"templateId": "academic-default", "instructionsText": "make headings red"},
        )
    finally:
        app.dependency_overrides.pop(get_ai_provider, None)

    assert response.status_code == 200
    assert response.json()["document"]["resolvedStyles"]["Heading 1"]["color"] == "red"


def test_format_document_with_instructions_file():
    document_id = _create_document()
    instruction_response = AIInstructionExtractionResponse(
        rules=[AIFormattingRule(target="Paragraph", property="fontFamily", value="Georgia")]
    )
    app.dependency_overrides[get_ai_provider] = lambda: FakeAIProvider([instruction_response])
    try:
        response = client.post(
            f"/api/documents/{document_id}/format",
            files={"instructionsFile": ("rules.txt", b"Use Georgia font for body text.", "text/plain")},
        )
    finally:
        app.dependency_overrides.pop(get_ai_provider, None)

    assert response.status_code == 200
    assert response.json()["document"]["resolvedStyles"]["Paragraph"]["font-family"] == "Georgia"


def test_format_document_rejects_unsupported_instructions_file_type():
    document_id = _create_document()

    response = client.post(
        f"/api/documents/{document_id}/format",
        files={
            "instructionsFile": (
                "rules.docx",
                b"not really a docx",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 400


def test_format_document_rejects_unknown_template():
    document_id = _create_document()

    response = client.post(f"/api/documents/{document_id}/format", data={"templateId": "does-not-exist"})

    assert response.status_code == 400


def test_format_unknown_document_returns_404():
    response = client.post("/api/documents/does-not-exist/format", data={"templateId": "academic-default"})

    assert response.status_code == 404


def test_set_and_clear_element_style():
    document_id = _create_document()
    formatted = client.post(f"/api/documents/{document_id}/format", data={"templateId": "academic-default"}).json()["document"]
    paragraph_id = formatted["elements"][1]["id"]  # index 0 is the heading

    set_response = client.patch(
        f"/api/documents/{document_id}/elements/{paragraph_id}/style",
        json={"property": "fontFamily", "value": "Georgia"},
    )
    assert set_response.status_code == 200
    body = set_response.json()
    assert body["resolvedStyles"][paragraph_id]["font-family"] == "Georgia"
    assert body["elements"][1]["styleRef"] == paragraph_id
    # The coarse "Paragraph" target (and any other paragraph) is unaffected.
    assert body["resolvedStyles"]["Paragraph"]["font-family"] == "Times New Roman"

    clear_response = client.delete(f"/api/documents/{document_id}/elements/{paragraph_id}/style/fontFamily")
    assert clear_response.status_code == 200
    cleared_body = clear_response.json()
    assert paragraph_id not in cleared_body["resolvedStyles"]
    assert cleared_body["elements"][1]["styleRef"] == "Paragraph"


def test_set_element_style_rejects_unknown_element():
    document_id = _create_document()

    response = client.patch(
        f"/api/documents/{document_id}/elements/does-not-exist/style",
        json={"property": "fontFamily", "value": "Georgia"},
    )

    assert response.status_code == 404


def test_set_element_style_unknown_document_returns_404():
    response = client.patch(
        "/api/documents/does-not-exist/elements/also-does-not-exist/style",
        json={"property": "fontFamily", "value": "Georgia"},
    )

    assert response.status_code == 404


def test_format_document_returns_409_when_conflict_exists():
    document_id = _create_document()
    formatted = client.post(f"/api/documents/{document_id}/format", data={"templateId": "academic-default"}).json()["document"]
    paragraph_id = formatted["elements"][1]["id"]
    client.patch(f"/api/documents/{document_id}/elements/{paragraph_id}/style", json={"property": "color", "value": "red"})

    custom = client.post(
        "/api/templates",
        json={
            "name": "Conflict Test Template",
            "category": "academic",
            "styleSystem": {"paragraph": {"color": "blue"}},
        },
    ).json()

    response = client.post(f"/api/documents/{document_id}/format", data={"templateId": custom["id"]})

    assert response.status_code == 409
    conflicts = response.json()["detail"]["conflicts"]
    assert len(conflicts) == 1
    assert conflicts[0]["elementId"] == paragraph_id
    assert conflicts[0]["property"] == "color"
    assert conflicts[0]["currentValue"] == "red"
    assert conflicts[0]["requiredValue"] == "blue"

    # Nothing was actually applied -- the document is unchanged.
    unchanged = client.get(f"/api/documents/{document_id}").json()
    assert unchanged["templateId"] == "academic-default"
    assert unchanged["resolvedStyles"][paragraph_id]["color"] == "red"


def test_format_document_with_resolutions_applies_correctly():
    document_id = _create_document()
    formatted = client.post(f"/api/documents/{document_id}/format", data={"templateId": "academic-default"}).json()["document"]
    paragraph_id = formatted["elements"][1]["id"]
    client.patch(f"/api/documents/{document_id}/elements/{paragraph_id}/style", json={"property": "color", "value": "red"})

    custom = client.post(
        "/api/templates",
        json={
            "name": "Conflict Test Template 2",
            "category": "academic",
            "styleSystem": {"paragraph": {"color": "blue"}},
        },
    ).json()

    resolutions = [{"elementId": paragraph_id, "property": "color", "resolution": "apply_recommended"}]
    response = client.post(
        f"/api/documents/{document_id}/format",
        data={"templateId": custom["id"], "resolutions": json.dumps(resolutions)},
    )

    assert response.status_code == 200
    body = response.json()["document"]
    assert body["templateId"] == custom["id"]
    assert body["resolvedStyles"]["Paragraph"]["color"] == "blue"
    # The override was dropped (resolution: apply_recommended) -- no more per-element entry.
    assert paragraph_id not in body["resolvedStyles"]


def test_undo_redo_round_trip_through_format_and_override():
    document_id = _create_document()
    original = client.get(f"/api/documents/{document_id}").json()
    assert original["templateId"] is None

    client.post(f"/api/documents/{document_id}/format", data={"templateId": "academic-default"})
    after_format = client.get(f"/api/documents/{document_id}").json()
    assert after_format["templateId"] == "academic-default"

    undo_response = client.post(f"/api/documents/{document_id}/undo")
    assert undo_response.status_code == 200
    assert undo_response.json()["templateId"] is None

    redo_response = client.post(f"/api/documents/{document_id}/redo")
    assert redo_response.status_code == 200
    assert redo_response.json()["templateId"] == "academic-default"


def test_add_page_creates_a_real_page_break_element():
    document_id = _create_document()

    response = client.post(f"/api/documents/{document_id}/pages", json={})

    assert response.status_code == 201
    body = response.json()
    assert body["elements"][-1]["type"] == "page_break"


def test_add_page_after_a_specific_element():
    document_id = _create_document()
    created = client.get(f"/api/documents/{document_id}").json()
    heading_id = created["elements"][0]["id"]

    response = client.post(f"/api/documents/{document_id}/pages", json={"afterElementId": heading_id})

    body = response.json()
    ordered = sorted(body["elements"], key=lambda el: el["order"])
    assert ordered[1]["type"] == "page_break"


def test_add_page_unknown_after_element_returns_404():
    document_id = _create_document()

    response = client.post(f"/api/documents/{document_id}/pages", json={"afterElementId": "does-not-exist"})

    assert response.status_code == 404


def test_add_page_unknown_document_returns_404():
    response = client.post("/api/documents/does-not-exist/pages", json={})

    assert response.status_code == 404


def test_add_page_is_undoable():
    document_id = _create_document()
    before = client.get(f"/api/documents/{document_id}").json()
    before_count = len(before["elements"])

    client.post(f"/api/documents/{document_id}/pages", json={})
    undo_response = client.post(f"/api/documents/{document_id}/undo")

    assert len(undo_response.json()["elements"]) == before_count


def test_add_element_creates_a_real_paragraph():
    document_id = _create_document()

    response = client.post(f"/api/documents/{document_id}/elements", json={"elementType": "paragraph", "text": "New paragraph"})

    assert response.status_code == 201
    body = response.json()
    assert body["elements"][-1]["type"] == "paragraph"
    assert body["elements"][-1]["content"] == "New paragraph"


def test_add_element_list_and_table_get_real_structure():
    document_id = _create_document()

    list_response = client.post(f"/api/documents/{document_id}/elements", json={"elementType": "list"})
    table_response = client.post(f"/api/documents/{document_id}/elements", json={"elementType": "table"})

    assert list_response.json()["elements"][-1]["listItems"] is not None
    assert table_response.json()["elements"][-1]["table"]["rows"]


def test_add_element_rejects_unsupported_type():
    document_id = _create_document()

    response = client.post(f"/api/documents/{document_id}/elements", json={"elementType": "image"})

    assert response.status_code == 422


def test_add_element_after_a_specific_element():
    document_id = _create_document()
    created = client.get(f"/api/documents/{document_id}").json()
    heading_id = created["elements"][0]["id"]

    response = client.post(f"/api/documents/{document_id}/elements", json={"elementType": "paragraph", "afterElementId": heading_id, "text": "Right after"})

    ordered = sorted(response.json()["elements"], key=lambda el: el["order"])
    assert ordered[1]["content"] == "Right after"


def test_add_element_unknown_after_element_returns_404():
    document_id = _create_document()

    response = client.post(f"/api/documents/{document_id}/elements", json={"elementType": "paragraph", "afterElementId": "does-not-exist"})

    assert response.status_code == 404


def test_add_element_unknown_document_returns_404():
    response = client.post("/api/documents/does-not-exist/elements", json={"elementType": "paragraph"})

    assert response.status_code == 404


def test_add_element_is_undoable():
    document_id = _create_document()
    before_count = len(client.get(f"/api/documents/{document_id}").json()["elements"])

    client.post(f"/api/documents/{document_id}/elements", json={"elementType": "paragraph", "text": "Undo me"})
    undo_response = client.post(f"/api/documents/{document_id}/undo")

    assert len(undo_response.json()["elements"]) == before_count


def test_update_content_replaces_element_text_and_preserves_style_ref():
    document_id = _create_document()
    client.post(f"/api/documents/{document_id}/format", data={"templateId": "academic-default"})
    current = client.get(f"/api/documents/{document_id}").json()
    heading = current["elements"][0]
    heading["content"] = "Edited Title"
    heading["inline"] = [{"text": "Edited Title", "marks": []}]

    response = client.put(f"/api/documents/{document_id}/content", json={"elements": current["elements"]})

    assert response.status_code == 200
    body = response.json()
    assert body["elements"][0]["content"] == "Edited Title"
    assert body["elements"][0]["styleRef"] == "Heading 1"  # formatting survives a pure content sync


def test_update_content_removing_an_element_prunes_its_override():
    document_id = _create_document()
    current = client.get(f"/api/documents/{document_id}").json()
    paragraph_id = current["elements"][1]["id"]
    client.patch(f"/api/documents/{document_id}/elements/{paragraph_id}/style", json={"property": "color", "value": "red"})

    remaining_elements = [current["elements"][0]]  # drop the paragraph entirely
    response = client.put(f"/api/documents/{document_id}/content", json={"elements": remaining_elements})

    assert response.status_code == 200
    body = response.json()
    assert len(body["elements"]) == 1
    assert paragraph_id not in body["resolvedStyles"]


def test_update_content_unknown_document_returns_404():
    response = client.put("/api/documents/does-not-exist/content", json={"elements": []})

    assert response.status_code == 404


def test_update_content_is_undoable():
    document_id = _create_document()
    current = client.get(f"/api/documents/{document_id}").json()
    original_content = current["elements"][0]["content"]
    current["elements"][0]["content"] = "Changed"
    current["elements"][0]["inline"] = [{"text": "Changed", "marks": []}]

    client.put(f"/api/documents/{document_id}/content", json={"elements": current["elements"]})
    undo_response = client.post(f"/api/documents/{document_id}/undo")

    assert undo_response.json()["elements"][0]["content"] == original_content


def test_formatted_document_is_persisted_in_the_database(api_db):
    """A restart can only lose what isn't in the database, so read the row back
    through a separate engine, independent of any in-process state."""
    document_id = _create_document()
    client.post(f"/api/documents/{document_id}/format", data={"templateId": "academic-default"})

    with OrmSession(api_db) as db:
        row = db.get(DocumentRow, document_id)
    assert row is not None
    assert row.data["templateId"] == "academic-default"


def test_undo_with_nothing_to_undo_returns_400():
    document_id = _create_document()

    response = client.post(f"/api/documents/{document_id}/undo")

    assert response.status_code == 400


def test_redo_with_nothing_to_redo_returns_400():
    document_id = _create_document()

    response = client.post(f"/api/documents/{document_id}/redo")

    assert response.status_code == 400


def test_undo_unknown_document_returns_404():
    response = client.post("/api/documents/does-not-exist/undo")

    assert response.status_code == 404


def test_export_docx_returns_downloadable_file():
    document_id = _create_document()

    response = client.get(f"/api/documents/{document_id}/export/docx")

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    assert "attachment" in response.headers["content-disposition"]
    assert response.content[:2] == b"PK"  # .docx is a zip archive


def test_export_docx_query_params_reach_the_builder():
    document_id = _create_document()
    client.patch(f"/api/documents/{document_id}/settings", json={"property": "header", "value": "Test Header"})

    default_response = client.get(f"/api/documents/{document_id}/export/docx")
    omitted_response = client.get(f"/api/documents/{document_id}/export/docx?includeHeaders=false")

    default_doc = DocxDocument(io.BytesIO(default_response.content))
    omitted_doc = DocxDocument(io.BytesIO(omitted_response.content))
    assert default_doc.sections[0].header.paragraphs[0].text == "Test Header"
    assert omitted_doc.sections[0].header.paragraphs[0].text == ""


def test_export_pdf_returns_downloadable_file():
    document_id = _create_document()

    response = client.get(f"/api/documents/{document_id}/export/pdf")

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert "attachment" in response.headers["content-disposition"]
    assert response.content.startswith(b"%PDF")


def test_export_docx_unknown_document_returns_404():
    response = client.get("/api/documents/does-not-exist/export/docx")

    assert response.status_code == 404


def test_export_pdf_unknown_document_returns_404():
    response = client.get("/api/documents/does-not-exist/export/pdf")

    assert response.status_code == 404


def test_style_analysis_returns_scored_result():
    document_id = _create_document()
    fake_response = AIStyleAnalysisResponse(
        consistency_score=0.9,
        tone="Formal and academic",
        summary="Consistent formal tone throughout.",
        flagged=[],
    )
    app.dependency_overrides[get_ai_provider] = lambda: FakeAIProvider([fake_response])
    try:
        response = client.post(f"/api/documents/{document_id}/style-analysis")
    finally:
        app.dependency_overrides.pop(get_ai_provider, None)

    assert response.status_code == 200
    body = response.json()
    assert body["status"] == "ok"
    assert body["consistencyScore"] == 0.9
    assert body["tone"] == "Formal and academic"


def test_style_analysis_flags_reference_real_elements_only():
    document_id = _create_document()
    document = client.get(f"/api/documents/{document_id}").json()
    real_element_id = document["elements"][0]["id"]
    fake_response = AIStyleAnalysisResponse(
        consistency_score=0.5,
        tone="Mixed",
        summary="One outlier paragraph.",
        flagged=[
            AIStyleFlag(element_id=real_element_id, reason="Breaks tone"),
            AIStyleFlag(element_id="hallucinated-id", reason="Should be dropped"),
        ],
    )
    app.dependency_overrides[get_ai_provider] = lambda: FakeAIProvider([fake_response])
    try:
        response = client.post(f"/api/documents/{document_id}/style-analysis")
    finally:
        app.dependency_overrides.pop(get_ai_provider, None)

    body = response.json()
    assert len(body["flagged"]) == 1
    assert body["flagged"][0]["elementId"] == real_element_id


def test_style_analysis_unknown_document_returns_404():
    response = client.post("/api/documents/does-not-exist/style-analysis")

    assert response.status_code == 404


def test_export_docx_with_non_latin1_title_does_not_crash():
    document_id = _create_document()
    client.patch(f"/api/documents/{document_id}", json={"title": "Курсова работа"})

    response = client.get(f"/api/documents/{document_id}/export/docx")

    assert response.status_code == 200
    assert "filename*=UTF-8''" in response.headers["content-disposition"]


def test_export_pdf_with_non_latin1_title_does_not_crash():
    document_id = _create_document()
    client.patch(f"/api/documents/{document_id}", json={"title": "Курсова работа"})

    response = client.get(f"/api/documents/{document_id}/export/pdf")

    assert response.status_code == 200
    assert "filename*=UTF-8''" in response.headers["content-disposition"]


def test_redo_unknown_document_returns_404():
    response = client.post("/api/documents/does-not-exist/redo")

    assert response.status_code == 404


def test_rename_document_updates_title():
    document_id = _create_document()

    response = client.patch(f"/api/documents/{document_id}", json={"title": "New Title"})

    assert response.status_code == 200
    assert response.json()["metadata"]["title"] == "New Title"
    assert client.get(f"/api/documents/{document_id}").json()["metadata"]["title"] == "New Title"


def test_rename_document_rejects_blank_title():
    document_id = _create_document()

    response = client.patch(f"/api/documents/{document_id}", json={"title": "   "})

    assert response.status_code == 422


def test_rename_document_unknown_document_returns_404():
    response = client.patch("/api/documents/does-not-exist", json={"title": "New Title"})

    assert response.status_code == 404


def test_rename_document_is_undoable():
    document_id = _create_document()
    original = client.get(f"/api/documents/{document_id}").json()["metadata"]["title"]

    client.patch(f"/api/documents/{document_id}", json={"title": "Renamed"})
    undo_response = client.post(f"/api/documents/{document_id}/undo")

    assert undo_response.json()["metadata"]["title"] == original


def test_set_page_setting_wins_over_template_and_survives_reformat():
    document_id = _create_document()

    set_response = client.patch(
        f"/api/documents/{document_id}/settings", json={"property": "pageSize", "value": "Legal"}
    )
    assert set_response.status_code == 200
    assert set_response.json()["settings"]["pageSize"] == "Legal"

    # Applying a template (academic-default sets pageSize A4) must not silently override the direct page setting.
    formatted = client.post(f"/api/documents/{document_id}/format", data={"templateId": "academic-default"})
    assert formatted.json()["document"]["settings"]["pageSize"] == "Legal"


def test_clear_page_setting_falls_back_to_template_default():
    document_id = _create_document()
    client.post(f"/api/documents/{document_id}/format", data={"templateId": "academic-default"})
    client.patch(f"/api/documents/{document_id}/settings", json={"property": "pageSize", "value": "Legal"})

    response = client.delete(f"/api/documents/{document_id}/settings/pageSize")

    assert response.status_code == 200
    assert response.json()["settings"]["pageSize"] == "A4"


def test_set_page_setting_unknown_document_returns_404():
    response = client.patch("/api/documents/does-not-exist/settings", json={"property": "pageSize", "value": "Legal"})

    assert response.status_code == 404


def test_set_page_setting_with_margin_and_unit():
    document_id = _create_document()

    response = client.patch(
        f"/api/documents/{document_id}/settings", json={"property": "marginTop", "value": "4", "unit": "cm"}
    )

    assert response.status_code == 200
    assert response.json()["settings"]["marginTopCm"] == 4.0


def test_a_new_document_has_the_render_specifications_look_from_the_start():
    """Resolved styles are worked out on creation and on every read, so a pasted
    document shows (and exports) the shared base look before any formatting."""
    response = client.post("/api/documents", json={"text": "# Title\n\nA body paragraph long enough to be real text."})

    assert response.status_code == 201
    styles = response.json()["resolvedStyles"]
    assert (styles["Heading 1"]["font-size"], styles["Heading 1"]["font-family"]) == ("20pt", "Arial")
    fetched = client.get(f"/api/documents/{response.json()['id']}").json()
    assert fetched["resolvedStyles"]["Paragraph"]["margin-bottom"] == "8pt"
    assert (fetched["settings"]["pageWidthMm"], fetched["settings"]["pageHeightMm"]) == (210.0, 297.0)
