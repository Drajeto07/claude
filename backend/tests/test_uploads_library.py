"""The editor's Uploads panel (S2-131a): the workspace's pictures listed, a page at a time, and put in a document again
by their id -- never stored twice, never another workspace's, and kept while any document shows them."""

import asyncio
import base64
import io
import json
from datetime import datetime, timedelta, timezone

import pytest
from docx import Document as DocxDocument
from fastapi.testclient import TestClient
from PIL import Image as PILImage
from pypdf import PdfReader
from sqlalchemy import func, select, update
from sqlalchemy.ext.asyncio import async_sessionmaker
from sqlalchemy.orm import Session as OrmSession
from sqlalchemy.pool import NullPool

from app.db.models import DocumentAsset
from app.db.session import make_engine
from app.main import app
from app.models.document import Document, DocumentMetadata, Element, ElementType, ImageContent
from app.services.asset_cleanup import sweep_unused_assets
from app.services.image_assets import drop_foreign_pictures
from app.storage.local_provider import LocalStorageProvider

_DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_LEFT_OUT = "A picture that isn't one of this workspace's was left out."


def _png(colour: str = "blue") -> bytes:
    buf = io.BytesIO()
    PILImage.new("RGB", (8, 8), colour).save(buf, format="PNG")
    return buf.getvalue()


def _docx_with_picture(png: bytes) -> bytes:
    doc = DocxDocument()
    doc.add_paragraph("Before the figure.")
    doc.add_picture(io.BytesIO(png))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _signed_in_client(email: str) -> TestClient:
    client = TestClient(app, base_url="https://testserver")
    assert client.post("/api/v1/auth/register", json={"email": email, "password": "long enough password"}).status_code == 201
    return client


@pytest.fixture
def alice(api_db):
    return _signed_in_client("alice@example.com")


def _upload_picture_document(client: TestClient, png: bytes) -> dict:
    response = client.post("/api/v1/documents/upload", files={"file": ("report.docx", _docx_with_picture(png), _DOCX_TYPE)})
    assert response.status_code == 201
    return response.json()


def _blank(client: TestClient, title: str = "Notes") -> dict:
    return client.post("/api/v1/documents", json={"text": f"# {title}\n\nSome text to start with here."}).json()


def _picture_ids(document: dict) -> list[str]:
    return [e["image"]["assetId"] for e in document["elements"] if e["type"] == "image"]


def _by_id(asset_id: str) -> dict:
    return {"type": "image", "content": "", "order": 99, "image": {"src": "", "assetId": asset_id}}


def _pasted(png: bytes) -> dict:
    return {"type": "image", "content": "", "order": 99, "image": {"src": "data:image/png;base64," + base64.b64encode(png).decode()}}


def _assets(api_db) -> list[tuple[str, str]]:
    with OrmSession(api_db) as db:
        return [(row.id, row.content_type) for row in db.execute(select(DocumentAsset)).scalars()]


def _asset_count(api_db) -> int:
    with OrmSession(api_db) as db:
        return db.execute(select(func.count()).select_from(DocumentAsset)).scalar_one()


# --- the list ------------------------------------------------------------------------------------------------------


def test_an_empty_workspace_lists_no_pictures_and_only_the_signed_in_may_ask(alice):
    assert alice.get("/api/v1/assets").json() == {"items": [], "next": None}
    assert TestClient(app, base_url="https://testserver").get("/api/v1/assets").status_code == 401


def test_the_list_is_the_workspace_s_pictures_only(api_db, alice):
    document = _upload_picture_document(alice, _png())
    [picture] = _picture_ids(document)
    # The Word file kept as the original is an asset of the workspace too -- but not a picture.
    assert sorted(kind for _, kind in _assets(api_db)) == [_DOCX_TYPE, "image/png"]
    bob = _signed_in_client("bob@example.com")
    _upload_picture_document(bob, _png("red"))

    [item] = alice.get("/api/v1/assets").json()["items"]

    assert item["id"] == picture and item["contentType"] == "image/png" and item["sizeBytes"] == len(_png())
    assert set(item) == {"id", "contentType", "sizeBytes", "createdAt"}
    assert picture not in {each["id"] for each in bob.get("/api/v1/assets").json()["items"]}
    assert len(bob.get("/api/v1/assets").json()["items"]) == 1


def test_the_list_comes_a_bounded_page_at_a_time_newest_first(api_db, alice):
    document = _blank(alice)
    saved = alice.put(f"/api/v1/documents/{document['id']}/content", json={"elements": document["elements"] + [_pasted(_png()) for _ in range(30)]})
    assert saved.status_code == 200
    stored = _picture_ids(saved.json())
    assert len(set(stored)) == 30
    # Two of them stored in the same instant: the id orders them (the cursor is the time and the id).
    with OrmSession(api_db) as db:
        same = datetime(2026, 10, 1, tzinfo=timezone.utc)
        db.execute(update(DocumentAsset).where(DocumentAsset.id.in_(stored[:2])).values(created_at=same))
        db.commit()

    first = alice.get("/api/v1/assets").json()
    second = alice.get("/api/v1/assets", params={"before": first["next"]}).json()

    assert len(first["items"]) == 24 and first["next"] == first["items"][-1]["id"]
    assert len(second["items"]) == 6 and second["next"] is None
    listed = [item["id"] for item in first["items"] + second["items"]]
    assert sorted(listed) == sorted(stored)  # every one, none twice
    times = [item["createdAt"] for item in first["items"] + second["items"]]
    assert times == sorted(times, reverse=True)
    assert sorted(listed[-2:], reverse=True) == listed[-2:]
    assert [i["id"] for i in alice.get("/api/v1/assets", params={"limit": 48}).json()["items"]] == listed
    assert alice.get("/api/v1/assets", params={"limit": 49}).status_code == 422
    assert alice.get("/api/v1/assets", params={"limit": 0}).status_code == 422


def test_a_cursor_that_isn_t_one_of_the_workspace_s_pictures_is_refused(api_db, alice):
    alices = _picture_ids(_upload_picture_document(alice, _png()))[0]
    bob = _signed_in_client("bob@example.com")
    docx_original = next(asset_id for asset_id, kind in _assets(api_db) if kind == _DOCX_TYPE)

    assert bob.get("/api/v1/assets", params={"before": alices}).status_code == 400
    assert alice.get("/api/v1/assets", params={"before": "does-not-exist"}).status_code == 400
    assert alice.get("/api/v1/assets", params={"before": docx_original}).status_code == 400
    assert alice.get("/api/v1/assets", params={"before": alices}).json() == {"items": [], "next": None}


# --- putting a listed picture in again --------------------------------------------------------------------------------


def test_a_picture_put_in_again_is_the_same_asset_not_a_new_one(api_db, alice):
    first = _upload_picture_document(alice, _png())
    [picture] = _picture_ids(first)
    second = _blank(alice, "Second")
    before = _asset_count(api_db)

    saved = alice.put(f"/api/v1/documents/{second['id']}/content", json={"elements": second["elements"] + [_by_id(picture)]})

    assert saved.status_code == 200
    assert _picture_ids(saved.json()) == [picture]
    assert _LEFT_OUT not in saved.json()["unsupportedFeatures"]
    assert _asset_count(api_db) == before
    assert alice.get(f"/api/v1/assets/{picture}").content == _png()
    # Saved again as the editor re-sends it: still the one asset.
    alice.put(f"/api/v1/documents/{second['id']}/content", json={"elements": saved.json()["elements"]})
    assert _asset_count(api_db) == before
    assert [item["id"] for item in alice.get("/api/v1/assets").json()["items"]] == [picture]


def test_a_picture_put_in_again_by_a_patch_is_kept_too(api_db, alice):
    [picture] = _picture_ids(_upload_picture_document(alice, _png()))
    second = _blank(alice, "Second")
    patch = {"added": [{"after": second["elements"][-1]["id"], "element": _by_id(picture)}]}

    answer = alice.patch(f"/api/v1/documents/{second['id']}/content", json=patch, headers={"If-Match": str(second["revision"])})

    assert answer.status_code == 200
    assert _picture_ids(alice.get(f"/api/v1/documents/{second['id']}").json()) == [picture]


def test_the_exports_of_both_documents_carry_the_one_picture(alice):
    [picture] = _picture_ids(_upload_picture_document(alice, _png()))
    second = _blank(alice, "Second")
    alice.put(f"/api/v1/documents/{second['id']}/content", json={"elements": second["elements"] + [_by_id(picture)]})

    docx = DocxDocument(io.BytesIO(alice.get(f"/api/v1/documents/{second['id']}/export/docx").content))
    pdf = PdfReader(io.BytesIO(alice.get(f"/api/v1/documents/{second['id']}/export/pdf").content))

    assert len(docx.inline_shapes) == 1
    [blob] = [part.blob for part in docx.part.package.parts if part.partname.startswith("/word/media/")]
    assert PILImage.open(io.BytesIO(blob)).size == (8, 8)
    assert sum(len(page.images) for page in pdf.pages) == 1


# --- what a save refuses to show ----------------------------------------------------------------------------------------


def test_another_workspace_s_picture_is_left_out_of_a_save(api_db, alice):
    [alices] = _picture_ids(_upload_picture_document(alice, _png()))
    bob = _signed_in_client("bob@example.com")
    bobs = _blank(bob, "Mine")
    before = _asset_count(api_db)

    saved = bob.put(f"/api/v1/documents/{bobs['id']}/content", json={"elements": bobs["elements"] + [_by_id(alices)]})

    assert saved.status_code == 200
    assert _picture_ids(saved.json()) == []
    assert _LEFT_OUT in saved.json()["unsupportedFeatures"]
    assert [e["order"] for e in saved.json()["elements"]] == list(range(len(bobs["elements"])))
    assert _asset_count(api_db) == before
    assert _picture_ids(bob.get(f"/api/v1/documents/{bobs['id']}").json()) == []
    export = DocxDocument(io.BytesIO(bob.get(f"/api/v1/documents/{bobs['id']}/export/docx").content))
    assert len(export.inline_shapes) == 0


def test_another_workspace_s_picture_is_left_out_of_a_patch_too(alice):
    [alices] = _picture_ids(_upload_picture_document(alice, _png()))
    bob = _signed_in_client("bob@example.com")
    bobs = _blank(bob, "Mine")
    patch = {"added": [{"after": bobs["elements"][-1]["id"], "element": _by_id(alices)}]}

    answer = bob.patch(f"/api/v1/documents/{bobs['id']}/content", json=patch, headers={"If-Match": str(bobs["revision"])})

    assert answer.status_code == 200
    stored = bob.get(f"/api/v1/documents/{bobs['id']}").json()
    assert _picture_ids(stored) == [] and _LEFT_OUT in stored["unsupportedFeatures"]


def test_a_missing_picture_or_a_kept_word_original_is_left_out(api_db, alice):
    _upload_picture_document(alice, _png())
    docx_original = next(asset_id for asset_id, kind in _assets(api_db) if kind == _DOCX_TYPE)
    document = _blank(alice)

    saved = alice.put(
        f"/api/v1/documents/{document['id']}/content",
        json={"elements": document["elements"] + [_by_id("does-not-exist"), _by_id(docx_original)]},
    ).json()

    assert _picture_ids(saved) == [] and _LEFT_OUT in saved["unsupportedFeatures"]


class _Pictures:
    def __init__(self, own: set[str]) -> None:
        self.own = own
        self.asked: list[set[str]] = []

    async def pictures_in(self, workspace_id: str, asset_ids) -> set[str]:
        self.asked.append(set(asset_ids))
        return self.own & set(asset_ids)


def _picture(asset_id: str, order: int = 0) -> Element:
    return Element(type=ElementType.IMAGE, content="", image=ImageContent(src="", assetId=asset_id), order=order)


def test_a_foreign_picture_nested_in_a_block_is_left_out_and_what_the_document_showed_stays():
    quote = Element(type=ElementType.QUOTE, content="", order=1, children=[_picture("mine", 0), _picture("foreign", 1)])
    document = Document(
        metadata=DocumentMetadata(title="T"),
        elements=[_picture("planted-long-ago", 0), quote, _picture("foreign", 2)],
    )
    pictures = _Pictures({"mine"})

    changed = asyncio.run(drop_foreign_pictures(document, {"planted-long-ago"}, pictures, "ws"))

    assert changed
    assert pictures.asked == [{"mine", "foreign"}]  # only what the save adds is looked up
    assert [e.image.assetId if e.image else e.type for e in document.elements] == ["planted-long-ago", ElementType.QUOTE]
    assert [e.image.assetId for e in document.elements[1].children] == ["mine"]
    assert [e.order for e in document.elements] == [0, 1] and document.elements[1].children[0].order == 0
    assert document.unsupportedFeatures == [_LEFT_OUT]


def test_a_save_that_adds_no_picture_looks_nothing_up():
    document = Document(metadata=DocumentMetadata(title="T"), elements=[_picture("a")])
    pictures = _Pictures(set())

    assert not asyncio.run(drop_foreign_pictures(document, {"a"}, pictures, "ws"))
    assert pictures.asked == [] and document.unsupportedFeatures == []


# --- kept while any document shows it -----------------------------------------------------------------------------------


def test_deleting_one_document_keeps_the_picture_another_one_shows(api_db, alice, tmp_path):
    first = _upload_picture_document(alice, _png())
    [picture] = _picture_ids(first)
    second = _blank(alice, "Second")
    alice.put(f"/api/v1/documents/{second['id']}/content", json={"elements": second["elements"] + [_by_id(picture)]})

    assert alice.delete(f"/api/v1/documents/{first['id']}").status_code == 204

    assert alice.get(f"/api/v1/assets/{picture}").content == _png()
    assert [item["id"] for item in alice.get("/api/v1/assets").json()["items"]] == [picture]
    docx = DocxDocument(io.BytesIO(alice.get(f"/api/v1/documents/{second['id']}/export/docx").content))
    assert len(docx.inline_shapes) == 1

    # The daily sweep (where it runs) keeps it too, once out of its grace period -- the other document shows it.
    with OrmSession(api_db) as db:
        db.execute(update(DocumentAsset).values(created_at=datetime.now(timezone.utc) - timedelta(days=3)))
        db.commit()

    async def sweep() -> int:
        engine = make_engine(f"sqlite+aiosqlite:///{tmp_path / 'api.db'}", poolclass=NullPool)
        try:
            return await sweep_unused_assets(async_sessionmaker(engine, expire_on_commit=False), LocalStorageProvider(tmp_path / "assets"))
        finally:
            await engine.dispose()

    # It ran: the deleted document's kept Word original, which nothing shows, went.
    assert asyncio.run(sweep()) == 1
    assert _assets(api_db) == [(picture, "image/png")]
    assert alice.get(f"/api/v1/assets/{picture}").content == _png()


# --- the real workspace of each document, not only the "one workspace per user" of today -------------------------------


def _user_id(api_db, email: str) -> str:
    from app.db.models import User

    with OrmSession(api_db) as db:
        return db.execute(select(User.id).where(User.email == email)).scalar_one()


def _second_workspace_for(api_db, email: str, document_ids: list[str]) -> str:
    """A second workspace the user is a member of (as a shared one would be), holding these of their documents and
    the pictures those documents stored."""
    from app.db.models import Workspace, WorkspaceMember
    from app.db.models import Document as DocumentRow

    with OrmSession(api_db) as db:
        workspace = Workspace(name="Shared", slug=f"shared-{len(document_ids)}-{email}")
        db.add(workspace)
        db.flush()
        db.add(WorkspaceMember(workspace_id=workspace.id, user_id=_user_id(api_db, email), role="member"))
        db.execute(update(DocumentRow).where(DocumentRow.id.in_(document_ids)).values(workspace_id=workspace.id))
        db.execute(update(DocumentAsset).where(DocumentAsset.document_id.in_(document_ids)).values(workspace_id=workspace.id))
        db.commit()
        return workspace.id


def test_a_document_takes_only_its_own_workspace_s_pictures_even_from_a_member_of_both(api_db, alice):
    [own] = _picture_ids(_upload_picture_document(alice, _png()))  # alice's own workspace
    shared_source = _upload_picture_document(alice, _png("red"))
    [shared_picture] = _picture_ids(shared_source)
    shared_target = _blank(alice, "Shared target")
    _second_workspace_for(api_db, "alice@example.com", [shared_source["id"], shared_target["id"]])

    saved = alice.put(
        f"/api/v1/documents/{shared_target['id']}/content",
        json={"elements": shared_target["elements"] + [_by_id(own), _by_id(shared_picture)]},
    ).json()

    # alice may see both pictures, but a document of the shared workspace may show only that workspace's.
    assert _picture_ids(saved) == [shared_picture]
    assert _LEFT_OUT in saved["unsupportedFeatures"]
    assert alice.get(f"/api/v1/assets/{own}").status_code == 200


def test_the_list_is_the_edited_document_s_workspace(api_db, alice):
    [own] = _picture_ids(_upload_picture_document(alice, _png()))
    shared_source = _upload_picture_document(alice, _png("red"))
    [shared_picture] = _picture_ids(shared_source)
    shared_target = _blank(alice, "Shared target")
    _second_workspace_for(api_db, "alice@example.com", [shared_source["id"], shared_target["id"]])
    own_document = _blank(alice, "Own")
    bob = _signed_in_client("bob@example.com")
    bobs = _blank(bob, "Bob's")

    def listed(**params) -> list[str]:
        return [item["id"] for item in alice.get("/api/v1/assets", params=params).json()["items"]]

    assert listed(documentId=shared_target["id"]) == [shared_picture]
    assert listed(documentId=own_document["id"]) == [own]
    assert listed() == [own]  # no document named: the user's own workspace
    # Another person's document, or none: 404 alike, nothing listed.
    refused = alice.get("/api/v1/assets", params={"documentId": bobs["id"]})
    missing = alice.get("/api/v1/assets", params={"documentId": "does-not-exist"})
    assert refused.status_code == missing.status_code == 404
    assert refused.json()["message"] == missing.json()["message"]
    # Nothing of a stored picture but its id, type, size and date is told: never where it is stored.
    item = alice.get("/api/v1/assets", params={"documentId": shared_target["id"]}).json()["items"][0]
    assert "storage" not in json.dumps(item).lower() and "/" not in item["id"]


def test_a_new_document_made_from_another_keeps_only_its_workspace_s_pictures(api_db, alice):
    """Every new document goes through DocumentService.create (imports, translated versions, clean copies): a copy
    that carries pictures by id keeps them only when they are of the workspace it goes into."""
    from app.db.session import make_engine as engine_for
    from app.services.document_service import DocumentService

    [own] = _picture_ids(_upload_picture_document(alice, _png()))
    bob = _signed_in_client("bob@example.com")
    [bobs] = _picture_ids(_upload_picture_document(bob, _png("red")))
    copy = Document(metadata=DocumentMetadata(title="Copy"), elements=[_picture(own, 0), _picture(bobs, 1)])

    async def create() -> Document:
        engine = engine_for(f"sqlite+aiosqlite:///{api_db.url.database}", poolclass=NullPool)
        try:
            async with async_sessionmaker(engine, expire_on_commit=False)() as session:
                service = DocumentService(session, user_id=_user_id(api_db, "alice@example.com"), storage=app.dependency_overrides[_storage_dependency()]())
                return await service.create(copy)
        finally:
            await engine.dispose()

    created = asyncio.run(create())

    assert [e.image.assetId for e in created.elements if e.image] == [own]
    assert _LEFT_OUT in created.unsupportedFeatures
    assert _picture_ids(alice.get(f"/api/v1/documents/{created.id}").json()) == [own]


def _storage_dependency():
    from app.api.deps import get_storage_provider

    return get_storage_provider


def test_restoring_a_version_brings_back_only_the_workspace_s_own_pictures(api_db, alice):
    """A version made before saves were checked could hold another workspace's picture id: restoring it brings back
    the document's own pictures and leaves that one out."""
    from app.db.models import Document as DocumentRow
    from app.db.models import DocumentVersion

    [own] = _picture_ids(_upload_picture_document(alice, _png()))
    bob = _signed_in_client("bob@example.com")
    [bobs] = _picture_ids(_upload_picture_document(bob, _png("red")))
    document = _blank(alice)
    with OrmSession(api_db) as db:
        row = db.get(DocumentRow, document["id"])
        version = db.execute(
            select(DocumentVersion).where(DocumentVersion.document_id == row.id).order_by(DocumentVersion.revision_number)
        ).scalars().first()
        data = version.data
        data["elements"] = data["elements"] + [
            {**_by_id(own), "id": "kept-own-picture", "order": len(data["elements"])},
            {**_by_id(bobs), "id": "planted-picture", "order": len(data["elements"]) + 1},
        ]
        version.data = data
        number = version.revision_number
        db.commit()

    restored = alice.post(f"/api/v1/documents/{document['id']}/versions/{number}/restore")

    assert restored.status_code == 200
    assert _picture_ids(restored.json()) == [own]
    assert _LEFT_OUT in restored.json()["unsupportedFeatures"]


def _stored_files(tmp_path) -> int:
    return sum(1 for path in (tmp_path / "assets").rglob("*") if path.is_file())


def test_reuse_stores_no_second_file_and_a_version_alone_keeps_the_picture_from_the_sweep(api_db, alice, tmp_path):
    first = _upload_picture_document(alice, _png())
    [picture] = _picture_ids(first)
    second = _blank(alice, "Second")
    files = _stored_files(tmp_path)

    shown = alice.put(f"/api/v1/documents/{second['id']}/content", json={"elements": second["elements"] + [_by_id(picture)]}).json()
    assert _stored_files(tmp_path) == files  # the same stored file, not a copy
    # A minute later (a save within a minute of the last step joins it: VersionHistory CONTENT_MERGE_WINDOW).
    from app.db.models import DocumentVersion

    with OrmSession(api_db) as db:
        db.execute(update(DocumentVersion).values(created_at=datetime.now(timezone.utc) - timedelta(minutes=2)))
        db.commit()
    # Taken out of the second document again, and the first document deleted: only the second's history shows it.
    alice.put(f"/api/v1/documents/{second['id']}/content", json={"elements": [e for e in shown["elements"] if e["type"] != "image"]})
    assert alice.delete(f"/api/v1/documents/{first['id']}").status_code == 204
    with OrmSession(api_db) as db:
        db.execute(update(DocumentAsset).values(created_at=datetime.now(timezone.utc) - timedelta(days=3)))
        db.commit()

    async def sweep() -> int:
        engine = make_engine(f"sqlite+aiosqlite:///{tmp_path / 'api.db'}", poolclass=NullPool)
        try:
            return await sweep_unused_assets(async_sessionmaker(engine, expire_on_commit=False), LocalStorageProvider(tmp_path / "assets"))
        finally:
            await engine.dispose()

    asyncio.run(sweep())

    # Undo could bring it back, so it stays: row, file and served bytes.
    assert picture in {asset_id for asset_id, _ in _assets(api_db)}
    assert alice.get(f"/api/v1/assets/{picture}").content == _png()
    undone = alice.post(f"/api/v1/documents/{second['id']}/undo")
    assert undone.status_code == 200 and _picture_ids(undone.json()) == [picture]
