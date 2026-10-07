"""Content controls in tables, lists and text boxes (tracker DOCX-023A): kept like the ones in
top-level paragraphs -- a control in a cell's paragraph, a list item or a text box as a fragment
around its text, one around a cell or a row with that cell or row -- so a block written anew
(changed here, restyled, or a file written from scratch) still has them; before, they were kept
only as their text, copied with an unchanged block and lost once it was written anew. And rows
inside a row-level control (a repeating section) are read at all: before, they were left out."""

import io

import pytest
from docx import Document as DocxDocument
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from fastapi.testclient import TestClient

from app.export.docx_export import build_docx
from app.export.package_check import package_problems
from app.formatting.engine import recompute_styles
from app.main import app
from app.models.document import ElementType
from app.services.ingestion_service import build_document_from_docx

_NS = nsdecls("w")
_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
client = TestClient(app, base_url="https://testserver")


def _control(tag: str, inner: str, number: int, *, plain: bool = True) -> str:
    """A content control: a plain-text one (in a paragraph), or a rich-text one -- the only kind
    Word takes around a cell or rows."""
    kind = "<w:text/>" if plain else ""
    return (
        f'<w:sdt {_NS}><w:sdtPr><w:alias w:val="{tag.title()}"/><w:tag w:val="{tag}"/><w:id w:val="{number}"/>{kind}</w:sdtPr>'
        f"<w:sdtContent>{inner}</w:sdtContent></w:sdt>"
    )


def _run(text: str) -> str:
    return f'<w:r><w:t xml:space="preserve">{text}</w:t></w:r>'


def _wrap(element, tag: str, number: int) -> None:
    """The element moved into a content control of its own, where it was."""
    control = parse_xml(_control(tag, "", number, plain=False))
    element.getparent().replace(element, control)
    control.find(qn("w:sdtContent")).append(element)


def _word_file() -> bytes:
    """A table: a control in a cell's paragraph, one in a cell of two paragraphs, a control
    around a cell, and two rows in one control (a repeating section's); a list item with a
    control; a text box whose paragraph has one."""
    doc = DocxDocument()
    doc.add_paragraph("Before the table.")
    table = doc.add_table(rows=4, cols=2)
    texts = [["Name", "Signed"], ["Client", "On date"], ["Row A", "A2"], ["Row B", "B2"]]
    for row, values in zip(table.rows, texts):
        for cell, value in zip(row.cells, values):
            cell.paragraphs[0].add_run(value)
    first = table.rows[0].cells[0].paragraphs[0]._p
    first.append(parse_xml(f"<w:p {_NS}>{_control('name', _run(' Ada Lovelace'), 101)}</w:p>")[0])
    two = table.rows[0].cells[1]._tc
    two.append(parse_xml(f"<w:p {_NS}>{_run('By: ')}{_control('signer', _run('Grace Hopper'), 102)}</w:p>"))
    _wrap(table.rows[1].cells[1]._tc, "datecell", 103)
    rows = [table.rows[2]._tr, table.rows[3]._tr]
    section = parse_xml(_control("items", "", 104, plain=False))
    rows[0].getparent().replace(rows[0], section)
    for number, row in enumerate(rows, start=1):  # each row in an item of its own, as a repeating section's
        item = parse_xml(_control(f"item{number}", "", 110 + number, plain=False))
        item.find(qn("w:sdtContent")).append(row)
        section.find(qn("w:sdtContent")).append(item)
    item = doc.add_paragraph("First item: ", style="List Bullet")._p
    item.append(parse_xml(f"<w:p {_NS}>{_control('choice', _run('yes'), 105)}</w:p>")[0])
    doc.add_paragraph("Second item.", style="List Bullet")
    anchor = doc.add_paragraph("Anchor.")._p
    anchor.append(
        parse_xml(
            f'<w:r {_NS} xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
            'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
            'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><w:drawing><wp:inline distT="0" distB="0" distL="0" distR="0">'
            '<wp:extent cx="2286000" cy="457200"/><wp:docPr id="7" name="Text Box 7"/><wp:cNvGraphicFramePr/>'
            '<a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"><wps:wsp>'
            '<wps:cNvSpPr txBox="1"/><wps:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="2286000" cy="457200"/></a:xfrm>'
            '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></wps:spPr><wps:txbx><w:txbxContent>'
            f"<w:p>{_run('Boxed: ')}{_control('boxed', _run('inside'), 106)}</w:p>"
            '</w:txbxContent></wps:txbx><wps:bodyPr/></wps:wsp></a:graphicData></a:graphic></wp:inline></w:drawing></w:r>'
        )
    )
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _controls(data: bytes) -> dict[str, str]:
    """Each content control in the body: its tag, and the text inside it."""
    body = DocxDocument(io.BytesIO(data)).element.body
    found: dict[str, str] = {}
    for sdt in body.iter(qn("w:sdt")):
        tag = sdt.find(f"{qn('w:sdtPr')}/{qn('w:tag')}")
        content = sdt.find(qn("w:sdtContent"))
        if tag is not None and content is not None:
            found[tag.get(qn("w:val"))] = "".join(text.text or "" for text in content.iter(qn("w:t")))
    return found


_EXPECTED = {
    "name": " Ada Lovelace",
    "signer": "Grace Hopper",
    "datecell": "On date",
    "items": "Row AA2Row BB2",
    "item1": "Row AA2",
    "item2": "Row BB2",
    "choice": "yes",
    "boxed": "inside",
}


def test_every_control_in_a_table_a_list_and_a_text_box_is_read_with_all_its_text():
    document = build_document_from_docx(_word_file(), "controls.docx", None)
    assert document.importReport.contentStatus == "verified"  # the repeating section's rows too
    table = next(element for element in document.elements if element.type == ElementType.TABLE).table
    assert [[cell.inline[0].text if cell.inline else "" for cell in row.cells] for row in table.rows][2:] == [["Row A", "A2"], ["Row B", "B2"]]
    features = {item.feature: item for item in document.importReport.items}
    assert "docx.content_control.nested" not in features and "docx.content_control" in features
    assert "docx.preserved.flattened" not in features  # none of them kept only as its text
    assert _controls(_word_file()) == _EXPECTED


@pytest.mark.parametrize("into_source", [False, True])
def test_a_file_written_anew_has_every_control_back_around_its_text_cell_and_rows(into_source):
    source = _word_file()
    document = build_document_from_docx(source, "controls.docx", None)
    recompute_styles(document)
    if into_source:  # every block changed here: each written anew into the original
        for element in document.elements:
            element.sourceBlocks = element.sourceHash = None
    exported = build_docx(document, source=source if into_source else None)
    assert package_problems(exported) == []
    assert _controls(exported) == _EXPECTED
    again = build_document_from_docx(exported, "again.docx", None)
    assert again.importReport.contentStatus == "verified" and _controls(build_docx(again)) == _EXPECTED


def test_a_control_whose_text_was_changed_here_goes_around_the_new_text():
    document = build_document_from_docx(_word_file(), "controls.docx", None)
    table = next(element for element in document.elements if element.type == ElementType.TABLE).table
    cell = table.rows[0].cells[0]
    cell.inline = [run.model_copy(update={"text": run.text.replace("Ada Lovelace", "Mary Somerville")}) for run in cell.inline]
    items = next(element for element in document.elements if element.type == ElementType.LIST).listItems
    items[0].inline = [run.model_copy(update={"text": run.text.replace("yes", "no")}) for run in items[0].inline]
    recompute_styles(document)
    found = _controls(build_docx(document))
    assert found["name"] == " Mary Somerville" and found["choice"] == "no"


@pytest.fixture
def signed_in(api_db):
    client.cookies.clear()
    assert client.post("/api/v1/auth/register", json={"email": "nested-controls@example.com", "password": "long enough password"}).status_code == 201
    yield
    client.cookies.clear()


def test_a_save_from_the_editor_keeps_them_and_can_t_add_one(signed_in):
    uploaded = client.post("/api/v1/documents/upload", files={"file": ("controls.docx", _word_file(), _DOCX)}).json()
    elements = uploaded["elements"]
    for element in elements:  # as the editor sends them: what the import kept is the server's
        for row in (element.get("table") or {}).get("rows", []):
            row.pop("preservedAttributes", None)
            for cell in row["cells"]:
                cell.pop("preservedAttributes", None)
        for item in element.get("listItems") or []:
            item["preservedAttributes"] = {"ooxml": [{"kind": "field", "instr": "DDEAUTO calc", "start": 0, "end": 1, "text": "F"}]}
    saved = client.put(f"/api/v1/documents/{uploaded['id']}/content", json={"elements": elements}, headers={"If-Match": str(uploaded["revision"])})
    assert saved.status_code == 200, saved.text
    lists = [element for element in saved.json()["elements"] if element["type"] == "list"]
    assert lists[0]["listItems"][0]["preservedAttributes"]["ooxml"][0]["kind"] == "control"  # the stored one, not the one sent
    assert lists[0]["listItems"][1]["preservedAttributes"] is None
    exported = client.get(f"/api/v1/documents/{uploaded['id']}/export/docx")
    assert exported.status_code == 200 and _controls(exported.content) == _EXPECTED
    assert "DDEAUTO" not in DocxDocument(io.BytesIO(exported.content)).element.xml
