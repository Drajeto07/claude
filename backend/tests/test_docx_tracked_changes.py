"""Tracked changes (tracker DOCX-022): never accepted without saying so, and never
lost without it. The import reads a file's tracked changes as accepted -- the editor
shows the document as it would be -- and a Word export into the file keeps them in
every block not changed here. Accepting them all instead is the user's choice. Before
this, every export accepted them (measured in Word: a07's 7 revisions, 0 after)."""

import io
import re
import zipfile
from pathlib import Path

import pytest
from docx import Document as DocxDocument
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from fastapi.testclient import TestClient

from app.export.docx_export import _self_contained, build_docx
from app.export.package_check import package_problems
from app.export.provenance import stamp
from app.fidelity.imports import TRACKED_ACCEPTED, TRACKED_KEPT
from app.fidelity.report import ReportBuilder
from app.formatting.engine import recompute_styles
from app.main import app
from app.models.document import Document, InlineRun
from app.parsers.docx import parse_docx

A07 = Path(__file__).parent / "fixtures" / "word" / "a07-review.docx"
_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_W = nsdecls("w")
client = TestClient(app, base_url="https://testserver")


def _marks(data: bytes) -> dict[str, int]:
    """How many of each mark Word keeps of a tracked change the body holds."""
    xml = zipfile.ZipFile(io.BytesIO(data)).read("word/document.xml").decode("utf-8")
    return {kind: len(re.findall(rf"<w:{kind}\b", xml)) for kind in ("ins", "del", "moveFrom", "moveTo", "rPrChange")}


_A07 = {"ins": 7, "del": 7, "moveFrom": 2, "moveTo": 2, "rPrChange": 1}


def _a07() -> tuple[bytes, Document]:
    source = A07.read_bytes()
    document = parse_docx(source, A07.name)
    recompute_styles(document)
    stamp(document)  # as an upload does
    return source, document


def _saved(word) -> bytes:
    buffer = io.BytesIO()
    word.save(buffer)
    return buffer.getvalue()


def test_the_import_reads_tracked_changes_as_accepted_and_keeps_them_for_word():
    source, document = _a07()
    assert _marks(source) == _A07

    table = next(element.table for element in document.elements if element.table)

    assert document.trackedChanges == "kept"
    assert [[cell.inline[0].text if cell.inline else "" for cell in row.cells] for row in table.rows] == [
        ["row two", ""],
        ["inserted row", ""],
    ]  # the row deleted while tracking is gone, as accepted (it came in as an empty row)
    assert "Tracked changes were imported as accepted (insertions kept, deletions removed)." in document.unsupportedFeatures


def test_an_unchanged_document_keeps_its_tracked_changes_in_a_word_export():
    source, document = _a07()

    exported = build_docx(document, source=source)

    assert _marks(exported) == _A07  # Word shows its 7 revisions again (measured; 0 before)
    assert package_problems(exported) == []


def test_a_block_changed_here_has_its_changes_accepted_and_the_export_says_so():
    source, document = _a07()
    edited = next(element for element in document.elements if element.content.startswith("Base sentence"))
    edited.inline = [*(edited.inline or []), InlineRun(text=" Edited.")]
    edited.content = f"{edited.content} Edited."
    report = ReportBuilder()

    exported = build_docx(document, source=source, report=report)

    assert _marks(exported) == {**_A07, "ins": 6}  # that paragraph's insertion accepted, the rest kept
    [rewritten] = [item for item in report.items() if item.feature == "export.docx.rewritten_blocks"]
    assert "tracked changes" in rewritten.reason
    assert package_problems(exported) == []


def test_tracked_changes_accepted_as_chosen_are_in_no_export():
    source, document = _a07()
    document.trackedChanges = "accepted"
    report = ReportBuilder()

    exported = build_docx(document, source=source, report=report)

    assert _marks(exported) == dict.fromkeys(_A07, 0)
    assert not any("tracked changes" in item.reason for item in report.items())  # chosen: not reported lost
    assert package_problems(exported) == []


def test_formatting_changes_alone_are_noted_and_kept():
    word = DocxDocument()
    paragraph = word.add_paragraph()
    paragraph._p.append(
        parse_xml(
            f'<w:r {_W}><w:rPr><w:b/><w:rPrChange w:id="1" w:author="Ana" w:date="2026-01-01T00:00:00Z"><w:rPr/></w:rPrChange>'
            "</w:rPr><w:t>Made bold while tracking.</w:t></w:r>"
        )
    )
    source = _saved(word)
    document = parse_docx(source, "bold.docx")
    recompute_styles(document)
    stamp(document)

    assert document.trackedChanges == "kept"
    assert "Tracked changes were imported as accepted (insertions kept, deletions removed)." in document.unsupportedFeatures
    assert _marks(build_docx(document, source=source))["rPrChange"] == 1


def test_a_file_without_tracked_changes_has_nothing_to_choose():
    word = DocxDocument()
    word.add_paragraph("Nothing tracked.")

    assert parse_docx(_saved(word), "plain.docx").trackedChanges is None


def _paragraph(xml: str):
    return parse_xml(f"<w:p {_W}>{xml}</w:p>")


def test_moved_text_is_copied_only_with_both_ends_of_its_range():
    started = _paragraph('<w:moveFromRangeStart w:id="7" w:name="move1"/><w:moveFrom w:id="8"><w:r><w:t>Moved.</w:t></w:r></w:moveFrom>')
    ended = _paragraph('<w:moveFromRangeEnd w:id="7"/>')

    assert _self_contained([started, ended], revisions=True)
    assert not _self_contained([started], revisions=True)  # the range would be left open
    assert not _self_contained([started, ended], revisions=False)  # accepted: never copied


# -- the choice, through the API ----------------------------------------------------------


@pytest.fixture
def uploaded(api_db):
    client.cookies.clear()
    assert client.post("/api/v1/auth/register", json={"email": "tracked@example.com", "password": "long enough password"}).status_code == 201
    response = client.post("/api/v1/documents/upload", files={"file": (A07.name, A07.read_bytes(), _DOCX)})
    assert response.status_code == 201, response.text
    yield response.json()
    client.cookies.clear()


def _tracked_item(document: dict) -> dict:
    return next(item for item in document["importReport"]["items"] if item["feature"] == "docx.tracked_changes")


def _exported_marks(document_id: str) -> dict[str, int]:
    response = client.get(f"/api/v1/documents/{document_id}/export/docx")
    assert response.status_code == 200, response.text
    return _marks(response.content)


def test_the_choice_is_the_users_and_the_report_says_which(uploaded):
    assert uploaded["trackedChanges"] == "kept"
    assert (_tracked_item(uploaded)["policy"], _tracked_item(uploaded)["reason"], _tracked_item(uploaded)["contentChanged"]) == (
        "detected_not_editable",
        TRACKED_KEPT,
        False,
    )
    assert _exported_marks(uploaded["id"]) == _A07

    accepted = client.put(f"/api/v1/documents/{uploaded['id']}/tracked-changes", json={"choice": "accepted"})
    assert accepted.status_code == 200, accepted.text
    assert accepted.json()["trackedChanges"] == "accepted"
    assert (_tracked_item(accepted.json())["policy"], _tracked_item(accepted.json())["reason"]) == ("lossy", TRACKED_ACCEPTED)
    assert accepted.json()["elements"] == uploaded["elements"]  # the content as it was: read as accepted either way
    assert _exported_marks(uploaded["id"]) == dict.fromkeys(_A07, 0)

    kept = client.put(f"/api/v1/documents/{uploaded['id']}/tracked-changes", json={"choice": "kept"})
    assert kept.status_code == 200 and kept.json()["trackedChanges"] == "kept"
    assert _exported_marks(uploaded["id"]) == _A07  # the original file still has them


def test_a_document_without_tracked_changes_has_no_choice_to_make(api_db):
    client.cookies.clear()
    assert client.post("/api/v1/auth/register", json={"email": "untracked@example.com", "password": "long enough password"}).status_code == 201
    word = DocxDocument()
    word.add_paragraph("Nothing tracked.")
    document = client.post("/api/v1/documents/upload", files={"file": ("plain.docx", _saved(word), _DOCX)}).json()

    refused = client.put(f"/api/v1/documents/{document['id']}/tracked-changes", json={"choice": "accepted"})
    not_rejected = client.put(f"/api/v1/documents/{document['id']}/tracked-changes", json={"choice": "rejected"})
    invalid = client.put(f"/api/v1/documents/{document['id']}/tracked-changes", json={"choice": "ignored"})
    client.cookies.clear()

    assert refused.status_code == 409 and not_rejected.status_code == 409
    assert invalid.status_code == 422
