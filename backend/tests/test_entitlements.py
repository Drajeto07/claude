"""Plan limits (корекции.docx §35/§36): every limit is enforced on the backend,
from the workspace plan's entitlements, whichever endpoint the work comes in by
-- and refused with 402 "plan_limit" before any of it is done."""

import base64
import io
import os
from datetime import datetime, timezone

import pytest
from docx import Document as DocxDocument
from fastapi.testclient import TestClient
from PIL import Image as PILImage
from sqlalchemy import func, select
from sqlalchemy.orm import Session as OrmSession

from app.ai.factory import get_ai_provider
from app.ai.schemas import AIBlock, AIBlockType, AIFormattingRule, AIInstructionExtractionResponse, AIStructureResponse
from app.billing.plans import FREE, PLANS
from app.config import get_settings
from app.db.models import DocumentAsset, ProcessingJob, Subscription, UsageRecord, WorkspaceMember
from app.jobs import queue as queue_module
from app.jobs.queue import PRIORITY_HEAD_START, get_job_queue
from app.main import app
from app.services.auth_service import AuthService
from app.services.entitlements_service import EntitlementsService, effective
from tests.fakes import FakeAIProvider

client = TestClient(app, base_url="https://testserver")
_MARKDOWN = "# Plans\n\nA paragraph that makes this a real document."
_PROSE = "Just some plain prose without any markdown in it at all, written the way people write an email."
_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@pytest.fixture(autouse=True)
def signed_in(api_db):
    client.cookies.clear()
    assert client.post("/api/auth/register", json={"email": "plan@example.com", "password": "long enough password"}).status_code == 201
    yield
    client.cookies.clear()
    for dependency in (get_ai_provider, get_job_queue):
        app.dependency_overrides.pop(dependency, None)


def _free_plan(monkeypatch, **limits) -> None:
    free = PLANS[FREE]
    monkeypatch.setitem(PLANS, FREE, free.model_copy(update={"entitlements": free.entitlements.model_copy(update=limits)}))


def _workspace_id(api_db) -> str:
    with OrmSession(api_db) as session:
        return session.scalars(select(WorkspaceMember.workspace_id)).one()


def _refused(response, entitlement: str) -> dict:
    assert response.status_code == 402, response.text
    body = response.json()
    assert (body["code"], body["details"]["entitlement"]) == ("plan_limit", entitlement)
    return body


def _document() -> str:
    response = client.post("/api/documents", json={"text": _MARKDOWN})
    assert response.status_code == 201, response.text
    return response.json()["id"]


def _docx(picture: bytes | None = None) -> bytes:
    doc = DocxDocument()
    doc.add_paragraph("A reference paragraph.")
    if picture is not None:
        doc.add_picture(io.BytesIO(picture))
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _noise_png(width: int, height: int) -> bytes:
    buffer = io.BytesIO()
    PILImage.frombytes("RGB", (width, height), os.urandom(width * height * 3)).save(buffer, format="PNG")
    return buffer.getvalue()


def _jobs(api_db) -> int:
    with OrmSession(api_db) as session:
        return session.scalar(select(func.count(ProcessingJob.id)))


def test_the_document_limit_stops_new_documents_however_they_are_made(monkeypatch, api_db):
    _free_plan(monkeypatch, maxDocuments=1)
    document_id = _document()

    body = _refused(client.post("/api/documents", json={"text": _MARKDOWN}), "maxDocuments")
    assert (body["details"]["limit"], body["details"]["used"]) == (1, 1)
    assert "1 document." in body["message"]
    _refused(client.post("/api/documents/upload", files={"file": ("a.txt", b"Some text.", "text/plain")}), "maxDocuments")
    _refused(client.post("/api/jobs/import-text", json={"text": _MARKDOWN}), "maxDocuments")
    _refused(client.post("/api/jobs/import-file", files={"file": ("a.docx", _docx(), _DOCX)}), "maxDocuments")
    assert _jobs(api_db) == 0  # refused before anything was queued

    assert client.delete(f"/api/documents/{document_id}").status_code == 204
    assert client.post("/api/jobs/import-text", json={"text": _MARKDOWN}).json()["status"] == "succeeded"


def test_an_export_the_plan_doesnt_include_is_refused(monkeypatch):
    _free_plan(monkeypatch, canExportPdf=False)
    document_id = _document()

    body = _refused(client.get(f"/api/documents/{document_id}/export/pdf"), "canExportPdf")
    assert "PDF" in body["message"]
    _refused(client.post("/api/jobs/export", json={"documentId": document_id, "format": "pdf"}), "canExportPdf")
    assert client.get(f"/api/documents/{document_id}/export/docx").status_code == 200
    assert client.post("/api/jobs/export", json={"documentId": document_id, "format": "docx"}).json()["status"] == "succeeded"


def test_a_file_over_the_plans_size_is_refused_before_it_is_read(monkeypatch):
    _free_plan(monkeypatch, maxDocumentSizeMb=1)
    big = b"word " * 300_000  # 1.5 MB, not a valid .docx either: the size is checked first

    body = _refused(client.post("/api/documents/upload", files={"file": ("big.txt", big, "text/plain")}), "maxDocumentSizeMb")
    assert (body["details"]["limit"], body["details"]["used"]) == (1, 2)
    _refused(client.post("/api/jobs/import-file", files={"file": ("big.docx", big, _DOCX)}), "maxDocumentSizeMb")
    _refused(client.post("/api/jobs/extract-reference", files={"file": ("big.docx", big, _DOCX)}), "maxDocumentSizeMb")
    _refused(client.post("/api/templates/extract", files={"file": ("big.docx", big, _DOCX)}), "maxDocumentSizeMb")
    assert client.post("/api/documents/upload", files={"file": ("small.txt", b"Small enough.", "text/plain")}).status_code == 201


def test_no_plan_takes_a_file_over_the_servers_own_upload_cap(monkeypatch):
    monkeypatch.setattr(get_settings(), "max_upload_size_mb", 10)
    pro = PLANS["pro"].entitlements.model_copy(update={"maxDocumentSizeMb": 50})

    assert effective(pro).maxDocumentSizeMb == 10
    assert effective(pro.model_copy(update={"maxDocumentSizeMb": 5})).maxDocumentSizeMb == 5


def test_used_up_ai_operations_refuse_work_that_is_all_ai(monkeypatch, api_db):
    _free_plan(monkeypatch, maxAiOperations=1)
    document_id = _document()
    answer = AIInstructionExtractionResponse(rules=[AIFormattingRule(target="Paragraph", property="bold", value="true")])
    app.dependency_overrides[get_ai_provider] = lambda: FakeAIProvider([answer])
    first = client.post("/api/jobs/format", data={"documentId": document_id, "instructionsText": "make the text bold"}).json()
    assert first["result"]["status"] == "applied"

    body = _refused(client.post("/api/jobs/format", data={"documentId": document_id, "instructionsText": "make it blue"}), "maxAiOperations")
    assert (body["details"]["limit"], body["details"]["used"]) == (1, 1)
    _refused(client.post(f"/api/documents/{document_id}/format", data={"instructionsText": "make it blue"}), "maxAiOperations")
    _refused(client.post(f"/api/documents/{document_id}/style-analysis"), "maxAiOperations")
    # A template needs no AI, so it still applies.
    assert client.post("/api/jobs/format", data={"documentId": document_id, "templateId": "academic-default"}).json()["status"] == "succeeded"


def test_used_up_ai_operations_make_the_optional_ai_steps_fall_back(monkeypatch):
    _free_plan(monkeypatch, maxAiOperations=0)
    fake = FakeAIProvider([])  # any call would fail the test: it has nothing to answer with
    app.dependency_overrides[get_ai_provider] = lambda: fake

    assert client.post("/api/jobs/import-text", json={"text": _PROSE}).json()["status"] == "succeeded"
    assert client.post("/api/documents", json={"text": _PROSE}).status_code == 201
    assert fake.calls == 0


@pytest.mark.parametrize("path", ["/api/jobs/import-text", "/api/documents"])
def test_the_allowance_counts_the_calls_made_for_the_same_request(monkeypatch, api_db, path):
    # One operation left: the first answer fails the fidelity check (a completed,
    # counted call), so the retry would be the second -- and is refused.
    _free_plan(monkeypatch, maxAiOperations=1)
    diverged = AIStructureResponse(
        document_type="note",
        document_type_confidence=0.9,
        blocks=[AIBlock(type=AIBlockType.PARAGRAPH, text="Something else entirely", confidence=0.9)],
    )
    fake = FakeAIProvider([diverged, diverged])
    app.dependency_overrides[get_ai_provider] = lambda: fake

    response = client.post(path, json={"text": _PROSE})

    assert response.status_code in (201, 202) and fake.calls == 1
    with OrmSession(api_db) as session:
        assert session.scalar(select(func.sum(UsageRecord.quantity)).where(UsageRecord.metric == "ai_operations")) == 1


def test_the_template_limit_stops_new_templates(monkeypatch):
    _free_plan(monkeypatch, maxTemplates=1)
    assert client.post("/api/templates", json={"name": "House style"}).status_code == 201

    body = _refused(client.post("/api/templates", json={"name": "Another"}), "maxTemplates")
    assert "1 template." in body["message"]


def test_storage_over_the_plan_is_refused_before_anything_is_stored(monkeypatch, api_db, tmp_path):
    _free_plan(monkeypatch, maxStorageMb=1)
    photo = _noise_png(700, 600)  # 1.2 MB that doesn't compress: over the plan's 1 MB on its own

    body = _refused(client.post("/api/documents/upload", files={"file": ("photo.docx", _docx(picture=photo), _DOCX)}), "maxStorageMb")
    assert body["details"]["limit"] == 1
    document = client.get(f"/api/documents/{_document()}").json()
    pasted = {"type": "image", "content": "", "order": 99, "image": {"src": "data:image/png;base64," + base64.b64encode(photo).decode()}}
    _refused(client.put(f"/api/documents/{document['id']}/content", json={"elements": document["elements"] + [pasted]}), "maxStorageMb")
    with OrmSession(api_db) as session:
        assert session.scalar(select(func.count(DocumentAsset.id))) == 0
    assert not any(path.is_file() for path in (tmp_path / "assets").rglob("*"))

    icon = {**pasted, "image": {"src": "data:image/png;base64," + base64.b64encode(_noise_png(8, 8)).decode()}}
    assert client.put(f"/api/documents/{document['id']}/content", json={"elements": document["elements"] + [icon]}).status_code == 200
    # Room for the photo once: it counts as stored, not also as the inline copy the import briefly holds.
    _free_plan(monkeypatch, maxStorageMb=2)
    assert client.post("/api/documents/upload", files={"file": ("photo.docx", _docx(picture=photo), _DOCX)}).status_code == 201


async def test_a_subscription_grants_its_plan_only_while_it_is_in_good_standing(db_session):
    user = await AuthService(db_session).register("paid@example.com", "long enough password", None)
    workspace_id = await AuthService(db_session).default_workspace_id(user.id)
    plans = EntitlementsService(db_session)
    assert (await plans.plan(workspace_id)).key == FREE

    subscription = Subscription(workspace_id=workspace_id, plan="business", status="active", current_period_end=datetime.now(timezone.utc))
    db_session.add(subscription)
    await db_session.commit()
    assert (await plans.plan(workspace_id)).key == "business"
    assert await plans.ai_remaining(workspace_id) == PLANS["business"].entitlements.maxAiOperations

    for status, plan in (("past_due", "business"), ("trialing", "business"), ("canceled", FREE), ("unpaid", FREE), ("incomplete", FREE)):
        subscription.status = status
        await db_session.commit()
        assert (await plans.plan(workspace_id)).key == plan, status


def test_jobs_of_a_plan_with_priority_processing_are_queued_ahead(monkeypatch, api_db):
    business = PLANS["business"]
    monkeypatch.setitem(
        PLANS, "business", business.model_copy(update={"entitlements": business.entitlements.model_copy(update={"priorityProcessing": True})})
    )
    queued: list[tuple[str, bool]] = []

    class Recording:
        async def enqueue(self, job_id, *, priority=False):
            queued.append((job_id, priority))

    app.dependency_overrides[get_job_queue] = lambda: Recording()
    client.post("/api/jobs/import-text", json={"text": _MARKDOWN})
    with OrmSession(api_db) as session:
        session.add(Subscription(workspace_id=_workspace_id(api_db), plan="business", status="active"))
        session.commit()
    client.post("/api/jobs/import-text", json={"text": _MARKDOWN})

    assert [priority for _, priority in queued] == [False, True]


async def test_the_arq_queue_puts_a_priority_job_ahead_of_those_queued_since(monkeypatch):
    queued = []

    class FakePool:
        async def enqueue_job(self, function, *args, **kwargs):
            queued.append(kwargs)

    monkeypatch.setattr(queue_module, "_arq_pool", FakePool())
    before = datetime.now(timezone.utc)

    await queue_module.ArqQueue("redis://unused").enqueue("job-1", priority=True)

    # Queued as if it had waited PRIORITY_HEAD_START already: arq takes the oldest score first.
    assert queued[0]["_job_id"] == "job-1"
    assert before - PRIORITY_HEAD_START <= queued[0]["_defer_until"] <= datetime.now(timezone.utc) - PRIORITY_HEAD_START
