"""A link's title -- its tooltip (ScreenTip) in Word, the title attribute in the
editor -- is kept from import to export (tracker EDIT-010)."""

import io

from docx import Document as DocxDocument
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

from app.export.docx_export import build_docx
from app.models.document import MarkType
from app.parsers.docx import parse_docx
from app.parsers.markdown import parse_markdown


def _save(document) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _links(document) -> list[tuple[str, str | None, str | None]]:
    return [
        (run.text, mark.href, mark.title)
        for element in document.elements
        for run in element.inline or []
        for mark in run.marks
        if mark.type == MarkType.LINK
    ]


def _word_with_links() -> bytes:
    document = DocxDocument()
    paragraph = document.add_paragraph("See ")
    rel_id = paragraph.part.relate_to("https://example.com/guide", RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    paragraph._p.append(
        parse_xml(f'<w:hyperlink {nsdecls("w", "r")} r:id="{rel_id}" w:tooltip="The  full\nguide"><w:r><w:t>the guide</w:t></w:r></w:hyperlink>')
    )
    field = document.add_paragraph("Or ")
    field._p.append(
        parse_xml(
            f'<w:fldSimple {nsdecls("w")} w:instr=\' HYPERLINK "https://example.com/faq" \\o "Questions and answers" \'>'
            "<w:r><w:t>the FAQ</w:t></w:r></w:fldSimple>"
        )
    )
    plain = document.add_paragraph("And ")
    plain_rel = plain.part.relate_to("https://example.com/", RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    plain._p.append(parse_xml(f'<w:hyperlink {nsdecls("w", "r")} r:id="{plain_rel}"><w:r><w:t>home</w:t></w:r></w:hyperlink>'))
    return _save(document)


def test_word_tooltips_are_imported_and_written_back():
    imported = parse_docx(_word_with_links(), "links.docx")

    assert _links(imported) == [
        ("the guide", "https://example.com/guide", "The full guide"),
        ("the FAQ", "https://example.com/faq", "Questions and answers"),
        ("home", "https://example.com/", None),
    ]

    exported = DocxDocument(io.BytesIO(build_docx(imported)))
    tooltips = [link.get(qn("w:tooltip")) for link in exported.element.body.iter(qn("w:hyperlink"))]
    assert tooltips == ["The full guide", "Questions and answers", None]
    assert _links(parse_docx(build_docx(imported), "again.docx")) == _links(imported)


def test_markdown_link_titles_are_kept():
    document = parse_markdown('Read [the guide](https://example.com/guide "The full guide") first.', "Notes")

    assert _links(document) == [("the guide", "https://example.com/guide", "The full guide")]
