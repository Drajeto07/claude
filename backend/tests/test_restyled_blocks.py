"""Blocks only restyled here keep their original XML where nothing conflicts (tracker DOCX-029).
A template, an instruction or a person restyling a kind of block used to write every such block
anew in a Word export into the original: what the model can't hold (stretched text, a text effect,
a bookmark) was lost from it. Now a paragraph or heading whose content is as imported, whose
Word style is the one the export writes for its kind (with the new look) and whose own formatting
sets none of what the new look changed is copied; the others are written anew, as before."""

import io
import zipfile

import pytest
from docx import Document as DocxDocument
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from fastapi.testclient import TestClient

from app.export.docx_export import build_docx
from app.export.package_check import package_problems
from app.export.provenance import look_changes, stamp
from app.formatting.engine import recompute_styles
from app.main import app
from app.models.document import FormattingProperty, FormattingRule
from app.services.ingestion_service import build_document_from_docx

_W = nsdecls("w")
_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
client = TestClient(app, base_url="https://testserver")


def _word_file() -> bytes:
    """Paragraphs holding what the model can't: stretched text (w:w) and an outline effect, a
    bookmark; one that sets its own font size; one in a style of its own; a heading."""
    word = DocxDocument()
    word.add_heading("Results", level=1)
    styled = word.add_paragraph("Marked ")
    styled._p.append(parse_xml(f'<w:r {_W}><w:rPr><w:w w:val="150"/></w:rPr><w:t>twice</w:t></w:r>'))
    styled._p.append(parse_xml(f'<w:r {_W}><w:rPr><w:outline/></w:rPr><w:t xml:space="preserve"> outlined</w:t></w:r>'))
    marked = word.add_paragraph()
    marked._p.append(parse_xml(f'<w:bookmarkStart {_W} w:id="7" w:name="Results"/>'))
    marked._p.append(parse_xml(f'<w:r {_W}><w:t>The results section.</w:t></w:r>'))
    marked._p.append(parse_xml(f'<w:bookmarkEnd {_W} w:id="7"/>'))
    sized = word.add_paragraph()
    sized._p.append(parse_xml(f'<w:r {_W}><w:rPr><w:sz w:val="30"/><w:w w:val="150"/></w:rPr><w:t>Large and stretched.</w:t></w:r>'))
    body_text = word.add_paragraph("In Body Text ", style="Body Text")  # a paragraph, in a style the export doesn't write
    body_text._p.append(parse_xml(f'<w:r {_W}><w:rPr><w:w w:val="80"/></w:rPr><w:t>narrow</w:t></w:r>'))
    out = io.BytesIO()
    word.save(out)
    return out.getvalue()


def _imported():
    source = _word_file()
    document = build_document_from_docx(source, "report.docx", None)
    recompute_styles(document)
    stamp(document)
    return source, document


def _restyle(document, target: str, prop: FormattingProperty, value: str, unit: str | None = None) -> None:
    document.formattingRules.append(FormattingRule(target=target, property=prop, value=value, unit=unit, priority=1, source="template"))
    recompute_styles(document)


def _body(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as package:
        return package.read("word/document.xml").decode("utf-8")


def _paragraph_of(body: str, text: str) -> str:
    """The XML of the paragraph holding `text`."""
    at = body.index(text)
    return body[body.rindex("<w:p>", 0, at) if "<w:p>" in body[:at] else body.rindex("<w:p ", 0, at) : body.index("</w:p>", at)]


def test_a_paragraph_only_restyled_is_copied_with_what_the_model_cant_hold():
    source, document = _imported()
    _restyle(document, "Paragraph", FormattingProperty.FONT_FAMILY, "Times New Roman")
    exported = build_docx(document, source=source)
    body = _body(exported)
    assert 'w:val="150"' in _paragraph_of(body, "twice") and "<w:outline/>" in body  # copied, not written anew
    assert 'w:name="Results"' in body  # the bookmark's paragraph too
    assert package_problems(exported) == []
    # The new look comes from the Word style the export writes for paragraphs.
    normal = DocxDocument(io.BytesIO(exported)).styles["Normal"]
    assert normal.font.name == "Times New Roman"


def test_one_whose_own_formatting_sets_what_changed_is_written_anew():
    source, document = _imported()
    _restyle(document, "Paragraph", FormattingProperty.FONT_SIZE, "12", "pt")
    body = _body(build_docx(document, source=source))
    assert 'w:val="150"' not in _paragraph_of(body, "Large and stretched.")  # its own size: written anew
    assert 'w:val="150"' in _paragraph_of(body, "twice")  # sets no size of its own: copied


def test_one_in_another_word_style_or_changed_here_is_written_anew():
    source, document = _imported()
    marked = next(element for element in document.elements if element.content.startswith("Marked"))
    marked.inline = [run.model_copy(update={"text": run.text.replace("twice", "once")}) for run in marked.inline]
    marked.content = marked.content.replace("twice", "once")
    _restyle(document, "Paragraph", FormattingProperty.FONT_FAMILY, "Times New Roman")
    assert look_changes(document, marked) is None  # what it holds changed
    body = _body(build_docx(document, source=source))
    assert 'w:val="150"' not in _paragraph_of(body, "Marked once")
    assert 'w:val="80"' not in _paragraph_of(body, "narrow")  # Body Text isn't the Word style written for paragraphs (Normal)


def test_a_heading_restyled_keeps_its_xml_and_takes_the_new_look_from_its_style():
    source, document = _imported()
    heading = next(element for element in document.elements if element.content == "Results")
    _restyle(document, "Heading 1", FormattingProperty.COLOR, "#C00000")
    assert look_changes(document, heading) == {"color"}
    exported = build_docx(document, source=source)
    assert DocxDocument(io.BytesIO(exported)).styles["Heading 1"].font.color.rgb is not None
    assert package_problems(exported) == []


def test_a_document_stamped_before_restyles_as_it_did():
    source, document = _imported()
    document.sourceStyles = None  # stamped before DOCX-029
    _restyle(document, "Paragraph", FormattingProperty.FONT_FAMILY, "Times New Roman")
    assert 'w:val="150"' not in _paragraph_of(_body(build_docx(document, source=source)), "twice")


@pytest.fixture
def uploaded(api_db):
    client.cookies.clear()
    assert client.post("/api/v1/auth/register", json={"email": "restyled@example.com", "password": "long enough password"}).status_code == 201
    response = client.post("/api/v1/documents/upload", files={"file": ("report.docx", _word_file(), _DOCX)})
    assert response.status_code == 201, response.text
    yield response.json()
    client.cookies.clear()


def test_after_a_template_the_word_export_keeps_what_the_model_cant_hold(uploaded):
    assert client.post(f"/api/v1/documents/{uploaded['id']}/format", data={"templateId": "academic-default"}).status_code == 200
    exported = client.get(f"/api/v1/documents/{uploaded['id']}/export/docx")
    assert exported.status_code == 200
    body = _body(exported.content)
    assert 'w:val="150"' in _paragraph_of(body, "twice") and "<w:outline/>" in body and 'w:name="Results"' in body
    assert 'w:val="150"' not in _paragraph_of(body, "Large and stretched.")  # sets its own size, which the template changed
    normal = DocxDocument(io.BytesIO(exported.content)).styles["Normal"]
    assert normal.font.name == "Times New Roman" and normal.font.size.pt == 12


@pytest.mark.parametrize("text", ["twice", "Large and stretched."])
def test_one_given_its_own_look_here_is_written_anew_with_it(text):
    """With no look of its own at import ("twice"), or one (its own size) it keeps."""
    source, document = _imported()
    element = next(element for element in document.elements if text in element.content)
    document.formattingRules.append(FormattingRule(target=element.id, property=FormattingProperty.ALIGNMENT, value="center", priority=0, source="user"))
    _restyle(document, "Paragraph", FormattingProperty.FONT_FAMILY, "Times New Roman")
    assert look_changes(document, element) is None  # set on it apart from its kind, here
    paragraph = _paragraph_of(_body(build_docx(document, source=source)), text)
    assert 'w:val="150"' not in paragraph and '<w:jc w:val="center"/>' in paragraph
