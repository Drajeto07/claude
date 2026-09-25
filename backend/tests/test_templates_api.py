import io

import pytest
from docx import Document as DocxDocument
from docx.shared import Pt
from fastapi.testclient import TestClient
from sqlalchemy import func, select
from sqlalchemy.orm import Session as OrmSession

from app.ai.base import AIStructuredOutputError
from app.ai.factory import get_ai_provider
from app.db.models import TemplateVersion
from app.formatting.priorities import Priority
from app.formatting.templates import BUILTIN_TEMPLATES
from app.main import app
from tests.fakes import FakeAIProvider

# https: the session cookie is Secure, and a cookie jar never returns it over plain http.
client = TestClient(app, base_url="https://testserver")

_REPORT = {
    "name": "Team report",
    "category": "business",
    "description": "Our house style",
    "styleSystem": {
        "page": {"size": "Letter", "marginLeftCm": 2.5},
        "paragraph": {"fontFamily": "Georgia", "fontSizePt": 11},
        "headings": {"h1": {"fontSizePt": 18, "bold": True}},
    },
}


@pytest.fixture(autouse=True)
def signed_in(api_db):
    client.cookies.clear()
    response = client.post("/api/v1/auth/register", json={"email": "owner@example.com", "password": "long enough password"})
    assert response.status_code == 201
    yield
    client.cookies.clear()


def _create(payload: dict | None = None) -> dict:
    response = client.post("/api/v1/templates", json=payload or _REPORT)
    assert response.status_code == 201, response.text
    return response.json()


def _update(template: dict, body: dict, version: int | None = None):
    headers = {"If-Match": str(version if version is not None else template["version"])}
    return client.put(f"/api/v1/templates/{template['id']}", json=body, headers=headers)


def _document() -> str:
    response = client.post("/api/v1/documents", json={"text": "# Title\n\nA body paragraph that is long enough."})
    assert response.status_code == 201
    return response.json()["id"]


def _format(document_id: str, template_id: str):
    return client.post(f"/api/v1/documents/{document_id}/format", data={"templateId": template_id})


_DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _reference_docx() -> bytes:
    doc = DocxDocument()
    doc.styles["Normal"].font.name = "Georgia"
    doc.styles["Normal"].font.size = Pt(13)
    doc.add_heading("Report", level=1)
    doc.add_paragraph("A body paragraph long enough to be what most of this reference's text looks like.")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def no_ai():
    """Reading a reference may ask the AI which paragraphs are headings; tests never reach a real one."""
    app.dependency_overrides[get_ai_provider] = lambda: FakeAIProvider([AIStructuredOutputError("no AI in tests")])
    yield
    app.dependency_overrides.pop(get_ai_provider, None)


def test_every_template_endpoint_needs_a_signed_in_user():
    client.cookies.clear()
    calls = [
        client.get("/api/v1/templates"),
        client.post("/api/v1/templates", json=_REPORT),
        client.get("/api/v1/templates/academic-default"),
        client.put("/api/v1/templates/x", json={"name": "y"}),
        client.delete("/api/v1/templates/x"),
        client.post("/api/v1/templates/academic-default/duplicate"),
        client.put("/api/v1/templates/default", json={"templateId": None}),
        client.post("/api/v1/templates/preview", json={}),
        client.get("/api/v1/templates/x/versions"),
        client.post("/api/v1/templates/x/versions/1/restore"),
        client.post("/api/v1/templates/extract", files={"file": ("ref.docx", _reference_docx(), _DOCX_TYPE)}),
    ]
    assert [response.status_code for response in calls] == [401] * len(calls)


def test_list_shows_the_builtins_first_read_only_with_engine_computed_previews():
    templates = client.get("/api/v1/templates").json()

    assert [t["id"] for t in templates] == list(BUILTIN_TEMPLATES)
    academic = templates[0]
    assert academic["builtin"] is True and academic["editable"] is False
    assert academic["styleSystem"]["paragraph"]["fontFamily"] == "Times New Roman"
    assert academic["previewStyles"]["Paragraph"]["font-family"] == "Times New Roman"
    assert academic["previewStyles"]["Heading 1"]["font-weight"] == "bold"


def test_create_a_template_and_find_it_after_the_builtins():
    created = _create()

    assert created["name"] == "Team report" and created["category"] == "business"
    assert created["builtin"] is False and created["editable"] is True
    assert created["version"] == 1 and created["visibility"] == "workspace"
    assert created["styleSystem"]["paragraph"]["fontFamily"] == "Georgia"
    assert created["notes"] == []
    assert [t["id"] for t in client.get("/api/v1/templates").json()][-1] == created["id"]
    assert client.get(f"/api/v1/templates/{created['id']}").json()["name"] == "Team report"


def test_an_empty_template_can_be_created_and_filled_in_later():
    created = _create({"name": "Blank"})
    assert created["category"] == "general"
    assert created["styleSystem"]["paragraph"]["fontFamily"] is None


@pytest.mark.parametrize(
    "payload",
    [
        {"name": "   "},
        {"name": "X", "category": " "},
        {"name": "X", "styleSystem": {"paragraph": {"fontSizePt": -3}}},
        {"name": "X", "styleSystem": {"paragraph": {"font": "Arial"}}},
        {"name": "X", "styleSystem": {}, "sourceDocumentId": "abc"},
        {"name": "X", "visibility": "everyone"},
    ],
)
def test_invalid_templates_are_rejected(payload):
    assert client.post("/api/v1/templates", json=payload).status_code == 422


def test_unknown_template_is_404():
    assert client.get("/api/v1/templates/does-not-exist").status_code == 404
    assert _update({"id": "does-not-exist", "version": 1}, {"name": "x"}).status_code == 404
    assert client.delete("/api/v1/templates/does-not-exist").status_code == 404


def test_formatting_with_a_workspace_template_applies_its_rules_at_the_custom_tier():
    template = _create()
    document_id = _document()

    response = _format(document_id, template["id"])

    assert response.status_code == 200
    document = response.json()["document"]
    assert document["templateId"] == template["id"]
    assert document["resolvedStyles"]["Paragraph"]["font-family"] == "Georgia"
    assert document["settings"]["pageSize"] == "Letter" and document["settings"]["marginLeftCm"] == 2.5
    georgia = next(r for r in document["formattingRules"] if r["value"] == "Georgia")
    assert (georgia["priority"], georgia["source"]) == (Priority.CUSTOM_TEMPLATE, "custom_template")


def test_editing_a_template_bumps_its_version_and_changes_what_it_applies():
    template = _create()

    response = _update(template, {"styleSystem": {"paragraph": {"fontFamily": "Garamond"}}})

    assert response.status_code == 200
    edited = response.json()
    assert edited["version"] == 2
    assert edited["styleSystem"]["paragraph"]["fontFamily"] == "Garamond"
    assert edited["styleSystem"]["headings"]["h1"]["fontSizePt"] is None  # the whole style system is replaced
    document = _format(_document(), template["id"]).json()["document"]
    assert document["resolvedStyles"]["Paragraph"]["font-family"] == "Garamond"


def test_a_stale_edit_is_refused_with_412_and_changes_nothing():
    template = _create()
    assert _update(template, {"name": "Renamed in another tab"}).status_code == 200

    response = _update(template, {"name": "My stale rename"}, version=1)

    assert response.status_code == 412
    assert response.json()["details"]["currentVersion"] == 2
    assert client.get(f"/api/v1/templates/{template['id']}").json()["name"] == "Renamed in another tab"


def test_saving_without_changes_adds_no_version(api_db):
    template = _create()

    response = _update(template, {"name": "Team report", "styleSystem": _REPORT["styleSystem"]})

    assert response.json()["version"] == 1
    with OrmSession(api_db) as session:
        assert session.scalar(select(func.count()).select_from(TemplateVersion)) == 1


def test_rename_keeps_the_style_and_is_recorded_in_the_history():
    template = _create()

    renamed = _update(template, {"name": "Quarterly report"}).json()

    assert renamed["name"] == "Quarterly report"
    assert renamed["styleSystem"] == template["styleSystem"]
    versions = client.get(f"/api/v1/templates/{template['id']}/versions").json()
    assert [(v["number"], v["name"], v["current"]) for v in versions] == [
        (2, "Quarterly report", True),
        (1, "Team report", False),
    ]
    assert versions[0]["author"] == "owner@example.com"


def test_restoring_an_old_version_brings_its_style_back_as_a_new_version():
    template = _create()
    edited = _update(template, {"styleSystem": {"paragraph": {"fontFamily": "Garamond"}}}).json()

    response = client.post(f"/api/v1/templates/{template['id']}/versions/1/restore", headers={"If-Match": str(edited["version"])})

    assert response.status_code == 200
    restored = response.json()
    assert restored["version"] == 3
    assert restored["styleSystem"] == template["styleSystem"]
    assert client.post(f"/api/v1/templates/{template['id']}/versions/99/restore").status_code == 404


def test_template_history_keeps_only_the_configured_number_of_versions(monkeypatch):
    from app.config import get_settings

    monkeypatch.setattr(get_settings(), "template_history_max_versions", 2)
    template = _create()
    for index in range(3):
        template = _update(template, {"name": f"Name {index}"}).json()

    assert [v["number"] for v in client.get(f"/api/v1/templates/{template['id']}/versions").json()] == [4, 3]


def test_builtins_cannot_be_edited_or_deleted():
    assert _update({"id": "academic-default", "version": 1}, {"name": "Mine now"}).status_code == 403
    assert client.delete("/api/v1/templates/academic-default").status_code == 403
    assert client.get("/api/v1/templates/academic-default/versions").json() == []


def test_deleting_a_template_leaves_documents_formatted_with_it_intact():
    template = _create()
    document_id = _document()
    _format(document_id, template["id"])

    assert client.delete(f"/api/v1/templates/{template['id']}").status_code == 204

    assert client.get(f"/api/v1/templates/{template['id']}").status_code == 404
    document = client.get(f"/api/v1/documents/{document_id}").json()
    assert document["resolvedStyles"]["Paragraph"]["font-family"] == "Georgia"
    # Formatting with the deleted template again is the usual unknown-template error.
    assert _format(document_id, template["id"]).status_code == 400


def test_duplicating_a_builtin_gives_an_editable_copy_that_leaves_the_original_alone():
    response = client.post("/api/v1/templates/academic-default/duplicate")

    assert response.status_code == 201
    copy = response.json()
    assert copy["name"] == f"{BUILTIN_TEMPLATES['academic-default'].name} (copy)"
    assert copy["editable"] is True and copy["builtin"] is False
    assert copy["styleSystem"] == client.get("/api/v1/templates/academic-default").json()["styleSystem"]

    _update(copy, {"styleSystem": {"paragraph": {"fontFamily": "Garamond"}}})
    assert client.get("/api/v1/templates/academic-default").json()["styleSystem"]["paragraph"]["fontFamily"] == "Times New Roman"
    assert BUILTIN_TEMPLATES["academic-default"].styleSystem.paragraph.fontFamily == "Times New Roman"


def test_duplicate_can_be_given_a_name():
    template = _create()
    copy = client.post(f"/api/v1/templates/{template['id']}/duplicate", json={"name": "Variant B"}).json()
    assert copy["name"] == "Variant B" and copy["id"] != template["id"]


def test_the_workspace_default_template_is_set_shown_and_cleared():
    template = _create()

    assert client.put("/api/v1/templates/default", json={"templateId": "modern-report"}).json() == {"templateId": "modern-report"}
    assert [t["id"] for t in client.get("/api/v1/templates").json() if t["isDefault"]] == ["modern-report"]

    client.put("/api/v1/templates/default", json={"templateId": template["id"]})
    assert client.get(f"/api/v1/templates/{template['id']}").json()["isDefault"] is True

    assert client.put("/api/v1/templates/default", json={"templateId": None}).json() == {"templateId": None}
    assert not any(t["isDefault"] for t in client.get("/api/v1/templates").json())


def test_deleting_the_default_template_clears_the_default():
    template = _create()
    client.put("/api/v1/templates/default", json={"templateId": template["id"]})

    client.delete(f"/api/v1/templates/{template['id']}")

    assert not any(t["isDefault"] for t in client.get("/api/v1/templates").json())


def test_an_unknown_template_cannot_become_the_default():
    assert client.put("/api/v1/templates/default", json={"templateId": "nope"}).status_code == 404


def test_preview_resolves_an_unsaved_style_system_with_the_real_engine():
    response = client.post(
        "/api/v1/templates/preview",
        json={"page": {"orientation": "landscape", "marginTopCm": 1}, "captions": {"italic": True}},
    )

    assert response.status_code == 200
    body = response.json()
    assert body["settings"]["orientation"] == "landscape" and body["settings"]["marginTopCm"] == 1
    assert body["resolvedStyles"]["Caption"]["font-style"] == "italic"
    assert body["resolvedStyles"]["Paragraph"]["font-family"] == "Arial"  # the engine's own default shows through
    assert client.post("/api/v1/templates/preview", json={"page": {"size": "B5"}}).status_code == 422


def test_a_document_can_be_saved_as_a_template_without_its_one_off_overrides():
    document_id = _document()
    formatted = _format(document_id, "academic-default").json()["document"]
    paragraph_id = formatted["elements"][1]["id"]
    client.patch(f"/api/v1/documents/{document_id}/settings", json={"property": "marginTop", "value": "4", "unit": "cm"})
    client.patch(f"/api/v1/documents/{document_id}/elements/{paragraph_id}/style", json={"property": "color", "value": "red"})

    response = client.post("/api/v1/templates", json={"name": "From my thesis", "sourceDocumentId": document_id})

    assert response.status_code == 201
    template = response.json()
    assert template["sourceDocumentId"] == document_id
    assert template["notes"] == []
    style = template["styleSystem"]
    assert style["paragraph"]["fontFamily"] == "Times New Roman"  # from the template it was formatted with
    assert style["paragraph"]["lineSpacing"] == 1.5
    assert style["page"]["marginTopCm"] == 4  # the document's own page setting wins, as it does in the document
    assert style["paragraph"]["color"] is None  # the red on one paragraph isn't the document's style


def test_saving_a_template_from_an_unknown_document_is_404():
    response = client.post("/api/v1/templates", json={"name": "X", "sourceDocumentId": "no-such-document"})
    assert response.status_code == 404


def test_another_accounts_templates_are_invisible_and_unusable():
    template = _create()
    document_id = _document()

    client.cookies.clear()
    client.post("/api/v1/auth/register", json={"email": "someone@example.com", "password": "another long password"})
    own_document = _document()

    assert template["id"] not in {t["id"] for t in client.get("/api/v1/templates").json()}
    assert client.get(f"/api/v1/templates/{template['id']}").status_code == 404
    assert _update(template, {"name": "Hijacked"}).status_code == 404
    assert client.delete(f"/api/v1/templates/{template['id']}").status_code == 404
    assert client.post(f"/api/v1/templates/{template['id']}/duplicate").status_code == 404
    assert client.put("/api/v1/templates/default", json={"templateId": template["id"]}).status_code == 404
    assert _format(own_document, template["id"]).status_code == 400
    assert client.post("/api/v1/templates", json={"name": "X", "sourceDocumentId": document_id}).status_code == 404


# -- Format by Example -------------------------------------------------------------


def test_a_reference_documents_style_is_read_then_saved_and_applied_like_any_template(no_ai):
    response = client.post("/api/v1/templates/extract", files={"file": ("Annual report.docx", _reference_docx(), _DOCX_TYPE)})

    assert response.status_code == 200, response.text
    extracted = response.json()
    assert extracted["suggestedName"] == "Annual report style"
    assert extracted["styleSystem"]["paragraph"]["fontFamily"] == "Georgia"
    assert extracted["previewStyles"]["Paragraph"]["font-family"] == "Georgia"
    assert (extracted["headingsFrom"], extracted["headingCounts"]) == ("styles", {"1": 1})
    assert client.get("/api/v1/templates").json()[-1]["id"] in BUILTIN_TEMPLATES  # reading saves nothing

    template = _create({"name": extracted["suggestedName"], "styleSystem": extracted["styleSystem"]})
    again = client.post("/api/v1/templates/extract", files={"file": ("Annual report.docx", _reference_docx(), _DOCX_TYPE)})
    assert again.json()["suggestedName"] == "Annual report style 2"  # never a second template of the same name
    formatted = _format(_document(), template["id"])

    assert formatted.status_code == 200
    styles = formatted.json()["document"]["resolvedStyles"]
    assert (styles["Paragraph"]["font-family"], styles["Paragraph"]["font-size"]) == ("Georgia", "13pt")


def test_the_reference_has_to_be_a_readable_word_file(no_ai):
    pdf = client.post("/api/v1/templates/extract", files={"file": ("notes.pdf", b"%PDF-1.4", "application/pdf")})
    broken = client.post("/api/v1/templates/extract", files={"file": ("broken.docx", b"not a zip file", _DOCX_TYPE)})

    assert pdf.status_code == 400 and "Word file" in pdf.json()["message"]
    assert broken.status_code == 400
