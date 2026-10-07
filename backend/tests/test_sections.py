"""Word sections as the model has them (tracker DOCX-015): a section break is an
element of its own -- how the next section starts, and the section it ends: its
page setup (size, orientation, margins, header and footer distances, columns, page
numbering) and its own headers and footers (main, first page, even pages; none of
a kind is the previous section's). Document.lastSection holds the last section's.
The importer makes one per section's end, a Word export writes each back as its
sectPr in the schema's order, and a PDF gives each section its pages, numbers,
headers and footers as Word does."""

import io
import zipfile

import pytest
from docx import Document as DocxDocument
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm
from fastapi.testclient import TestClient
from pydantic import ValidationError

from app.export.docx_export import build_docx
from app.export.package_check import package_problems
from app.export.pdf_export import build_pdf
from app.fidelity.report import ReportBuilder
from app.main import app
from app.models.document import Element, ElementType
from app.parsers.docx import parse_docx

client = TestClient(app, base_url="https://testserver")
_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_W = nsdecls("w")
# CT_SectPr's children, in the schema's order.
_SECT_PR_ORDER = (
    "headerReference footerReference footnotePr endnotePr type pgSz pgMar paperSrc pgBorders lnNumType pgNumType cols "
    "formProt vAlign noEndnote titlePg textDirection bidi rtlGutter docGrid printerSettings"
).split()


def _word_file() -> bytes:
    word = DocxDocument()
    first = word.sections[0]
    first.orientation, first.page_width, first.page_height = WD_ORIENT.LANDSCAPE, first.page_height, first.page_width
    first._sectPr.append(parse_xml(f'<w:pgNumType {_W} w:fmt="lowerRoman" w:start="1"/>'))
    columns = first._sectPr.find(qn("w:cols"))
    columns.set(qn("w:num"), "2")
    word.add_paragraph("A wide section in two columns.")
    last = word.add_section(WD_SECTION.ODD_PAGE)
    last.orientation, last.page_width, last.page_height = WD_ORIENT.PORTRAIT, last.page_height, last.page_width
    for extra in last._sectPr.findall(qn("w:pgNumType")):
        last._sectPr.remove(extra)
    last._sectPr.find(qn("w:cols")).set(qn("w:num"), "1")
    word.add_paragraph("A narrow one, on an odd page.")
    buffer = io.BytesIO()
    word.save(buffer)
    return buffer.getvalue()


def _breaks(document) -> list[Element]:
    return [element for element in document.elements if element.type == ElementType.SECTION_BREAK]


def test_a_section_break_holds_how_the_next_section_starts_and_the_setup_of_the_one_it_ends():
    document = parse_docx(_word_file(), "sections.docx")

    [section] = _breaks(document)
    settings = section.sectionBreak
    assert [element.type for element in document.elements] == [ElementType.PARAGRAPH, ElementType.SECTION_BREAK, ElementType.PARAGRAPH]
    assert settings.start == "oddPage"
    assert (settings.orientation, round(settings.pageWidthMm), round(settings.pageHeightMm)) == ("landscape", 279, 216)
    assert (settings.columns, settings.pageNumberStart, settings.pageNumberFormat) == (2, 1, "lowerRoman")
    assert settings.marginTopCm == 2.54 and settings.headerDistanceCm == 1.27


def test_a_word_export_writes_each_section_back_in_the_schemas_order():
    document = parse_docx(_word_file(), "sections.docx")

    exported = build_docx(document)

    assert package_problems(exported) == []
    with zipfile.ZipFile(io.BytesIO(exported)) as package:
        body = parse_xml(package.read("word/document.xml"))
    sections = list(body.iter(qn("w:sectPr")))
    assert len(sections) == 2
    for sect_pr in sections:
        tags = [child.tag.split("}")[1] for child in sect_pr]
        assert tags == sorted(tags, key=_SECT_PR_ORDER.index), tags
    last = sections[-1].find(qn("w:type"))
    assert last is not None and last.get(qn("w:val")) == "oddPage"  # how the last section starts
    again = parse_docx(exported, "again.docx")
    assert [element.sectionBreak for element in _breaks(again)] == [element.sectionBreak for element in _breaks(document)]


def test_without_page_breaks_sections_run_on():
    document = parse_docx(_word_file(), "sections.docx")

    exported = build_docx(document, include_page_breaks=False)

    with zipfile.ZipFile(io.BytesIO(exported)) as package:
        body = parse_xml(package.read("word/document.xml"))
    starts = [(sect_pr.find(qn("w:type")).get(qn("w:val")) if sect_pr.find(qn("w:type")) is not None else "nextPage") for sect_pr in body.iter(qn("w:sectPr"))]
    assert starts[-1] == "continuous"


def test_a_pdf_gives_each_section_its_own_pages_and_numbers():
    from pypdf import PdfReader

    document = parse_docx(_word_file(), "sections.docx")
    document.settings.showPageNumbers = True
    report = ReportBuilder()

    pages = PdfReader(io.BytesIO(build_pdf(document, report=report))).pages

    sizes = [(float(page.mediabox.width), float(page.mediabox.height)) for page in pages]
    assert len(pages) == 3  # the wide section, a blank page, the narrow one on an odd page
    assert sizes[0][0] > sizes[0][1] and sizes[2][0] < sizes[2][1]
    texts = [page.extract_text() for page in pages]
    assert "A wide section" in texts[0] and "Page i" in texts[0]  # lower-case roman numbering, from 1
    assert "A narrow one" in texts[2] and "Page 3" in texts[2]
    assert "export.pdf.sections" not in {item.feature for item in report.items()}


def test_only_a_section_break_has_section_settings():
    with pytest.raises(ValidationError):
        Element.model_validate({"type": "paragraph", "content": "x", "order": 0, "sectionBreak": {"start": "continuous"}})
    assert Element.model_validate({"type": "section_break", "content": "", "order": 0}).sectionBreak.start == "nextPage"
    with pytest.raises(ValidationError):
        Element.model_validate({"type": "section_break", "content": "", "order": 0, "sectionBreak": {"start": "sideways"}})


def test_a_section_written_anew_keeps_its_own_headers(api_db):
    word = DocxDocument()
    word.sections[0].header.paragraphs[0].text = "Chapter one"
    ending = word.add_paragraph("The end of chapter one.")
    word.add_section(WD_SECTION.NEW_PAGE)
    holder = next(p for p in word.element.body.iterchildren(qn("w:p")) if p.find(f"{qn('w:pPr')}/{qn('w:sectPr')}") is not None)
    ending._p.get_or_add_pPr().append(holder.find(f"{qn('w:pPr')}/{qn('w:sectPr')}"))  # the section ends in the text
    word.element.body.remove(holder)
    later = word.sections[-1]
    later.header.is_linked_to_previous = False
    later.header.paragraphs[0].text = "Chapter two"
    word.add_paragraph("Chapter two.")
    buffer = io.BytesIO()
    word.save(buffer)
    client.cookies.clear()
    assert client.post("/api/v1/auth/register", json={"email": "sections@example.com", "password": "long enough password"}).status_code == 201
    document = client.post("/api/v1/documents/upload", files={"file": ("chapters.docx", buffer.getvalue(), _DOCX)}).json()
    elements = document["elements"]
    edited = next(element for element in elements if element["content"] == "The end of chapter one.")
    edited["content"], edited["inline"] = "The end of chapter one, revised.", [{"text": "The end of chapter one, revised.", "marks": []}]
    assert client.put(f"/api/v1/documents/{document['id']}/content", json={"elements": elements}).status_code == 200

    job = client.post("/api/v1/jobs/export", json={"documentId": document["id"], "format": "docx"}).json()
    items = {item["feature"]: item for item in client.get(f"/api/v1/jobs/{job['id']}").json()["result"]["fidelity"]["items"]}
    exported = client.get(f"/api/v1/jobs/{job['id']}/file").content
    client.cookies.clear()

    assert "export.docx.section_lost" not in items  # the section is still there, from its section break
    assert "export.docx.rewritten_blocks" not in items  # and its own header, from the model (DOCX-015)
    assert [section.header.paragraphs[0].text for section in DocxDocument(io.BytesIO(exported)).sections] == ["Chapter one", "Chapter two"]
    assert package_problems(exported) == []


def test_a_picture_fits_its_sections_column_and_page_in_a_pdf():
    from PIL import Image as PILImage

    image = io.BytesIO()
    PILImage.new("RGB", (400, 1200), "navy").save(image, format="PNG")  # tall: the page's full width would be too high
    word = DocxDocument()
    first = word.sections[0]
    first.orientation, first.page_width, first.page_height = WD_ORIENT.LANDSCAPE, first.page_height, first.page_width
    first._sectPr.find(qn("w:cols")).set(qn("w:num"), "2")
    word.add_picture(io.BytesIO(image.getvalue()), width=first.page_width - first.left_margin - first.right_margin)
    word.add_section(WD_SECTION.NEW_PAGE)
    word.add_paragraph("After.")
    buffer = io.BytesIO()
    word.save(buffer)

    pdf = build_pdf(parse_docx(buffer.getvalue(), "tall.docx"))  # reportlab refuses a picture bigger than its frame

    assert pdf.startswith(b"%PDF")


# -- headers and footers per section (DOCX-015 part 2) ---------------------------------------


def _chapters_file(*, second_linked: bool = False) -> bytes:
    word = DocxDocument()
    first = word.sections[0]
    first.different_first_page_header_footer = True
    first.first_page_header.is_linked_to_previous = False  # an own, empty first-page header: a cover without one
    first.header.paragraphs[0].text = "Chapter one"
    word.add_paragraph("One.")
    last = word.add_section(WD_SECTION.NEW_PAGE)
    last.different_first_page_header_footer = False
    if not second_linked:
        last.header.is_linked_to_previous = False
        last.header.paragraphs[0].text = "Chapter two"
    last.footer.is_linked_to_previous = False
    last.footer.paragraphs[0].text = "Footer two"
    word.add_paragraph("Two.")
    buffer = io.BytesIO()
    word.save(buffer)
    return buffer.getvalue()


def test_each_section_has_its_own_headers_and_footers_or_the_previous_ones():
    document = parse_docx(_chapters_file(), "chapters.docx")

    [section] = _breaks(document)
    first = section.sectionBreak
    assert (first.header, first.firstHeader, first.differentFirstPage, first.footer) == ("Chapter one", "", True, None)
    assert (document.settings.header, document.settings.footer) == ("Chapter two", "Footer two")
    assert parse_docx(_chapters_file(second_linked=True), "linked.docx").settings.header is None  # linked to the previous


def test_a_word_export_writes_each_sections_headers_and_links_the_rest():
    document = parse_docx(_chapters_file(), "chapters.docx")

    exported = build_docx(document)

    assert package_problems(exported) == []
    first, last = DocxDocument(io.BytesIO(exported)).sections
    assert first.header.paragraphs[0].text == "Chapter one" and first.different_first_page_header_footer
    assert first.first_page_header.paragraphs[0].text == "" and not first.first_page_header.is_linked_to_previous
    assert first.footer.is_linked_to_previous  # none of its own
    assert (last.header.paragraphs[0].text, last.footer.paragraphs[0].text) == ("Chapter two", "Footer two")
    again = parse_docx(exported, "again.docx")
    assert [element.sectionBreak for element in _breaks(again)] == [element.sectionBreak for element in _breaks(document)]


def test_a_pdf_shows_each_page_its_sections_headers():
    from pypdf import PdfReader

    texts = [page.extract_text() for page in PdfReader(io.BytesIO(build_pdf(parse_docx(_chapters_file(), "chapters.docx")))).pages]
    linked = [page.extract_text() for page in PdfReader(io.BytesIO(build_pdf(parse_docx(_chapters_file(second_linked=True), "linked.docx")))).pages]

    assert "Chapter" not in texts[0]  # the cover: its own first-page header is empty
    assert "Chapter two" in texts[1] and "Footer two" in texts[1] and "Chapter one" not in texts[1]
    assert "Chapter one" in linked[1]  # no header of its own: the previous section's, as Word shows it


def _continuous_file() -> bytes:
    word = DocxDocument()
    word.sections[0].header.paragraphs[0].text = "Alpha header"
    word.add_paragraph("The first section.")
    last = word.add_section(WD_SECTION.CONTINUOUS)
    last.header.is_linked_to_previous = False
    last.header.paragraphs[0].text = "Beta header"
    for index in range(70):
        word.add_paragraph(f"Line {index} of the second section.")
    buffer = io.BytesIO()
    word.save(buffer)
    return buffer.getvalue()


def test_a_continuous_section_shows_its_own_header_from_the_next_page():
    from pypdf import PdfReader

    document = parse_docx(_continuous_file(), "continuous.docx")

    texts = [page.extract_text() for page in PdfReader(io.BytesIO(build_pdf(document))).pages]

    assert len(texts) >= 2 and "Line 0 of the second section." in texts[0]  # it begins on the first page
    assert "Alpha header" in texts[0] and "Beta header" not in texts[0]  # which stays the first section's, as in Word
    assert "Beta header" in texts[1] and "Alpha header" not in texts[1]


def _renumbered_file() -> bytes:
    word = DocxDocument()
    word.add_paragraph("Front matter.")
    last = word.add_section(WD_SECTION.NEW_PAGE)
    last._sectPr.find(qn("w:cols")).addprevious(parse_xml(f'<w:pgNumType {_W} w:fmt="upperRoman" w:start="4"/>'))
    word.add_paragraph("The last section.")
    buffer = io.BytesIO()
    word.save(buffer)
    return buffer.getvalue()


def test_the_last_sections_own_page_numbering_is_kept_in_both_exports():
    from pypdf import PdfReader

    document = parse_docx(_renumbered_file(), "renumbered.docx")
    document.settings.showPageNumbers = True
    assert (document.lastSection.pageNumberStart, document.lastSection.pageNumberFormat) == (4, "upperRoman")

    texts = [page.extract_text() for page in PdfReader(io.BytesIO(build_pdf(document))).pages]
    exported = build_docx(document)

    assert "Page 1" in texts[0] and "Page IV" in texts[1]
    assert package_problems(exported) == []
    with zipfile.ZipFile(io.BytesIO(exported)) as package:
        body = parse_xml(package.read("word/document.xml"))
    last = body.find(qn("w:body")).find(qn("w:sectPr"))
    numbering = last.find(qn("w:pgNumType"))
    assert (numbering.get(qn("w:fmt")), numbering.get(qn("w:start"))) == ("upperRoman", "4")
    order = [child.tag.split("}")[1] for child in last]
    assert order == sorted(order, key=_SECT_PR_ORDER.index)
    again = parse_docx(exported, "again.docx")
    assert (again.lastSection.pageNumberStart, again.lastSection.pageNumberFormat) == (4, "upperRoman")


def test_numbers_are_counted_as_word_counts_them():
    from app.parsers.docx_styles import format_number

    assert [format_number(value, "lowerLetter") for value in (1, 26, 27, 28, 53)] == ["a", "z", "aa", "bb", "aaa"]
    assert (format_number(28, "upperLetter"), format_number(1994, "lowerRoman"), format_number(4, "upperRoman")) == ("BB", "mcmxciv", "IV")
    assert (format_number(0, "lowerRoman"), format_number(4000, "upperRoman"), format_number(0, "lowerLetter")) == ("0", "4000", "0")


_TEXT_BOX = (
    '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
    'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" xmlns:v="urn:schemas-microsoft-com:vml">'
    "<mc:AlternateContent><mc:Choice Requires=\"wps\"><w:drawing><wps:wsp><wps:txbx><w:txbxContent>"
    "<w:p><w:r><w:t>Boxed words</w:t></w:r></w:p></w:txbxContent></wps:txbx></wps:wsp></w:drawing></mc:Choice>"
    "<mc:Fallback><w:pict><v:shape><v:textbox><w:txbxContent><w:p><w:r><w:t>Boxed words</w:t></w:r></w:p>"
    "</w:txbxContent></v:textbox></v:shape></w:pict></mc:Fallback></mc:AlternateContent></w:r>"
)


def test_a_text_box_in_a_header_is_read_once():
    word = DocxDocument()
    paragraph = word.sections[0].header.paragraphs[0]
    paragraph.text = "Header words"
    paragraph._p.append(parse_xml(_TEXT_BOX))
    word.add_paragraph("Body.")
    buffer = io.BytesIO()
    word.save(buffer)

    document = parse_docx(buffer.getvalue(), "boxed.docx")

    assert document.settings.header == "Header words Boxed words"


def _back_cover_file() -> bytes:
    word = DocxDocument()
    word.sections[0].header.paragraphs[0].text = "Chapter one"
    word.add_paragraph("One.")
    last = word.add_section(WD_SECTION.NEW_PAGE)
    last.header.is_linked_to_previous = False  # its own header, left empty: a back cover without one
    word.add_paragraph("The back cover.")
    buffer = io.BytesIO()
    word.save(buffer)
    return buffer.getvalue()


def test_a_last_section_with_its_own_empty_header_shows_none():
    from pypdf import PdfReader

    document = parse_docx(_back_cover_file(), "cover.docx")

    texts = [page.extract_text() for page in PdfReader(io.BytesIO(build_pdf(document))).pages]
    exported = DocxDocument(io.BytesIO(build_docx(document)))

    assert (document.settings.header, document.lastSection.header) == (None, "")
    assert "Chapter one" in texts[0] and "Chapter one" not in texts[1]  # not the previous section's
    last = exported.sections[-1]
    assert not last.header.is_linked_to_previous and last.header.paragraphs[0].text == ""
    assert parse_docx(_chapters_file(second_linked=True), "linked.docx").lastSection.header is None  # linked: none of its own


def _two_column_file() -> bytes:
    word = DocxDocument()
    section = word.sections[0]
    section._sectPr.find(qn("w:cols")).set(qn("w:num"), "2")
    section._sectPr.find(qn("w:cols")).set(qn("w:space"), "709")  # 1.25 cm
    section.header_distance = Cm(2)
    section.header.paragraphs[0].text = "Top words"
    for index in range(80):
        word.add_paragraph(f"Line {index} of a two-column page.")
    buffer = io.BytesIO()
    word.save(buffer)
    return buffer.getvalue()


def _text_at(pdf: bytes, page: int = 0) -> list[tuple[str, float, float]]:
    from pypdf import PdfReader

    found: list[tuple[str, float, float]] = []

    def visit(text, cm_matrix, tm_matrix, font_dict, font_size):
        if text.strip():
            found.append((text.strip(), cm_matrix[4] + tm_matrix[4], cm_matrix[5] + tm_matrix[5]))

    PdfReader(io.BytesIO(pdf)).pages[page].extract_text(visitor_text=visit)
    return found


def test_the_last_sections_columns_and_header_distance_are_kept_in_both_exports():
    from docx.shared import Cm as WordCm

    document = parse_docx(_two_column_file(), "columns.docx")
    last = document.lastSection
    assert (last.columns, last.columnSpacingCm, last.headerDistanceCm) == (2, 1.25, 2.0)

    pdf = build_pdf(document)
    placed = _text_at(pdf)
    width = float(__import__("pypdf").PdfReader(io.BytesIO(pdf)).pages[0].mediabox.width)
    height = float(__import__("pypdf").PdfReader(io.BytesIO(pdf)).pages[0].mediabox.height)
    lines = [(x, y) for text, x, y in placed if text.startswith("Line ")]
    assert any(x < width / 2 for x, _ in lines) and any(x > width / 2 for x, _ in lines)  # two columns side by side
    [(_, _, top)] = [entry for entry in placed if entry[0] == "Top words"]
    assert abs((height - top) - (2 * 72 / 2.54 + 9)) < 1  # the header at its own distance from the page's top

    exported = build_docx(document)
    assert package_problems(exported) == []
    section = DocxDocument(io.BytesIO(exported)).sections[-1]
    columns = section._sectPr.find(qn("w:cols"))
    assert (columns.get(qn("w:num")), columns.get(qn("w:space"))) == ("2", "709")
    assert abs(section.header_distance - WordCm(2)) < WordCm(0.01)


def _custom_size_file(width_mm: float = 200, height_mm: float = 250) -> bytes:
    """One section on a paper size the app doesn't list."""
    from docx.shared import Mm

    word = DocxDocument()
    section = word.sections[0]
    section.page_width, section.page_height = Mm(width_mm), Mm(height_mm)
    word.add_paragraph("On paper of its own size.")
    buffer = io.BytesIO()
    word.save(buffer)
    return buffer.getvalue()


def test_a_last_section_on_an_unlisted_paper_size_keeps_it_in_both_exports():
    # DOCX-015A: it used to be shown and exported on A4.
    document = parse_docx(_custom_size_file(), "custom.docx")
    assert (round(document.lastSection.pageWidthMm), round(document.lastSection.pageHeightMm)) == (200, 250)
    assert any("it is kept for this document" in note for note in document.unsupportedFeatures)

    word = DocxDocument(io.BytesIO(build_docx(document)))
    assert (round(word.sections[-1].page_width.mm), round(word.sections[-1].page_height.mm)) == (200, 250)

    from pypdf import PdfReader

    box = PdfReader(io.BytesIO(build_pdf(document))).pages[0].mediabox
    assert (round(float(box.width) / 72 * 25.4), round(float(box.height) / 72 * 25.4)) == (200, 250)

    listed = parse_docx(_word_file(), "listed.docx")  # A4 and Letter are the app's: DocumentSettings has them
    assert listed.lastSection is None or listed.lastSection.pageWidthMm is None


def test_a_page_size_chosen_here_replaces_the_unlisted_one(api_db):
    client.cookies.clear()
    assert client.post("/api/v1/auth/register", json={"email": "paper@example.com", "password": "long enough password"}).status_code == 201
    uploaded = client.post("/api/v1/documents/upload", files={"file": ("custom.docx", _custom_size_file(), _DOCX)})
    assert uploaded.status_code == 201, uploaded.text[:300]
    document = uploaded.json()
    assert round(document["lastSection"]["pageWidthMm"]) == 200
    size_item = next(item for item in document["importReport"]["items"] if item["feature"] == "docx.page_setup.size")
    assert size_item["policy"] == "detected_preserved" and not size_item["contentChanged"]

    chosen = client.patch(f"/api/v1/documents/{document['id']}/settings", json={"property": "pageSize", "value": "Letter"})
    assert chosen.status_code == 200, chosen.text[:300]
    after = chosen.json()
    assert after["settings"]["pageSize"] == "Letter"
    assert after["lastSection"] is None or after["lastSection"]["pageWidthMm"] is None  # the choice is the page now
    pdf = client.get(f"/api/v1/documents/{document['id']}/export/pdf").content
    from pypdf import PdfReader

    box = PdfReader(io.BytesIO(pdf)).pages[0].mediabox
    assert (round(float(box.width)), round(float(box.height))) == (612, 792)  # Letter
    client.cookies.clear()
