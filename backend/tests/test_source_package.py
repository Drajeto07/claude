"""An imported Word file is kept as it was (tracker DOCX-010), and a Word export
is written into it (DOCX-011): the body is the document's, everything else --
custom styles, headers and footers of every kind, columns, custom properties,
the theme -- is the original file's, rewritten only where the document changed
it. Without the file, or with one that isn't the one kept, the export is built
without it and says so."""

import hashlib
import io
import zipfile

import pytest
from docx import Document as DocxDocument
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from fastapi.testclient import TestClient
from PIL import Image as PILImage
from sqlalchemy import text

from app.export.docx_export import build_docx
from app.export.package_check import package_problems
from app.fidelity.content import compare_words, document_words, words
from app.fidelity.docx_source import read_docx_source
from app.fidelity.report import ReportBuilder
from app.main import app
from app.models.document import Document, Element, ElementType, InlineRun

client = TestClient(app, base_url="https://testserver")
_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_CUSTOM = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties" '
    'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
    '<property fmtid="{D5CDD505-2E9C-101B-9397-08002B2CF9AE}" pid="2" name="Client"><vt:lpwstr>Acme Holdings</vt:lpwstr></property>'
    "</Properties>"
)


def _png() -> bytes:
    buffer = io.BytesIO()
    PILImage.new("RGB", (40, 20), "green").save(buffer, format="PNG")
    return buffer.getvalue()


def _with_custom_properties(data: bytes) -> bytes:
    """docProps/custom.xml as Word writes it: the part, its content type and the package relationship."""
    source, target = io.BytesIO(data), io.BytesIO()
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as after:
        for item in before.infolist():
            content = before.read(item.filename)
            if item.filename == "[Content_Types].xml":
                content = content.replace(
                    b"</Types>",
                    b'<Override PartName="/docProps/custom.xml" ContentType="application/vnd.openxmlformats-officedocument.custom-properties+xml"/></Types>',
                )
            elif item.filename == "_rels/.rels":
                content = content.replace(
                    b"</Relationships>",
                    b'<Relationship Id="rIdCustom" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/custom-properties" Target="docProps/custom.xml"/></Relationships>',
                )
            after.writestr(item, content)
        after.writestr("docProps/custom.xml", _CUSTOM)
    return target.getvalue()


def _word_file() -> bytes:
    document = DocxDocument()
    document.styles["Normal"].font.name = "Georgia"
    document.styles.add_style("Memo Note", WD_STYLE_TYPE.PARAGRAPH).font.italic = True
    section = document.sections[0]
    section.different_first_page_header_footer = True
    section.first_page_header.paragraphs[0].text = "CONFIDENTIAL cover page"
    section.header.paragraphs[0].text = "Quarterly report"
    section._sectPr.find(qn("w:cols")).set(qn("w:num"), "2")
    document.add_heading("Findings", level=1)
    document.add_paragraph("The memo text stays in its own style.", style="Memo Note")
    reviewed = document.add_paragraph("A reviewed sentence.")
    document.add_comment(runs=reviewed.runs, text="Check this figure.", author="Ana Petrova", initials="AP")
    document.add_picture(io.BytesIO(_png()))
    buffer = io.BytesIO()
    document.save(buffer)
    return _with_custom_properties(buffer.getvalue())


@pytest.fixture
def uploaded(api_db):
    client.cookies.clear()
    assert client.post("/api/v1/auth/register", json={"email": "source@example.com", "password": "long enough password"}).status_code == 201
    original = _word_file()
    response = client.post("/api/v1/documents/upload", files={"file": ("report.docx", original, _DOCX)})
    assert response.status_code == 201, response.text
    yield original, response.json()
    client.cookies.clear()


def _export(document_id: str) -> bytes:
    response = client.get(f"/api/v1/documents/{document_id}/export/docx")
    assert response.status_code == 200, response.text
    return response.content


def _parts(data: bytes) -> dict[str, str]:
    with zipfile.ZipFile(io.BytesIO(data)) as package:
        return {name: package.read(name).decode("utf-8", "replace") for name in package.namelist()}


def test_the_imported_word_file_is_kept_as_it_was(uploaded):
    original, document = uploaded

    source = document["sourcePackage"]
    assert (source["sha256"], source["size"], source["format"]) == (hashlib.sha256(original).hexdigest(), len(original), "docx")
    pasted = client.post("/api/v1/documents", json={"text": "Pasted text."}).json()
    assert pasted["sourcePackage"] is None


def test_a_word_export_keeps_what_the_document_model_doesnt_hold(uploaded):
    _, document = uploaded

    exported = _export(document["id"])

    parts = _parts(exported)
    assert "Acme Holdings" in parts["docProps/custom.xml"]  # custom properties
    assert "Memo Note" in parts["word/styles.xml"] and "Georgia" in parts["word/styles.xml"]  # the file's own styles, unchanged
    body = parts["word/document.xml"]
    assert 'w:num="2"' in body and "<w:titlePg" in body  # two columns, a first page of its own
    headers = [xml for name, xml in parts.items() if name.startswith("word/header")]
    assert any("CONFIDENTIAL cover page" in xml for xml in headers) and any("Quarterly report" in xml for xml in headers)
    assert parts["word/comments.xml"].count("<w:comment ") == 1  # written again, not twice
    assert len([name for name in parts if name.startswith("word/media/")]) == 1  # the old body's picture isn't left behind
    reopened = compare_words(document_words(Document.model_validate(document).elements), words(read_docx_source(exported).body), method="t")
    assert reopened.verified  # the body is the document's
    assert package_problems(exported) == []  # and Word gets a sound package (TEST-023)


def test_a_header_changed_in_the_app_is_rewritten_and_the_rest_kept(uploaded):
    _, document = uploaded
    assert client.patch(f"/api/v1/documents/{document['id']}/settings", json={"property": "header", "value": "Annual report"}).status_code == 200

    headers = [xml for name, xml in _parts(_export(document["id"])).items() if name.startswith("word/header")]

    assert any("Annual report" in xml for xml in headers)
    assert not any("Quarterly report" in xml for xml in headers)
    assert any("CONFIDENTIAL cover page" in xml for xml in headers)


def test_a_template_restyles_the_file_and_nothing_else_does(uploaded):
    _, document = uploaded
    assert client.post(f"/api/v1/documents/{document['id']}/format", data={"templateId": "academic-default"}).status_code == 200

    styles = _parts(_export(document["id"]))["word/styles.xml"]

    assert "Times New Roman" in styles  # the template's body font, in the file's Normal style
    assert "Memo Note" in styles


def test_a_stored_file_that_isnt_the_one_kept_is_not_used_and_the_export_says_so(uploaded, api_db, tmp_path):
    _, document = uploaded
    with api_db.connect() as connection:
        key = connection.execute(text("SELECT storage_key FROM document_assets WHERE id = :id"), {"id": document["sourcePackage"]["assetId"]}).scalar_one()
    (tmp_path / "assets" / key).write_bytes(b"not the original file")

    job = client.post("/api/v1/jobs/export", json={"documentId": document["id"], "format": "docx"}).json()
    finished = client.get(f"/api/v1/jobs/{job['id']}").json()

    assert finished["status"] == "succeeded", finished
    reasons = {item["feature"]: item["reason"] for item in finished["result"]["fidelity"]["items"]}
    assert "isn't the one this document came from" in reasons["export.docx.source_missing"]
    assert "export.docx.source_package" not in reasons


def test_an_unreadable_file_is_left_out_of_the_export_and_named():
    document = Document(elements=[Element(type=ElementType.PARAGRAPH, content="Body.", inline=[InlineRun(text="Body.")], order=0)])
    report = ReportBuilder()

    exported = build_docx(document, source=b"not a zip", report=report)

    assert [item.feature for item in report.items()] == ["export.docx.source_unreadable"]
    assert DocxDocument(io.BytesIO(exported)).paragraphs[0].text == "Body."


def test_the_import_report_says_what_the_word_export_keeps(uploaded):
    _, document = uploaded

    items = {item["feature"]: item for item in document["importReport"]["items"]}

    assert "docx.header_footer.variants" not in items  # the first-page header is the document's own now (DOCX-015)
    assert (document["lastSection"]["firstHeader"], document["lastSection"]["differentFirstPage"]) == ("CONFIDENTIAL cover page", True)
    assert items["docx.metadata.custom_properties"]["policy"] == "detected_not_editable"
    assert "the Word export keeps the columns" in next(item["reason"] for item in document["importReport"]["items"] if item["feature"] == "docx.layout")
    assert "docx.header_footer.text" not in items  # the first-page header's words aren't lost: the Word export keeps them
    assert document["importReport"]["contentLossCount"] == 0

    job = client.post("/api/v1/jobs/export", json={"documentId": document["id"], "format": "pdf"}).json()
    pdf = client.get(f"/api/v1/jobs/{job['id']}").json()
    assert "export.pdf.word_only" in {item["feature"] for item in pdf["result"]["fidelity"]["items"]}


def test_headers_of_earlier_sections_are_their_sections_own(api_db):
    from docx.enum.section import WD_SECTION

    client.cookies.clear()
    assert client.post("/api/v1/auth/register", json={"email": "sections@example.com", "password": "long enough password"}).status_code == 201
    word = DocxDocument()
    word.sections[0].header.paragraphs[0].text = "Chapter one running head"
    word.add_paragraph("Chapter one.")
    later = word.add_section(WD_SECTION.NEW_PAGE)
    later.header.is_linked_to_previous = False
    later.header.paragraphs[0].text = "Quarterly report"
    word.add_paragraph("Chapter two.")
    buffer = io.BytesIO()
    word.save(buffer)

    document = client.post("/api/v1/documents/upload", files={"file": ("chapters.docx", buffer.getvalue(), _DOCX)}).json()
    exported = _export(document["id"])
    client.cookies.clear()

    assert "docx.header_footer.text" not in {item["feature"] for item in document["importReport"]["items"]}  # nothing left out
    section_break = next(element for element in document["elements"] if element["type"] == "section_break")
    assert section_break["sectionBreak"]["header"] == "Chapter one running head"  # the first section's own (DOCX-015)
    assert document["settings"]["header"] == "Quarterly report"
    sections = DocxDocument(io.BytesIO(exported)).sections
    assert [section.header.paragraphs[0].text for section in sections] == ["Chapter one running head", "Quarterly report"]


def test_sections_own_properties_are_kept_and_named(api_db):
    from docx.enum.section import WD_ORIENT, WD_SECTION
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    client.cookies.clear()
    assert client.post("/api/v1/auth/register", json={"email": "borders@example.com", "password": "long enough password"}).status_code == 201
    word = DocxDocument()
    first = word.sections[0]
    first.orientation, first.page_width, first.page_height = WD_ORIENT.LANDSCAPE, first.page_height, first.page_width
    word.add_paragraph("A wide first section.")
    last = word.add_section(WD_SECTION.NEW_PAGE)
    last.orientation, last.page_width, last.page_height = WD_ORIENT.PORTRAIT, last.page_height, last.page_width
    last._sectPr.append(parse_xml(f'<w:pgBorders {nsdecls("w")}><w:top w:val="single" w:sz="4" w:space="24" w:color="auto"/></w:pgBorders>'))
    word.add_paragraph("The last section, with a border.")
    buffer = io.BytesIO()
    word.save(buffer)

    document = client.post("/api/v1/documents/upload", files={"file": ("borders.docx", buffer.getvalue(), _DOCX)}).json()
    exported = _export(document["id"])
    client.cookies.clear()

    items = {(item["feature"], item["policy"]) for item in document["importReport"]["items"]}
    assert ("docx.sections.page_borders", "detected_not_editable") in items  # the last section's: kept in the Word export
    assert ("docx.sections.page_setup", "detected_not_editable") in items  # the landscape first section's, while unchanged
    body = _parts(exported)["word/document.xml"]
    assert "<w:pgBorders" in body and 'w:orient="landscape"' in body
