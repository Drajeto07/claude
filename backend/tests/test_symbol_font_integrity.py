"""B-032: a character whose look depends on its font -- a symbol font's (Wingdings, Symbol, Webdings...), text in
the Private Use Area, a w:sym -- is never changed by a safe repair, and SafeEdit never calls a file verified when
one was.

Found by the Validation Pack (S2-100, r05-fonts.docx: the font fix set a Wingdings tick in Arial; SafeEdit compared
the Unicode words, which hadn't changed, and said verified). These tests check the Word files' XML; they don't
render the pages, so they prove the font and the characters are kept, not how a given viewer draws them.
"""

import io
import re
import zipfile

import pytest
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.doctor import safe_edit
from app.doctor.corpus import imported, run
from app.doctor.doctor import analyse
from app.doctor.safe_edit import verify
from app.formatting import health_fixes
from scripts.doctor_validation import FIXTURES, differences, kept
from scripts.make_doctor_fixtures import _fixed_dates, new_document, save, text_paragraphs

R05 = FIXTURES / "r05-fonts.docx"
TICK = ""  # Wingdings: a tick


def _with_document_xml(data: bytes, change) -> bytes:
    source = zipfile.ZipFile(io.BytesIO(data))
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            content = source.read(item.filename)
            target.writestr(item, change(content) if item.filename == "word/document.xml" else content)
    return buffer.getvalue()


def _document(build) -> bytes:
    document = new_document()
    text_paragraphs(document, 3)
    build(document)
    text_paragraphs(document, 2, 3)
    return _fixed_dates(save(document))


def _run_in(font: str, text: str):
    def build(document):
        paragraph = document.add_paragraph("Approved ")
        paragraph.add_run(text).font.name = font
        paragraph.add_run(" by the board on 2026-10-01.")

    return build


def _sym(char: str):
    def build(document):
        paragraph = document.add_paragraph("Done ")
        symbol = OxmlElement("w:sym")
        symbol.set(qn("w:font"), "Wingdings")
        symbol.set(qn("w:char"), char)
        paragraph.add_run()._r.append(symbol)
        paragraph.add_run(" and ").font.name = "Comic Sans MS"  # a font repair that would rewrite this block
        paragraph.add_run("checked.")

    return build


def _character_style(document):
    style = document.styles.add_style("Tick", WD_STYLE_TYPE.CHARACTER)
    style.font.name = "Wingdings"
    paragraph = document.add_paragraph("Ticked ")
    paragraph.add_run(TICK, style="Tick")
    paragraph.add_run(" through a style.")


def _mixed(document):
    paragraph = document.add_paragraph("Totals are final ")
    paragraph.add_run("except travel").font.name = "Comic Sans MS"
    paragraph.add_run(" and ")
    paragraph.add_run(TICK).font.name = "Wingdings"
    paragraph.add_run(" approved.")


SYMBOL_CASES = {
    "wingdings-private-use": _run_in("Wingdings", TICK),
    "symbol-ascii-alpha": _run_in("Symbol", "a"),  # Symbol's "a" is an alpha
    "webdings": _run_in("Webdings", ""),
    "icon-font-private-use": _run_in("Font Awesome 5 Free", ""),
    "w-sym-known-tick": _sym("F0FC"),
    "w-sym-unknown": _sym("F028"),
    "character-style": _character_style,
    "symbol-next-to-a-real-off-font": _mixed,
}


def _old_font_fix(monkeypatch):
    """The font fix as it was before B-032: every font other than the body's main one cleared."""
    monkeypatch.setattr(health_fixes, "font_carries_meaning", lambda font, text: False)


# -- The original corruption, and the guard ----------------------------------------------------------------------


def test_the_original_wingdings_corruption_reproduces_with_the_old_font_fix_and_is_now_caught(monkeypatch):
    data = R05.read_bytes()
    _old_font_fix(monkeypatch)
    outcome = run(data, R05.name)
    assert outcome.by_repair == {"fonts": 2}  # the Comic Sans words and the Wingdings tick
    assert differences(data, outcome.file)["symbols"]  # read independently: the tick is no longer in Wingdings
    assert kept({"symbols": [[TICK, "Wingdings"]]}, outcome.file) == ["Wingdings character"]
    assert not outcome.verified  # before B-032 this file was verified
    assert [check.name for check in outcome.result.failed()] == ["preserved"]
    assert "symbol-font texts changed or set in another font" in outcome.result.failed()[0].detail


def test_the_font_fix_now_leaves_the_tick_alone_and_still_repairs_the_off_font_words():
    data = R05.read_bytes()
    outcome = run(data, R05.name)
    assert outcome.by_repair == {"fonts": 1}  # the Comic Sans words only
    assert outcome.verified, [check.detail for check in outcome.result.failed()]
    assert not any(differences(data, outcome.file).values())
    assert kept({"symbols": [[TICK, "Wingdings"]]}, outcome.file) == []
    after = imported(outcome.file, R05.name)
    fonts = {mark.fontFamily for element in after.elements for run_ in element.inline or [] for mark in run_.marks if mark.fontFamily}
    assert fonts == {"Wingdings"}


# -- Each kind of font-dependent text ---------------------------------------------------------------------------


@pytest.mark.parametrize("case", sorted(SYMBOL_CASES), ids=sorted(SYMBOL_CASES))
def test_font_dependent_text_survives_doctor_and_its_file_is_verified(case):
    data = _document(SYMBOL_CASES[case])
    outcome = run(data, f"{case}.docx")
    assert outcome.refusals == []
    assert outcome.verified, [check.detail for check in outcome.result.failed()]
    assert not any(differences(data, outcome.file).values())  # text, numbers, tables, structure, objects, symbols
    assert safe_edit._font_dependent(data) == safe_edit._font_dependent(outcome.file) != []


@pytest.mark.parametrize("case", sorted(SYMBOL_CASES), ids=sorted(SYMBOL_CASES))
def test_with_the_old_font_fix_no_changed_symbol_is_verified(case, monkeypatch):
    data = _document(SYMBOL_CASES[case])
    _old_font_fix(monkeypatch)
    outcome = run(data, f"{case}.docx")
    if safe_edit._font_dependent(outcome.file) != safe_edit._font_dependent(data):
        assert not outcome.verified
        assert "preserved" in [check.name for check in outcome.result.failed()]
    else:
        assert outcome.verified  # nothing font-dependent changed (the block was refused, or the font came from a style)


@pytest.mark.parametrize(("char", "becomes"), [("F0FC", "✓"), ("F028", None)], ids=["known-tick-becomes-text", "unknown-is-dropped"])
def test_a_w_sym_rewritten_the_old_way_is_caught_by_safeedit(char, becomes, monkeypatch):
    """Before B-032 a block holding a w:sym could be rewritten by another repair: the import reads a known Wingdings
    sign as its Unicode look-alike (a tick becomes the text U+2713) and leaves an unknown one out."""
    data = _document(_sym(char))
    monkeypatch.delitem(safe_edit.PRESERVED, "symbol characters")  # the per-block refusal and its count, as before
    outcome = run(data, "sym.docx")
    assert outcome.by_repair == {"fonts": 1}
    with zipfile.ZipFile(io.BytesIO(outcome.file)) as package:
        xml = package.read("word/document.xml").decode("utf-8")
    assert "<w:sym " not in xml
    assert (becomes in xml) if becomes else True
    assert not outcome.verified
    assert any(check.name == "preserved" and "symbol-font" in check.detail for check in outcome.result.failed())


def test_a_block_holding_a_symbol_character_is_refused_per_block():
    data = _document(_sym("F028"))
    document = imported(data, "sym.docx")
    analysis = analyse(document, data, name="sym.docx")
    assert [refusal.checkId for refusal in analysis.blockRefusals] == ["fonts"]
    assert "symbol characters" in analysis.blockRefusals[0].reason


# -- SafeEdit on its own: deliberate changes ---------------------------------------------------------------------

_WINGDINGS_FONTS = b'<w:rFonts w:ascii="Wingdings" w:hAnsi="Wingdings"/>'
DELIBERATE = {
    "font removed": lambda xml: xml.replace(_WINGDINGS_FONTS, b"", 1),
    "set in Arial": lambda xml: xml.replace(_WINGDINGS_FONTS, b'<w:rFonts w:ascii="Arial" w:hAnsi="Arial"/>', 1),
    "set in another symbol font": lambda xml: xml.replace(b"Wingdings", b"Webdings"),
    "the character replaced": lambda xml: xml.replace(TICK.encode(), "".encode(), 1),  # a cross
    "the character removed": lambda xml: xml.replace(TICK.encode(), b"", 1),
    "ordinary text set in a symbol font": lambda xml: xml.replace(
        b'<w:r><w:t xml:space="preserve">Totals are final </w:t>', b"<w:r><w:rPr>" + _WINGDINGS_FONTS + b'</w:rPr><w:t xml:space="preserve">Totals are final </w:t>', 1
    ),
}


@pytest.mark.parametrize("change", sorted(DELIBERATE), ids=sorted(DELIBERATE))
def test_safeedit_does_not_verify_a_deliberate_symbol_change(change):
    data = R05.read_bytes()
    changed = _with_document_xml(data, DELIBERATE[change])
    assert changed != data
    result = verify(data, changed, set(range(len(safe_edit._body_children(data)))))  # every block "touched": only the guard can catch it
    assert not result.verified
    assert any(check.name == "preserved" and "symbol-font" in check.detail for check in result.failed())


@pytest.mark.parametrize("change", ["removed", "made text"], ids=["removed", "made-text"])
def test_safeedit_does_not_verify_a_lost_or_converted_w_sym(change):
    data = _document(_sym("F0FC"))
    replacement = b"" if change == "removed" else "<w:t>✓</w:t>".encode()
    changed = _with_document_xml(data, lambda xml: re.sub(rb"<w:sym [^>]*/>", replacement, xml, count=1))
    result = verify(data, changed, set(range(len(safe_edit._body_children(data)))))
    assert not result.verified
    assert any(check.name == "preserved" for check in result.failed())


def test_runs_split_or_joined_in_one_symbol_font_are_not_a_change():
    data = _document(_run_in("Wingdings", TICK + ""))
    split = _with_document_xml(
        data,
        lambda xml: xml.replace(
            ("<w:t>" + TICK + "</w:t>").encode(),
            ("<w:t>" + TICK + "</w:t></w:r><w:r><w:rPr>").encode() + _WINGDINGS_FONTS + ("</w:rPr><w:t></w:t>").encode(),
        ),
    )
    assert split != data
    assert verify(data, split, set(range(len(safe_edit._body_children(data))))).verified


def test_styles_safeedit_cannot_read_are_not_verified():
    data = R05.read_bytes()
    source = zipfile.ZipFile(io.BytesIO(data))
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            target.writestr(item, b"<not xml" if item.filename == "word/styles.xml" else source.read(item.filename))
    result = verify(data, buffer.getvalue(), set())
    assert not result.verified


def test_an_unchanged_file_is_verified():
    data = R05.read_bytes()
    assert verify(data, data, set()).verified


# -- The two lists agree, and ordinary fonts are repaired -------------------------------------------------------


def test_the_repair_and_safeedit_know_the_same_symbol_fonts():
    assert health_fixes.SYMBOL_FONTS == safe_edit.SYMBOL_FONTS


@pytest.mark.parametrize(
    ("font", "text", "carries"),
    [
        ("Wingdings", TICK, True),
        ("Symbol", "a", True),
        ("Webdings", "x", True),
        (" wingdings 2 ", "x", True),
        ("Font Awesome 5 Free", "", True),
        ("Arial", "\U000f0001", True),
        ("Comic Sans MS", "except travel", False),
        ("Times New Roman", "✓ and α", False),  # real Unicode signs: any font with them shows them
        (None, "text", False),
    ],
)
def test_which_text_carries_its_meaning_in_its_font(font, text, carries):
    assert health_fixes.font_carries_meaning(font, text) is carries


def test_ordinary_off_font_words_are_still_repaired():
    data = _document(_run_in("Comic Sans MS", "except the travel budget"))
    outcome = run(data, "ordinary.docx")
    assert outcome.by_repair == {"fonts": 1}
    assert outcome.verified
    assert not any(differences(data, outcome.file).values())
    after = imported(outcome.file, "ordinary.docx")
    assert not any(mark.fontFamily for element in after.elements for run_ in element.inline or [] for mark in run_.marks)
