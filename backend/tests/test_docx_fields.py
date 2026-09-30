"""Fields that run across paragraphs -- a table of contents, a bibliography -- stay
fields through a Word export (tracker DOCX-020): the import keeps where each starts on
the block it begins in and where it ends on the block it ends in (Word ends them in a
paragraph of their own, which the import leaves out); a Word export writes them back
around their entries, and into the original it copies an unchanged one whole. One
whose first or last paragraph was deleted is written as its text, and the export says
so."""

import io
import zipfile
from pathlib import Path

from docx import Document as DocxDocument
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from lxml import etree

from app.export.docx_export import build_docx
from app.export.package_check import package_problems
from app.export.provenance import stamp
from app.fidelity.report import ReportBuilder
from app.formatting.engine import apply_formatting
from app.formatting.templates import BUILTIN_TEMPLATES
from app.parsers.docx import parse_docx

WORD = Path(__file__).parent / "fixtures" / "word"
_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_NS = nsdecls("w")
_TOC = (
    f'<w:p {_NS}><w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText xml:space="preserve"> TOC \\o "1-3" \\h </w:instrText></w:r>'
    '<w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>Chapter one</w:t></w:r></w:p>',
    f'<w:p {_NS}><w:r><w:t>Chapter two</w:t></w:r></w:p>',
    f'<w:p {_NS}><w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>',  # where Word ends one: a paragraph of its own
)


def _word_file() -> bytes:
    word = DocxDocument()
    body = word.element.body
    for xml in _TOC:
        body.insert(len(body) - 1, parse_xml(xml))
    word.add_heading("Chapter one", level=1)
    word.add_paragraph("The text.")
    buffer = io.BytesIO()
    word.save(buffer)
    return buffer.getvalue()


def _paragraphs(data: bytes) -> list:
    with zipfile.ZipFile(io.BytesIO(data)) as package:
        return list(etree.fromstring(package.read("word/document.xml")).find(f"{_W}body").iter(f"{_W}p"))


def _field_chars(paragraph) -> list[str]:
    return [node.get(f"{_W}fldCharType") for node in paragraph.iter(f"{_W}fldChar")]


def _balanced(data: bytes) -> bool:
    depth = 0
    for paragraph in _paragraphs(data):
        for kind in _field_chars(paragraph):
            depth += {"begin": 1, "end": -1}.get(kind, 0)
            if depth < 0:
                return False
    return depth == 0


def _instructions(data: bytes) -> str:
    return " ".join("".join(node.itertext()) for paragraph in _paragraphs(data) for node in paragraph.iter(f"{_W}instrText"))


def test_a_table_of_contents_is_kept_where_it_starts_and_ends():
    document = parse_docx(_word_file(), "toc.docx")
    first, second = document.elements[0], document.elements[1]

    [start] = first.preservedAttributes["ooxml"]
    [end] = second.preservedAttributes["ooxml"]
    assert (start["kind"], start["instr"].split()[0], start["start"]) == ("field_open", "TOC", 0)
    assert (end["kind"], end["region"], end["paragraph"], end["start"]) == ("field_close", start["region"], True, len("Chapter two"))
    items = {item.feature: item for item in document.importReport.items}
    assert items["docx.toc"].policy == "detected_not_editable"
    assert "docx.preserved.flattened" not in items and "docx.empty_paragraph" not in items  # the end's paragraph isn't spacing


def test_written_anew_a_table_of_contents_goes_back_around_its_entries():
    exported = build_docx(parse_docx(_word_file(), "toc.docx"))
    paragraphs = _paragraphs(exported)

    assert _field_chars(paragraphs[0]) == ["begin", "separate"] and "TOC" in _instructions(exported)
    assert ["".join(paragraph.itertext()) for paragraph in paragraphs[1:3]] == ["Chapter two", ""]
    assert _field_chars(paragraphs[2]) == ["end"]  # in a paragraph of its own, as Word ends one
    assert _balanced(exported) and package_problems(exported) == []


def test_into_the_original_an_unchanged_table_of_contents_is_copied_whole():
    data = _word_file()
    document = parse_docx(data, "toc.docx")
    stamp(document)
    noted = ReportBuilder()

    exported = build_docx(document, source=data, report=noted)

    assert _instructions(exported).count("TOC") == 1 and _balanced(exported)
    assert "export.docx.rewritten_blocks" not in {item.feature for item in noted.items()}


def test_a_table_of_contents_missing_its_last_paragraph_is_written_as_its_text():
    document = parse_docx(_word_file(), "toc.docx")
    document.elements = [document.elements[0], *document.elements[2:]]  # its last entry, and its end, deleted
    noted = ReportBuilder()

    exported = build_docx(document, report=noted)

    assert "TOC" not in _instructions(exported) and _balanced(exported)
    assert "export.docx.field_region" in {item.feature for item in noted.items()}
    assert "Chapter one" in "".join(_paragraphs(exported)[0].itertext())


def test_word_s_own_table_of_contents_and_bibliography_survive_a_template():
    template = BUILTIN_TEMPLATES["academic-default"]
    for name, fields in (("a06-fields.docx", ("TOC", "BIBLIOGRAPHY", "CITATION")), ("r01-university-paper.docx", ("TOC",))):
        data = (WORD / name).read_bytes()
        document = parse_docx(data, name)
        stamp(document)
        apply_formatting(document, template_id=template.id, template_rules=template.rules, instruction_rules=[])

        exported = build_docx(document, source=data)

        for field in fields:
            assert field in _instructions(exported), (name, field)
        assert _balanced(exported) and package_problems(exported) == [], name
