"""Deleting an account (tracker ACCT-005, brief §73): the password first, then every row of
the user's in one transaction, their files from storage afterwards, a goodbye e-mail and
the cookie cleared. What the policy keeps or refuses is services/account_deletion.py's."""

import base64
import io
import json
import logging

import pytest
from docx import Document as DocxDocument
from fastapi.testclient import TestClient
from PIL import Image as PILImage
from sqlalchemy import insert, select, update

from app.ai.factory import get_ai_provider
from app.config import get_settings
from app.db.models import Base, Subscription, User, WorkspaceMember, WorkspaceRole
from app.main import app
from app.services.account_deletion import delete_files
from app.storage.local_provider import LocalStorageProvider
from tests.fakes import FakeAIProvider
from tests.helpers import error_body

pytestmark = pytest.mark.security  # the security regression suite (TEST-030)

_PASSWORD = "long enough password"
_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _client() -> TestClient:
    return TestClient(app, base_url="https://testserver")


def _docx(text: str) -> bytes:
    doc = DocxDocument()
    doc.add_heading("Report", level=1)
    doc.add_paragraph(text)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _png() -> str:
    buffer = io.BytesIO()
    PILImage.new("RGB", (4, 3), (10, 120, 200)).save(buffer, format="PNG")
    return "data:image/png;base64," + base64.b64encode(buffer.getvalue()).decode()


def _user_with_everything(email: str) -> tuple[TestClient, dict]:
    """A user with a document holding a picture, an imported Word file (its original kept),
    a finished export, a template, a second session and a used reset link."""
    client = _client()
    me = client.post("/api/v1/auth/register", json={"email": email, "password": _PASSWORD}).json()
    document = client.post("/api/v1/documents", json={"text": f"# {email}\n\nA private paragraph."}).json()
    picture = {"type": "image", "content": "", "order": 9, "image": {"src": _png()}}
    assert client.put(f"/api/v1/documents/{document['id']}/content", json={"elements": document["elements"] + [picture]}).status_code == 200
    imported = client.post("/api/v1/jobs/import-file", files={"file": ("report.docx", _docx(f"Text of {email}."), _DOCX)}).json()
    exported = client.post("/api/v1/jobs/export", json={"documentId": document["id"], "format": "docx"}).json()
    assert client.post("/api/v1/templates", json={"name": f"{email}'s style"}).status_code == 201
    assert _client().post("/api/v1/auth/login", json={"email": email, "password": _PASSWORD}).status_code == 200
    assert client.post("/api/v1/auth/password-reset", json={"email": email}).status_code == 202
    assert imported["status"] == exported["status"] == "succeeded", (imported, exported)
    return client, me


@pytest.fixture
def two_users(api_db):
    app.dependency_overrides[get_ai_provider] = lambda: FakeAIProvider([])
    frank = _user_with_everything("frank@example.com")
    grace = _user_with_everything("grace@example.com")
    yield api_db, frank, grace
    app.dependency_overrides.pop(get_ai_provider, None)


def _everything(engine) -> dict[str, list[dict]]:
    """Every row of every table the models know -- a table added later included."""
    with engine.connect() as connection:
        return {table.name: [dict(row._mapping) for row in connection.execute(table.select())] for table in Base.metadata.sorted_tables}


def _mentions(row: dict, needles: set[str]) -> bool:
    for value in row.values():
        if isinstance(value, (dict, list)):
            value = json.dumps(value, ensure_ascii=False)
        if isinstance(value, str) and any(needle in value for needle in needles):
            return True
    return False


def _ids_of(rows: dict[str, list[dict]], me: dict) -> set[str]:
    """The user's id, address and workspace, and the id of every row that mentions any of
    them, or a row that does, and so on."""
    found = {me["id"], me["email"], me["workspaceId"]}
    while True:
        more = {row["id"] for table in rows.values() for row in table if row["id"] not in found and _mentions(row, found)}
        if not more:
            return found
        found |= more


def _files(storage_root, rows: dict[str, list[dict]], ids: set[str]) -> list:
    """The files in storage that the rows of `ids` name."""
    keys = [row["storage_key"] for row in rows["document_assets"] if row["id"] in ids]
    for job in rows["processing_jobs"]:
        if job["id"] in ids:
            keys += [key for key in (job["input_key"], (job["result"] or {}).get("key")) if key]
    # A finished job's upload is gone already; its row still names it.
    return [storage_root / key for key in keys if (storage_root / key).exists()]


def _delete(client: TestClient, password: str = _PASSWORD):
    return client.request("DELETE", "/api/v1/auth/account", json={"password": password})


def test_every_row_and_file_of_the_user_goes_and_everything_of_another_stays(two_users, tmp_path, sent_mail, caplog):
    engine, (frank, frank_me), (grace, grace_me) = two_users
    before = _everything(engine)
    franks, graces = _ids_of(before, frank_me), _ids_of(before, grace_me)
    assert not franks & graces
    franks_files, graces_files = _files(tmp_path / "assets", before, franks), _files(tmp_path / "assets", before, graces)
    # A picture, a kept original and an export's file each.
    assert len(franks_files) >= 3 and len(graces_files) >= 3
    tables_with_frank = {name for name, rows in before.items() if any(row["id"] in franks for row in rows)}
    assert {"users", "sessions", "account_tokens", "known_browsers", "workspaces", "workspace_members", "documents",
            "document_versions", "document_assets", "processing_jobs", "templates", "usage_records"} <= tables_with_frank

    with caplog.at_level(logging.INFO, logger="app.audit"):
        deleted = _delete(frank)

    assert deleted.status_code == 204
    assert 'smartdoc_session=""' in deleted.headers["set-cookie"]
    after = _everything(engine)
    for table, rows in after.items():
        assert not [row for row in rows if _mentions(row, franks)], f"{table} still holds something of the deleted user"
        assert [row for row in before[table] if row["id"] not in franks] == rows, f"{table} lost or changed another user's rows"
    assert not [path for path in franks_files if path.exists()]
    assert all(path.exists() for path in graces_files)
    [goodbye] = [message for message in sent_mail if message.kind == "account_deleted"]
    assert goodbye.to == "frank@example.com"
    [line] = [record for record in caplog.records if record.getMessage() == "auth.account_deleted"]
    assert not _mentions({key: str(value) for key, value in vars(line).items()}, franks | {"frank"})
    assert frank.get("/api/v1/auth/me").status_code == 401
    assert _client().post("/api/v1/auth/login", json={"email": "frank@example.com", "password": _PASSWORD}).status_code == 401
    assert grace.get("/api/v1/auth/me").status_code == 200


def test_a_wrong_password_deletes_nothing_and_counts_with_sign_ins(two_users, sent_mail, monkeypatch):
    engine, (frank, _), _ = two_users
    before = _everything(engine)
    monkeypatch.setattr(get_settings(), "rate_limit_login_account", "3/minute")  # one used by the fixture's sign-in

    refused = [_delete(frank, "not my password") for _ in range(3)]

    assert (refused[0].status_code, error_body(refused[0])["code"]) == (400, "wrong_password")
    assert [response.status_code for response in refused] == [400, 400, 429]
    assert _everything(engine) == before
    assert not [message for message in sent_mail if message.kind == "account_deleted"]
    assert _delete(_client()).status_code == 401  # signed in only


def test_a_renewing_paid_plan_is_cancelled_first(two_users):
    engine, (frank, frank_me), _ = two_users
    with engine.begin() as connection:
        connection.execute(insert(Subscription).values(id="sub-frank", workspace_id=frank_me["workspaceId"], plan="pro", status="active"))

    refused = _delete(frank)

    assert (refused.status_code, error_body(refused)["code"]) == (409, "subscription_active")
    assert frank.get("/api/v1/auth/me").status_code == 200
    with engine.begin() as connection:  # cancelled at the period's end: it may go now
        connection.execute(update(Subscription).where(Subscription.id == "sub-frank").values(cancel_at_period_end=True))
    assert _delete(frank).status_code == 204


def test_a_workspace_others_are_members_of_needs_another_owner_first(two_users):
    engine, (frank, frank_me), (grace, grace_me) = two_users
    with engine.begin() as connection:
        connection.execute(
            insert(WorkspaceMember).values(id="grace-in-frank", workspace_id=frank_me["workspaceId"], user_id=grace_me["id"], role=WorkspaceRole.MEMBER.value)
        )
    before = _everything(engine)

    refused = _delete(frank)

    assert (refused.status_code, error_body(refused)["code"]) == (409, "workspace_has_members")
    assert _everything(engine) == before
    with engine.begin() as connection:  # Grace owns it too now: Frank may go, the workspace and its documents stay
        connection.execute(update(WorkspaceMember).where(WorkspaceMember.id == "grace-in-frank").values(role=WorkspaceRole.OWNER.value))
    assert _delete(frank).status_code == 204
    with engine.connect() as connection:
        assert connection.scalar(select(User.id).where(User.id == frank_me["id"])) is None
        assert connection.scalar(select(WorkspaceMember.workspace_id).where(WorkspaceMember.id == "grace-in-frank")) == frank_me["workspaceId"]


async def test_a_file_that_cant_be_deleted_is_logged_by_its_key_and_the_rest_go(tmp_path, caplog):
    storage = LocalStorageProvider(tmp_path)
    for key in ("ws/one", "ws/two"):
        await storage.put(key, b"x", "application/octet-stream")

    class Flaky(LocalStorageProvider):
        async def delete(self, key: str) -> None:
            if key == "ws/one":
                raise OSError("storage unreachable")
            await super().delete(key)

    with caplog.at_level(logging.WARNING):
        assert await delete_files(Flaky(tmp_path), ["ws/one", "ws/two"]) == 1

    assert (tmp_path / "ws" / "one").exists() and not (tmp_path / "ws" / "two").exists()
    assert "ws/one" in caplog.text
