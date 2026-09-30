"""Numbered headings (tracker DOCX-016A): Word's numbers for headings stay numbering.
The import keeps the headings' own text and the numbering (Document.headingNumbering);
the editor, the PDF and a Word export number them again, counted over the headings in
order, so they follow when headings move. Measured in Word before this: a03 written
anew had its headings' numbers typed into their text, and a heading edited here showed
two numbers ("2.1 2.1 Scope and aims"): its style's and the typed one."""

import io
import zipfile
from pathlib import Path

from docx import Document as DocxDocument
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from lxml import etree
from pypdf import PdfReader

from app.export.docx_export import build_docx
from app.export.package_check import package_problems
from app.export.pdf_export import build_pdf
from app.export.provenance import stamp
from app.formatting.engine import recompute_styles
from app.formatting.list_numbering import heading_labels
from app.models.document import Document, ElementType, HeadingNumbering, InlineRun, ListLevel
from app.parsers.docx import parse_docx

A03 = Path(__file__).parent / "fixtures" / "word" / "a03-lists.docx"
_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _a03() -> tuple[bytes, Document]:
    source = A03.read_bytes()
    document = parse_docx(source, A03.name)
    recompute_styles(document)
    stamp(document)
    return source, document


def _headings(document: Document) -> list[tuple[int, str]]:
    return [(element.level, element.content) for element in document.elements if element.type == ElementType.HEADING]


def _labels(document: Document) -> list[str | None]:
    headings = [element for element in document.elements if element.type == ElementType.HEADING]
    labels = heading_labels([(element.id, element.level or 1, element.numbered is not False) for element in headings], document.headingNumbering)
    return [labels.get(element.id) for element in headings]


def _saved(word) -> bytes:
    buffer = io.BytesIO()
    word.save(buffer)
    return buffer.getvalue()


def test_word_heading_numbering_stays_numbering():
    _, document = _a03()

    assert _headings(document) == [(1, "Lists and numbering"), (1, "Introduction"), (2, "Scope"), (1, "Method")]
    assert [(level.format, level.text) for level in document.headingNumbering.levels[:3]] == [("decimal", "%1"), ("decimal", "%1.%2"), ("decimal", "%1.%2.%3")]
    assert _labels(document) == ["1", "2", "2.1", "3"]
    assert any(note.startswith("The numbers Word gives the headings are kept as numbering") for note in document.unsupportedFeatures)


def test_the_numbers_follow_when_headings_move():
    _, document = _a03()
    headings = [element for element in document.elements if element.type == ElementType.HEADING]
    method = headings[-1]
    document.elements.remove(method)
    document.elements.insert(document.elements.index(headings[1]), method)  # Method before Introduction now

    assert [label for label in _labels(document)] == ["1", "2", "3", "3.1"]


def test_a_title_isnt_numbered_and_counts_for_nothing():
    word = DocxDocument()
    word.add_paragraph("A report", style="Title")
    for text in ("Introduction", "Method"):
        word.add_paragraph(text, style="Heading 1")
    numbering = word.part.numbering_part.element
    numbering.append(
        parse_xml(
            f'<w:abstractNum {nsdecls("w")} w:abstractNumId="90"><w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/>'
            '<w:lvlText w:val="%1."/></w:lvl></w:abstractNum>'
        )
    )
    numbering.append(parse_xml(f'<w:num {nsdecls("w")} w:numId="90"><w:abstractNumId w:val="90"/></w:num>'))
    style = word.styles["Heading 1"].element
    style.get_or_add_pPr().append(parse_xml(f'<w:numPr {nsdecls("w")}><w:numId w:val="90"/></w:numPr>'))

    document = parse_docx(_saved(word), "report.docx")

    assert _headings(document) == [(1, "A report"), (1, "Introduction"), (1, "Method")]
    assert [element.numbered for element in document.elements if element.type == ElementType.HEADING] == [False, None, None]
    assert _labels(document) == [None, "1.", "2."]


def _numbers_in_word(data: bytes) -> list[tuple[str | None, str]]:
    """Each heading of the exported file, read again: the number its numbering gives it, and its text."""
    again = parse_docx(data, "again.docx")
    return list(zip(_labels(again), [text for _, text in _headings(again)], strict=True))


def test_a_new_word_file_numbers_the_headings_again():
    _, document = _a03()

    exported = build_docx(document)

    assert _numbers_in_word(exported) == [("1", "Lists and numbering"), ("2", "Introduction"), ("2.1", "Scope"), ("3", "Method")]
    with zipfile.ZipFile(io.BytesIO(exported)) as package:
        styles, body = etree.fromstring(package.read("word/styles.xml")), etree.fromstring(package.read("word/document.xml"))
    used = body.find(f".//{_W}p/{_W}pPr/{_W}numPr/{_W}numId").get(f"{_W}val")  # the first heading's numbering
    for style_id in ("Heading1", "Heading2"):  # the styles number with it: a heading made in Word later is numbered too
        assert styles.find(f"{_W}style[@{_W}styleId='{style_id}']/{_W}pPr/{_W}numPr/{_W}numId").get(f"{_W}val") == used
    assert package_problems(exported) == []


def test_a_heading_edited_here_gets_one_number_into_the_original():
    source, document = _a03()
    scope = next(element for element in document.elements if element.content == "Scope")
    scope.inline = [*(scope.inline or []), InlineRun(text=" and aims")]
    scope.content = "Scope and aims"

    exported = build_docx(document, source=source)

    assert _numbers_in_word(exported) == [("1", "Lists and numbering"), ("2", "Introduction"), ("2.1", "Scope and aims"), ("3", "Method")]
    body = zipfile.ZipFile(io.BytesIO(exported)).read("word/document.xml").decode("utf-8")
    assert f'<w:numId w:val="{document.headingNumbering.sourceNumId}"/>' in body  # the original's numbering: it counts on with the others
    assert "2.1 Scope" not in body
    assert package_problems(exported) == []


def test_a_pdf_prints_the_numbers():
    _, document = _a03()

    text = "\n".join(page.extract_text() for page in PdfReader(io.BytesIO(build_pdf(document))).pages)

    lines = text.splitlines()
    assert all(heading in lines for heading in ("1 Lists and numbering", "2 Introduction", "2.1 Scope", "3 Method"))


def test_a_level_numbers_in_its_own_format_and_label():
    level = ListLevel(format="upperRoman", text="Глава %1", start=3)
    numbering = HeadingNumbering(levels=[level])

    labels = heading_labels([("a", 1, True), ("b", 1, True), ("c", 2, True), ("d", 1, False)], numbering)

    assert labels == {"a": "Глава III", "b": "Глава IV"}  # level 2 isn't numbered; d not at all


def test_headings_numbered_by_two_lists_keep_their_numbers_in_their_text():
    """Not one numbering: as before, each heading's number is written into its text --
    and a Word export into the original never numbers it a second time."""
    word = DocxDocument()
    numbering = word.part.numbering_part.element
    for num in ("91", "92"):
        numbering.append(
            parse_xml(
                f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{num}"><w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/>'
                '<w:lvlText w:val="%1."/></w:lvl></w:abstractNum>'
            )
        )
        numbering.append(parse_xml(f'<w:num {nsdecls("w")} w:numId="{num}"><w:abstractNumId w:val="{num}"/></w:num>'))
    for text, num in (("Introduction", "91"), ("Appendix", "92")):
        paragraph = word.add_paragraph(text, style="Heading 1")
        paragraph._p.get_or_add_pPr().append(parse_xml(f'<w:numPr {nsdecls("w")}><w:ilvl w:val="0"/><w:numId w:val="{num}"/></w:numPr>'))
    source = _saved(word)
    document = parse_docx(source, "two.docx")
    recompute_styles(document)
    stamp(document)
    assert document.headingNumbering is None
    assert _headings(document) == [(1, "1. Introduction"), (1, "1. Appendix")]
    appendix = document.elements[-1]
    appendix.inline = [*(appendix.inline or []), InlineRun(text=" A")]
    appendix.content = "1. Appendix A"

    exported = build_docx(document, source=source)

    body = zipfile.ZipFile(io.BytesIO(exported)).read("word/document.xml").decode("utf-8")
    assert '<w:numId w:val="0"/>' in body  # written anew: its number is in its text, not a second one from its list
