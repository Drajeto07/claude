"""Monthly usage reserved before the work that uses it (tracker PLAN-003, PLAN-001):
an AI call's operation and an export are taken under the workspace's hold in a short
transaction of their own, committed before the call or the export is made, and
given back when it doesn't happen -- consistent with the jobs' retries, timeouts
and cancels (JOB-001), so a job counts once however often it runs. A PDF's pages
are counted with the document it becomes. The race itself (two jobs, room for one
call) is in tests/test_plan_limits_atomic.py."""

import asyncio
import io

import pytest
import pytest_asyncio
from fastapi.testclient import TestClient
from reportlab.pdfgen import canvas
from sqlalchemy import create_engine, event, func, select
from sqlalchemy.ext.asyncio import async_sessionmaker
from sqlalchemy.orm import Session as OrmSession
from sqlalchemy.pool import NullPool

from app.ai.base import AIProvider
from app.ai.budget import AIBudgetExceededError, BudgetedAIProvider
from app.ai.factory import get_ai_provider
from app.ai.schemas import AIInstructionExtractionResponse
from app.billing import units
from app.billing.plans import FREE, PLANS
from app.db.base import Base
from app.db.models import Document as DocumentRow
from app.db.models import JobStatus, ProcessingJob, UsageRecord
from app.db.session import make_engine
from app.jobs import runner as runner_module
from app.jobs.runner import EXPORT, IMPORT_TEXT, JobRunner
from app.main import app
from app.services.auth_service import AuthService
from app.services.entitlements_service import AILimitReachedError, EntitlementsService, UsageReservations, metered
from app.storage.local_provider import LocalStorageProvider
from tests.conftest import _enable_sqlite_fk
from tests.fakes import FakeAIProvider

client = TestClient(app, base_url="https://testserver")
_ANSWER = AIInstructionExtractionResponse(rules=[])


def _free_plan(monkeypatch, **limits) -> None:
    free = PLANS[FREE]
    monkeypatch.setitem(PLANS, FREE, free.model_copy(update={"entitlements": free.entitlements.model_copy(update=limits)}))


@pytest_asyncio.fixture
async def workspace(tmp_path):
    """Sessions on a SQLite file, a new connection each (as in production, where a
    reservation's session is not the job's), and a signed-up user's workspace."""
    path = tmp_path / "reservations.db"
    sync_engine = create_engine(f"sqlite:///{path}")
    Base.metadata.create_all(sync_engine)
    sync_engine.dispose()
    engine = make_engine(f"sqlite+aiosqlite:///{path}", poolclass=NullPool)
    event.listen(engine.sync_engine, "connect", _enable_sqlite_fk)
    sessions = async_sessionmaker(engine, expire_on_commit=False)
    async with sessions() as session:
        user = await AuthService(session).register("reserve@example.com", "long enough password", None)
        workspace_id = await AuthService(session).default_workspace_id(user.id)
        await session.commit()
    yield sessions, workspace_id, user.id
    await engine.dispose()


async def _counted(sessions, workspace_id: str, metric: str = units.AI_OPERATIONS) -> int:
    async with sessions() as session:
        return await EntitlementsService(session).used_this_month(workspace_id, metric)


async def _job(sessions, workspace_id: str, user_id: str, **fields) -> str:
    async with sessions() as session:
        job = ProcessingJob(**{"workspace_id": workspace_id, "created_by": user_id, "job_type": IMPORT_TEXT, "payload": {"text": "x"}} | fields)
        session.add(job)
        await session.commit()
        return job.id


async def _status(sessions, job_id: str) -> str:
    async with sessions() as session:
        return (await session.get(ProcessingJob, job_id)).status


def _one_structured_call(seen: list | None = None):
    """A job kind making one AI call, noting what it gave back."""

    async def kind(ctx):
        try:
            await ctx.provider.complete_structured("Format this.", response_model=AIInstructionExtractionResponse)
        except AILimitReachedError:
            if seen is not None:
                seen.append("refused")
        return {}

    return kind


class _Hanging(AIProvider):
    """A call that doesn't come back until it is cut off."""

    def __init__(self) -> None:
        self.started = asyncio.Event()

    def provider_name(self) -> str:
        return "hanging"

    async def complete(self, prompt: str, *, max_tokens: int = 256, system: str | None = None) -> str:
        self.started.set()
        await asyncio.sleep(3600)
        return ""

    async def complete_structured(self, prompt, *, response_model, max_tokens=8192, system=None):
        self.started.set()
        await asyncio.sleep(3600)
        raise AssertionError("never answers")


# -- the reservation itself ------------------------------------------------------------


async def test_an_ai_call_is_counted_before_it_is_made_and_committed_on_its_own(workspace):
    sessions, workspace_id, _ = workspace
    seen_during_the_call: list[int] = []

    class Looking(FakeAIProvider):
        async def complete_structured(self, prompt, *, response_model, max_tokens=8192, system=None):
            # From another connection: the reservation is already committed.
            seen_during_the_call.append(await _counted(sessions, workspace_id))
            return await super().complete_structured(prompt, response_model=response_model, max_tokens=max_tokens, system=system)

    provider = metered(Looking([_ANSWER]), UsageReservations(sessions, workspace_id))
    await provider.complete_structured("Format this.", response_model=AIInstructionExtractionResponse)

    assert seen_during_the_call == [1] and await _counted(sessions, workspace_id) == 1


async def test_a_call_that_fails_gives_its_operation_back(workspace):
    sessions, workspace_id, _ = workspace
    provider = metered(FakeAIProvider([ConnectionError("down"), _ANSWER]), UsageReservations(sessions, workspace_id))

    with pytest.raises(ConnectionError):
        await provider.complete_structured("Format this.", response_model=AIInstructionExtractionResponse)
    assert await _counted(sessions, workspace_id) == 0
    await provider.complete_structured("Format this.", response_model=AIInstructionExtractionResponse)
    assert await _counted(sessions, workspace_id) == 1


async def test_a_call_the_ai_budget_cuts_off_is_given_back(workspace):
    sessions, workspace_id, _ = workspace
    provider = BudgetedAIProvider(metered(_Hanging(), UsageReservations(sessions, workspace_id)), calls=5, seconds=0.2)

    with pytest.raises(AIBudgetExceededError):
        await provider.complete("Which are headings?")

    assert await _counted(sessions, workspace_id) == 0


async def test_used_up_operations_refuse_the_call_before_it_is_made(monkeypatch, workspace):
    sessions, workspace_id, _ = workspace
    _free_plan(monkeypatch, maxAiOperations=1)
    fake = FakeAIProvider([_ANSWER])
    provider = metered(fake, UsageReservations(sessions, workspace_id))
    await provider.complete_structured("first", response_model=AIInstructionExtractionResponse)

    with pytest.raises(AILimitReachedError):
        await provider.complete_structured("second", response_model=AIInstructionExtractionResponse)
    assert fake.calls == 1 and await _counted(sessions, workspace_id) == 1


async def test_a_reservation_that_cant_be_given_back_stays_counted(workspace, caplog):
    sessions, workspace_id, _ = workspace
    reservations = UsageReservations(sessions, workspace_id)
    reservation = await reservations.take(units.AI)

    def broken():
        raise ConnectionError("database gone")

    await UsageReservations(broken, workspace_id).give_back(reservation)

    assert await _counted(sessions, workspace_id) == 1 and "Could not give back" in caplog.text


# -- with the jobs' retries, timeouts and cancels (JOB-001) ------------------------------------


async def test_a_retried_job_whose_call_failed_transiently_counts_its_call_once(monkeypatch, workspace, tmp_path):
    sessions, workspace_id, user_id = workspace
    monkeypatch.setitem(runner_module.KINDS, IMPORT_TEXT, _one_structured_call())
    fake = FakeAIProvider([ConnectionError("AI unreachable"), _ANSWER])
    runner = JobRunner(sessions, LocalStorageProvider(tmp_path), fake)
    job_id = await _job(sessions, workspace_id, user_id)

    wait = await runner.run(job_id)  # reserved, failed transiently, given back
    assert wait is not None and await _counted(sessions, workspace_id) == 0
    assert await runner.run(job_id) is None  # the retry: reserved, made

    assert fake.calls == 2 and await _status(sessions, job_id) == JobStatus.SUCCEEDED.value
    assert await _counted(sessions, workspace_id) == 1


async def test_a_call_cut_off_by_the_jobs_timeout_is_not_counted(monkeypatch, workspace, tmp_path):
    sessions, workspace_id, user_id = workspace
    monkeypatch.setitem(runner_module.KINDS, IMPORT_TEXT, _one_structured_call())
    monkeypatch.setattr(runner_module, "timeout_for", lambda job_type: 0.3)
    job_id = await _job(sessions, workspace_id, user_id)

    await JobRunner(sessions, LocalStorageProvider(tmp_path), _Hanging()).run(job_id)

    assert await _status(sessions, job_id) == JobStatus.FAILED.value
    assert await _counted(sessions, workspace_id) == 0


async def test_a_job_cancelled_during_its_call_keeps_the_completed_call_counted(monkeypatch, workspace, tmp_path):
    # A cancel stops a job at its next check, so a call under way completes: it was
    # made (and paid for), so it counts, and the job stops right after it.
    sessions, workspace_id, user_id = workspace
    job_id = await _job(sessions, workspace_id, user_id)

    class CancelledMeanwhile(FakeAIProvider):
        async def complete_structured(self, prompt, *, response_model, max_tokens=8192, system=None):
            async with sessions() as session:
                (await session.get(ProcessingJob, job_id)).status = JobStatus.CANCELLED.value
                await session.commit()
            return await super().complete_structured(prompt, response_model=response_model, max_tokens=max_tokens, system=system)

    async def kind(ctx):
        await _one_structured_call()(ctx)
        await ctx.report("formatting", 50)
        raise AssertionError("the cancel is noticed at the check")

    monkeypatch.setitem(runner_module.KINDS, IMPORT_TEXT, kind)
    await JobRunner(sessions, LocalStorageProvider(tmp_path), CancelledMeanwhile([_ANSWER])).run(job_id)

    assert await _status(sessions, job_id) == JobStatus.CANCELLED.value
    assert await _counted(sessions, workspace_id) == 1


async def test_a_refused_call_lets_the_job_go_on_without_the_ai(monkeypatch, workspace, tmp_path):
    sessions, workspace_id, user_id = workspace
    _free_plan(monkeypatch, maxAiOperations=0)
    refused: list[str] = []
    monkeypatch.setitem(runner_module.KINDS, IMPORT_TEXT, _one_structured_call(refused))
    fake = FakeAIProvider([])
    job_id = await _job(sessions, workspace_id, user_id)

    await JobRunner(sessions, LocalStorageProvider(tmp_path), fake).run(job_id)

    assert (refused, fake.calls, await _status(sessions, job_id)) == (["refused"], 0, JobStatus.SUCCEEDED.value)


# -- exports --------------------------------------------------------------------------------


async def _document(sessions, workspace_id: str, user_id: str) -> str:
    from app.models.document import Document, DocumentMetadata, Element, ElementType
    from app.repositories.document_repository import DocumentRepository

    document = Document(metadata=DocumentMetadata(title="Exported"), elements=[Element(type=ElementType.PARAGRAPH, content="Words.", order=0)])
    async with sessions() as session:
        row = await DocumentRepository(session).create(workspace_id, document, created_by=user_id)
        await session.commit()
        return row.id


async def test_an_export_job_that_is_retried_counts_one_export(monkeypatch, workspace, tmp_path):
    sessions, workspace_id, user_id = workspace
    storage = LocalStorageProvider(tmp_path / "files")
    real_put = storage.put
    failures = [ConnectionError("storage unreachable")]

    async def flaky_put(key, data, content_type):
        if failures:
            raise failures.pop()
        await real_put(key, data, content_type)

    monkeypatch.setattr(storage, "put", flaky_put)
    document_id = await _document(sessions, workspace_id, user_id)
    job_id = await _job(sessions, workspace_id, user_id, job_type=EXPORT, document_id=document_id, payload={"format": "docx"})
    runner = JobRunner(sessions, storage, FakeAIProvider([]))

    assert await runner.run(job_id) is not None  # failed transiently: its export given back
    assert await _counted(sessions, workspace_id, units.EXPORTS) == 0
    await runner.run(job_id)

    assert await _status(sessions, job_id) == JobStatus.SUCCEEDED.value
    assert await _counted(sessions, workspace_id, units.EXPORTS) == 1


async def test_an_export_job_the_month_has_no_room_for_makes_nothing(monkeypatch, workspace, tmp_path):
    sessions, workspace_id, user_id = workspace
    _free_plan(monkeypatch, maxExports=0)
    document_id = await _document(sessions, workspace_id, user_id)
    job_id = await _job(sessions, workspace_id, user_id, job_type=EXPORT, document_id=document_id, payload={"format": "pdf"})

    await JobRunner(sessions, LocalStorageProvider(tmp_path), FakeAIProvider([])).run(job_id)

    async with sessions() as session:
        job = await session.get(ProcessingJob, job_id)
    assert (job.status, job.result, job.failure_reason) == (JobStatus.FAILED.value, None, "user:PlanLimitError")
    assert "exports" in job.error_message and await _counted(sessions, workspace_id, units.EXPORTS) == 0


@pytest.fixture
def signed_in(api_db):
    client.cookies.clear()
    assert client.post("/api/v1/auth/register", json={"email": "units@example.com", "password": "long enough password"}).status_code == 201
    yield api_db
    client.cookies.clear()
    app.dependency_overrides.pop(get_ai_provider, None)


def _exports(api_db) -> int:
    with OrmSession(api_db) as session:
        return session.scalar(select(func.coalesce(func.sum(UsageRecord.quantity), 0)).where(UsageRecord.metric == units.EXPORTS))


def test_the_months_exports_are_a_limit_on_every_way_to_export(monkeypatch, signed_in):
    _free_plan(monkeypatch, maxExports=2)
    document_id = client.post("/api/v1/documents", json={"text": "# Exports\n\nA paragraph."}).json()["id"]

    assert client.get(f"/api/v1/documents/{document_id}/export/docx").status_code == 200
    assert client.post("/api/v1/jobs/export", json={"documentId": document_id, "format": "pdf"}).json()["status"] == "succeeded"
    for refused in (
        client.get(f"/api/v1/documents/{document_id}/export/pdf"),
        client.post("/api/v1/jobs/export", json={"documentId": document_id, "format": "docx"}),
    ):
        assert (refused.status_code, refused.json()["code"], refused.json()["details"]["entitlement"]) == (402, "plan_limit", "maxExports")
    assert _exports(signed_in) == 2


def test_a_direct_export_that_fails_to_build_is_given_back(monkeypatch, signed_in):
    from app.api import documents as documents_api

    document_id = client.post("/api/v1/documents", json={"text": "# Exports\n\nA paragraph."}).json()["id"]

    def broken(*args, **kwargs):
        raise RuntimeError("renderer crashed")

    monkeypatch.setattr(documents_api, "build_pdf", broken)
    failing = TestClient(app, base_url="https://testserver", raise_server_exceptions=False, cookies=client.cookies)
    assert failing.get(f"/api/v1/documents/{document_id}/export/pdf").status_code == 500

    assert _exports(signed_in) == 0


# -- PDF pages, counted at import --------------------------------------------------------------


def _pdf(pages: int) -> bytes:
    buffer = io.BytesIO()
    document = canvas.Canvas(buffer, invariant=1)
    for number in range(pages):
        document.drawString(72, 720, f"Page {number + 1} of a short PDF.")
        document.showPage()
    document.save()
    return buffer.getvalue()


def _pdf_pages(api_db) -> int:
    with OrmSession(api_db) as session:
        return session.scalar(select(func.coalesce(func.sum(UsageRecord.quantity), 0)).where(UsageRecord.metric == units.PDF_PAGES))


def test_a_pdf_import_counts_its_pages_with_the_document(signed_in):
    upload = client.post("/api/v1/documents/upload", files={"file": ("three.pdf", _pdf(3), "application/pdf")})
    job = client.post("/api/v1/jobs/import-file", files={"file": ("two.pdf", _pdf(2), "application/pdf")}).json()

    assert upload.status_code == 201 and job["status"] == "succeeded"
    assert _pdf_pages(signed_in) == 5
    pdf_unit = next(unit for unit in client.get("/api/v1/usage").json()["units"] if unit["key"] == "pdfPages")
    assert (pdf_unit["used"], pdf_unit["limit"]) == (5, PLANS[FREE].entitlements.maxPdfPages)


def test_a_pdf_with_more_pages_than_the_month_has_left_is_refused_before_it_is_read(monkeypatch, signed_in):
    _free_plan(monkeypatch, maxPdfPages=4)
    assert client.post("/api/v1/documents/upload", files={"file": ("three.pdf", _pdf(3), "application/pdf")}).status_code == 201
    fake = FakeAIProvider([])  # nothing may be asked about a refused file
    app.dependency_overrides[get_ai_provider] = lambda: fake

    refused = client.post("/api/v1/documents/upload", files={"file": ("two.pdf", _pdf(2), "application/pdf")})
    job = client.post("/api/v1/jobs/import-file", files={"file": ("two.pdf", _pdf(2), "application/pdf")}).json()

    assert (refused.status_code, refused.json()["details"]) == (402, {"entitlement": "maxPdfPages", "limit": 4, "used": 3})
    assert "1 of its 4 left" in refused.json()["message"]
    assert (job["status"], "PDF pages" in job["error"]) == ("failed", True)
    assert _pdf_pages(signed_in) == 3 and fake.calls == 0
    with OrmSession(signed_in) as session:
        assert session.scalar(select(func.count(DocumentRow.id))) == 1


async def test_a_documents_pdf_pages_are_checked_again_under_its_hold(monkeypatch, workspace, tmp_path):
    # What another import counted after this one's early check: the check made with
    # the document, under the hold it takes for the document count, still refuses.
    from app.models.document import Document, DocumentMetadata
    from app.services.document_service import DocumentService
    from app.services.entitlements_service import PlanLimitError
    from app.services.usage_service import usage_row

    sessions, workspace_id, user_id = workspace
    _free_plan(monkeypatch, maxPdfPages=4)
    async with sessions() as session:
        session.add(usage_row(workspace_id, units.PDF_PAGES, 2))
        await session.commit()

    async with sessions() as session:
        service = DocumentService(session, user_id=user_id, storage=LocalStorageProvider(tmp_path))
        with pytest.raises(PlanLimitError):
            await service.create(Document(metadata=DocumentMetadata(title="Three pages"), elements=[]), pdf_pages=3)
        await service.create(Document(metadata=DocumentMetadata(title="Two pages"), elements=[]), pdf_pages=2)

    assert await _counted(sessions, workspace_id, units.PDF_PAGES) == 4
    async with sessions() as session:
        assert (await session.execute(select(DocumentRow.title))).scalars().all() == ["Two pages"]


def test_a_pdf_that_fills_the_month_exactly_is_taken_and_word_files_count_no_pages(monkeypatch, signed_in):
    from docx import Document as DocxDocument

    _free_plan(monkeypatch, maxPdfPages=2)
    word = io.BytesIO()
    document = DocxDocument()
    document.add_paragraph("A Word paragraph.")
    document.save(word)
    docx_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

    assert client.post("/api/v1/documents/upload", files={"file": ("two.pdf", _pdf(2), "application/pdf")}).status_code == 201
    assert client.post("/api/v1/documents/upload", files={"file": ("word.docx", word.getvalue(), docx_type)}).status_code == 201
    assert _pdf_pages(signed_in) == 2
