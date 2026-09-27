"""Word's character formatting beyond bold, italic and plain lines (tracker
DOCX-013): underline styles, a double strikethrough, capitals, small capitals,
character spacing and raised or lowered text are read as Word resolves them
(the run, its character style, its paragraph's style, the defaults), kept in the
model, written back into Word in the schema's order, and drawn in a PDF as far as
it can, which names what it can't."""

import io
import zipfile

import pytest
from docx import Document as DocxDocument
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from fastapi.testclient import TestClient
from pydantic import ValidationError
from pypdf import PdfReader

from app.export.docx_export import _RPR_ORDER, build_docx
from app.export.package_check import package_problems
from app.export.pdf_export import build_pdf
from app.fidelity.docx_detect import detect_docx_features
from app.fidelity.exports import export_report
from app.fidelity.report import ReportBuilder
from app.main import app
from app.models.document import Mark, MarkType
from app.parsers.docx import parse_docx

client = TestClient(app, base_url="https://testserver")
_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_W = nsdecls("w")


def _word_file() -> bytes:
    word = DocxDocument()
    label = word.styles.add_style("Label", WD_STYLE_TYPE.CHARACTER)
    label.element.get_or_add_rPr().append(parse_xml(f"<w:smallCaps {_W}/>"))
    ruled = word.styles.add_style("Ruled", WD_STYLE_TYPE.PARAGRAPH)
    ruled.element.get_or_add_rPr().append(parse_xml(f'<w:u {_W} w:val="double"/>'))
    paragraph = word.add_paragraph("Plain ")
    for properties, text in (
        ('<w:u w:val="double"/>', "double"),
        ('<w:u w:val="dotDash"/>', "dotdash"),
        ("<w:dstrike/>", "twice"),
        ("<w:caps/>", "shouting"),
        ('<w:spacing w:val="40"/>', "spaced"),
        ('<w:position w:val="-4"/>', "lowered"),
        ('<w:caps/><w:smallCaps/>', "both"),  # all caps wins, as in Word
    ):
        paragraph._p.append(parse_xml(f'<w:r {_W}><w:rPr>{properties}</w:rPr><w:t xml:space="preserve">{text} </w:t></w:r>'))
    paragraph.add_run("quiet label", style="Label")
    word.add_paragraph("A ruled paragraph.", style="Ruled")
    shown = word.add_paragraph("Ruled, but not this", style="Ruled")
    shown.runs[0]._r.get_or_add_rPr().append(parse_xml(f'<w:u {_W} w:val="none"/>'))
    buffer = io.BytesIO()
    word.save(buffer)
    return buffer.getvalue()


def _marks(document) -> dict[str, dict[str, object]]:
    return {
        run.text.strip(): {
            f"{mark.type.value}.{key}" if key != "type" else mark.type.value: value
            for mark in run.marks
            for key, value in mark.model_dump(exclude_none=True).items()
        }
        for element in document.elements
        for run in element.inline or []
        if run.marks
    }


def test_run_formatting_is_read_as_word_resolves_it():
    marks = _marks(parse_docx(_word_file(), "chars.docx"))

    assert marks["double"] == {"underline": MarkType.UNDERLINE, "underline.lineStyle": "double"}
    assert marks["dotdash"]["underline.lineStyle"] == "dashed"  # the closest style the app has
    assert marks["twice"]["strike.lineStyle"] == "double"
    assert marks["shouting"]["textStyle.caps"] is True
    assert marks["spaced"]["textStyle.letterSpacingPt"] == 2.0
    assert marks["lowered"]["textStyle.baselineShiftPt"] == -2.0
    assert marks["both"]["textStyle.caps"] is True and "textStyle.smallCaps" not in marks["both"]
    assert marks["quiet label"]["textStyle.smallCaps"] is True  # from its character style
    assert marks["A ruled paragraph."]["underline.lineStyle"] == "double"  # from its paragraph's style
    assert "Ruled, but not this" not in marks  # the run turns the style's underline off


def test_only_the_approximated_underline_is_reported():
    items = {item.feature: item for item in detect_docx_features(_word_file())}

    assert items["docx.underline_variant"].count == 1 and "dotdash" in (items["docx.underline_variant"].sourceState or "")
    assert "docx.caps" not in items


def test_a_word_export_writes_them_back_in_the_schemas_order():
    document = parse_docx(_word_file(), "chars.docx")

    exported = build_docx(document)

    assert package_problems(exported) == []
    again = parse_docx(exported, "again.docx")
    assert [[(run.text, run.marks) for run in element.inline or []] for element in again.elements] == [
        [(run.text, run.marks) for run in element.inline or []] for element in document.elements
    ]
    with zipfile.ZipFile(io.BytesIO(exported)) as package:
        body = parse_xml(package.read("word/document.xml"))
    for properties in body.iter(qn("w:rPr")):
        positions = [_RPR_ORDER.index(child.tag) for child in properties if child.tag in _RPR_ORDER]
        assert positions == sorted(positions), [child.tag for child in properties]


def test_a_highlight_with_everything_else_stays_in_order():
    document = parse_docx(_word_file(), "chars.docx")
    run = document.elements[0].inline[1]  # "double "
    run.marks.append(Mark(type=MarkType.TEXT_STYLE, backgroundColor="#FFFF00", caps=True, letterSpacingPt=1, baselineShiftPt=2))

    with zipfile.ZipFile(io.BytesIO(build_docx(document))) as package:
        body = package.read("word/document.xml").decode("utf-8")

    start = body.index("<w:rPr>", body.index("double") - 400)
    properties = body[start : body.index("</w:rPr>", start)]
    order = [properties.index(tag) for tag in ("<w:caps/>", "<w:spacing", "<w:position", "<w:highlight", "<w:u ")]
    assert order == sorted(order)


def test_a_pdf_prints_capitals_and_notes_what_it_cant_draw():
    document = parse_docx(_word_file(), "chars.docx")
    report = ReportBuilder()

    pdf = build_pdf(document, report=report)

    text = " ".join(page.extract_text() for page in PdfReader(io.BytesIO(pdf)).pages)
    assert "SHOUTING" in text and "QUIET LABEL" in text and "shouting" not in text
    features = {item.feature for item in report.items()}
    assert {"export.pdf.character_spacing", "export.pdf.underline_style"} <= features
    assert export_report(document, pdf, "pdf", report.items()).content.verified  # capitals counted as printed


@pytest.mark.parametrize(
    "mark",
    [
        {"type": "bold", "lineStyle": "double"},
        {"type": "strike", "lineStyle": "wavy"},
        {"type": "textStyle", "letterSpacingPt": 500},
        {"type": "textStyle", "baselineShiftPt": -101},
    ],
)
def test_the_model_refuses_what_it_cant_mean(mark):
    with pytest.raises(ValidationError):
        Mark.model_validate(mark)


def test_unset_and_zero_mean_the_same():
    mark = Mark.model_validate({"type": "textStyle", "caps": False, "smallCaps": False, "letterSpacingPt": 0, "baselineShiftPt": 0})

    assert (mark.caps, mark.smallCaps, mark.letterSpacingPt, mark.baselineShiftPt) == (None, None, None, None)


def test_an_editor_save_keeps_them(api_db):
    client.cookies.clear()
    assert client.post("/api/v1/auth/register", json={"email": "chars@example.com", "password": "long enough password"}).status_code == 201
    document = client.post("/api/v1/documents/upload", files={"file": ("chars.docx", _word_file(), _DOCX)}).json()

    saved = client.put(f"/api/v1/documents/{document['id']}/content", json={"elements": document["elements"]})
    client.cookies.clear()

    assert saved.status_code == 200
    runs = {run["text"].strip(): run["marks"] for run in saved.json()["elements"][0]["inline"]}
    assert runs["double"][0]["lineStyle"] == "double" and runs["spaced"][0]["letterSpacingPt"] == 2.0
    assert document["importReport"]["content"]["verified"]


def test_a_run_that_turns_off_its_styles_bold_leaves_the_rest_bold():
    from app.formatting.engine import recompute_styles

    word = DocxDocument()
    heading = word.add_heading("Mostly bold, ", level=1)  # Heading 1 is bold by its style
    heading.add_run("except this").bold = False
    buffer = io.BytesIO()
    word.save(buffer)

    document = parse_docx(buffer.getvalue(), "heading.docx")
    recompute_styles(document)

    element = document.elements[0]
    assert [(run.text, [mark.type for mark in run.marks]) for run in element.inline] == [("Mostly bold, ", [MarkType.BOLD]), ("except this", [])]
    assert document.resolvedStyles[element.styleRef].get("font-weight") != "bold"  # the heading's look no longer claims it
    again = parse_docx(build_docx(document), "again.docx")
    assert [(run.text, [mark.type for mark in run.marks]) for run in again.elements[0].inline] == [("Mostly bold, ", [MarkType.BOLD]), ("except this", [])]
