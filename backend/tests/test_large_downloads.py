"""B-035B: files larger than one hosted response leave the API as a stream (security/serving.py).

On Vercel a function's response body stops at 4.5 MB (500 FUNCTION_RESPONSE_PAYLOAD_TOO_LARGE); a streamed one is
let through (Vercel's docs). Here, with deterministic synthetic files: a picture of 4.9 and 8.8 MB, a DOCX and a PDF
over 4.5 MB (the PDF about 11 MB), an export job's file, a batch ZIP, a Doctor download and the PDF tools' results
come out byte for byte (SHA-256) and as valid packages, with their types, names and privacy headers, streamed in
parts without a Content-Length, while small files leave in one body as before. Access is checked before anything is
sent; a storage error is a plain 500, never a short file; a client that goes away mid-download leaves the stored
file as it was. What this can't show: that Vercel's container runtime and proxy pass such a stream on -- that is
the owner's Preview test.
"""

import asyncio
import base64
import hashlib
import io
import logging
import random
import socket
import threading
import time
import zipfile

import pytest
import uvicorn
from docx import Document as DocxDocument
from docx.shared import Cm
from fastapi import FastAPI
from fastapi.testclient import TestClient
from PIL import Image
from pypdf import PdfReader

from app.billing.plans import FREE, PLANS
from app.main import app
from app.security import serving
from app.security.serving import STREAM_ABOVE, STREAM_PART, file_response
from app.storage.factory import get_storage_provider
from app.storage.local_provider import LocalStorageProvider

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PASSWORD = "long enough password"
LIMIT = 4_500_000  # Vercel's response body limit


def _client(email: str) -> TestClient:
    client = TestClient(app, base_url="https://testserver")
    assert client.post("/api/v1/auth/register", json={"email": email, "password": PASSWORD}).status_code == 201
    return client


@pytest.fixture
def alice(api_db) -> TestClient:
    return _client("alice-downloads@example.com")


@pytest.fixture
def bob(api_db) -> TestClient:
    return _client("bob-downloads@example.com")


def _plan(monkeypatch, **limits) -> None:
    free = PLANS[FREE]
    monkeypatch.setitem(PLANS, FREE, free.model_copy(update={"entitlements": free.entitlements.model_copy(update=limits)}))


def _noise_png(size: int, seed: int) -> bytes:
    """A PNG of about `size` bytes that doesn't compress, the same every time."""
    side = int((size / 3) ** 0.5)
    buffer = io.BytesIO()
    Image.frombytes("RGB", (side, side), random.Random(seed).randbytes(side * side * 3)).save(buffer, "PNG", compress_level=0)
    return buffer.getvalue()


def _docx(picture_mb: float, seed: int, title: str = "Отчет за 2026") -> tuple[bytes, bytes]:
    document = DocxDocument()
    document.core_properties.title = title
    document.add_paragraph(f"{title}: revenue 1,250 EUR.")
    picture = _noise_png(int(picture_mb * 1024 * 1024), seed)
    document.add_picture(io.BytesIO(picture), width=Cm(12))
    document.add_paragraph("Text after the picture.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue(), picture


def _upload(client: TestClient, data: bytes, name: str = "report.docx") -> dict:
    response = client.post("/api/v1/documents/upload", files={"file": (name, data, DOCX)})
    assert response.status_code in (200, 201), response.text
    return response.json()


def _image_id(document: dict) -> str:
    return next(element for element in document["elements"] if element["type"] == "image")["image"]["assetId"]


def _sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _streamed(response) -> None:
    """Sent in parts: no Content-Length (chunked on the wire), the API's privacy headers as on any file."""
    assert "content-length" not in response.headers, response.headers
    assert response.headers["x-content-type-options"] == "nosniff"
    assert response.headers["content-security-policy"] == "default-src 'none'; sandbox"
    assert response.headers["cache-control"].startswith("private")


def _media(docx: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(io.BytesIO(docx)) as package:
        assert package.testzip() is None
        return {name: package.read(name) for name in package.namelist() if name.startswith("word/media/")}


@pytest.fixture
def large(alice) -> dict:
    """One document with an 8.8 MB picture, uploaded whole (the test server takes 10 MB in one request)."""
    data, picture = _docx(8.8, seed=88)
    document = _upload(alice, data)
    return {"id": document["id"], "document": document, "picture": picture, "source": data}


# -- file_response itself ------------------------------------------------------------------------------------


def test_a_small_file_leaves_in_one_body_as_before_and_a_large_one_in_parts_with_the_same_bytes_and_headers():
    small = random.Random(1).randbytes(STREAM_ABOVE)
    big = random.Random(2).randbytes(STREAM_ABOVE + 1)
    whole = file_response(small, "application/pdf", filename="Отчет.pdf")
    assert type(whole).__name__ == "Response" and whole.body == small

    streamed = file_response(big, "application/pdf", filename="Отчет.pdf")
    assert type(streamed).__name__ == "StreamingResponse"
    assert {name: value for name, value in streamed.headers.items() if name != "content-length"} == {
        name: value for name, value in whole.headers.items() if name != "content-length"
    }
    assert "content-length" not in streamed.headers

    async def collect() -> list[bytes]:
        return [part async for part in streamed.body_iterator]

    parts = asyncio.run(collect())
    assert all(len(part) <= STREAM_PART for part in parts) and len(parts) == -(-len(big) // STREAM_PART)
    assert b"".join(parts) == big
    assert STREAM_ABOVE < LIMIT


# -- Stored pictures ------------------------------------------------------------------------------------------


@pytest.mark.parametrize("picture_mb", [4.9, 8.8])
def test_a_stored_picture_over_the_limit_is_served_byte_for_byte(alice, picture_mb):
    data, picture = _docx(picture_mb, seed=int(picture_mb * 10))
    asset = alice.get(f"/api/v1/assets/{_image_id(_upload(alice, data))}")
    assert asset.status_code == 200 and len(asset.content) > LIMIT
    assert _sha(asset.content) == _sha(picture)
    assert asset.headers["content-type"] == "image/png" and asset.headers["content-disposition"] == "inline"
    assert asset.headers["cache-control"] == "private, max-age=31536000, immutable"
    _streamed(asset)


def test_a_small_picture_is_served_whole_as_before(alice):
    data, picture = _docx(0.5, seed=5)
    asset = alice.get(f"/api/v1/assets/{_image_id(_upload(alice, data))}")
    assert asset.content == picture and asset.headers["content-length"] == str(len(picture))
    assert asset.headers["content-disposition"] == "inline"


# -- Exports ---------------------------------------------------------------------------------------------------


def test_a_direct_docx_export_over_the_limit_is_a_whole_valid_word_file(alice, large):
    exported = alice.get(f"/api/v1/documents/{large['id']}/export/docx")
    assert exported.status_code == 200 and len(exported.content) > LIMIT
    assert exported.headers["content-type"] == DOCX
    # The document's own title, in Cyrillic: an ASCII fallback and the RFC 6266 UTF-8 name.
    assert exported.headers["content-disposition"] == (
        "attachment; filename=\"_____ __ 2026.docx\"; filename*=UTF-8''%D0%9E%D1%82%D1%87%D0%B5%D1%82%20%D0%B7%D0%B0%202026.docx"
    )
    _streamed(exported)
    assert [_sha(media) for media in _media(exported.content).values()] == [_sha(large["picture"])]
    with zipfile.ZipFile(io.BytesIO(exported.content)) as package:
        assert b"<w:drawing>" in package.read("word/document.xml") and b"Text after the picture." in package.read("word/document.xml")


def test_a_direct_pdf_export_of_about_11_mb_is_a_whole_valid_pdf(alice, large):
    exported = alice.get(f"/api/v1/documents/{large['id']}/export/pdf")
    assert exported.status_code == 200 and 9 * 1024 * 1024 < len(exported.content) < 13 * 1024 * 1024
    assert exported.headers["content-type"] == "application/pdf" and exported.headers["content-disposition"].endswith("%202026.pdf")
    _streamed(exported)
    assert exported.content.startswith(b"%PDF-") and exported.content.rstrip().endswith(b"%%EOF")
    reader = PdfReader(io.BytesIO(exported.content))
    assert len(reader.pages) >= 1
    assert any(page.images for page in reader.pages)
    assert "Text after the picture." in "".join(page.extract_text() for page in reader.pages)


def test_a_small_export_leaves_in_one_body_with_its_length(alice):
    document = alice.post("/api/v1/documents", json={"text": "# Отчет\n\nA small document."}).json()
    exported = alice.get(f"/api/v1/documents/{document['id']}/export/docx")
    assert exported.status_code == 200 and exported.headers["content-length"] == str(len(exported.content))
    assert "filename*=UTF-8''%D0%9E%D1%82%D1%87%D0%B5%D1%82.docx" in exported.headers["content-disposition"]


@pytest.mark.parametrize("file_format", ["docx", "pdf"])
def test_an_export_jobs_file_over_the_limit_is_exactly_the_stored_file(alice, large, file_format):
    job = alice.post("/api/v1/jobs/export", json={"documentId": large["id"], "format": file_format})
    assert job.status_code == 202 and job.json()["status"] == "succeeded", job.text
    download = alice.get(f"/api/v1/jobs/{job.json()['id']}/file")
    assert download.status_code == 200 and len(download.content) > LIMIT
    _streamed(download)
    stored = _stored_job_file(job.json()["id"])
    assert _sha(download.content) == _sha(stored)  # exactly what the job kept
    assert alice.get(f"/api/v1/jobs/{job.json()['id']}/file").content == stored  # downloading again changes nothing
    if file_format == "docx":
        assert [_sha(media) for media in _media(download.content).values()] == [_sha(large["picture"])]
    else:
        assert len(PdfReader(io.BytesIO(download.content)).pages) >= 1


def _stored_job_file(job_id: str) -> bytes:
    storage: LocalStorageProvider = app.dependency_overrides[get_storage_provider]()
    keys = [key for key, _ in asyncio.run(storage.list("jobs/")) if job_id in key and not key.endswith("input")]
    assert len(keys) == 1, keys
    return asyncio.run(storage.get(keys[0]))


def test_a_batch_zip_over_the_limit_holds_every_document_whole(alice, monkeypatch):
    _plan(monkeypatch, maxBatchJobs=5)
    pictures = []
    ids = []
    for seed in (41, 42):
        data, picture = _docx(3.0, seed=seed, title=f"Part {seed}")
        ids.append(_upload(alice, data, name=f"part-{seed}.docx")["id"])
        pictures.append(picture)
    job = alice.post("/api/v1/jobs/batch-export", json={"documentIds": ids, "format": "docx"})
    assert job.json()["status"] == "succeeded", job.text
    download = alice.get(f"/api/v1/jobs/{job.json()['id']}/file")
    assert download.status_code == 200 and download.headers["content-type"] == "application/zip" and len(download.content) > LIMIT
    _streamed(download)
    with zipfile.ZipFile(io.BytesIO(download.content)) as package:
        assert package.testzip() is None
        inner = [package.read(name) for name in package.namelist()]
    assert sorted(_sha(media) for part in inner for media in _media(part).values()) == sorted(_sha(picture) for picture in pictures)


def test_a_doctor_download_over_the_limit_is_a_whole_word_file(alice, large):
    download = alice.post(f"/api/v1/documents/{large['id']}/doctor/download", json={"wordRepairs": []})
    assert download.status_code == 200, download.text[:300]
    assert download.headers["x-safeedit"] == "verified" and "repaired" in download.headers["content-disposition"]
    assert len(download.content) > LIMIT
    _streamed(download)
    assert [_sha(media) for media in _media(download.content).values()] == [_sha(large["picture"])]


# -- PDF tools ---------------------------------------------------------------------------------------------------


def _pdf(client: TestClient, seed: int) -> bytes:
    data, _ = _docx(2.6, seed=seed, title=f"Pdf {seed}")
    return client.get(f"/api/v1/documents/{_upload(client, data, name=f'pdf-{seed}.docx')['id']}/export/pdf").content


def test_the_pdf_tools_results_over_the_limit_are_whole(alice):
    first, second = _pdf(alice, 61), _pdf(alice, 62)
    assert len(first) + len(second) > LIMIT
    merged = alice.post("/api/v1/pdf/merge", files=[("files", ("a.pdf", first, "application/pdf")), ("files", ("b.pdf", second, "application/pdf"))])
    assert merged.status_code == 200 and len(merged.content) > LIMIT and merged.headers["content-type"] == "application/pdf"
    _streamed(merged)
    pages = len(PdfReader(io.BytesIO(first)).pages) + len(PdfReader(io.BytesIO(second)).pages)
    assert len(PdfReader(io.BytesIO(merged.content)).pages) == pages

    split = alice.post("/api/v1/pdf/split", files={"file": ("merged.pdf", merged.content, "application/pdf")}, data={"ranges": f"[[1, {len(PdfReader(io.BytesIO(first)).pages)}], [{len(PdfReader(io.BytesIO(first)).pages) + 1}, {pages}]]"})
    assert split.status_code == 200 and split.headers["content-type"] == "application/zip" and len(split.content) > LIMIT
    _streamed(split)
    with zipfile.ZipFile(io.BytesIO(split.content)) as package:
        assert package.testzip() is None and len(package.namelist()) == 2

    rotated = alice.post("/api/v1/pdf/pages", files={"file": ("merged.pdf", merged.content, "application/pdf")}, data={"operations": '[{"op": "rotate", "pages": [1], "degrees": 90}]'})
    assert rotated.status_code == 200 and len(rotated.content) > LIMIT
    _streamed(rotated)
    assert PdfReader(io.BytesIO(rotated.content)).pages[0].rotation == 90


# -- Access, missing files, errors -------------------------------------------------------------------------------


def test_another_user_and_no_user_get_the_existing_refusals_and_nothing_of_the_file(alice, bob, large):
    job = alice.post("/api/v1/jobs/export", json={"documentId": large["id"], "format": "docx"}).json()
    asset_id = _image_id(large["document"])
    for path in (f"/api/v1/assets/{asset_id}", f"/api/v1/jobs/{job['id']}/file", f"/api/v1/documents/{large['id']}/export/docx", f"/api/v1/documents/{large['id']}/export/pdf"):
        refused = bob.get(path)
        assert refused.status_code == 404, path
        assert len(refused.content) < 1000 and large["picture"][:64] not in refused.content
        anonymous = TestClient(app, base_url="https://testserver").get(path)
        assert anonymous.status_code == 401, path
    assert bob.post(f"/api/v1/documents/{large['id']}/doctor/download", json={"wordRepairs": []}).status_code == 404


def test_a_missing_or_expired_export_is_a_plain_404(alice, large):
    assert alice.get("/api/v1/jobs/00000000-0000-4000-8000-000000000000/file").status_code == 404
    job = alice.post("/api/v1/jobs/export", json={"documentId": large["id"], "format": "docx"}).json()
    storage: LocalStorageProvider = app.dependency_overrides[get_storage_provider]()
    for key, _ in asyncio.run(storage.list("jobs/")):
        if job["id"] in key and not key.endswith("input"):
            asyncio.run(storage.delete(key))  # what the expiry sweep does
    gone = alice.get(f"/api/v1/jobs/{job['id']}/file")
    assert gone.status_code == 404 and "expired" in gone.text


def test_a_storage_error_is_a_plain_500_before_any_byte_of_the_file(alice, large, monkeypatch):
    asset_id = _image_id(large["document"])

    async def broken(self, key):
        raise ConnectionError("storage unreachable")

    monkeypatch.setattr(LocalStorageProvider, "get", broken)
    failed = alice.get(f"/api/v1/assets/{asset_id}")
    assert failed.status_code == 500 and failed.headers["content-type"].startswith("application/json")
    assert large["picture"][:64] not in failed.content


def test_nothing_of_a_large_file_or_where_it_is_stored_goes_into_the_logs(alice, large, caplog):
    asset_id = _image_id(large["document"])
    with caplog.at_level(logging.DEBUG):
        assert alice.get(f"/api/v1/assets/{asset_id}").status_code == 200
        assert alice.get(f"/api/v1/documents/{large['id']}/export/docx").status_code == 200
    logged = caplog.text
    picture = large["picture"]
    for fragment in (picture[100:140].hex(), base64.b64encode(picture[99:150]).decode(), "Signature=", "X-Amz-"):
        assert fragment not in logged
    storage: LocalStorageProvider = app.dependency_overrides[get_storage_provider]()
    keys = [key for key, _ in asyncio.run(storage.list(""))]
    assert any(asset_id in key for key in keys)
    assert not [key for key in keys if key in logged]  # no storage key, so nowhere a file lives


# -- On the wire: a real server, a real socket -------------------------------------------------------------------


@pytest.fixture
def served():
    """file_response behind a real uvicorn and the same kind of http middleware the app runs (BaseHTTPMiddleware)."""
    big = random.Random(7).randbytes(9 * 1024 * 1024)
    probe = FastAPI()

    @probe.middleware("http")
    async def passthrough(request, call_next):
        response = await call_next(request)
        response.headers["X-Request-Id"] = "probe"
        return response

    @probe.get("/file")
    async def download():
        return file_response(big, "application/pdf", filename="big.pdf")

    with socket.socket() as finder:
        finder.bind(("127.0.0.1", 0))
        port = finder.getsockname()[1]
    server = uvicorn.Server(uvicorn.Config(probe, host="127.0.0.1", port=port, log_level="warning"))
    thread = threading.Thread(target=server.run, daemon=True)
    thread.start()
    for _ in range(100):
        if server.started:
            break
        time.sleep(0.05)
    yield port, big
    server.should_exit = True
    thread.join(timeout=10)


def _get(port: int) -> socket.socket:
    connection = socket.create_connection(("127.0.0.1", port), timeout=10)
    connection.sendall(b"GET /file HTTP/1.1\r\nHost: 127.0.0.1\r\nConnection: close\r\n\r\n")
    return connection


def _read_all(connection: socket.socket) -> bytes:
    chunks = []
    while chunk := connection.recv(65536):
        chunks.append(chunk)
    return b"".join(chunks)


def _dechunk(body: bytes) -> tuple[bytes, bool]:
    """The body of a chunked response, and whether it ended with the closing zero-length chunk."""
    out, rest = [], body
    while b"\r\n" in rest:
        size_line, rest = rest.split(b"\r\n", 1)
        size = int(size_line, 16)
        if size == 0:
            return b"".join(out), True
        out.append(rest[:size])
        rest = rest[size + 2 :]
    return b"".join(out), False


def test_on_the_wire_a_large_file_is_chunked_whole_and_a_client_that_leaves_harms_nothing(served):
    port, big = served
    # A client that reads a little and goes away: the server goes on serving others.
    for _ in range(3):
        leaving = _get(port)
        assert leaving.recv(1024).startswith(b"HTTP/1.1 200")
        leaving.close()
    raw = _read_all(_get(port))
    head, body = raw.split(b"\r\n\r\n", 1)
    headers = head.decode("latin-1").lower()
    assert "transfer-encoding: chunked" in headers and "content-length" not in headers
    assert "x-request-id: probe" in headers and "content-disposition: attachment" in headers
    content, finished = _dechunk(body)
    assert finished and _sha(content) == _sha(big)


def test_on_the_wire_a_stream_cut_short_is_never_a_finished_file(served):
    port, big = served
    connection = _get(port)
    received = b""
    while len(received) < 2 * 1024 * 1024:
        received += connection.recv(65536)
    connection.close()
    head, body = received.split(b"\r\n\r\n", 1)
    content, finished = _dechunk(body)
    # What arrived is the start of the file, without the closing chunk: a browser shows a failed download, not a file.
    assert not finished
    assert 1024 * 1024 < len(content) < len(big) and big.startswith(content)


def test_on_the_wire_several_large_downloads_at_once_each_come_whole(served):
    port, big = served
    results: list[str] = []

    def one() -> None:
        head, body = _read_all(_get(port)).split(b"\r\n\r\n", 1)
        content, finished = _dechunk(body)
        results.append(_sha(content) if finished else "unfinished")

    threads = [threading.Thread(target=one) for _ in range(6)]
    for thread in threads:
        thread.start()
    for thread in threads:
        thread.join(timeout=60)
    assert results == [_sha(big)] * 6


@pytest.fixture
def real_app_server(api_db):
    """The real app -- every middleware it runs (CORS, request context, rate limit, isolation gate, cross-site check,
    body-size cap) -- behind a real uvicorn on a real socket, on the test's database and storage."""
    with socket.socket() as finder:
        finder.bind(("127.0.0.1", 0))
        port = finder.getsockname()[1]
    server = uvicorn.Server(uvicorn.Config(app, host="127.0.0.1", port=port, log_level="warning", lifespan="off"))
    thread = threading.Thread(target=server.run, daemon=True)
    thread.start()
    for _ in range(100):
        if server.started:
            break
        time.sleep(0.05)
    yield port
    server.should_exit = True
    thread.join(timeout=10)


def _raw_get(port: int, path: str, cookie: str) -> tuple[str, bytes]:
    connection = socket.create_connection(("127.0.0.1", port), timeout=30)
    connection.sendall(f"GET {path} HTTP/1.1\r\nHost: 127.0.0.1\r\nCookie: {cookie}\r\nConnection: close\r\n\r\n".encode())
    head, body = _read_all(connection).split(b"\r\n\r\n", 1)
    return head.decode("latin-1").lower(), body


def test_through_the_real_app_on_the_wire_large_files_are_chunked_and_small_ones_keep_their_length(alice, large, real_app_server):
    from app.api.deps import SESSION_COOKIE

    cookie = f"{SESSION_COOKIE}={alice.cookies.get(SESSION_COOKIE)}"
    for path, expected in (
        (f"/api/v1/assets/{_image_id(large['document'])}", large["picture"]),
        (f"/api/v1/documents/{large['id']}/export/docx", None),
    ):
        headers, body = _raw_get(real_app_server, path, cookie)
        assert headers.startswith("http/1.1 200"), headers
        assert "transfer-encoding: chunked" in headers and "content-length" not in headers, headers
        assert "x-request-id:" in headers and "x-content-type-options: nosniff" in headers  # the app's own middleware ran
        content, finished = _dechunk(body)
        assert finished and len(content) > LIMIT
        if expected is not None:
            assert _sha(content) == _sha(expected)
        else:
            assert [_sha(media) for media in _media(content).values()] == [_sha(large["picture"])]
    small = alice.post("/api/v1/documents", json={"text": "# Small\n\nA small document."}).json()
    headers, body = _raw_get(real_app_server, f"/api/v1/documents/{small['id']}/export/docx", cookie)
    assert "content-length: " in headers and "transfer-encoding" not in headers
    assert zipfile.ZipFile(io.BytesIO(body)).testzip() is None
    # Without a valid session: 401 over the wire, and nothing of the file.
    headers, body = _raw_get(real_app_server, f"/api/v1/assets/{_image_id(large['document'])}", f"{SESSION_COOKIE}=not-a-session")
    assert headers.startswith("http/1.1 401") and len(body) < 2000
