"""Word's hidden text stays hidden (tracker DOCX-025): imported as the Document
Model's hidden mark -- set on the run, by its character style, by its
paragraph's style or by the document's defaults, as Word resolves it -- kept
through the editor, hidden again in a Word export and left out of a PDF, which
says so. The content checks count it where it is: in the Word file and the
document on both sides, never in a printed page."""

import io
import zipfile

from docx import Document as DocxDocument
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from fastapi.testclient import TestClient
from pypdf import PdfReader

from app.export.docx_export import build_docx
from app.export.pdf_export import build_pdf
from app.fidelity.content import document_words, hidden_words
from app.fidelity.docx_detect import detect_docx_features
from app.main import app
from app.models.document import MarkType
from app.parsers.docx import parse_docx

client = TestClient(app, base_url="https://testserver")
_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_W = nsdecls("w")


def _word_file(*, hidden_normal: bool = False) -> bytes:
    word = DocxDocument()
    secret = word.styles.add_style("Secret", WD_STYLE_TYPE.CHARACTER)
    secret.element.get_or_add_rPr().append(parse_xml(f"<w:vanish {_W}/>"))
    notes = word.styles.add_style("Speaker notes", WD_STYLE_TYPE.PARAGRAPH)
    notes.element.get_or_add_rPr().append(parse_xml(f"<w:vanish {_W}/>"))
    if hidden_normal:
        word.styles["Normal"].element.get_or_add_rPr().append(parse_xml(f"<w:vanish {_W}/>"))

    direct = word.add_paragraph("The answer is ")
    direct._p.append(parse_xml(f'<w:r {_W}><w:rPr><w:vanish/></w:rPr><w:t xml:space="preserve">forty-two, </w:t></w:r>'))
    direct._p.append(parse_xml(f"<w:r {_W}><w:t>as agreed.</w:t></w:r>"))
    styled = word.add_paragraph("Styled ")
    styled.add_run("secretly", style="Secret")
    web = styled.add_run(" and web-only")
    web._r.get_or_add_rPr().append(parse_xml(f"<w:webHidden {_W}/>"))  # hidden in Word's web view only: shown
    notes_paragraph = word.add_paragraph("Remember to smile.", style="Speaker notes")
    notes_paragraph.add_run(" Shown anyway.")._r.get_or_add_rPr().append(parse_xml(f'<w:vanish {_W} w:val="0"/>'))
    buffer = io.BytesIO()
    word.save(buffer)
    return buffer.getvalue()


def _runs(document) -> list[list[tuple[str, list[MarkType]]]]:
    return [[(run.text, [mark.type for mark in run.marks]) for run in element.inline or []] for element in document.elements]


_EXPECTED = [
    [("The answer is ", []), ("forty-two, ", [MarkType.HIDDEN]), ("as agreed.", [])],
    [("Styled ", []), ("secretly", [MarkType.HIDDEN]), (" and web-only", [])],
    [("Remember to smile.", [MarkType.HIDDEN]), (" Shown anyway.", [])],
]


def test_hidden_text_is_read_from_the_run_its_style_the_paragraphs_style_as_word_does():
    assert _runs(parse_docx(_word_file(), "hidden.docx")) == _EXPECTED


def test_a_hidden_default_paragraph_style_hides_paragraphs_without_a_style():
    data = _word_file(hidden_normal=True)

    document = parse_docx(data, "hidden.docx")
    detected = {item.feature: item for item in detect_docx_features(data)}

    assert _runs(document) == [
        [("The answer is forty-two, as agreed.", [MarkType.HIDDEN])],
        [("Styled secretly and web-only", [MarkType.HIDDEN])],
        _EXPECTED[2],  # its own style, and a run that shows itself
    ]
    assert detected["docx.hidden_text"].count == 7  # the report counts the same runs: 3 + 3 + 1


def test_the_import_report_says_it_is_kept_hidden():
    item = next(item for item in detect_docx_features(_word_file()) if item.feature == "docx.hidden_text")

    assert item.policy == "detected_preserved" and not item.contentChanged and item.count == 3
    assert "kept hidden" in item.reason


def test_an_upload_keeps_it_hidden_and_its_content_check_counts_it_on_both_sides(api_db):
    client.cookies.clear()
    assert client.post("/api/v1/auth/register", json={"email": "hidden@example.com", "password": "long enough password"}).status_code == 201
    document = client.post("/api/v1/documents/upload", files={"file": ("hidden.docx", _word_file(), _DOCX)}).json()

    runs = document["elements"][0]["inline"]
    assert [(run["text"], [mark["type"] for mark in run["marks"]]) for run in runs][1] == ("forty-two, ", ["hidden"])
    assert document["importReport"]["content"]["verified"]

    saved = client.put(f"/api/v1/documents/{document['id']}/content", json={"elements": document["elements"]})  # through the editor
    client.cookies.clear()
    assert saved.status_code == 200
    [mark] = saved.json()["elements"][0]["inline"][1]["marks"]
    assert mark["type"] == "hidden" and all(value is None for key, value in mark.items() if key != "type")


def test_a_word_export_hides_it_again():
    document = parse_docx(_word_file(), "hidden.docx")

    exported = build_docx(document)

    with zipfile.ZipFile(io.BytesIO(exported)) as package:
        body = package.read("word/document.xml").decode("utf-8")
    assert body.count("<w:vanish/>") == 3
    assert _runs(parse_docx(exported, "again.docx")) == _EXPECTED


def test_a_pdf_leaves_it_out_and_says_so(api_db):
    client.cookies.clear()
    assert client.post("/api/v1/auth/register", json={"email": "print@example.com", "password": "long enough password"}).status_code == 201
    document = client.post("/api/v1/documents/upload", files={"file": ("hidden.docx", _word_file(), _DOCX)}).json()

    job = client.post("/api/v1/jobs/export", json={"documentId": document["id"], "format": "pdf"}).json()
    finished = client.get(f"/api/v1/jobs/{job['id']}").json()
    pdf = client.get(f"/api/v1/jobs/{job['id']}/file")
    client.cookies.clear()

    text = " ".join(page.extract_text() for page in PdfReader(io.BytesIO(pdf.content)).pages)
    assert "forty-two" not in text and "secretly" not in text and "smile" not in text
    assert "as agreed" in text and "Shown anyway" in text and "web-only" in text
    fidelity = finished["result"]["fidelity"]
    note = next(item for item in fidelity["items"] if item["feature"] == "export.pdf.hidden_text")
    assert note["policy"] == "detected_preserved" and "6 words" in note["reason"]
    assert fidelity["content"]["verified"]  # checked against the words a page shows


def test_the_pdf_builder_prints_only_what_shows():
    document = parse_docx(_word_file(), "hidden.docx")

    text = " ".join(page.extract_text() for page in PdfReader(io.BytesIO(build_pdf(document))).pages)

    assert hidden_words(document.elements) == 6
    assert document_words(document.elements, visible_only=True) == ["The", "answer", "is", "as", "agreed", "Styled", "and", "web", "only", "Shown", "anyway"]
    assert "forty" not in text
