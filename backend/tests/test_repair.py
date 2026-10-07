"""Repair document (tracker REV-004, brief §60): what is broken, by kind -- numbering, styles,
tables, links, what the app doesn't hold, malformed input -- with each fix proposed for review
(detected issue, proposed fix, preview, apply). Broken tables are new here: rows short of the
table's width, a cell merged down past the last row, empty rows."""

import io

import pytest
from docx import Document as DocxDocument
from fastapi.testclient import TestClient

from app.fidelity.report import ContentCheck, FidelityReport, FidelityStage
from app.formatting.engine import recompute_styles
from app.formatting.health_fixes import fixes
from app.formatting.repair import repair_report
from app.formatting.table_repair import repaired, table_problems
from app.main import app
from app.models.document import Document, Element, ElementType, InlineRun, TableCell, TableContent, TableRow

_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _cell(text: str = "", **extra) -> TableCell:
    return TableCell(inline=[InlineRun(text=text)] if text else [], **extra)


def _table(rows: list[list[TableCell]]) -> Element:
    table = TableContent(rows=[TableRow(cells=cells) for cells in rows])
    content = "\n".join(" | ".join(cell.inline[0].text if cell.inline else "" for cell in row) for row in rows)
    return Element(type=ElementType.TABLE, content=content, table=table, order=0)


def test_a_short_row_a_span_past_the_end_and_an_empty_row_are_found_and_put_right():
    element = _table(
        [
            [_cell("Name"), _cell("Qty"), _cell("Price")],
            [_cell("Tea"), _cell("2")],  # one short
            [_cell(), _cell(), _cell()],  # empty
            [_cell("Total", colspan=2), _cell("9", rowspan=3)],  # merged down past the end
        ]
    )
    problems = table_problems(element.table)
    assert (problems.width, problems.short, problems.empty, problems.overflowing) == (3, [1], [2], [3])
    fixed = repaired(element)
    assert [len(row.cells) for row in fixed.table.rows] == [3, 3, 2]
    assert fixed.table.rows[2].cells[1].rowspan == 1
    assert fixed.content.splitlines() == ["Name | Qty | Price", "Tea | 2 | ", "Total | 9"]
    assert not table_problems(fixed.table)
    assert element.table.rows[1].cells[-1].inline[0].text == "2"  # the original is left as it was


def test_cells_merged_down_fill_the_rows_below_and_aren_t_broken():
    element = _table([[_cell("A", rowspan=2), _cell("B")], [_cell("C")], [_cell(), _cell()]])
    problems = table_problems(element.table)
    assert problems.short == [] and problems.overflowing == [] and problems.empty == [2]
    assert repaired(_table([[_cell("A", rowspan=2), _cell("B")], [_cell("C")]])) is None
    blank = _table([[_cell(), _cell()], [_cell(), _cell()]])
    assert repaired(blank) is None  # nothing but empty rows: blank, not broken


def _document(*elements: Element) -> Document:
    document = Document(elements=[element.model_copy(update={"order": index}) for index, element in enumerate(elements)])
    recompute_styles(document)
    return document


def test_the_repair_report_groups_what_health_finds_by_kind_with_its_fixes():
    broken = _table([[_cell("a"), _cell("b")], [_cell("c")]])
    link = Element(type=ElementType.PARAGRAPH, content="See here", inline=[InlineRun(text="See here", marks=[{"type": "link", "href": "javascript:alert(1)"}])], order=0)
    document = _document(broken, link)
    report = repair_report(document)
    kinds = {issue.kind: issue for issue in report.issues}
    assert kinds["tables"].checkId == "broken_tables" and kinds["tables"].fixes == 1 and kinds["tables"].status == "fail"
    assert kinds["tables"].elementIds == [document.elements[0].id]
    assert [issue.kind for issue in report.issues] == sorted((issue.kind for issue in report.issues), key=["numbering", "styles", "tables", "links", "structures", "input"].index)
    [proposal] = fixes(document, ["broken_tables"])
    assert proposal.category == "structure" and proposal.replacement.table.rows[1].cells[1].inline == []


def test_text_the_import_read_wrong_is_malformed_input_with_no_fix():
    document = _document(Element(type=ElementType.PARAGRAPH, content="Text.", inline=[InlineRun(text="Text.")], order=0))
    document.importReport = FidelityReport(stage=FidelityStage.IMPORT, sourceType="pdf", content=ContentCheck(method="pdf-text", verified=False, sourceWords=10, resultWords=7, missing=3))
    [issue] = [issue for issue in repair_report(document).issues if issue.kind == "input"]
    assert issue.fixes == 0 and issue.status == "fail" and "3 words missing" in issue.summary


@pytest.fixture
def client(api_db):
    test_client = TestClient(app, base_url="https://testserver")
    assert test_client.post("/api/v1/auth/register", json={"email": "repair@example.com", "password": "long enough password"}).status_code == 201
    return test_client


def test_detected_proposed_previewed_and_applied_through_the_api(client):
    word = DocxDocument()
    table = word.add_table(rows=2, cols=3)
    for column, text in enumerate(["Name", "Qty", "Price"]):
        table.cell(0, column).text = text
    table.cell(1, 0).text = "Tea"
    row = table.rows[1]._tr
    row.remove(row.tc_lst[-1])  # a row Word left a cell short
    out = io.BytesIO()
    word.save(out)
    uploaded = client.post("/api/v1/documents/upload", files={"file": ("table.docx", out.getvalue(), _DOCX)}).json()

    report = client.get(f"/api/v1/documents/{uploaded['id']}/repair").json()
    [tables] = [issue for issue in report["issues"] if issue["checkId"] == "broken_tables"]
    assert tables["fixes"] == 1 and tables["kind"] == "tables"

    proposed = client.post(f"/api/v1/documents/{uploaded['id']}/health/fixes", json={"checkIds": ["broken_tables"]}).json()
    [proposal] = [item for item in proposed["document"]["proposals"] if item["checkId"] == "broken_tables"]
    assert [len(row["cells"]) for row in proposal["replacement"]["table"]["rows"]] == [3, 3]  # the preview
    table_before = next(element for element in proposed["document"]["elements"] if element["type"] == "table")
    assert len(table_before["table"]["rows"][1]["cells"]) == 2  # nothing applied yet

    accepted = client.post(f"/api/v1/documents/{uploaded['id']}/proposals/{proposal['id']}/accept", json={})
    assert accepted.status_code == 200, accepted.text
    table_after = next(element for element in accepted.json()["elements"] if element["type"] == "table")
    assert len(table_after["table"]["rows"][1]["cells"]) == 3
    assert not [issue for issue in client.get(f"/api/v1/documents/{uploaded['id']}/repair").json()["issues"] if issue["checkId"] == "broken_tables"]


def test_a_row_merged_into_from_above_isn_t_an_empty_row():
    above = _table([[_cell("A", rowspan=2), _cell("B")], [_cell()], [_cell("C"), _cell("D")]])
    assert table_problems(above.table).empty == []
    down = _table([[_cell("A"), _cell("B")], [_cell(rowspan=2), _cell()], [_cell("C")]])
    assert table_problems(down.table).empty == []


def test_the_report_goes_by_kind_then_failures_first(monkeypatch):
    from app.formatting import repair
    from app.formatting.health import HealthCheck, HealthReport

    def check(check_id: str, status: str) -> HealthCheck:
        return HealthCheck(id=check_id, title=check_id, status=status, summary="", issues=[], weight=1)

    shuffled = [check("links", "warn"), check("page_breaks", "warn"), check("fonts", "warn"), check("hierarchy", "warn"), check("numbering", "fail"), check("captions", "fail")]
    monkeypatch.setattr(repair, "check_health", lambda document: HealthReport(score=50, rating="fair", checks=shuffled))
    report = repair.repair_report(Document())
    assert [issue.checkId for issue in report.issues] == ["numbering", "hierarchy", "fonts", "links", "page_breaks"]  # captions isn't a repair


def test_an_import_whose_words_checked_out_isn_t_malformed():
    document = _document(Element(type=ElementType.PARAGRAPH, content="Text.", inline=[InlineRun(text="Text.")], order=0))
    document.importReport = FidelityReport(stage=FidelityStage.IMPORT, sourceType="pdf", content=ContentCheck(method="pdf-text", verified=True, sourceWords=1, resultWords=1))
    assert not [issue for issue in repair_report(document).issues if issue.kind == "input"]
