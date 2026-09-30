"""Pictures past the limits (tracker SEC-012): judged by their header before anything
decodes them, on import (left out and said to be), on the editor's save (removed, or the
save refused before anything is stored) and on export (a picture stored before the
limits is left out) -- with only PNG, JPEG, GIF, WebP and BMP ever opened."""

import base64
import io
import time

import pytest
from docx import Document as DocxDocument
from docx.shared import Inches
from fastapi.testclient import TestClient
from PIL import Image as PILImage
from sqlalchemy import func, select
from sqlalchemy.orm import Session as OrmSession

from app.ai.base import AIStructuredOutputError
from app.ai.factory import get_ai_provider
from app.db.models import DocumentAsset
from app.main import app
from app.parsers.docx import parse_docx
from app.security import files as limits
from app.security.files import picture_problem
from tests.fakes import FakeAIProvider
from tests.malformed_pictures import EPS, PICTURES, jpeg_claiming, png_claiming, real

pytestmark = pytest.mark.security  # the security regression suite (TEST-030)

client = TestClient(app, base_url="https://testserver")


@pytest.fixture
def signed_in(api_db):
    client.cookies.clear()
    app.dependency_overrides[get_ai_provider] = lambda: FakeAIProvider([AIStructuredOutputError("no AI in tests")] * 3)
    assert client.post("/api/v1/auth/register", json={"email": "pictures@example.com", "password": "long enough password"}).status_code == 201
    yield api_db
    client.cookies.clear()
    app.dependency_overrides.pop(get_ai_provider, None)


def _assets(api_db) -> int:
    with OrmSession(api_db) as session:
        return session.scalar(select(func.count(DocumentAsset.id)))


def _docx(*pictures: bytes) -> bytes:
    document = DocxDocument()
    document.add_paragraph("Before the pictures.")
    for picture in pictures:
        document.add_picture(io.BytesIO(picture), width=Inches(1))
    document.add_paragraph("After the pictures.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _notes(document) -> dict[str, int]:
    return {item.feature: item.count for item in document.importReport.items}


def _inline(picture: bytes, content_type: str = "image/png") -> dict:
    return {"type": "image", "content": "", "order": 0, "image": {"src": f"data:{content_type};base64," + base64.b64encode(picture).decode()}}


@pytest.mark.parametrize("name", list(PICTURES))
def test_a_picture_is_judged_by_its_header_before_anything_decodes_it(name):
    data, content_type, expected = PICTURES[name]
    start = time.perf_counter()
    problem = picture_problem(data, content_type)
    assert (problem.kind if problem else "ok") == expected
    assert time.perf_counter() - start < 0.5  # a claimed 10 gigapixels costs nothing to judge


def test_pillow_itself_opens_nothing_past_the_limit_anywhere():
    # The backstop for any decode that doesn't ask first -- reportlab's included -- and
    # whatever warning filters are in force (pytest's own, for one).
    for claimed in ((8_000, 7_000), (100_000, 100_000)):
        with pytest.raises(PILImage.DecompressionBombError):
            PILImage.open(io.BytesIO(png_claiming(*claimed)))
    assert PILImage.open(io.BytesIO(png_claiming(7_000, 7_000))).size == (7_000, 7_000)  # just under: a picture


def test_only_the_allowed_formats_are_ever_opened():
    assert PILImage.open(io.BytesIO(EPS)).format == "EPS"  # what Pillow would otherwise open (and run Ghostscript for)
    assert picture_problem(EPS).kind == "unreadable"
    assert picture_problem(real("TIFF")).kind == "unreadable"


def test_a_word_file_s_pictures_past_the_limits_are_left_out_and_said_to_be():
    document = parse_docx(_docx(real("PNG"), png_claiming(100_000, 100_000), jpeg_claiming(60_000, 60_000)), "pictures.docx")

    assert len([element for element in document.elements if element.type == "image"]) == 1
    assert _notes(document)["docx.image.too_large"] == 2
    assert [element.content for element in document.elements if element.type == "paragraph"] == ["Before the pictures.", "After the pictures."]


def test_a_word_file_with_more_pictures_than_a_document_holds(monkeypatch):
    monkeypatch.setattr(limits, "MAX_PICTURES", 2)
    document = parse_docx(_docx(real("PNG"), real("PNG"), real("PNG")), "many.docx")
    assert len([element for element in document.elements if element.type == "image"]) == 2
    assert _notes(document)["docx.image.too_many"] == 1

    monkeypatch.setattr(limits, "MAX_PICTURES", 1000)
    monkeypatch.setattr(limits, "MAX_PICTURE_TOTAL_BYTES", len(real("PNG")) + 1)
    document = parse_docx(_docx(real("PNG"), real("PNG")), "heavy.docx")
    assert len([element for element in document.elements if element.type == "image"]) == 1
    assert _notes(document)["docx.image.too_many"] == 1


def test_a_saved_picture_past_the_limits_is_removed_and_said_to_be(signed_in):
    document = client.post("/api/v1/documents", json={"text": "# Pictures\n\nSome text."}).json()
    elements = document["elements"] + [_inline(png_claiming(100_000, 100_000)), _inline(real("PNG"))]

    saved = client.put(f"/api/v1/documents/{document['id']}/content", json={"elements": elements})

    assert saved.status_code == 200, saved.text[:300]
    kept = [element for element in saved.json()["elements"] if element["type"] == "image"]
    assert len(kept) == 1 and kept[0]["image"]["assetId"]
    assert any("has more than 50 megapixels" in note for note in saved.json()["unsupportedFeatures"])
    assert _assets(signed_in) == 1


def test_a_save_past_a_document_s_pictures_is_refused_before_anything_is_stored(signed_in, monkeypatch):
    document = client.post("/api/v1/documents", json={"text": "# Pictures\n\nSome text."}).json()
    first = client.put(f"/api/v1/documents/{document['id']}/content", json={"elements": document["elements"] + [_inline(real("PNG"))]}).json()
    assert _assets(signed_in) == 1

    monkeypatch.setattr(limits, "MAX_PICTURES", 1)  # the stored one counts: a second is one too many
    refused = client.put(f"/api/v1/documents/{document['id']}/content", json={"elements": first["elements"] + [_inline(real("GIF"), "image/gif")]})
    body = refused.json()
    assert (refused.status_code, body["code"], body["message"]) == (413, "too_large", "A document can hold at most 1 pictures.")
    assert _assets(signed_in) == 1

    monkeypatch.setattr(limits, "MAX_PICTURES", 1000)
    monkeypatch.setattr(limits, "MAX_PICTURE_TOTAL_BYTES", len(real("PNG")) + 10)
    refused = client.put(f"/api/v1/documents/{document['id']}/content", json={"elements": first["elements"] + [_inline(real("PNG"))]})
    assert (refused.status_code, refused.json()["code"]) == (413, "too_large")
    assert "pictures can take at most" in refused.json()["message"]
    assert _assets(signed_in) == 1
    stored = client.get(f"/api/v1/documents/{document['id']}").json()
    assert len([element for element in stored["elements"] if element["type"] == "image"]) == 1  # the document as it was


def test_an_export_leaves_out_a_picture_from_before_the_limits(signed_in, monkeypatch):
    document = client.post("/api/v1/documents", json={"text": "# Pictures\n\nSome text."}).json()
    client.put(f"/api/v1/documents/{document['id']}/content", json={"elements": document["elements"] + [_inline(real("PNG"))]})
    monkeypatch.setattr(limits, "MAX_PICTURE_PIXELS", 10)  # the stored 6 x 4 picture is now past them

    for extension in ("docx", "pdf"):
        job = client.post("/api/v1/jobs/export", json={"documentId": document["id"], "format": extension}).json()
        assert job["status"] == "succeeded", job
        features = {item["feature"] for item in job["result"]["fidelity"]["items"]}
        assert "export.image.too_large" in features, (extension, features)
