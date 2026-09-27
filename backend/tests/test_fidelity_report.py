"""The Document Fidelity Report (tracker FID-001, FID-005): items classified by
policy, and a content check that says "verified" only when the source's words
and the document's are the same, in the same order."""

import asyncio
import io
import time
from pathlib import Path

import pytest
from docx import Document as DocxDocument
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from fastapi.testclient import TestClient

from app.ai.schemas import AIBlock, AIBlockType, AIStructureResponse
from app.fidelity.content import compare_words, words
from app.fidelity.report import FidelityPolicy, FidelityReport, FidelityStage, ReportBuilder
from app.main import app
from app.services.ingestion_service import build_document_from_docx, build_document_from_text
from tests.fakes import FakeAIProvider

FIXTURES = Path(__file__).parent / "fixtures" / "documents"


def _docx(build) -> bytes:
    document = DocxDocument()
    build(document)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# -- the comparison ----------------------------------------------------------------------


def test_the_same_words_are_verified_whatever_the_spacing_and_punctuation():
    check = compare_words(words("Dose: 5 mg,\ntwice   a day."), words("Dose 5 mg - twice a day"), method="test")
    assert check.verified and check.sourceWords == 6


@pytest.mark.parametrize(
    "result, kind",
    [
        ("Take 50 mg twice a day.", "changed"),  # a number
        ("Take 5 mg a day.", "missing"),  # a word gone
        ("Take 5 mg twice a day. Never stop.", "added"),  # a sentence invented
        ("Take 5 mg twice a day, always.", "added"),
        ("Please take 5 mg twice a day.", "changed"),  # "Take" became "Please take"
    ],
)
def test_any_word_that_differs_is_reported(result, kind):
    check = compare_words(words("Take 5 mg twice a day."), words(result), method="test")
    assert not check.verified
    assert [sample.kind for sample in check.samples] == [kind]


def test_negation_removed_is_caught():
    check = compare_words(words("Do not keep it away from children."), words("Do keep it away from children."), method="test")
    assert not check.verified and check.missing == 1
    assert check.samples[0].source == "not" and check.samples[0].context == "Do"


def test_a_moved_passage_is_reported_as_moved_not_lost():
    check = compare_words(words("alpha beta gamma delta epsilon"), words("alpha delta epsilon beta gamma"), method="test")
    assert not check.verified
    assert check.moved == 2 and check.missing == 0 and check.added == 0


def test_a_long_document_with_one_change_is_compared_quickly():
    source = [f"word{i}" for i in range(200_000)]
    result = source.copy()
    result[123_456] = "changed"
    started = time.perf_counter()
    check = compare_words(source, result, method="test")
    assert time.perf_counter() - started < 2
    assert check.samples[0].source == "word123456" and check.samples[0].result == "changed"


def test_items_repeat_as_a_count_and_policies_decide_what_needs_review():
    builder = ReportBuilder()
    builder.add("docx.image.linked", FidelityPolicy.UNSUPPORTED, "A linked image was not imported.", element_id="a", content_changed=True)
    builder.add("docx.image.linked", FidelityPolicy.UNSUPPORTED, "A linked image was not imported.", element_id="b", content_changed=True)
    builder.add("docx.bookmark", FidelityPolicy.DETECTED_NOT_EDITABLE, "Bookmarks aren't shown in the editor.")
    report = FidelityReport(stage=FidelityStage.IMPORT, sourceType="docx", items=builder.items())

    [linked, bookmark] = report.items
    assert linked.count == 2 and linked.elementIds == ["a", "b"]
    assert report.reviewCount == 1
    assert report.contentStatus == "unverified"  # nothing compared: nothing claimed


# -- Word imports --------------------------------------------------------------------------


@pytest.mark.parametrize("path", sorted(FIXTURES.glob("*.docx")), ids=lambda path: path.name)
def test_every_golden_document_imports_with_its_content_verified(path):
    report = build_document_from_docx(path.read_bytes(), path.name, None).importReport

    assert report.contentStatus == "verified", report.content.samples


def test_text_the_importer_skips_is_reported_missing():
    """Ruby text (furigana) sits in a container the importer doesn't read: the
    check, which reads all of the file's text, notices."""

    def build(document):
        paragraph = document.add_paragraph("Before the ruby ")
        paragraph._p.append(
            parse_xml(
                f"<w:r {nsdecls('w')}><w:ruby><w:rubyPr/><w:rt><w:r><w:t>kan</w:t></w:r></w:rt>"
                "<w:rubyBase><w:r><w:t>漢字</w:t></w:r></w:rubyBase></w:ruby></w:r>"
            )
        )
        document.add_paragraph("after.")

    report = build_document_from_docx(_docx(build), "ruby.docx", None).importReport

    assert report.contentStatus == "changed"
    assert report.content.missing == 1  # the ruby text and its base run together: one word
    assert [sample.source for sample in report.content.samples] == ["kan漢字"]


def test_deleted_revisions_and_field_codes_are_not_content_but_field_results_are():
    def build(document):
        paragraph = document.add_paragraph("Kept ")
        paragraph._p.append(parse_xml(f'<w:del {nsdecls("w")} w:id="1" w:author="x"><w:r><w:delText>removed </w:delText></w:r></w:del>'))
        paragraph._p.append(parse_xml(f'<w:ins {nsdecls("w")} w:id="2" w:author="x"><w:r><w:t>inserted </w:t></w:r></w:ins>'))
        for xml in (
            '<w:r><w:fldChar w:fldCharType="begin"/></w:r>',
            '<w:r><w:instrText xml:space="preserve"> DATE \\@ "yyyy" </w:instrText></w:r>',
            '<w:r><w:fldChar w:fldCharType="separate"/></w:r>',
            "<w:r><w:t>2026</w:t></w:r>",
            '<w:r><w:fldChar w:fldCharType="end"/></w:r>',
        ):
            paragraph._p.append(parse_xml(xml.replace("<w:r>", f"<w:r {nsdecls('w')}>", 1)))

    report = build_document_from_docx(_docx(build), "revisions.docx", None).importReport

    assert report.contentStatus == "verified", report.content.samples
    assert report.content.sourceWords == 3  # Kept inserted 2026
    assert any(item.feature == "docx.tracked_changes" and item.contentChanged for item in report.items)


def test_header_text_that_is_left_out_is_reported():
    def build(document):
        section = document.sections[0]
        section.different_first_page_header_footer = True
        section.header.paragraphs[0].text = "Main header"
        section.first_page_header.paragraphs[0].text = "Cover page header"
        document.add_paragraph("Body.")

    report = build_document_from_docx(_docx(build), "headers.docx", None).importReport

    [lost] = [item for item in report.items if item.feature == "docx.header_footer.text"]
    assert lost.policy == FidelityPolicy.UNSUPPORTED and lost.contentChanged
    assert lost.sourceState == "Cover page header"
    assert any(item.feature == "docx.header_footer.variants" for item in report.items)


# -- text imports ---------------------------------------------------------------------------


def test_markdown_syntax_and_link_addresses_are_not_words():
    text = "# Title\n\nSee [the docs](https://example.com/very/long/path) and **bold** `code`.\n\n- [x] done\n- [ ] todo\n"
    document = asyncio.run(build_document_from_text(text, None, FakeAIProvider([])))

    report = document.importReport
    assert report.sourceType == "paste" and report.content.method == "markdown-text"
    assert report.contentStatus == "verified", report.content.samples


def test_an_ai_answer_that_drops_a_sentence_never_reaches_the_document():
    text = "Take the tablet with water. Do not take more than two a day. Keep away from children."
    answer = AIStructureResponse(
        document_type="general",
        document_type_confidence=0.9,
        blocks=[AIBlock(type=AIBlockType.PARAGRAPH, text="Take the tablet with water. Keep away from children.", confidence=0.9)],
    )
    provider = FakeAIProvider([answer, answer])

    document = asyncio.run(build_document_from_text(text, None, provider))

    assert provider.calls == 2  # the answer and the retry were both refused (app/fidelity/text_check.py)
    assert document.importReport.contentStatus == "verified"
    assert "Do not take more than two a day." in " ".join(element.content for element in document.elements)


def test_the_import_check_still_catches_a_dropped_sentence():
    source = words("Take the tablet with water. Do not take more than two a day. Keep away from children.")
    check = compare_words(source, words("Take the tablet with water. Keep away from children."), method="source-text")

    assert not check.verified
    assert [sample.source for sample in check.samples] == ["Do not take more than two a day"]


# -- through the API ------------------------------------------------------------------------


@pytest.fixture
def alice(api_db):
    client = TestClient(app, base_url="https://testserver")
    assert client.post("/api/v1/auth/register", json={"email": "alice@example.com", "password": "long enough password"}).status_code == 201
    return client


def test_an_uploaded_word_file_comes_with_its_report(alice):
    data = (FIXTURES / "12-complex.docx").read_bytes()
    response = alice.post(
        "/api/v1/documents/upload",
        files={"file": ("complex.docx", data, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert response.status_code == 201
    report = response.json()["importReport"]
    assert report["stage"] == "import" and report["sourceType"] == "docx"
    assert report["contentStatus"] == "verified" and report["content"]["sourceWords"] > 0
    # Stored with the document, and not touched by the editor's saves.
    document = response.json()
    alice.put(f"/api/v1/documents/{document['id']}/content", json={"elements": document["elements"]})
    assert alice.get(f"/api/v1/documents/{document['id']}").json()["importReport"] == report


def test_a_pdf_upload_says_only_its_text_was_imported(alice):
    data = (Path(__file__).parent / "fixtures" / "sample.pdf").read_bytes()
    response = alice.post("/api/v1/documents/upload", files={"file": ("sample.pdf", data, "application/pdf")})
    assert response.status_code == 201
    report = response.json()["importReport"]
    assert report["sourceType"] == "pdf" and report["content"]["method"] == "pdf-extracted-text"
    assert any(item["feature"] == "pdf.layout" for item in report["items"])
