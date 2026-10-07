"""Tables text flows around (tracker DOCX-017B): a Word table that floats (w:tblpPr) floats to its
side on the pages here and in a PDF, with the text beside it -- the side worked out at import from
its position and width, as a floating picture's -- and keeps exactly where it floats in a Word export.
One centred, as wide as the text or too tall to stay beside its text is drawn in line."""

import io

import pytest
from docx import Document as DocxDocument
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

from app.export.docx_export import build_docx
from app.export.pdf_export import build_pdf
from app.formatting.engine import recompute_styles
from app.models.document import Document, Element, ElementType, InlineRun, TableCell, TableContent, TableFloat, TableRow
from app.parsers.docx import parse_docx
from app.parsers.docx_tables import table_side

_W = nsdecls("w")
TEXT = "Text that flows beside the table, line after line, as Word lays it out around a table that floats. " * 5


@pytest.mark.parametrize(
    ("floating", "width", "side"),
    [
        ({"xAlign": "right"}, 6, "right"),
        ({"xAlign": "outside"}, 6, "right"),
        ({"xAlign": "left"}, 6, "left"),
        ({"xAlign": "inside"}, 6, "left"),
        ({"xAlign": "center"}, 6, None),
        ({}, 6, "left"),  # no position: where it stands, at the left
        ({"horizontalAnchor": "margin", "xCm": 10}, 6, "right"),  # its middle at 13 of 17
        ({"horizontalAnchor": "margin", "xCm": 1}, 6, "left"),
        ({"horizontalAnchor": "page", "xCm": 6}, 6, "left"),  # 4 into the column: its middle at 7 of 17
        ({"horizontalAnchor": "page", "xCm": 12}, 6, "right"),
        ({"xAlign": "right"}, 15, None),  # as wide as the text: nothing beside it
        ({"xAlign": "right"}, None, None),  # its width unknown
    ],
)
def test_the_side_a_floating_table_floats_to(floating, width, side):
    # In a 17 cm column 2 cm in from the page's edge.
    assert table_side(floating, width, 2.0, 17.0) == side


def _word(tblpPr: str, widths=(1701, 1701)) -> bytes:
    """A two-column table (3 cm columns by default) floating as `tblpPr` says, with text after it."""
    word = DocxDocument()
    table = word.add_table(rows=2, cols=2)
    for column, twips in zip(table._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        column.set(qn("w:w"), str(twips))
    for index, cell in enumerate(table._cells):
        cell.text = f"cell {index}"
    table._tbl.tblPr.insert(1, parse_xml(f"<w:tblpPr {_W} {tblpPr}/>"))
    word.add_paragraph(TEXT)
    buffer = io.BytesIO()
    word.save(buffer)
    return buffer.getvalue()


def _table(document):
    return next(element.table for element in document.elements if element.type == ElementType.TABLE)


def test_a_floating_table_gets_its_side_at_import_and_keeps_where_it_floats_in_word():
    data = _word('w:rightFromText="284" w:vertAnchor="text" w:horzAnchor="margin" w:tblpXSpec="right" w:tblpY="113"')
    document = parse_docx(data, "floating.docx")

    floating = _table(document).floating
    assert (floating.side, floating.xAlign, floating.rightFromTextCm) == ("right", "right", 0.5)
    [note] = [item for item in document.importReport.items if item.feature == "docx.table.floating"]
    assert "float at their side" in note.reason
    again = _table(parse_docx(build_docx(document), "again.docx")).floating
    assert again == floating  # where it floats, exactly; its side worked out again the same


def test_a_centred_or_full_width_floating_table_is_drawn_in_line_and_said_so():
    centred = parse_docx(_word('w:horzAnchor="margin" w:tblpXSpec="center"'), "centred.docx")
    wide = parse_docx(_word('w:horzAnchor="margin" w:tblpXSpec="right"', widths=(4800, 4800)), "wide.docx")  # 17 cm

    for document in (centred, wide):
        assert _table(document).floating.side is None
        [note] = [item for item in document.importReport.items if item.feature == "docx.table.floating"]
        assert "in line" in note.reason


def _cell(text: str) -> TableCell:
    return TableCell(inline=[InlineRun(text=text)])


def _document(side: str | None, rows: int = 2) -> Document:
    table = Element(
        type=ElementType.TABLE, content="cells", order=0,
        table=TableContent(
            rows=[TableRow(cells=[_cell(f"r{row}a"), _cell(f"r{row}b")]) for row in range(rows)],
            columnWidthsCm=[2.5, 2.5],
            hasHeaderRow=False,
            floating=TableFloat(xAlign=side, side=side) if side else None,
        ),
    )
    paragraph = Element(type=ElementType.PARAGRAPH, content=TEXT, inline=[InlineRun(text=TEXT)], order=1)
    document = Document(elements=[table, paragraph])
    recompute_styles(document)
    return document


def _lines(pdf: bytes) -> list:
    from pdfminer.high_level import extract_pages
    from pdfminer.layout import LTTextBox, LTTextLine

    return [line for page in extract_pages(io.BytesIO(pdf)) for box in page if isinstance(box, LTTextBox) for line in box if isinstance(line, LTTextLine)]


def _text_beside(pdf: bytes) -> tuple[float, float, float]:
    """The table's first cell's top, and the left edge and top of the paragraph's first line."""
    lines = _lines(pdf)
    cell = next(line for line in lines if line.get_text().startswith("r0a"))
    text = max((line for line in lines if line.get_text().startswith("Text that flows")), key=lambda line: line.y1)
    return cell.y1, text.x0, text.y1


def test_a_pdf_wraps_the_text_around_a_floating_table_at_its_side():
    left_margin = 2 * 72 / 2.54
    cell_top, text_left, text_top = _text_beside(build_pdf(_document("left")))
    assert text_left > left_margin + 5 * 72 / 2.54  # right of the 5 cm table
    assert abs(text_top - cell_top) < 20  # beside it, not below

    cell_top, text_left, text_top = _text_beside(build_pdf(_document("right")))
    assert text_left - left_margin < 10 and abs(text_top - cell_top) < 20  # at the margin, the table on the right

    cell_top, text_left, text_top = _text_beside(build_pdf(_document(None)))
    assert text_top < cell_top - 20  # in line: the text below it


def test_a_pdf_draws_a_floating_table_too_tall_to_stay_beside_its_text_in_line():
    pdf = build_pdf(_document("left", rows=60))  # longer than a page

    lines = _lines(pdf)
    assert any(line.get_text().startswith("r59a") for line in lines)  # every row drawn, over pages
    first = max((line for line in lines if line.get_text().startswith("Text that flows")), key=lambda line: line.y1)
    assert first.x0 - 2 * 72 / 2.54 < 10  # the text after it at the margin, below it


def test_the_text_keeps_the_tables_own_distance_from_it():
    from app.export.pdf_export import _placement_of

    document = _document("right")
    document.elements[0].table.floating = TableFloat(xAlign="right", side="right", leftFromTextCm=0.8, rightFromTextCm=0.1, bottomFromTextCm=0.4)

    placement = _placement_of(document.elements[0])

    assert (placement.side, placement.distanceLeftCm, placement.distanceRightCm, placement.distanceBottomCm) == ("right", 0.8, 0.1, 0.4)
