"""Clean copy (tracker REV-005, brief §59): a new document with what was chosen taken out --
comments, tracked changes accepted, hidden text, the Word file's metadata, formatting set apart from
the styles -- each action only when asked, the original left as it is."""

import io

import pytest
from docx import Document as DocxDocument
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Pt
from fastapi.testclient import TestClient
from PIL import Image

from app.main import app
from app.models.document import MarkType

_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_W = nsdecls("w")


def _png() -> bytes:
    out = io.BytesIO()
    Image.new("RGB", (40, 30), (200, 30, 30)).save(out, format="PNG")
    return out.getvalue()


def _word_file(*, tracked: bool = False) -> bytes:
    word = DocxDocument()
    word.core_properties.author = "Ana Ivanova"
    word.core_properties.keywords = "draft, internal"
    noted = word.add_paragraph("A paragraph with a comment.")
    word.add_comment(noted.runs[0], text="Is this right?", author="Reviewer", initials="RV")
    hidden = word.add_paragraph("Shown text.")
    hidden._p.append(parse_xml(f'<w:r {_W}><w:rPr><w:vanish/></w:rPr><w:t xml:space="preserve"> secret note</w:t></w:r>'))
    styled = word.add_paragraph()
    run = styled.add_run("Set in its own font.")
    run.font.name, run.font.size = "Georgia", Pt(15)  # the whole paragraph's: kept as the block's own look
    mixed = word.add_paragraph("Mostly plain, ")
    mixed.add_run("one word").font.name = "Verdana"  # one run's: kept as that run's (a monospace one would be code)
    word.add_paragraph().add_run().add_picture(io.BytesIO(_png()))
    if tracked:
        changed = word.add_paragraph("Tracked ")
        changed._p.append(parse_xml(f'<w:ins {_W} w:id="9" w:author="Ana" w:date="2026-01-01T00:00:00Z"><w:r><w:t>insertion</w:t></w:r></w:ins>'))
    out = io.BytesIO()
    word.save(out)
    return out.getvalue()


@pytest.fixture
def client(api_db):
    test_client = TestClient(app, base_url="https://testserver")
    assert test_client.post("/api/v1/auth/register", json={"email": "clean@example.com", "password": "long enough password"}).status_code == 201
    return test_client


def _upload(client, data: bytes) -> dict:
    response = client.post("/api/v1/documents/upload", files={"file": ("report.docx", data, _DOCX)})
    assert response.status_code == 201, response.text
    return response.json()


def _fragments(document: dict, kind: str) -> list:
    found = []

    def walk(elements):
        for element in elements or []:
            found.extend(f for f in ((element.get("preservedAttributes") or {}).get("ooxml") or []) if f.get("kind") == kind)
            for item in element.get("listItems") or []:
                found.extend(f for f in ((item.get("preservedAttributes") or {}).get("ooxml") or []) if f.get("kind") == kind)
            walk(element.get("children"))

    walk(document["elements"])
    return found


def _marks(document: dict, kind: str) -> int:
    return sum(1 for element in document["elements"] for run in element.get("inline") or [] for mark in run["marks"] if mark["type"] == kind)


def test_every_action_taken_out_of_a_new_document_and_the_original_left_as_it_was(client):
    original = _upload(client, _word_file())
    assert _fragments(original, "comment") and _marks(original, MarkType.HIDDEN.value) and _marks(original, MarkType.TEXT_STYLE.value)
    assert original["metadata"]["sourceProperties"]["author"] == "Ana Ivanova"
    options = {"removeComments": True, "removeHiddenText": True, "removeMetadata": True, "normaliseFormatting": True}

    response = client.post(f"/api/v1/documents/{original['id']}/clean-copy", json=options)

    assert response.status_code == 201, response.text
    copy, summary = response.json()["document"], response.json()["summary"]
    assert copy["id"] != original["id"] and copy["metadata"]["title"].endswith("(clean copy)")
    assert summary["commentsRemoved"] == 1 and summary["hiddenRunsRemoved"] == 1 and summary["originalFileLeftOut"] is True
    assert set(summary["metadataRemoved"]) >= {"author", "keywords"} and summary["formattingRemoved"] > 0
    assert not _fragments(copy, "comment") and not _marks(copy, MarkType.HIDDEN.value) and not _marks(copy, MarkType.TEXT_STYLE.value)
    assert "secret note" not in " ".join(element["content"] for element in copy["elements"])
    assert copy["metadata"]["sourceProperties"] is None and copy["sourcePackage"] is None
    ids = {element["id"] for element in copy["elements"]}
    assert not [rule for rule in copy["formattingRules"] if rule["target"] in ids]

    [picture] = [element for element in copy["elements"] if element["type"] == "image"]
    [before] = [element for element in original["elements"] if element["type"] == "image"]
    assert picture["image"]["assetId"] and picture["image"]["assetId"] != before["image"]["assetId"]  # its own

    exported = client.get(f"/api/v1/documents/{copy['id']}/export/docx")
    assert exported.status_code == 200
    word = DocxDocument(io.BytesIO(exported.content))
    xml = word.element.xml
    assert "commentReference" not in xml and "<w:vanish/>" not in xml and word.core_properties.author == ""
    assert len(word.inline_shapes) == 1

    again = client.get(f"/api/v1/documents/{original['id']}").json()
    assert again["revision"] == original["revision"] and _fragments(again, "comment") and _marks(again, MarkType.HIDDEN.value)


def test_only_what_is_chosen_is_taken_out(client):
    original = _upload(client, _word_file())
    copy = client.post(f"/api/v1/documents/{original['id']}/clean-copy", json={"removeComments": True}).json()["document"]
    assert not _fragments(copy, "comment")
    assert _marks(copy, MarkType.HIDDEN.value) and _marks(copy, MarkType.TEXT_STYLE.value)  # kept: not asked
    ids = {element["id"] for element in copy["elements"]}
    assert [rule for rule in copy["formattingRules"] if rule["target"] in ids]  # the blocks' own looks too
    assert copy["metadata"]["sourceProperties"]["author"] == "Ana Ivanova"


def test_nothing_chosen_or_tracked_changes_left_kept_is_refused(client):
    plain = _upload(client, _word_file())
    nothing = client.post(f"/api/v1/documents/{plain['id']}/clean-copy", json={})
    assert nothing.status_code == 422 and nothing.json()["code"] == "nothing_chosen"

    tracked = _upload(client, _word_file(tracked=True))
    assert tracked["trackedChanges"] == "kept"
    refused = client.post(f"/api/v1/documents/{tracked['id']}/clean-copy", json={"removeComments": True})
    assert refused.status_code == 422 and refused.json()["code"] == "tracked_changes"
    accepted = client.post(f"/api/v1/documents/{tracked['id']}/clean-copy", json={"acceptTrackedChanges": True})
    assert accepted.status_code == 201 and accepted.json()["document"]["trackedChanges"] is None
    assert accepted.json()["summary"]["trackedChangesAccepted"] is True
    assert "insertion" in " ".join(element["content"] for element in accepted.json()["document"]["elements"])
    listed = client.get("/api/v1/documents").json()
    assert len(listed["items"]) == 3  # the two originals and the one copy
