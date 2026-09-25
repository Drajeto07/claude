import base64
import io
import json

import pytest
from docx import Document as DocxDocument
from fastapi.testclient import TestClient
from PIL import Image as PILImage
from pypdf import PdfReader
from sqlalchemy import func, select
from sqlalchemy.orm import Session as OrmSession

from app.db.models import Document as DocumentRow
from app.db.models import DocumentAsset
from app.main import app
from tests.helpers import error_body

_DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _png() -> bytes:
    buf = io.BytesIO()
    PILImage.new("RGB", (8, 8), "blue").save(buf, format="PNG")
    return buf.getvalue()


def _docx_with_picture(png: bytes) -> bytes:
    doc = DocxDocument()
    doc.add_paragraph("Before the figure.")
    doc.add_picture(io.BytesIO(png))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _signed_in_client(email: str) -> TestClient:
    client = TestClient(app, base_url="https://testserver")
    assert client.post("/api/auth/register", json={"email": email, "password": "long enough password"}).status_code == 201
    return client


@pytest.fixture
def alice(api_db):
    return _signed_in_client("alice@example.com")


def _upload_picture_document(client: TestClient, png: bytes) -> dict:
    response = client.post("/api/documents/upload", files={"file": ("report.docx", _docx_with_picture(png), _DOCX_TYPE)})
    assert response.status_code == 201
    return response.json()


def _image_elements(document: dict) -> list[dict]:
    return [e for e in document["elements"] if e["type"] == "image"]


def _asset_count(api_db) -> int:
    with OrmSession(api_db) as db:
        return db.execute(select(func.count()).select_from(DocumentAsset)).scalar_one()


def test_uploaded_docx_picture_is_stored_as_an_asset_and_served_back(alice):
    png = _png()

    document = _upload_picture_document(alice, png)

    [image] = [e["image"] for e in _image_elements(document)]
    assert image["assetId"] and image["src"] == ""
    served = alice.get(f"/api/assets/{image['assetId']}")
    assert served.status_code == 200
    assert served.content == png
    assert served.headers["content-type"] == "image/png"
    assert served.headers["x-content-type-options"] == "nosniff"
    assert "sandbox" in served.headers["content-security-policy"]
    assert served.headers["cache-control"].startswith("private")


def test_no_image_bytes_are_left_in_the_stored_document(api_db, alice):
    document = _upload_picture_document(alice, _png())

    with OrmSession(api_db) as db:
        stored = json.dumps(db.get(DocumentRow, document["id"]).data)
    assert "base64" not in stored


def test_assets_are_private_to_their_workspace(api_db, alice):
    asset_id = _image_elements(_upload_picture_document(alice, _png()))[0]["image"]["assetId"]
    bob = _signed_in_client("bob@example.com")
    anonymous = TestClient(app, base_url="https://testserver")

    assert bob.get(f"/api/assets/{asset_id}").status_code == 404
    assert error_body(bob.get(f"/api/assets/{asset_id}")) == error_body(bob.get("/api/assets/does-not-exist"))
    assert anonymous.get(f"/api/assets/{asset_id}").status_code == 401


def test_pasted_image_is_stored_once_not_on_every_autosave(api_db, alice):
    document = alice.post("/api/documents", json={"text": "# Notes\n\nSome text to start with here."}).json()
    pasted = {"type": "image", "content": "", "order": 99, "image": {"src": "data:image/png;base64," + base64.b64encode(_png()).decode()}}

    saved = alice.put(f"/api/documents/{document['id']}/content", json={"elements": document["elements"] + [pasted]})

    assert saved.status_code == 200
    [image] = [e["image"] for e in _image_elements(saved.json())]
    assert image["assetId"] and image["src"] == ""
    # The editor re-sends what the server returned; that must not store it again.
    alice.put(f"/api/documents/{document['id']}/content", json={"elements": saved.json()["elements"]})
    assert _asset_count(api_db) == 1


def test_inline_image_that_is_not_a_web_image_is_removed_and_reported(api_db, alice):
    document = alice.post("/api/documents", json={"text": "# Notes\n\nSome text to start with here."}).json()
    junk = {"type": "image", "content": "", "order": 99, "image": {"src": "data:text/html;base64,PHNjcmlwdD4="}}

    saved = alice.put(f"/api/documents/{document['id']}/content", json={"elements": document["elements"] + [junk]}).json()

    assert _image_elements(saved) == []
    assert any("was removed" in note for note in saved["unsupportedFeatures"])
    assert _asset_count(api_db) == 0


def test_docx_and_pdf_exports_embed_the_stored_picture(alice):
    document = _upload_picture_document(alice, _png())

    docx_bytes = alice.get(f"/api/documents/{document['id']}/export/docx").content
    pdf_bytes = alice.get(f"/api/documents/{document['id']}/export/pdf").content

    assert len(DocxDocument(io.BytesIO(docx_bytes)).inline_shapes) == 1
    assert sum(len(page.images) for page in PdfReader(io.BytesIO(pdf_bytes)).pages) == 1


def test_a_planted_foreign_asset_id_does_not_leak_into_an_export(api_db, alice):
    alices_asset = _image_elements(_upload_picture_document(alice, _png()))[0]["image"]["assetId"]
    bob = _signed_in_client("bob@example.com")
    bobs_document = bob.post("/api/documents", json={"text": "# Mine\n\nBob's own paragraph of text."}).json()
    planted = {"type": "image", "content": "", "order": 99, "image": {"src": "", "assetId": alices_asset}}
    bob.put(f"/api/documents/{bobs_document['id']}/content", json={"elements": bobs_document["elements"] + [planted]})

    export = bob.get(f"/api/documents/{bobs_document['id']}/export/docx")

    assert export.status_code == 200
    assert len(DocxDocument(io.BytesIO(export.content)).inline_shapes) == 0
