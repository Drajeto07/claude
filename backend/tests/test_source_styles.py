"""A Word file's own look beats the app's defaults (tracker FMT-004, audit AUD-03): what a
style leaves unset is what Word draws there -- no bold, no spacing, single lines, no
indent -- not what the render specification gives a document of its own. A template
still restyles the document."""

import io

from docx import Document as DocxDocument
from docx.oxml.ns import qn

from app.export.docx_export import build_docx
from app.formatting.engine import apply_formatting
from app.formatting.templates import BUILTIN_TEMPLATES
from app.parsers.docx import parse_docx


def _save(word) -> bytes:
    buffer = io.BytesIO()
    word.save(buffer)
    return buffer.getvalue()


def _unset(style, *tags: str) -> None:
    """Takes what `tags` set out of a Word style, so Word falls back on its own defaults."""
    for holder in (style.element.find(qn("w:pPr")), style.element.find(qn("w:rPr"))):
        for tag in tags:
            for child in holder.findall(qn(tag)) if holder is not None else []:
                holder.remove(child)


def _plain_word_file() -> bytes:
    """A heading style that isn't bold and has no spacing; body text, a quote and a
    table with no spacing either -- the document defaults set none."""
    word = DocxDocument()
    defaults = word.styles.element.find(qn("w:docDefaults")).find(qn("w:pPrDefault")).find(qn("w:pPr"))
    for spacing in defaults.findall(qn("w:spacing")):
        defaults.remove(spacing)
    _unset(word.styles["Heading 1"], "w:b", "w:bCs", "w:spacing")
    word.add_heading("A plain heading", level=1)
    word.add_paragraph("Body text.")
    word.add_paragraph("A quotation.", style="Quote")
    word.add_table(rows=1, cols=1).cell(0, 0).text = "A cell."
    word.add_paragraph("After the table.")
    return _save(word)


def test_a_heading_style_that_isnt_bold_stays_regular_and_unspaced():
    document = parse_docx(_plain_word_file(), "plain.docx")
    heading = document.resolvedStyles["Heading 1"]

    assert heading["font-weight"] == "normal"  # not the app's bold heading
    assert (heading["margin-top"], heading["margin-bottom"]) == ("0pt", "0pt")  # not 18 and 6 pt

    again = parse_docx(build_docx(document), "again.docx")
    assert again.resolvedStyles["Heading 1"]["font-weight"] == "normal"  # and so in the Word export


def test_what_the_file_leaves_unset_is_what_word_draws():
    styles = parse_docx(_plain_word_file(), "plain.docx").resolvedStyles

    assert styles["Paragraph"]["margin-bottom"] == "0pt"  # no spacing after, not the app's 8 pt
    assert styles["Quote"]["font-style"] == "italic"  # the file's Quote style
    assert styles["Quote"].get("margin-left", "0cm") == "0cm"  # without the app's 1 cm indent
    assert styles["Table"]["margin-bottom"] == "0pt"  # Word leaves no gap after a table


def test_a_template_still_restyles_an_imported_document():
    document = parse_docx(_plain_word_file(), "plain.docx")
    template = BUILTIN_TEMPLATES["academic-default"]

    apply_formatting(document, template_id=template.id, template_rules=template.rules, instruction_rules=[])

    assert document.resolvedStyles["Paragraph"]["font-family"] == "Times New Roman"
    assert document.resolvedStyles["Heading 1"]["font-weight"] == "bold"
