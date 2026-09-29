"""Word list numbering (tracker DOCX-016, FID-002): a list keeps its top level's
number format, where it starts, where it restarts and how it goes on counting after
an interruption, and each of its levels -- format, label, start, indent, legal
numbering, restart, suffix, bullet -- through a Word export; what isn't shown yet,
or can't be kept, is reported instead of silently renumbered."""

import io
import zipfile

from docx import Document as DocxDocument
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from app.export.docx_export import build_docx
from app.export.package_check import package_problems
from app.fidelity.report import FidelityPolicy
from app.formatting.list_numbering import Counters, format_number, level_label
from app.models.document import Element, ElementType, InlineRun, ListItem, ListLevel, ListNumbering
from app.parsers.docx import parse_docx


def _numbering(document, abstract_id: int, levels: list[tuple[str, str, int]]) -> None:
    """Adds a multi-level definition: (numFmt, lvlText, start) per level."""
    xml = "".join(
        f'<w:lvl w:ilvl="{ilvl}"><w:start w:val="{start}"/><w:numFmt w:val="{fmt}"/><w:lvlText w:val="{text}"/></w:lvl>'
        for ilvl, (fmt, text, start) in enumerate(levels)
    )
    numbering = document.part.numbering_part.element
    abstract = parse_xml(f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{abstract_id}">{xml}</w:abstractNum>')
    first_num = numbering.find(f"{{{numbering.nsmap['w']}}}num")
    if first_num is not None:
        first_num.addprevious(abstract)
    else:
        numbering.append(abstract)


def _num(document, num_id: int, abstract_id: int, restart_at: int | None = None) -> None:
    override = f'<w:lvlOverride w:ilvl="0"><w:startOverride w:val="{restart_at}"/></w:lvlOverride>' if restart_at is not None else ""
    document.part.numbering_part.element.append(
        parse_xml(f'<w:num {nsdecls("w")} w:numId="{num_id}"><w:abstractNumId w:val="{abstract_id}"/>{override}</w:num>')
    )


def _item(document, text: str, num_id: int, level: int = 0) -> None:
    paragraph = document.add_paragraph(text, style="List Paragraph")
    paragraph._p.get_or_add_pPr().append(
        parse_xml(f'<w:numPr {nsdecls("w")}><w:ilvl w:val="{level}"/><w:numId w:val="{num_id}"/></w:numPr>')
    )


def _import(build):
    document = DocxDocument()
    build(document)
    buffer = io.BytesIO()
    document.save(buffer)
    return parse_docx(buffer.getvalue(), "numbering.docx")


def _lists(document):
    return [element for element in document.elements if element.type == "list"]


def _features(document):
    return {item.feature for item in document.importReport.items}


def test_the_top_levels_format_and_start_are_kept():
    def build(document):
        _numbering(document, 90, [("upperRoman", "%1.", 1)])
        _numbering(document, 91, [("lowerLetter", "%1.", 5)])
        _num(document, 90, 90)
        _num(document, 91, 91)
        _item(document, "first roman", 90)
        document.add_paragraph("between")
        _item(document, "e is the fifth letter", 91)

    roman, letters = _lists(_import(build))
    assert (roman.numbering.start, roman.numbering.format) == (1, "upperRoman")
    assert (letters.numbering.start, letters.numbering.format) == (5, "lowerLetter")


def test_numbering_goes_on_after_an_interruption_and_restarts_where_word_restarts_it():
    def build(document):
        _numbering(document, 92, [("decimal", "%1.", 1)])
        _num(document, 92, 92)
        _num(document, 93, 92, restart_at=1)
        for text in ("a1", "a2", "a3"):
            _item(document, text, 92)
        document.add_paragraph("An interrupting paragraph.")
        _item(document, "a4", 92)  # the same list: Word shows 4
        document.add_paragraph("Another paragraph.")
        _item(document, "b1", 93)  # a new instance restarting at 1

    first, continued, restarted = _lists(_import(build))
    assert first.numbering is None  # 1, 2, 3 needs no numbering of its own
    assert continued.numbering.start == 4
    assert restarted.numbering is None


def test_an_empty_numbered_item_still_takes_its_number():
    def build(document):
        _numbering(document, 94, [("decimal", "%1.", 1)])
        _num(document, 94, 94)
        _item(document, "one", 94)
        _item(document, "", 94)  # Word shows "2." with nothing after it
        _item(document, "three", 94)

    document = _import(build)
    assert _lists(document)[1].numbering.start == 3
    assert "docx.list_numbering.empty_item" in _features(document)


def test_what_isnt_shown_yet_or_cant_be_kept_is_reported():
    def build(document):
        _numbering(document, 95, [("decimal", "Чл. %1.", 1), ("decimal", "%1.%2.", 1)])
        _numbering(document, 96, [("ordinal", "%1", 1)])
        _num(document, 95, 95)
        _num(document, 96, 96)
        _item(document, "Article", 95)
        _item(document, "Clause", 95, level=1)
        document.add_paragraph("between")
        _item(document, "first", 96)

    document = _import(build)
    items = {item.feature: item for item in document.importReport.items}
    for kept in ("docx.list_numbering.label", "docx.list_numbering.multilevel"):  # in a Word export; not shown here yet
        assert items[kept].policy == FidelityPolicy.DETECTED_NOT_EDITABLE and not items[kept].contentChanged
    assert items["docx.list_numbering.format"].contentChanged  # "1st" becomes "1"
    labels = [level.text for level in _lists(document)[0].numbering.levels]
    assert labels == ["Чл. %1.", "%1.%2."]


def test_a_list_keeps_its_format_start_and_continuation_through_a_round_trip():
    def build(document):
        _numbering(document, 97, [("lowerRoman", "%1.", 3)])
        _num(document, 97, 97)
        _item(document, "iii", 97)
        _item(document, "iv", 97)
        document.add_paragraph("interruption")
        _item(document, "v", 97)

    imported = _import(build)
    again = parse_docx(build_docx(imported), "again.docx")

    assert [(el.numbering.start, el.numbering.format) for el in _lists(again)] == [(3, "lowerRoman"), (5, "lowerRoman")]


def _level(ilvl: int, fmt: str, text: str, *, start: int = 1, left: int | None = None, hanging: int | None = None, extra: str = "", font: str = "") -> str:
    indent = f'<w:pPr><w:ind w:left="{left}" w:hanging="{hanging}"/></w:pPr>' if left is not None else ""
    fonts = f'<w:rPr><w:rFonts w:ascii="{font}" w:hAnsi="{font}"/></w:rPr>' if font else ""
    return f'<w:lvl w:ilvl="{ilvl}"><w:start w:val="{start}"/><w:numFmt w:val="{fmt}"/>{extra}<w:lvlText w:val="{text}"/>{indent}{fonts}</w:lvl>'


def _definition(document, abstract_id: int, levels: list[str], link: str = "") -> None:
    numbering = document.part.numbering_part.element
    abstract = parse_xml(f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{abstract_id}">{link}{"".join(levels)}</w:abstractNum>')
    first_num = numbering.find(f"{{{numbering.nsmap['w']}}}num")
    if first_num is not None:
        first_num.addprevious(abstract)
    else:
        numbering.append(abstract)


_ARTICLES = [
    _level(0, "decimal", "Чл. %1.", left=720, hanging=720),
    _level(1, "russianLower", "%2)", left=1440, hanging=360),
    _level(2, "decimal", "%1.%2.%3.", left=2160, hanging=720, extra='<w:lvlRestart w:val="1"/><w:isLgl/><w:suff w:val="space"/>'),
]


def _articles(document) -> None:
    _definition(document, 110, _ARTICLES)
    _num(document, 110, 110)
    _item(document, "Scope", 110)
    _item(document, "first point", 110, level=1)
    _item(document, "a sub-clause", 110, level=2)
    _item(document, "Terms", 110)


def test_each_level_keeps_its_format_label_start_and_indent():
    [articles] = _lists(_import(_articles))

    first, second, third = articles.numbering.levels[:3]
    assert (first.format, first.text, first.indentCm, first.hangingCm) == ("decimal", "Чл. %1.", 1.27, 1.27)
    assert (second.format, second.text, second.indentCm, second.hangingCm) == ("russianLower", "%2)", 2.54, 0.64)
    assert (third.text, third.legal, third.restartAfter, third.suffix) == ("%1.%2.%3.", True, 1, "space")


def test_a_lists_own_levels_come_back_from_a_word_export():
    imported = _import(_articles)

    exported = build_docx(imported)
    [again] = _lists(parse_docx(exported, "again.docx"))

    assert package_problems(exported) == []
    assert again.numbering.levels[:3] == _lists(imported)[0].numbering.levels[:3]
    assert [item.level for item in again.listItems] == [0, 1, 2, 0]


def test_a_list_taking_its_levels_from_a_numbering_style_has_them():
    def build(document):
        _definition(document, 111, [_level(0, "upperRoman", "Part %1", left=720, hanging=360)], link='<w:styleLink w:val="PartsList"/>')
        _definition(document, 112, [], link='<w:numStyleLink w:val="PartsList"/>')
        _num(document, 112, 112)
        _item(document, "One part", 112)

    [parts] = _lists(_import(build))

    assert (parts.numbering.format, parts.numbering.levels[0].text) == ("upperRoman", "Part %1")


def test_bullets_drawn_from_symbol_fonts_are_the_characters_they_show():
    def build(document):
        _definition(
            document,
            113,
            [
                _level(0, "bullet", "\uf0d8", left=720, hanging=360, font="Wingdings"),
                _level(1, "bullet", "–", left=1440, hanging=360, font="Arial"),
                _level(2, "bullet", "\uf0e0", left=2160, hanging=360, font="Wingdings"),
            ],
        )
        _num(document, 113, 113)
        _item(document, "arrow", 113)
        _item(document, "dash", 113, level=1)
        _item(document, "unknown", 113, level=2)

    document = _import(build)
    [bullets] = _lists(document)

    assert not bullets.ordered and [level.text for level in bullets.numbering.levels] == ["➢", "–", "•"]
    features = _features(document)
    assert "docx.list_numbering.bullet" in features and "docx.list_numbering.bullet_font" in features
    exported = build_docx(document)
    with zipfile.ZipFile(io.BytesIO(exported)) as package:
        assert "➢" in package.read("word/numbering.xml").decode("utf-8")


def test_the_usual_levels_need_no_numbering_of_their_own():
    def build(document):
        _definition(document, 114, [_level(0, "bullet", "•"), _level(1, "bullet", "◦")])
        _num(document, 114, 114)
        _item(document, "plain", 114)
        _item(document, "plainer", 114, level=1)

    [plain] = _lists(_import(build))

    assert plain.numbering is None


def test_a_label_is_written_as_text_never_as_markup():
    label = 'Art. "%1" <&>'
    element = Element(
        type=ElementType.LIST,
        content="One",
        ordered=True,
        order=0,
        listItems=[ListItem(inline=[InlineRun(text="One")])],
        numbering=ListNumbering(levels=[ListLevel(text=label)]),
    )
    from app.models.document import Document

    exported = build_docx(Document(elements=[element]))

    assert package_problems(exported) == []
    [again] = _lists(parse_docx(exported, "again.docx"))
    assert again.numbering.levels[0].text == label


def test_numbers_count_and_restart_as_word_counts_them():
    assert [format_number(value, "russianLower") for value in (1, 9, 10, 28, 29)] == ["а", "и", "к", "я", "аа"]
    assert (format_number(3, "russianUpper"), format_number(7, "decimalZero"), format_number(12, "decimalZero")) == ("В", "07", "12")
    counters = Counters([1, 1, 1], [None, None, 1])
    assert counters.advance(0) == [1] and counters.advance(1) == [1, 1] and counters.advance(2) == [1, 1, 1]
    assert counters.advance(1) == [1, 2]
    assert counters.advance(2) == [1, 2, 2]  # restarts only after level 1
    assert counters.advance(0) == [2] and counters.advance(2) == [2, 0, 1]  # a skipped level shows 0, as Word's 1.0.1
    assert level_label("%1.%2.", [2, 3], ["upperRoman", "lowerLetter"]) == "II.c."
    assert level_label("%1.%2.", [2, 3], ["upperRoman", "lowerLetter"], legal=True) == "2.3."


def test_instances_of_one_definition_number_on_together_unless_one_restarts():
    def build(document):
        _numbering(document, 115, [("decimal", "%1)", 1)])
        _num(document, 115, 115)
        _num(document, 116, 115)  # another instance, no restart: Word numbers on
        _num(document, 117, 115, restart_at=1)  # "Restart numbering"
        _item(document, "one", 115)
        _item(document, "two", 115)
        document.add_paragraph("between")
        _item(document, "three", 116)
        document.add_paragraph("between")
        _item(document, "one again", 117)
        _item(document, "two again", 117)

    starts = [element.numbering.start for element in _lists(_import(build))]

    assert starts == [1, 3, 1]
