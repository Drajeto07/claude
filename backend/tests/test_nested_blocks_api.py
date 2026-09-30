"""Blocks nested inside table cells, list items and quotes, through the API the
editor saves with (tracker CORE-003, DOCX-001, PDF-001): stored as sent,
pictures in them stored as assets, and every word of them in both exports."""

import base64
import io

import pytest
from docx import Document as DocxDocument
from docx.oxml.ns import qn
from fastapi.testclient import TestClient
from PIL import Image as PILImage
from pypdf import PdfReader
from sqlalchemy import func, select
from sqlalchemy.orm import Session as OrmSession

from app.db.models import DocumentAsset
from app.main import app
from app.models.document import MAX_BLOCK_DEPTH


def _png() -> bytes:
    buffer = io.BytesIO()
    PILImage.new("RGB", (40, 20), "green").save(buffer, format="PNG")
    return buffer.getvalue()


PNG_URI = "data:image/png;base64," + base64.b64encode(_png()).decode()


def _run(text: str) -> dict:
    return {"text": text, "marks": []}


def _paragraph(text: str, order: int = 0) -> dict:
    return {"type": "paragraph", "content": text, "inline": [_run(text)], "order": order}


def _list(texts: list[str], order: int = 0, **extra) -> dict:
    return {"type": "list", "content": "\n".join(texts), "order": order, "listItems": [{"inline": [_run(t)], "level": 0} for t in texts], **extra}


def _nested_elements() -> list[dict]:
    """The shapes the editor sends: a table whose cells hold a list, code and a
    picture; a list whose item holds code and a numbered sub-list; a quote
    holding a list; a numbered list starting at 5 in lower-case roman."""
    cell_blocks = [
        _paragraph("Items:", 0),
        _list(["alpha", "beta"], 1),
        {"type": "code_block", "content": "make build", "order": 2},
        {"type": "image", "content": "", "order": 3, "image": {"src": PNG_URI, "alt": "Build chart"}},
    ]
    table = {
        "type": "table",
        "content": "Items: alpha beta make build | plain",
        "order": 0,
        "table": {
            "rows": [
                {"cells": [{"inline": [_run("Items:\nalpha\nbeta\nmake build")], "blocks": cell_blocks}, {"inline": [_run("plain cell")]}]}
            ]
        },
    }
    steps = _list(["Install", "Run"], 1)
    steps["listItems"][0]["blocks"] = [
        {"type": "code_block", "content": "npm ci", "order": 0},
        _list(["first sub", "second sub"], 1, ordered=True),
    ]
    quote = {"type": "quote", "content": "Remember\nback up", "order": 2, "children": [_paragraph("Remember", 0), _list(["back up"], 1)]}
    numbered = _list(["fifth", "sixth"], 3, ordered=True, numbering={"start": 5, "format": "lowerRoman"})
    return [table, steps, quote, numbered]


WORDS = ["Items:", "alpha", "beta", "make build", "plain cell", "Install", "npm ci", "first sub", "second sub", "Run", "Remember", "back up", "fifth", "sixth"]


@pytest.fixture
def alice(api_db):
    client = TestClient(app, base_url="https://testserver")
    assert client.post("/api/v1/auth/register", json={"email": "alice@example.com", "password": "long enough password"}).status_code == 201
    return client


def _saved(client: TestClient, elements: list[dict]) -> dict:
    document = client.post("/api/v1/documents", json={"text": "# Start\n\nSome text to start with here."}).json()
    response = client.put(f"/api/v1/documents/{document['id']}/content", json={"elements": elements})
    assert response.status_code == 200, response.text
    return response.json()


def _asset_count(api_db) -> int:
    with OrmSession(api_db) as db:
        return db.execute(select(func.count()).select_from(DocumentAsset)).scalar_one()


def test_nested_blocks_are_stored_as_sent_and_their_picture_as_an_asset(api_db, alice):
    saved = _saved(alice, _nested_elements())

    table, steps, quote, numbered = saved["elements"]
    cell = table["table"]["rows"][0]["cells"][0]
    assert [block["type"] for block in cell["blocks"]] == ["paragraph", "list", "code_block", "image"]
    image = cell["blocks"][3]["image"]
    assert image["assetId"] and image["src"] == "" and image["alt"] == "Build chart"
    assert [block["type"] for block in steps["listItems"][0]["blocks"]] == ["code_block", "list"]
    assert [child["type"] for child in quote["children"]] == ["paragraph", "list"]
    assert numbered["numbering"] == {"start": 5, "format": "lowerRoman", "levels": None}
    # Reloaded from the database, the same.
    assert alice.get(f"/api/v1/documents/{saved['id']}").json()["elements"] == saved["elements"]
    # The editor sends back what it got: the picture is not stored again.
    alice.put(f"/api/v1/documents/{saved['id']}/content", json={"elements": saved["elements"]})
    assert _asset_count(api_db) == 1


def test_a_nested_inline_image_that_is_not_a_web_image_is_removed_and_reported(api_db, alice):
    item = _list(["with junk"], 0)
    item["listItems"][0]["blocks"] = [{"type": "image", "content": "", "order": 0, "image": {"src": "data:text/html;base64,PHNjcmlwdD4="}}]

    saved = _saved(alice, [item])

    assert saved["elements"][0]["listItems"][0]["blocks"] is None
    assert any("was removed" in note for note in saved["unsupportedFeatures"])
    assert _asset_count(api_db) == 0


def test_nesting_deeper_than_the_cap_is_refused(alice):
    element = _paragraph("deepest")
    for _ in range(MAX_BLOCK_DEPTH + 1):
        element = {"type": "table", "content": "", "order": 0, "table": {"rows": [{"cells": [{"inline": [], "blocks": [element]}]}]}}
    document = alice.post("/api/v1/documents", json={"text": "# Start\n\nSome text to start with here."}).json()

    response = alice.put(f"/api/v1/documents/{document['id']}/content", json={"elements": [element]})

    assert response.status_code == 422


def _docx_text(docx_bytes: bytes) -> str:
    body = DocxDocument(io.BytesIO(docx_bytes)).element.body
    return "\n".join("".join(t.text or "" for t in p.iter(qn("w:t"))) for p in body.iter(qn("w:p")))


def test_the_word_export_keeps_every_nested_block(alice):
    saved = _saved(alice, _nested_elements())

    docx_bytes = alice.get(f"/api/v1/documents/{saved['id']}/export/docx").content

    text = _docx_text(docx_bytes)
    positions = [text.find(word) for word in WORDS]
    assert -1 not in positions, [word for word, at in zip(WORDS, positions) if at == -1]
    assert positions == sorted(positions), "blocks come out in reading order"
    word = DocxDocument(io.BytesIO(docx_bytes))
    [picture] = word.inline_shapes
    assert picture._inline.docPr.get("descr") == "Build chart"
    cell = word.tables[0].cell(0, 0)
    styles = [p.style.name for p in cell.paragraphs]
    assert styles[:4] == ["Table Text", "List Bullet", "List Bullet", "Code"], styles
    assert cell.paragraphs[0].text == "Items:", "no empty paragraph left at the top of the cell"


def test_the_word_export_numbers_a_list_from_its_start_in_its_format(alice):
    saved = _saved(alice, [_list(["fifth", "sixth"], 0, ordered=True, numbering={"start": 5, "format": "lowerRoman"})])

    word = DocxDocument(io.BytesIO(alice.get(f"/api/v1/documents/{saved['id']}/export/docx").content))

    paragraph = next(p for p in word.paragraphs if p.text == "fifth")
    num_id = paragraph._p.pPr.numPr.numId.val
    numbering = word.part.numbering_part.element
    num = next(n for n in numbering.findall(qn("w:num")) if n.get(qn("w:numId")) == str(num_id))
    start = num.find(qn("w:lvlOverride")).find(qn("w:startOverride")).get(qn("w:val"))
    abstract_id = num.find(qn("w:abstractNumId")).get(qn("w:val"))
    abstract = next(a for a in numbering.findall(qn("w:abstractNum")) if a.get(qn("w:abstractNumId")) == abstract_id)
    level_zero = abstract.find(qn("w:lvl"))
    assert start == "5"
    assert level_zero.find(qn("w:numFmt")).get(qn("w:val")) == "lowerRoman"


def test_the_pdf_export_keeps_every_nested_block(alice):
    saved = _saved(alice, _nested_elements())

    pdf = PdfReader(io.BytesIO(alice.get(f"/api/v1/documents/{saved['id']}/export/pdf").content))

    text = "\n".join(page.extract_text() for page in pdf.pages)
    missing = [word for word in WORDS if word not in text]
    assert not missing
    assert "v. fifth" in text and "vi. sixth" in text
    assert sum(len(page.images) for page in pdf.pages) == 1
