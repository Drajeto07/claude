"""Export, then import again: what the app can represent must come back the
same (корекции.docx: nothing is "supported" without a round trip). Plus the
PDF-only checks: Cyrillic text, page-number fields, spans."""

import base64
import io

from docx import Document as DocxDocument
from docx.oxml.ns import qn
from PIL import Image as PILImage
from pypdf import PdfReader

from app.export.docx_export import build_docx
from app.export.pdf_export import build_pdf
from app.formatting.engine import recompute_styles
from app.models.document import (
    Document,
    DocumentMetadata,
    Element,
    ElementType,
    FormattingProperty,
    FormattingRule,
    ImageContent,
    InlineRun,
    ListItem,
    Mark,
    MarkType,
    TableCell,
    TableContent,
    TableRow,
)
from app.parsers.docx import parse_docx


def _png() -> bytes:
    buffer = io.BytesIO()
    PILImage.new("RGB", (60, 30), "green").save(buffer, format="PNG")
    return buffer.getvalue()


def _rich_document() -> Document:
    image = ImageContent(src="data:image/png;base64," + base64.b64encode(_png()).decode())
    elements = [
        Element(type=ElementType.HEADING, content="Отчет", inline=[InlineRun(text="Отчет")], level=1, order=0),
        Element(
            type=ElementType.PARAGRAPH,
            content="H2O e=mc2 важно виж https://example.com",
            inline=[
                InlineRun(text="H"),
                InlineRun(text="2", marks=[Mark(type=MarkType.SUBSCRIPT)]),
                InlineRun(text="O e=mc"),
                InlineRun(text="2", marks=[Mark(type=MarkType.SUPERSCRIPT)]),
                InlineRun(text=" "),
                InlineRun(
                    text="важно",
                    marks=[Mark(type=MarkType.TEXT_STYLE, fontFamily="Georgia", fontSizePt=14, color="#C00000", backgroundColor="#FFFF00")],
                ),
                InlineRun(text=" виж "),
                InlineRun(text="https://example.com", marks=[Mark(type=MarkType.LINK, href="https://example.com")]),
            ],
            order=1,
        ),
        Element(type=ElementType.HORIZONTAL_RULE, content="", order=2),
        Element(
            type=ElementType.LIST,
            content="",
            listItems=[
                ListItem(inline=[InlineRun(text="Направено")], checked=True),
                ListItem(inline=[InlineRun(text="Предстои")], checked=False),
            ],
            order=3,
        ),
        Element(
            type=ElementType.TABLE,
            content="",
            table=TableContent(
                rows=[
                    TableRow(
                        cells=[
                            TableCell(inline=[InlineRun(text="Категория")], header=True, background="#D9EAF7"),
                            TableCell(inline=[InlineRun(text="Общо")], header=True, colspan=2, background="#D9EAF7"),
                        ]
                    ),
                    TableRow(
                        cells=[
                            TableCell(inline=[InlineRun(text="Компютри")], rowspan=2),
                            TableCell(inline=[InlineRun(text="10")]),
                            TableCell(inline=[InlineRun(text="20")]),
                        ]
                    ),
                    TableRow(cells=[TableCell(inline=[InlineRun(text="5")]), TableCell(inline=[InlineRun(text="7")])]),
                ],
                hasHeaderRow=True,
                alignments=["left", "center", "center"],
            ),
            order=4,
        ),
        Element(type=ElementType.CODE_BLOCK, content="def f():\n    return 1", order=5),
        Element(type=ElementType.PAGE_BREAK, content="", order=6),
        Element(type=ElementType.IMAGE, content="", image=image, order=7),
        Element(type=ElementType.PARAGRAPH, content="Край", inline=[InlineRun(text="Край")], order=8),
    ]
    document = Document(metadata=DocumentMetadata(title="Round trip"), elements=elements)
    document.settings.footer = "Страница {PAGE} от {NUMPAGES}"
    image_id = elements[7].id
    document.formattingRules = [
        FormattingRule(target="Document", property=FormattingProperty.FOOTER, value="Страница {PAGE} от {NUMPAGES}"),
        FormattingRule(target="Document", property=FormattingProperty.PAGE_SIZE, value="A4"),
        FormattingRule(target="Document", property=FormattingProperty.MARGIN_LEFT, value="3", unit="cm"),
        FormattingRule(target=image_id, property=FormattingProperty.IMAGE_ALIGNMENT, value="center"),
        FormattingRule(target=image_id, property=FormattingProperty.IMAGE_WIDTH, value="50", unit="%"),
    ]
    recompute_styles(document)
    return document


def _round_trip(document: Document) -> Document:
    return parse_docx(build_docx(document), "round-trip.docx")


def _marks(run) -> dict:
    return {mark.type: mark for mark in run.marks}


def test_structure_survives_a_docx_round_trip():
    back = _round_trip(_rich_document())

    assert [element.type for element in back.elements] == [
        ElementType.HEADING,
        ElementType.PARAGRAPH,
        ElementType.HORIZONTAL_RULE,
        ElementType.LIST,
        ElementType.TABLE,
        ElementType.CODE_BLOCK,
        ElementType.PAGE_BREAK,
        ElementType.IMAGE,
        ElementType.PARAGRAPH,
    ]
    assert back.elements[5].content == "def f():\n    return 1"


def test_character_formatting_survives_a_docx_round_trip():
    paragraph = _round_trip(_rich_document()).elements[1]
    by_text: dict[str, list[dict]] = {}
    for run in paragraph.inline:
        by_text.setdefault(run.text, []).append(_marks(run))

    assert [set(marks) for marks in by_text["2"]] == [{MarkType.SUBSCRIPT}, {MarkType.SUPERSCRIPT}]
    style = by_text["важно"][0][MarkType.TEXT_STYLE]
    assert (style.fontFamily, style.fontSizePt, style.color, style.backgroundColor) == ("Georgia", 14, "#C00000", "#FFFF00")
    assert by_text["https://example.com"][0][MarkType.LINK].href == "https://example.com"


def test_a_checklist_survives_a_docx_round_trip_as_word_checkboxes():
    document = _rich_document()
    exported = build_docx(document)
    assert b"w14:checkbox" in exported or "w14:checkbox" in _document_xml(exported)

    checklist = parse_docx(exported, "x.docx").elements[3]

    assert [(item.checked, item.inline[0].text) for item in checklist.listItems] == [(True, "Направено"), (False, "Предстои")]


def _list(items: list[tuple[str, int]], *, ordered: bool, order: int, checked: bool | None = None) -> Element:
    return Element(
        type=ElementType.LIST,
        content="",
        ordered=ordered,
        listItems=[ListItem(inline=[InlineRun(text=text)], level=level, checked=checked) for text, level in items],
        order=order,
    )


def _lists_document(*elements: Element) -> Document:
    document = Document(metadata=DocumentMetadata(title="Lists"), elements=list(elements))
    recompute_styles(document)
    return document


def test_nested_list_levels_survive_a_docx_round_trip():
    levels = [("Едно", 0), ("Под", 1), ("Още по-навътре", 2), ("Обратно", 1), ("Две", 0)]
    back = _round_trip(
        _lists_document(
            _list(levels, ordered=False, order=0),
            _list(levels, ordered=True, order=1),
            _list(levels, ordered=False, order=2, checked=False),
        )
    )

    # Three lists right after one another stay three lists.
    assert [(element.type, element.ordered) for element in back.elements] == [
        (ElementType.LIST, False),
        (ElementType.LIST, True),
        (ElementType.LIST, False),
    ]
    for element in back.elements:
        assert [(item.inline[0].text, item.level) for item in element.listItems] == levels
    assert {item.checked for item in back.elements[2].listItems} == {False}


def test_each_numbered_list_starts_again_at_one_in_word():
    exported = build_docx(
        _lists_document(
            _list([("А", 0), ("Б", 0)], ordered=True, order=0),
            Element(type=ElementType.PARAGRAPH, content="между", inline=[InlineRun(text="между")], order=1),
            _list([("В", 0), ("Г", 0)], ordered=True, order=2),
        )
    )
    readback = DocxDocument(io.BytesIO(exported))

    num_ids = [p._p.pPr.numPr.numId.val for p in readback.paragraphs if p._p.pPr is not None and p._p.pPr.numPr is not None]
    assert len(num_ids) == 4 and num_ids[0] == num_ids[1] != num_ids[2] == num_ids[3]
    numbering = readback.part.numbering_part.element
    for num_id in (num_ids[0], num_ids[2]):
        num = numbering.num_having_numId(num_id)
        assert num.xpath('./w:lvlOverride[@w:ilvl="0"]/w:startOverride/@w:val') == ["1"]
        abstract_id = num.abstractNumId.val
        abstract = numbering.xpath(f'./w:abstractNum[@w:abstractNumId="{abstract_id}"]')[0]
        assert len(abstract.findall(qn("w:lvl"))) == 9  # real levels, not python-docx's single-level lists


def test_table_spans_shading_and_alignment_survive_a_docx_round_trip():
    table = _round_trip(_rich_document()).elements[4].table

    first, second, third = table.rows
    assert [(cell.colspan, cell.rowspan, cell.background) for cell in first.cells] == [(1, 1, "#D9EAF7"), (2, 1, "#D9EAF7")]
    assert [(cell.colspan, cell.rowspan) for cell in second.cells] == [(1, 2), (1, 1), (1, 1)]
    assert [cell.inline[0].text for cell in third.cells] == ["5", "7"]
    assert table.alignments == ["left", "center", "center"]


def test_page_setup_footer_fields_and_picture_placement_survive_a_docx_round_trip():
    back = _round_trip(_rich_document())

    assert (back.settings.pageSize, back.settings.marginLeftCm) == ("A4", 3.0)
    assert back.settings.footer == "Страница {PAGE} от {NUMPAGES}"
    image = back.elements[7]
    css = back.resolvedStyles[image.styleRef]
    assert (css["margin-left"], css["margin-right"], css["width"]) == ("auto", "auto", "50%")


def test_pdf_has_real_cyrillic_text_and_filled_in_page_numbers():
    reader = PdfReader(io.BytesIO(build_pdf(_rich_document())))

    assert len(reader.pages) == 2
    first_page = reader.pages[0].extract_text()
    assert "Отчет" in first_page and "Компютри" in first_page
    assert "Страница 1 от 2" in first_page
    assert "Страница 2 от 2" in reader.pages[1].extract_text()


def test_pdf_leaves_out_a_page_number_footer_when_page_numbers_are_off():
    reader = PdfReader(io.BytesIO(build_pdf(_rich_document(), include_page_numbers=False)))

    assert "Страница" not in reader.pages[0].extract_text()


def _document_xml(docx_bytes: bytes) -> str:
    import zipfile

    return zipfile.ZipFile(io.BytesIO(docx_bytes)).read("word/document.xml").decode("utf-8")
