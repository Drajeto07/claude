"""Comment threads (tracker DOCX-021): which comment answers which, and which are
resolved, go back into Word with the comments -- copied from the original, written
anew into it, or into a new file. Word keeps them apart from the comments, in
commentsExtended, naming each comment by its last paragraph's paraId: the export
writes that part again for the comments the file has, since comments written anew
get new ids and paragraphs (measured in Word: before this, a reply came back as a
comment of its own and a resolved comment as open)."""

import io
import re
import zipfile
from pathlib import Path

import pytest
from docx import Document as DocxDocument
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml.ns import qn

from app.export.docx_export import build_docx
from app.export.package_check import package_problems
from app.export.provenance import stamp
from app.fidelity.report import ReportBuilder
from app.formatting.engine import recompute_styles
from app.models.document import Document, InlineRun
from app.parsers.docx import parse_docx
from app.parsers.docx_comments import comment_threads

A07 = Path(__file__).parent / "fixtures" / "word" / "a07-review.docx"
_EXTENDED = "http://schemas.microsoft.com/office/2011/relationships/commentsExtended"
_EXTENDED_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.commentsExtended+xml"
_W14 = "{http://schemas.microsoft.com/office/word/2010/wordml}paraId"


def _kept_comments(document: Document) -> list[dict]:
    return [
        fragment
        for element in document.elements
        for fragment in (element.preservedAttributes or {}).get("ooxml") or []
        if fragment.get("kind") == "comment"
    ]


def _threads(data: bytes) -> list[tuple[str, str | None, bool]]:
    """Each comment as Word would list it: its author, whose comment it answers, and whether resolved."""
    word = DocxDocument(io.BytesIO(data))
    authors = {str(comment.comment_id): comment.author for comment in word.comments}
    threads = comment_threads(word.part)
    return [
        (author, authors.get(threads.get(comment_id, (None, False))[0]), threads.get(comment_id, (None, False))[1])
        for comment_id, author in authors.items()
    ]


def _ids(data: bytes) -> set[str]:
    return {str(comment.comment_id) for comment in DocxDocument(io.BytesIO(data)).comments}


_A07_THREADS = [("Reviewer A", None, False), ("Author C", "Reviewer A", False), ("Reviewer B", None, True)]


def _a07() -> tuple[bytes, Document]:
    source = A07.read_bytes()
    document = parse_docx(source, A07.name)
    recompute_styles(document)
    stamp(document)  # as an upload does
    return source, document


def _retype_commented(document: Document) -> None:
    """What typing in the editor does to the paragraphs with comments: text added at their end."""
    for element in document.elements:
        if any(fragment.get("kind") == "comment" for fragment in (element.preservedAttributes or {}).get("ooxml") or []):
            element.inline = [*(element.inline or []), InlineRun(text=" Edited.")]
            element.content = f"{element.content} Edited."


def test_the_import_keeps_which_comment_answers_which_and_which_are_resolved():
    _, document = _a07()

    kept = {fragment["author"]: fragment for fragment in _kept_comments(document)}

    assert set(kept) == {"Reviewer A", "Author C", "Reviewer B"}
    assert kept["Author C"]["replyTo"] == kept["Reviewer A"]["commentId"]
    assert (kept["Reviewer A"]["replyTo"], kept["Reviewer B"]["replyTo"]) == (None, None)
    assert [kept[name]["done"] for name in ("Reviewer A", "Author C", "Reviewer B")] == [False, False, True]
    assert any("with their replies and which are resolved" in note for note in document.unsupportedFeatures)


def test_a_thread_copied_into_the_original_is_as_it_was():
    source, document = _a07()
    assert _threads(source) == _A07_THREADS

    exported = build_docx(document, source=source)

    assert _ids(exported) == _ids(source)  # copied, ids and all
    assert _threads(exported) == _A07_THREADS
    assert package_problems(exported) == []


def test_a_thread_written_anew_into_the_original_keeps_its_replies_and_resolved_state():
    source, document = _a07()
    _retype_commented(document)

    exported = build_docx(document, source=source)

    assert _ids(exported).isdisjoint(_ids(source))  # written anew: new ids, new paragraphs
    assert sorted(_threads(exported), key=str) == sorted(_A07_THREADS, key=str)
    assert package_problems(exported) == []


def test_a_thread_goes_into_a_new_word_file_together():
    _, document = _a07()

    exported = build_docx(document)

    assert _threads(exported) == _A07_THREADS  # a reply right after the comment it answers, as Word lists them
    assert package_problems(exported) == []


def _references(data: bytes) -> list[str]:
    """Whose comment each reference mark in the body is, in the body's order."""
    word = DocxDocument(io.BytesIO(data))
    authors = {str(comment.comment_id): comment.author for comment in word.comments}
    return [authors[node.get(qn("w:id"))] for node in word.element.body.iter(qn("w:commentReference"))]


def test_a_replys_mark_comes_after_its_comments_as_word_writes_them():
    """Word takes a reply whose reference comes before its comment's for a comment of its
    own (measured): python-docx puts a later comment's end and reference on the same run
    before an earlier one's, so the export puts them back in the comments' order."""
    source, document = _a07()
    assert _references(source) == ["Reviewer A", "Author C", "Reviewer B"]  # Word's own

    assert _references(build_docx(document)) == ["Reviewer A", "Author C", "Reviewer B"]
    _retype_commented(document)
    assert _references(build_docx(document, source=source)) == ["Reviewer A", "Author C", "Reviewer B"]


def test_a_reply_is_written_after_the_comment_it_answers():
    _, document = _a07()
    for element in document.elements:  # the reply kept first, as a file written by another program might have it
        kept = (element.preservedAttributes or {}).get("ooxml") or []
        element.preservedAttributes = {**(element.preservedAttributes or {}), "ooxml": sorted(kept, key=lambda f: f.get("replyTo") is None)}

    exported = build_docx(document)

    assert _threads(exported) == _A07_THREADS


def _two_paragraph_thread() -> bytes:
    """A resolved comment on one paragraph and its reply on the next -- Word files can say so
    (Word then shows the reply on its comment's text)."""
    word = DocxDocument()
    first = word.add_paragraph("The first claim.")
    second = word.add_paragraph("The second claim.")
    question = word.add_comment(first.runs[0], text="Source?", author="Reviewer", initials="RV")
    answer = word.add_comment(second.runs[0], text="Added.", author="Author", initials="AU")
    question._comment_elm.findall(qn("w:p"))[-1].set(_W14, "0A000001")
    answer._comment_elm.findall(qn("w:p"))[-1].set(_W14, "0A000002")
    xml = (
        '<w15:commentsEx xmlns:w15="http://schemas.microsoft.com/office/word/2012/wordml">'
        '<w15:commentEx w15:paraId="0A000001" w15:done="1"/>'
        '<w15:commentEx w15:paraId="0A000002" w15:paraIdParent="0A000001" w15:done="0"/></w15:commentsEx>'
    )
    word.part.relate_to(Part(PackURI("/word/commentsExtended.xml"), _EXTENDED_TYPE, xml.encode(), word.part.package), _EXTENDED)
    buffer = io.BytesIO()
    word.save(buffer)
    return buffer.getvalue()


def test_a_reply_copied_as_it_was_answers_its_comment_written_anew():
    source = _two_paragraph_thread()
    document = parse_docx(source, "thread.docx")
    recompute_styles(document)
    stamp(document)
    assert _threads(source) == [("Reviewer", None, True), ("Author", "Reviewer", False)]
    first = document.elements[0]
    first.inline = [InlineRun(text="The first claim. Reworded.")]
    first.content = "The first claim. Reworded."

    exported = build_docx(document, source=source)

    ids = {comment.author: str(comment.comment_id) for comment in DocxDocument(io.BytesIO(exported)).comments}
    assert ids["Reviewer"] not in _ids(source) and ids["Author"] in _ids(source)  # one written anew, one copied
    assert sorted(_threads(exported), key=str) == sorted([("Reviewer", None, True), ("Author", "Reviewer", False)], key=str)
    assert package_problems(exported) == []


def test_a_reply_whose_comment_is_gone_is_written_on_its_own():
    _, document = _a07()
    for element in document.elements:  # Reviewer A's comment deleted with its text, say
        kept = (element.preservedAttributes or {}).get("ooxml") or []
        element.preservedAttributes = {**(element.preservedAttributes or {}), "ooxml": [f for f in kept if f.get("author") != "Reviewer A"]}

    exported = build_docx(document)

    assert _threads(exported) == [("Author C", None, False), ("Reviewer B", None, True)]
    assert package_problems(exported) == []  # no entry naming a comment that isn't there


def test_comments_without_replies_or_resolved_ones_add_no_part():
    _, document = _a07()
    for fragment in _kept_comments(document):
        fragment.update(replyTo=None, done=False)

    exported = build_docx(document)

    assert not [name for name in zipfile.ZipFile(io.BytesIO(exported)).namelist() if "commentsExtended" in name]
    assert len(_threads(exported)) == 3


def test_comments_kept_before_threads_were_go_back_as_before():
    """A document stored before DOCX-021: its comments carry no thread."""
    _, document = _a07()
    for fragment in _kept_comments(document):
        for name in ("commentId", "replyTo", "done"):
            fragment.pop(name)

    exported = build_docx(document)

    assert sorted(_threads(exported)) == sorted([("Reviewer A", None, False), ("Author C", None, False), ("Reviewer B", None, False)])
    assert package_problems(exported) == []


def test_malformed_thread_data_is_left_out():
    """preservedAttributes comes back from the browser; nothing in it is trusted."""
    _, document = _a07()
    for fragment, (name, value) in zip(_kept_comments(document), (("replyTo", "../0"), ("done", "yes"), ("commentId", 7)), strict=True):
        fragment[name] = value

    exported = build_docx(document)

    assert _threads(exported) == []
    assert package_problems(exported) == []


def test_the_package_check_names_a_thread_entry_that_names_no_comment():
    source = A07.read_bytes()
    with zipfile.ZipFile(io.BytesIO(source)) as package:
        parts = {name: package.read(name) for name in package.namelist()}
    parts["word/commentsExtended.xml"] = parts["word/commentsExtended.xml"].replace(b'w15:paraIdParent="', b'w15:paraIdParent="7F', 1)
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as package:
        for name, data in parts.items():
            package.writestr(name, data)

    assert [problem for problem in package_problems(buffer.getvalue()) if "names no comment" in problem]
    assert package_problems(source) == []


# -- a comment over more than one paragraph ------------------------------------------


def _spanning(*texts: str, end_in_empty: bool = False) -> bytes:
    """A comment from the first paragraph's text to the last's (or to an empty paragraph after it)."""
    word = DocxDocument()
    paragraphs = [word.add_paragraph(text) for text in texts]
    if end_in_empty:
        paragraphs.append(word.add_paragraph())
        paragraphs[-1].add_run("")
    word.add_comment([paragraphs[0].runs[0], paragraphs[-1].runs[0]], text="All of this.", author="Reviewer", initials="RV")
    buffer = io.BytesIO()
    word.save(buffer)
    return buffer.getvalue()


def _marks(data: bytes) -> list[list[str]]:
    """Each paragraph of the body: its comment marks and text, in order."""
    xml = zipfile.ZipFile(io.BytesIO(data)).read("word/document.xml").decode("utf-8")
    paragraphs = re.findall(r"<w:p[ >].*?</w:p>|<w:p/>", xml, re.S)
    found = [re.findall(r'<w:(commentRangeStart|commentRangeEnd|commentReference) w:id="\d+"|<w:t(?: [^>]*)?>([^<]*)</w:t>', p) for p in paragraphs]
    return [[mark or text for mark, text in marks] for marks in found]


_ACROSS_THREE = [
    ["commentRangeStart", "One."],
    ["Two."],
    ["Three.", "commentRangeEnd", "commentReference"],
]


@pytest.mark.parametrize("edited", [None, 0, 2, "new file"])
def test_a_comment_over_several_paragraphs_keeps_its_range(edited):
    """Measured: before, every export ended it where its first paragraph does."""
    source = _spanning("One.", "Two.", "Three.")
    document = parse_docx(source, "span.docx")
    recompute_styles(document)
    stamp(document)
    if isinstance(edited, int):
        element = document.elements[edited]
        element.inline = [*(element.inline or []), InlineRun(text=" Edited.")]
        element.content = f"{element.content} Edited."

    exported = build_docx(document) if edited == "new file" else build_docx(document, source=source)

    marks = _marks(exported)
    assert [[mark for mark in paragraph if mark.startswith("comment")] for paragraph in marks] == [
        [mark for mark in paragraph if mark.startswith("comment")] for paragraph in _ACROSS_THREE
    ]
    assert [(c.author, c.text) for c in DocxDocument(io.BytesIO(exported)).comments] == [("Reviewer", "All of this.")]
    assert package_problems(exported) == []


def test_a_comment_ending_in_an_empty_paragraph_ends_where_the_text_does():
    source = _spanning("One.", "Two.", end_in_empty=True)
    document = parse_docx(source, "span.docx")

    exported = build_docx(document)

    assert _marks(exported) == [["commentRangeStart", "One."], ["Two.", "commentRangeEnd", "commentReference"]]


def test_a_comment_running_into_a_table_stays_on_the_text_before_it_and_the_report_says_so():
    word = DocxDocument()
    before = word.add_paragraph("Before the table.")
    cell = word.add_table(rows=1, cols=1).cell(0, 0)
    cell.paragraphs[0].add_run("In the cell.")
    word.add_comment([before.runs[0], cell.paragraphs[0].runs[0]], text="Into the table.", author="Reviewer", initials="RV")
    buffer = io.BytesIO()
    word.save(buffer)

    document = parse_docx(buffer.getvalue(), "table.docx")
    exported = build_docx(document)

    assert "Comments running on into a list, a table or code cover only the text before it." in document.unsupportedFeatures
    assert [(c.author, c.text) for c in DocxDocument(io.BytesIO(exported)).comments] == [("Reviewer", "Into the table.")]
    assert _marks(exported)[0] == ["commentRangeStart", "Before the table.", "commentRangeEnd", "commentReference"]


def test_a_comment_whose_last_paragraph_is_deleted_covers_its_first_and_the_export_says_so():
    document = parse_docx(_spanning("One.", "Two."), "span.docx")
    document.elements = document.elements[:1]
    report = ReportBuilder()

    exported = build_docx(document, report=report)

    assert _marks(exported) == [["commentRangeStart", "One.", "commentRangeEnd", "commentReference"]]
    assert "export.docx.comment_range" in {item.feature for item in report.items()}
    assert package_problems(exported) == []
