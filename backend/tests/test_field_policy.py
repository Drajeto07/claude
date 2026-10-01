"""Which Word fields a document may hold (tracker SEC-015): a field that runs a program or
pulls in outside content keeps its last result as text -- in the document, in the Word
file kept as its original (body, headers, footers, notes), and in what an export writes
back -- and no save from the browser can add one."""

import io
import zipfile

import pytest
from docx import Document as DocxDocument
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from fastapi.testclient import TestClient
from lxml import etree

from app.ai.base import AIStructuredOutputError
from app.ai.factory import get_ai_provider
from app.export.docx_export import build_docx
from app.export.package_check import package_problems
from app.formatting.engine import recompute_styles
from app.main import app
from app.models.document import Document, DocumentMetadata, Element, ElementType, InlineRun
from app.parsers.docx import parse_docx
from app.security.fields import field_allowed, neutralize_fields
from app.security.package import Cleaned, clean_package
from app.services.ingestion_service import UNSAFE_FIELDS
from tests.fakes import FakeAIProvider

pytestmark = pytest.mark.security  # the security regression suite (TEST-030)

client = TestClient(app, base_url="https://testserver")
_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_NS = nsdecls("w")
B = "\\"


@pytest.fixture
def signed_in(api_db):
    client.cookies.clear()
    app.dependency_overrides[get_ai_provider] = lambda: FakeAIProvider([AIStructuredOutputError("no AI in tests")] * 3)
    assert client.post("/api/v1/auth/register", json={"email": "fields@example.com", "password": "long enough password"}).status_code == 201
    yield
    client.cookies.clear()
    app.dependency_overrides.pop(get_ai_provider, None)


@pytest.mark.parametrize(
    ("instr", "allowed"),
    [
        (" PAGE " + B + "* MERGEFORMAT ", True),
        ("page", True),
        ("=SUM(ABOVE)", True),
        ('DATE ' + B + '@ "d.M.yyyy"', True),
        ('HYPERLINK "https://example.com" ' + B + 'o "tip"', True),
        ("HYPERLINK " + B + 'l "_Toc123"', True),  # a table of contents' own entries
        ('TOC ' + B + 'o "1-3" ' + B + "h " + B + "z", True),
        ("CITATION Smi20 " + B + "l 1033", True),
        ('ADDIN ZOTERO_ITEM CSL_CITATION {"citationID":"a"}', True),
        ("FORMCHECKBOX", True),
        (" MERGEFIELD Name ", True),
        ("DDEAUTO c:" + B + "windows" + B + 'system32' + B + 'cmd.exe "/k calc"', False),
        ("dde excel sheet1", False),
        ('INCLUDETEXT "c:' + B + 'secret.docx"', False),
        ('includepicture "http://tracker.example/p.png" ' + B + "d", False),
        ('INCLUDE "a.docx"', False),
        ('IMPORT "a.png"', False),
        ('LINK Excel.Sheet.12 "c:' + B + 'x.xlsx" "Sheet1"', False),
        ('RD "c:' + B + 'chapter.docx"', False),
        ("DATABASE " + B + 'd "c:' + B + 'x.mdb"', False),
        ("MACROBUTTON AcceptAllChangesInDoc Click here", False),
        ('PRINT "' + B + 'p page"', False),
        ("AUTOTEXT Signature", False),  # a template's building block: content from outside the document
        ('AUTOTEXTLIST "Pick one" ' + B + "s Normal", False),
        ("GLOSSARY Signature " + B + "* MERGEFORMAT", False),
        ('STYLEREF "Heading 1"', True),
        ('HYPERLINK "javascript:alert(1)"', False),
        ("HYPERLINK file:///c:/secret.docx", False),
        ('HYPERLINK "' + B + B + "server" + B + 'share"', False),
        ("", False),
        ("   ", False),
    ],
)
def test_only_fields_that_show_what_the_document_holds_stay_fields(instr, allowed):
    assert field_allowed(instr) is allowed


def _story(inner: str) -> etree._Element:
    return etree.fromstring(f"<w:document {_NS}><w:body>{inner}</w:body></w:document>".encode())


def _complex(instr: str, result: str) -> str:
    return (
        '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
        f'<w:r><w:instrText xml:space="preserve">{instr}</w:instrText></w:r>'
        '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
        f"<w:r><w:t>{result}</w:t></w:r>"
        '<w:r><w:fldChar w:fldCharType="end"/></w:r>'
    )


def test_a_refused_field_keeps_its_result_and_an_allowed_one_stays_a_field():
    root = _story(
        '<w:p><w:fldSimple w:instr=" DDEAUTO cmd /k calc "><w:r><w:t>dde result</w:t></w:r></w:fldSimple>'
        '<w:fldSimple w:instr=" PAGE "><w:r><w:t>1</w:t></w:r></w:fldSimple></w:p>'
        f'<w:p>{_complex(" INCLUDEPICTURE &quot;http://tracker.example/x.png&quot; ", "picture result")}</w:p>'
        # IF is allowed; the INCLUDETEXT nested in its instruction isn't.
        '<w:p><w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText xml:space="preserve"> IF </w:instrText></w:r>'
        f'{_complex(" INCLUDETEXT &quot;x&quot; ", "inner")}'
        '<w:r><w:instrText xml:space="preserve"> = "" "a" "b" </w:instrText></w:r><w:r><w:fldChar w:fldCharType="separate"/></w:r>'
        '<w:r><w:t>if result</w:t></w:r><w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>'
        # Deleted with tracked changes: rejecting the deletion would bring it back.
        '<w:p><w:del w:id="1" w:author="a"><w:r><w:fldChar w:fldCharType="begin"/></w:r>'
        '<w:r><w:delInstrText xml:space="preserve"> DDE excel x </w:delInstrText></w:r><w:r><w:fldChar w:fldCharType="separate"/></w:r>'
        '<w:r><w:delText>deleted result</w:delText></w:r><w:r><w:fldChar w:fldCharType="end"/></w:r></w:del></w:p>'
        # A refused field whose instruction holds an allowed one: all of that is instruction.
        '<w:p><w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText xml:space="preserve"> INCLUDETEXT "</w:instrText></w:r>'
        f'{_complex(" MERGEFIELD path ", "nested path")}'
        '<w:r><w:instrText xml:space="preserve">" </w:instrText></w:r><w:r><w:fldChar w:fldCharType="separate"/></w:r>'
        '<w:r><w:t>included result</w:t></w:r><w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>'
        # An allowed field deleted with tracked changes stays one.
        '<w:p><w:del w:id="2" w:author="a"><w:r><w:fldChar w:fldCharType="begin"/></w:r>'
        '<w:r><w:delInstrText xml:space="preserve"> PAGE </w:delInstrText></w:r><w:r><w:fldChar w:fldCharType="separate"/></w:r>'
        '<w:r><w:delText>7</w:delText></w:r><w:r><w:fldChar w:fldCharType="end"/></w:r></w:del></w:p>'
        # Never calculated: no result, nothing to keep but the text around it.
        '<w:p><w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText xml:space="preserve"> INCLUDETEXT "y" </w:instrText></w:r>'
        '<w:r><w:fldChar w:fldCharType="end"/></w:r><w:r><w:t xml:space="preserve"> after</w:t></w:r></w:p>'
    )

    assert neutralize_fields(root) == 6
    xml = etree.tostring(root).decode()
    assert "DDE" not in xml and "INCLUDE" not in xml and "MERGEFIELD" not in xml and "nested path" not in xml
    for text in ("dde result", "picture result", "inner", "if result", "deleted result", "included result", " after"):
        assert text in xml
    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    assert [instr.text for instr in root.iter(f"{w}instrText")] == [" IF ", ' = "" "a" "b" ']
    assert [simple.get(f"{w}instr") for simple in root.iter(f"{w}fldSimple")] == [" PAGE "]
    assert [instr.text for instr in root.iter(f"{w}delInstrText")] == [" PAGE "]
    kinds = [mark.get(f"{w}fldCharType") for mark in root.iter(f"{w}fldChar")]
    assert kinds == ["begin", "separate", "end"] * 2  # the IF field's and the deleted PAGE field's: nothing left open


def _word_file() -> bytes:
    """A body with a DDEAUTO field, an INCLUDEPICTURE field and a DATE field; a header with
    an INCLUDETEXT field; a footer with a PAGE field."""
    doc = DocxDocument()
    body = doc.element.body
    paragraphs = [
        f'<w:p {_NS}><w:r><w:t xml:space="preserve">Run: </w:t></w:r><w:fldSimple w:instr=" DDEAUTO cmd /k calc "><w:r><w:t>Figures</w:t></w:r></w:fldSimple></w:p>',
        f'<w:p {_NS}><w:r><w:t xml:space="preserve">Logo: </w:t></w:r>{_complex(" INCLUDEPICTURE &quot;http://tracker.example/x.png&quot; " + B + "d ", "logo")}</w:p>',
        f'<w:p {_NS}><w:r><w:t xml:space="preserve">Printed on </w:t></w:r>{_complex(" DATE " + B + "@ &quot;d.M.yyyy&quot; ", "25.9.2026")}</w:p>',
    ]
    for xml in paragraphs:
        body.insert(len(body) - 1, parse_xml(xml))
    header = doc.sections[0].header.paragraphs[0]._p
    header.append(parse_xml(f'<w:fldSimple {_NS} w:instr=" INCLUDETEXT &quot;c:{B}{B}secret.docx&quot; "><w:r><w:t>Header text</w:t></w:r></w:fldSimple>'))
    footer = doc.sections[0].footer.paragraphs[0]._p
    footer.append(parse_xml(f'<w:fldSimple {_NS} w:instr=" PAGE "><w:r><w:t>1</w:t></w:r></w:fldSimple>'))
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _parts(data: bytes) -> dict[str, str]:
    with zipfile.ZipFile(io.BytesIO(data)) as package:
        return {name: package.read(name).decode("utf-8", "replace") for name in package.namelist() if name.endswith(".xml")}


_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _instructions(data: bytes) -> dict[str, list[str]]:
    """Every field instruction in every part, by part."""
    found: dict[str, list[str]] = {}
    for name, xml in _parts(data).items():
        root = etree.fromstring(xml.encode("utf-8"))
        instructions = [simple.get(f"{_W}instr") for simple in root.iter(f"{_W}fldSimple")]
        instructions += [text.text or "" for text in root.iter(f"{_W}instrText", f"{_W}delInstrText")]
        if instructions:
            found[name] = [instruction.split()[0] for instruction in instructions if instruction.strip()]
    return found


def test_a_word_file_s_unsafe_fields_are_kept_as_their_result_everywhere(signed_in):
    source = _word_file()
    assert _instructions(source) == {"word/document.xml": ["DDEAUTO", "INCLUDEPICTURE", "DATE"], "word/header1.xml": ["INCLUDETEXT"], "word/footer1.xml": ["PAGE"]}
    result = clean_package(source)
    cleaned = result.data
    assert (result.fields, result.links, result.external) == (3, 0, 0)
    assert _instructions(cleaned) == {"word/document.xml": ["DATE"], "word/footer1.xml": ["PAGE"]}
    parts = _parts(cleaned)
    untouched = set(_parts(source)) - {"word/document.xml", "word/header1.xml"}
    assert all(_parts(source)[name] == parts[name] for name in untouched)
    assert clean_package(cleaned) == Cleaned(cleaned)  # nothing more to do: the same bytes

    response = client.post("/api/v1/documents/upload", files={"file": ("fields.docx", source, _DOCX)})
    assert response.status_code == 201, response.text[:300]
    document = response.json()
    item = next(item for item in document["importReport"]["items"] if item["feature"] == "docx.field.unsafe")
    assert (item["count"], item["policy"]) == (3, "lossy") and UNSAFE_FIELDS in document["unsupportedFeatures"]
    texts = [element["content"] for element in document["elements"]]
    assert "Run: Figures" in texts and "Logo: logo" in texts and "Printed on 25.9.2026" in texts

    # The Word export is written into the kept original: none of it may come back.
    exported = client.get(f"/api/v1/documents/{document['id']}/export/docx")
    assert exported.status_code == 200
    names = [name for part in _instructions(exported.content).values() for name in part]
    assert all(field_allowed(name) for name in names) and "DATE" in names and "PAGE" in names
    assert package_problems(exported.content) == []


def _fragment(instr: str, text: str = "Some") -> dict:
    return {"kind": "field", "instr": instr, "start": 0, "end": len(text), "text": text}


def test_a_save_can_t_add_a_field_to_the_next_word_export(signed_in):
    upload = client.post("/api/v1/documents/upload", files={"file": ("fields.docx", _word_file(), _DOCX)}).json()
    dated = next(element for element in upload["elements"] if element["content"].startswith("Printed on"))
    kept = dated["preservedAttributes"]
    assert [fragment["instr"].split()[0] for fragment in kept["ooxml"]] == ["DATE"]

    # One the server knows, sent back with a DDE field for its DATE field; and a new one with a DDE field.
    dated["preservedAttributes"] = {"ooxml": [_fragment(" DDEAUTO cmd /k calc ", "Printed")]}
    new = {"type": "paragraph", "content": "Some text.", "order": 99, "inline": [{"text": "Some text.", "marks": []}], "preservedAttributes": {"ooxml": [_fragment(" DDEAUTO cmd ")]}}
    saved = client.put(f"/api/v1/documents/{upload['id']}/content", json={"elements": [*upload["elements"], new]})

    assert saved.status_code == 200, saved.text[:300]
    elements = saved.json()["elements"]
    assert next(element for element in elements if element["id"] == dated["id"])["preservedAttributes"] == kept
    assert elements[-1]["preservedAttributes"] is None
    exported = client.get(f"/api/v1/documents/{upload['id']}/export/docx")
    assert _instructions(exported.content)["word/document.xml"] == ["DATE"]


def test_an_export_writes_no_field_a_document_may_not_hold():
    """As if a fragment got past the rest: the export itself checks each field."""

    def element(text: str, fragments: list) -> Element:
        return Element(type=ElementType.PARAGRAPH, content=text, order=0, inline=[InlineRun(text=text)], preservedAttributes={"ooxml": fragments})

    first = element("Run this", [_fragment(" DDEAUTO cmd /k calc ", "this") | {"start": 4, "end": 8}])
    # A field running across paragraphs: its start refused, its end must go too.
    opening = element("Start", [{"kind": "field_open", "instr": " INCLUDETEXT x ", "region": "f1", "start": 0, "end": 0, "text": ""}])
    closing = element("End", [{"kind": "field_close", "region": "f1", "start": 3, "end": 3, "text": ""}])
    document = Document(metadata=DocumentMetadata(title="Fields"), elements=[first, opening, closing])
    for order, block in enumerate(document.elements):
        block.order = order
    recompute_styles(document)

    exported = build_docx(document)
    xml = _parts(exported)["word/document.xml"]
    assert _instructions(exported) == {} and "fldChar" not in xml
    assert [paragraph.text for paragraph in DocxDocument(io.BytesIO(exported)).paragraphs] == ["Run this", "Start", "End"]
    assert package_problems(exported) == []


def test_the_importer_itself_keeps_no_field_a_document_may_not_hold():
    # Whoever calls it with a file that wasn't made safe first: such a field is its result.
    document = parse_docx(_word_file(), "fields.docx")
    kept = [fragment["instr"].split()[0] for element in document.elements for fragment in (element.preservedAttributes or {}).get("ooxml", [])]
    assert kept == ["DATE"]
    assert [element.content for element in document.elements][:3] == ["Run: Figures", "Logo: logo", "Printed on 25.9.2026"]
