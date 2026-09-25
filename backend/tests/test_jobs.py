"""Background jobs (корекции.docx §52/§53): heavy work queued with 202, its real
stage and progress polled from GET /api/jobs/{id}, its result read back. Over
the API the jobs run eagerly (inside the request); the in-process queue, a
restart's cut-off jobs and the runner's own rules are tested directly."""

import asyncio
import io
from datetime import datetime, timedelta, timezone

import pytest
from docx import Document as DocxDocument
from docx.shared import Pt
from fastapi.testclient import TestClient
from sqlalchemy import select, update
from sqlalchemy.orm import Session as OrmSession

from app.ai.base import AIStructuredOutputError
from app.ai.factory import get_ai_provider
from app.db.models import JobStatus, ProcessingJob
from app.jobs import queue as queue_module
from app.jobs import runner as runner_module
from app.jobs.files import sweep_job_files
from app.jobs.queue import BackgroundQueue, fail_interrupted_jobs
from app.jobs.runner import JobRunner
from app.main import app
from app.models.document import Document, DocumentMetadata
from app.repositories.document_repository import DocumentRepository
from app.services.auth_service import AuthService
from app.storage.local_provider import LocalStorageProvider
from tests.fakes import FakeAIProvider

client = TestClient(app, base_url="https://testserver")
_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_MARKDOWN = "# Годишен отчет\n\nПърви абзац с достатъчно текст за истински документ.\n\n## Раздел\n\nОще текст."


@pytest.fixture(autouse=True)
def signed_in(api_db):
    client.cookies.clear()
    app.dependency_overrides[get_ai_provider] = lambda: FakeAIProvider([AIStructuredOutputError("no AI in tests")] * 3)
    assert client.post("/api/auth/register", json={"email": "jobs@example.com", "password": "long enough password"}).status_code == 201
    yield
    client.cookies.clear()
    app.dependency_overrides.pop(get_ai_provider, None)


def _docx(font: str = "Georgia") -> bytes:
    doc = DocxDocument()
    doc.styles["Normal"].font.name = font
    doc.styles["Normal"].font.size = Pt(12)
    doc.add_heading("Report", level=1)
    doc.add_paragraph("A body paragraph long enough to be the reference's body text.")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _document() -> str:
    job = client.post("/api/jobs/import-text", json={"text": _MARKDOWN, "title": "Jobs test"}).json()
    return job["result"]["documentId"]


def test_every_job_endpoint_needs_a_signed_in_user():
    client.cookies.clear()
    calls = [
        client.post("/api/jobs/import-text", json={"text": "x"}),
        client.post("/api/jobs/import-file", files={"file": ("a.docx", _docx(), _DOCX)}),
        client.post("/api/jobs/format", data={"documentId": "x"}),
        client.post("/api/jobs/export", json={"documentId": "x", "format": "pdf"}),
        client.post("/api/jobs/extract-reference", files={"file": ("a.docx", _docx(), _DOCX)}),
        client.get("/api/jobs/x"),
        client.get("/api/jobs/x/file"),
    ]
    assert [response.status_code for response in calls] == [401] * len(calls)


def test_pasted_text_becomes_a_document_through_a_job():
    response = client.post("/api/jobs/import-text", json={"text": _MARKDOWN, "title": "Jobs test"})

    assert response.status_code == 202
    job = response.json()
    assert (job["type"], job["status"], job["stage"], job["progress"], job["error"]) == ("import_text", "succeeded", "complete", 100, None)
    assert job["startedAt"] and job["finishedAt"]
    document = client.get(f"/api/documents/{job['result']['documentId']}").json()
    assert document["metadata"]["title"] == "Jobs test"
    assert client.get(f"/api/jobs/{job['id']}").json() == job


def test_an_uploaded_file_becomes_a_document_and_its_upload_is_not_kept(api_db, tmp_path):
    job = client.post("/api/jobs/import-file", files={"file": ("Отчет.docx", _docx(), _DOCX)}).json()

    assert job["status"] == "succeeded"
    document = client.get(f"/api/documents/{job['result']['documentId']}").json()
    assert document["resolvedStyles"]["Paragraph"]["font-family"] == "Georgia"
    assert not (tmp_path / "assets" / "jobs" / job["id"] / "input").exists()


def test_a_file_that_cant_be_read_fails_its_job_with_the_reason():
    job = client.post("/api/jobs/import-file", files={"file": ("broken.docx", b"not a zip file", _DOCX)}).json()

    assert (job["status"], job["stage"]) == ("failed", "failed")
    assert "not a valid .docx" in job["error"]


def test_what_can_be_checked_up_front_is_refused_before_queuing(monkeypatch):
    assert client.post("/api/jobs/import-file", files={"file": ("notes.odt", b"x", "application/octet-stream")}).status_code == 400
    assert client.post("/api/jobs/extract-reference", files={"file": ("ref.pdf", b"%PDF", "application/pdf")}).status_code == 400
    assert client.post("/api/jobs/format", data={"documentId": "missing"}).status_code == 404
    assert client.post("/api/jobs/export", json={"documentId": "missing", "format": "docx"}).status_code == 404
    assert client.post("/api/jobs/import-text", json={"text": "   "}).status_code == 422

    monkeypatch.setattr("app.api.uploads.get_settings", lambda: type("S", (), {"max_upload_size_mb": 0})())
    assert client.post("/api/jobs/import-file", files={"file": ("big.docx", _docx(), _DOCX)}).status_code == 413


def test_formatting_runs_as_a_job_and_reports_conflicts_to_resolve():
    document_id = _document()
    applied = client.post("/api/jobs/format", data={"documentId": document_id, "templateId": "academic-default"}).json()

    assert applied["status"] == "succeeded" and applied["result"]["status"] == "applied"
    styles = client.get(f"/api/documents/{document_id}").json()["resolvedStyles"]
    assert styles["Paragraph"]["font-family"] == "Times New Roman"

    heading = client.get(f"/api/documents/{document_id}").json()["elements"][0]
    client.patch(f"/api/documents/{document_id}/elements/{heading['id']}/style", json={"property": "fontFamily", "value": "Georgia"})
    conflicted = client.post("/api/jobs/format", data={"documentId": document_id, "templateId": "professional-cv"}).json()
    assert conflicted["result"]["status"] == "conflicts"
    assert conflicted["result"]["conflicts"][0]["elementId"] == heading["id"]

    resolutions = '[{"elementId": "%s", "property": "fontFamily", "resolution": "apply_recommended"}]' % heading["id"]
    resolved = client.post(
        "/api/jobs/format", data={"documentId": document_id, "templateId": "professional-cv", "resolutions": resolutions}
    ).json()
    assert resolved["result"]["status"] == "applied"


def test_formatting_a_document_changed_elsewhere_fails_with_that_reason():
    document_id = _document()

    job = client.post("/api/jobs/format", data={"documentId": document_id, "templateId": "academic-default"}, headers={"If-Match": "999"}).json()

    assert job["status"] == "failed" and "changed in another tab" in job["error"]


@pytest.mark.parametrize("extension, content_type, magic", [("pdf", "application/pdf", b"%PDF"), ("docx", _DOCX, b"PK")])
def test_an_export_is_rendered_by_a_job_and_downloaded_from_it(extension, content_type, magic):
    document_id = _document()

    job = client.post("/api/jobs/export", json={"documentId": document_id, "format": extension}).json()

    assert job["status"] == "succeeded"
    assert set(job["result"]) == {"filename", "contentType", "size", "expired"}  # where it's stored stays private
    assert job["result"]["expired"] is False
    assert (job["result"]["filename"], job["result"]["contentType"]) == (f"Jobs test.{extension}", content_type)
    download = client.get(f"/api/jobs/{job['id']}/file")
    assert download.status_code == 200 and download.content.startswith(magic)
    assert len(download.content) == job["result"]["size"]
    assert "Jobs%20test" in download.headers["content-disposition"]


def test_an_export_expires_and_its_file_goes(api_db):
    document_id = _document()
    job = client.post("/api/jobs/export", json={"documentId": document_id, "format": "pdf"}).json()
    with OrmSession(api_db) as session:
        session.execute(update(ProcessingJob).where(ProcessingJob.id == job["id"]).values(finished_at=datetime.now(timezone.utc) - timedelta(days=2)))
        session.commit()

    client.post("/api/jobs/export", json={"documentId": document_id, "format": "docx"})  # any new job tidies up

    assert client.get(f"/api/jobs/{job['id']}/file").status_code == 404
    assert client.get(f"/api/jobs/{job['id']}").json()["result"]["expired"] is True


def test_deleting_a_document_deletes_its_export_files(tmp_path):
    document_id = _document()
    job = client.post("/api/jobs/export", json={"documentId": document_id, "format": "docx"}).json()
    assert (tmp_path / "assets" / "jobs" / job["id"] / "output").exists()

    assert client.delete(f"/api/documents/{document_id}").status_code == 204

    assert not (tmp_path / "assets" / "jobs" / job["id"] / "output").exists()
    assert client.get(f"/api/jobs/{job['id']}/file").status_code == 404
    assert client.get(f"/api/jobs/{job['id']}").json()["result"]["expired"] is True


def test_a_finished_job_keeps_no_copy_of_the_users_text(api_db):
    imported = client.post("/api/jobs/import-text", json={"text": _MARKDOWN, "title": "Kept title"}).json()
    formatted = client.post(
        "/api/jobs/format", data={"documentId": imported["result"]["documentId"], "instructionsText": "make the title bold"}
    ).json()

    with OrmSession(api_db) as session:
        payloads = {job.id: job.payload for job in session.scalars(select(ProcessingJob))}
    assert payloads[imported["id"]] == {"title": "Kept title"}
    assert "instructionsText" not in payloads[formatted["id"]]


def test_a_job_that_cant_be_queued_fails_instead_of_waiting_forever(api_db, tmp_path):
    from app.jobs.queue import get_job_queue

    class Unreachable:
        async def enqueue(self, job_id, *, priority=False):
            raise ConnectionError("Redis is down")

    app.dependency_overrides[get_job_queue] = lambda: Unreachable()
    try:
        response = client.post("/api/jobs/import-file", files={"file": ("a.docx", _docx(), _DOCX)})
    finally:
        app.dependency_overrides.pop(get_job_queue, None)

    assert response.status_code == 503
    with OrmSession(api_db) as session:
        job = session.scalars(select(ProcessingJob)).one()
        assert (job.status, job.stage) == ("failed", "failed") and job.finished_at is not None
        assert not (tmp_path / "assets" / job.input_key).exists()


def test_nobody_else_sees_a_job_or_its_file():
    document_id = _document()
    job = client.post("/api/jobs/export", json={"documentId": document_id, "format": "pdf"}).json()
    client.cookies.clear()
    client.post("/api/auth/register", json={"email": "other@example.com", "password": "long enough password"})

    assert client.get(f"/api/jobs/{job['id']}").status_code == 404
    assert client.get(f"/api/jobs/{job['id']}/file").status_code == 404
    assert client.post("/api/jobs/export", json={"documentId": document_id, "format": "pdf"}).status_code == 404


def test_a_reference_documents_style_is_read_by_a_job():
    job = client.post("/api/jobs/extract-reference", files={"file": ("House style.docx", _docx("Garamond"), _DOCX)}).json()

    assert job["status"] == "succeeded"
    assert job["result"]["suggestedName"] == "House style style"
    assert job["result"]["styleSystem"]["paragraph"]["fontFamily"] == "Garamond"


def test_an_unexpected_failure_is_reported_without_details(monkeypatch):
    async def broken(ctx):
        raise RuntimeError("internal detail that must not reach the user")

    monkeypatch.setitem(runner_module.KINDS, runner_module.IMPORT_TEXT, broken)

    job = client.post("/api/jobs/import-text", json={"text": _MARKDOWN}).json()

    assert job["status"] == "failed"
    assert job["error"] == "Something went wrong while processing this. Please try again."


# -- the runner and the in-process queue, directly --------------------------------


async def _user_and_job(session_factory, **job) -> str:
    async with session_factory() as session:
        auth = AuthService(session)
        user = await auth.register("runner@example.com", "long enough password", None)
        row = ProcessingJob(
            workspace_id=await auth.default_workspace_id(user.id), created_by=user.id, job_type="import_text", payload={"text": _MARKDOWN}, **job
        )
        session.add(row)
        await session.commit()
        return row.id


async def test_the_in_process_queue_runs_a_job_after_the_request(db_session_factory, tmp_path):
    job_id = await _user_and_job(db_session_factory)
    runner = JobRunner(db_session_factory, LocalStorageProvider(tmp_path), FakeAIProvider([]))

    await BackgroundQueue(runner).enqueue(job_id)
    await asyncio.gather(*queue_module._running)

    async with db_session_factory() as session:
        job = await session.get(ProcessingJob, job_id)
        assert (job.status, job.attempts) == (JobStatus.SUCCEEDED.value, 1)


async def test_a_finished_job_is_never_run_again(db_session_factory, tmp_path):
    job_id = await _user_and_job(db_session_factory, status=JobStatus.SUCCEEDED.value, result={"documentId": "kept"})

    await JobRunner(db_session_factory, LocalStorageProvider(tmp_path), FakeAIProvider([])).run(job_id)

    async with db_session_factory() as session:
        job = await session.get(ProcessingJob, job_id)
        assert (job.result, job.attempts) == ({"documentId": "kept"}, 0)


async def test_jobs_a_restart_cut_off_are_failed_so_nobody_waits_forever(db_session_factory, tmp_path):
    storage = LocalStorageProvider(tmp_path)
    await storage.put("jobs/cut-off/input", b"uploaded", "application/octet-stream")
    job_id = await _user_and_job(db_session_factory, status=JobStatus.RUNNING.value, input_key="jobs/cut-off/input")

    assert await fail_interrupted_jobs(db_session_factory, storage) == 1

    async with db_session_factory() as session:
        job = (await session.scalars(select(ProcessingJob).where(ProcessingJob.id == job_id))).one()
        assert job.status == JobStatus.FAILED.value and "restarted" in job.error_message
        assert job.payload == {}  # the pasted text is not kept
    assert not (tmp_path / "jobs" / "cut-off" / "input").exists()


async def test_the_sweep_expires_old_or_orphaned_export_files_and_removes_old_jobs(db_session_factory, tmp_path):
    storage = LocalStorageProvider(tmp_path)
    now = datetime.now(timezone.utc)
    ages = {"recent": timedelta(hours=1), "orphaned": timedelta(hours=1), "day-old": timedelta(hours=30), "month-old": timedelta(days=30)}
    async with db_session_factory() as session:
        auth = AuthService(session)
        user = await auth.register("sweep@example.com", "long enough password", None)
        workspace_id = await auth.default_workspace_id(user.id)
        document = Document(metadata=DocumentMetadata(title="Swept"))
        await DocumentRepository(session).create(workspace_id, document)
        jobs = {}
        for name, age in ages.items():
            await storage.put(f"jobs/{name}/output", b"PK", "application/pdf")
            jobs[name] = ProcessingJob(
                workspace_id=workspace_id,
                created_by=user.id,
                document_id=None if name == "orphaned" else document.id,  # its document was deleted
                job_type="export",
                status=JobStatus.SUCCEEDED.value,
                result={"key": f"jobs/{name}/output", "filename": "a.pdf"},
                created_at=now - age,
                finished_at=now - age,
            )
            session.add(jobs[name])
        await session.commit()
        ids = {name: job.id for name, job in jobs.items()}

    assert await sweep_job_files(db_session_factory, storage) == (3, 1)

    async with db_session_factory() as session:
        assert (await session.get(ProcessingJob, ids["recent"])).result["key"] == "jobs/recent/output"
        assert (await session.get(ProcessingJob, ids["orphaned"])).result == {"filename": "a.pdf", "expired": True}
        assert (await session.get(ProcessingJob, ids["day-old"])).result == {"filename": "a.pdf", "expired": True}
        assert await session.get(ProcessingJob, ids["month-old"]) is None
    assert [(tmp_path / "jobs" / name / "output").exists() for name in ages] == [True, False, False, False]


async def test_the_arq_worker_runs_the_same_runner(db_session_factory, tmp_path):
    from app import worker

    job_id = await _user_and_job(db_session_factory)

    await worker.run_job({"runner": JobRunner(db_session_factory, LocalStorageProvider(tmp_path), FakeAIProvider([]))}, job_id)

    assert worker.WorkerSettings.functions == [worker.run_job]
    assert [job.coroutine for job in worker.WorkerSettings.cron_jobs] == [worker.sweep, worker.sweep_assets]
    async with db_session_factory() as session:
        assert (await session.get(ProcessingJob, job_id)).status == JobStatus.SUCCEEDED.value


async def test_the_arq_queue_hands_the_job_to_redis_under_its_own_id(monkeypatch):
    queued = []

    class FakePool:
        async def enqueue_job(self, function, *args, **kwargs):
            queued.append((function, args, kwargs))

    monkeypatch.setattr(queue_module, "_arq_pool", FakePool())

    await queue_module.ArqQueue("redis://unused").enqueue("job-1")

    assert queued == [("run_job", ("job-1",), {"_job_id": "job-1"})]
