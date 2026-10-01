"""Background job safety (JOB-001, docs/architecture/jobs.md): idempotency keys, retries
with backoff for transient errors only, a timeout per job type, cancellation, stuck-job
recovery and the dead letter. Over the API the jobs run eagerly; the runner and the
recovery sweep are tested directly with the in-process queue (no Redis)."""

import asyncio
import functools
import io
import threading
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timedelta, timezone

import pytest
from docx import Document as DocxDocument
from fastapi.testclient import TestClient
from sqlalchemy import event, func, select
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session as OrmSession

from app.ai.base import AIStructuredOutputError
from app.ai.factory import get_ai_provider
from app.config import get_settings
from app.db.base import Base
from app.db.models import Document as DocumentRow
from app.db.models import JobStatus, JobType, ProcessingJob
from app.jobs import policy
from app.jobs import runner as runner_module
from app.jobs.files import sweep_job_files
from app.jobs.queue import EagerQueue
from app.jobs.recovery import recover_stuck_jobs
from app.jobs.runner import JobCancelled, JobError, JobRunner
from app.main import app
from app.parsers.docx import DocxParseError
from app.parsers.pdf import PdfParseError
from app.security.files import UnsafeFileError
from app.services.auth_service import AuthService
from app.services.entitlements_service import PlanLimitError
from app.services.ingestion_service import UnsupportedFileTypeError
from app.services.job_service import JobService
from app.storage.local_provider import LocalStorageProvider
from tests.conftest import _enable_sqlite_fk
from tests.fakes import FakeAIProvider
from tests.helpers import error_body

client = TestClient(app, base_url="https://testserver")
_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_TEXT = "# Report\n\nA paragraph with enough text to be a real document."
_PASSWORD = "long enough password"
IMPORT_TEXT = JobType.IMPORT_TEXT.value


@pytest.fixture(autouse=True)
def signed_in(api_db):
    client.cookies.clear()
    app.dependency_overrides[get_ai_provider] = lambda: FakeAIProvider([AIStructuredOutputError("no AI in tests")] * 3)
    assert client.post("/api/v1/auth/register", json={"email": "safety@example.com", "password": _PASSWORD}).status_code == 201
    yield
    client.cookies.clear()
    app.dependency_overrides.pop(get_ai_provider, None)


@pytest.fixture(autouse=True)
def no_waiting(monkeypatch):
    """Retries wait no time here; what they would wait is checked on policy.backoff_seconds."""
    monkeypatch.setattr(get_settings(), "job_retry_base_seconds", 0.0)


@functools.cache  # python-docx stamps the save time, so two calls give different bytes
def _docx() -> bytes:
    doc = DocxDocument()
    doc.add_heading("Report", level=1)
    doc.add_paragraph("A body paragraph long enough to be the reference's body text.")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _import(text: str = _TEXT, key: str | None = None, http: TestClient = client):
    return http.post("/api/v1/jobs/import-text", json={"text": text}, headers={"Idempotency-Key": key} if key else {})


def _rows(api_db) -> list[ProcessingJob]:
    with OrmSession(api_db) as session:
        return list(session.scalars(select(ProcessingJob).order_by(ProcessingJob.created_at)))


def _documents(api_db) -> int:
    with OrmSession(api_db) as session:
        return session.scalar(select(func.count()).select_from(DocumentRow))


# -- (1) idempotency keys ---------------------------------------------------------


def test_the_same_key_and_request_gets_the_same_job_and_makes_no_second_one(api_db):
    first = _import(key="import-1")
    again = _import(key="import-1")

    assert first.status_code == again.status_code == 202
    assert again.json() == first.json()
    assert len(_rows(api_db)) == 1 and _documents(api_db) == 1  # nothing ran twice


def test_without_a_key_every_request_is_a_new_job(api_db):
    assert _import().json()["id"] != _import().json()["id"]
    assert len(_rows(api_db)) == 2


def test_a_key_belongs_to_one_user(api_db):
    mine = _import(key="shared-key").json()
    other = TestClient(app, base_url="https://testserver")
    assert other.post("/api/v1/auth/register", json={"email": "other-key@example.com", "password": _PASSWORD}).status_code == 201

    theirs = _import(key="shared-key", http=other)

    assert theirs.status_code == 202 and theirs.json()["id"] != mine["id"]
    assert client.get(f"/api/v1/jobs/{theirs.json()['id']}").status_code == 404


def test_the_same_key_with_another_request_is_refused_not_answered_with_the_other_job(api_db):
    _import(key="one-key")

    reused = _import(text="# Another\n\nA completely different text.", key="one-key")

    assert reused.status_code == 422
    assert reused.json()["code"] == "idempotency_key_reused"
    assert "different request" in reused.json()["message"]
    assert len(_rows(api_db)) == 1


@pytest.mark.parametrize("key", ["", "x" * 129, "has a space", "semi;colon", "slash/slash", "ключ"])
def test_a_key_that_isnt_valid_is_a_422_through_the_error_envelope_and_makes_no_job(api_db, key):
    # A non-ASCII header value is sent as UTF-8 bytes, as a browser would.
    response = client.post("/api/v1/jobs/import-text", json={"text": _TEXT}, headers={"Idempotency-Key": key.encode()})

    assert response.status_code == 422
    assert set(response.json()) == {"code", "message", "details", "request_id"}
    assert response.json()["code"] == "invalid_idempotency_key"
    assert _rows(api_db) == []


@pytest.mark.parametrize("key", ["a", "x" * 128, "550e8400-e29b-41d4-a716-446655440000", "order:42_retry.1-b"])
def test_keys_of_the_allowed_shape_are_taken(api_db, key):
    assert _import(key=key).status_code == 202


def test_every_job_creating_route_answers_a_repeat_with_its_job(api_db):
    document_id = _import().json()["result"]["documentId"]
    requests = {
        "import-text": lambda k: client.post("/api/v1/jobs/import-text", json={"text": _TEXT + k}, headers={"Idempotency-Key": "t" + k}),
        "import-file": lambda k: client.post(
            "/api/v1/jobs/import-file", files={"file": ("a.docx", _docx(), _DOCX)}, headers={"Idempotency-Key": "f" + k}
        ),
        "format": lambda k: client.post(
            "/api/v1/jobs/format", data={"documentId": document_id, "instructionsText": ""}, headers={"Idempotency-Key": "o" + k}
        ),
        "export": lambda k: client.post(
            "/api/v1/jobs/export", json={"documentId": document_id, "format": "pdf"}, headers={"Idempotency-Key": "e" + k}
        ),
        "extract-reference": lambda k: client.post(
            "/api/v1/jobs/extract-reference", files={"file": ("ref.docx", _docx(), _DOCX)}, headers={"Idempotency-Key": "r" + k}
        ),
    }

    for name, send in requests.items():
        first, again = send("1"), send("1")
        assert first.status_code == again.status_code == 202, name
        assert first.json()["id"] == again.json()["id"], name

    assert len(_rows(api_db)) == 1 + len(requests)  # the document's import, plus one job per route


def test_a_repeat_with_a_different_file_or_option_is_refused(api_db):
    headers = {"Idempotency-Key": "upload"}
    first = client.post("/api/v1/jobs/import-file", files={"file": ("a.docx", _docx(), _DOCX)}, headers=headers)
    other_file = DocxDocument()
    other_file.add_paragraph("Different bytes")
    buffer = io.BytesIO()
    other_file.save(buffer)

    assert client.post("/api/v1/jobs/import-file", files={"file": ("a.docx", buffer.getvalue(), _DOCX)}, headers=headers).status_code == 422
    assert client.post("/api/v1/jobs/import-file", files={"file": ("a.docx", _docx(), _DOCX)}, data={"autolink": "true"}, headers=headers).status_code == 422
    same = client.post("/api/v1/jobs/import-file", files={"file": ("a.docx", _docx(), _DOCX)}, headers=headers)
    assert same.json()["id"] == first.json()["id"]


def test_a_repeat_is_answered_even_when_the_plan_would_refuse_a_new_job(api_db, monkeypatch):
    from app.services.entitlements_service import EntitlementsService

    first = _import(key="before-the-limit")

    async def full(self, workspace_id):
        raise PlanLimitError("maxDocuments", "Your plan allows 1 document.", limit=1, used=1)

    monkeypatch.setattr(EntitlementsService, "check_new_document", full)

    assert _import(key="before-the-limit").json()["id"] == first.json()["id"]
    assert _import(text="# New\n\nAnother document.", key="after-the-limit").status_code != 202


def test_the_database_refuses_two_jobs_with_one_key_for_one_user(api_db):
    _import(key="unique")
    first = _rows(api_db)[0]

    with OrmSession(api_db) as session:
        session.add(
            ProcessingJob(
                workspace_id=first.workspace_id, created_by=first.created_by, job_type=IMPORT_TEXT, idempotency_key="unique", payload={}
            )
        )
        with pytest.raises(IntegrityError):
            session.commit()


def test_jobs_without_a_key_never_collide_in_the_index(api_db):
    _import(), _import(), _import()
    assert len(_rows(api_db)) == 3


def test_a_duplicate_that_gets_past_the_lookup_still_gets_the_first_job(api_db, monkeypatch):
    """Two requests with one key arrive together: both look, neither finds a job, both insert.
    The unique index lets one in; the other is answered with that job."""
    first = _import(key="race")
    monkeypatch.setattr(JobService, "find_replay", _blind_to_the_first_lookup())

    second = _import(key="race")

    assert second.status_code == 202 and second.json()["id"] == first.json()["id"]
    assert len(_rows(api_db)) == 1 and _documents(api_db) == 1


def test_a_race_with_a_different_request_is_refused_too(api_db, monkeypatch):
    _import(key="race-2")
    monkeypatch.setattr(JobService, "find_replay", _blind_to_the_first_lookup())

    assert _import(text="# Other\n\nDifferent text entirely.", key="race-2").status_code == 422


def _blind_to_the_first_lookup():
    original = JobService.find_replay
    calls = []

    async def find(self, key, job_type, **job):
        calls.append(1)
        return None if len(calls) == 1 else await original(self, key, job_type, **job)

    return find


def test_two_requests_sent_at_once_with_one_key_make_one_job(api_db):
    other = TestClient(app, base_url="https://testserver")
    other.cookies.update(client.cookies)
    barrier = threading.Barrier(2)

    def send(http):
        barrier.wait()
        return _import(key="together", http=http)

    with ThreadPoolExecutor(2) as pool:
        responses = list(pool.map(send, [client, other]))

    assert [response.status_code for response in responses] == [202, 202]
    assert responses[0].json()["id"] == responses[1].json()["id"]
    assert len(_rows(api_db)) == 1


async def test_a_key_is_forgotten_with_its_job(db_session_factory, tmp_path):
    storage = LocalStorageProvider(tmp_path)
    user_id, _ = await _user(db_session_factory)
    async with db_session_factory() as session:
        jobs = JobService(session, user_id=user_id, storage=storage)
        old, _new = await jobs.create(IMPORT_TEXT, payload={"text": "a"}, idempotency_key="forgotten")
        old.created_at = datetime.now(timezone.utc) - timedelta(days=30)
        old.status, old.finished_at = JobStatus.SUCCEEDED.value, old.created_at
        await session.commit()

    assert await sweep_job_files(db_session_factory, storage) == (0, 1)

    async with db_session_factory() as session:
        again, created = await JobService(session, user_id=user_id, storage=storage).create(IMPORT_TEXT, payload={"text": "a"}, idempotency_key="forgotten")
        assert created and again.id != old.id


# -- helpers for the runner tests -------------------------------------------------


async def _user(factory, email: str = "runner@example.com") -> tuple[str, str]:
    async with factory() as session:
        auth = AuthService(session)
        user = await auth.register(email, _PASSWORD, None)
        return user.id, await auth.default_workspace_id(user.id)


async def _add(factory, user: tuple[str, str], **fields) -> str:
    async with factory() as session:
        row = ProcessingJob(
            **{"workspace_id": user[1], "created_by": user[0], "job_type": IMPORT_TEXT, "payload": {"text": _TEXT}} | fields
        )
        session.add(row)
        await session.commit()
        return row.id


async def _load(factory, job_id: str) -> ProcessingJob:
    async with factory() as session:
        return await session.get(ProcessingJob, job_id)


def _runner(factory, tmp_path) -> JobRunner:
    return JobRunner(factory, LocalStorageProvider(tmp_path), FakeAIProvider([]))


def _failing_with(*errors, then=None):
    """A job kind that raises each of `errors` in turn, then returns `then`."""
    pending = list(errors)

    async def kind(ctx):
        if pending:
            raise pending.pop(0)
        return then or {"documentId": "made"}

    return kind


# -- (2) retries with backoff, for transient errors only --------------------------


async def test_a_transient_failure_is_tried_again_after_a_backoff_with_its_input_kept(db_session_factory, tmp_path, monkeypatch):
    storage = LocalStorageProvider(tmp_path)
    await storage.put("jobs/j/input", b"upload", "application/octet-stream")
    job_id = await _add(db_session_factory, await _user(db_session_factory), input_key="jobs/j/input")
    monkeypatch.setitem(runner_module.KINDS, IMPORT_TEXT, _failing_with(ConnectionError("network down")))
    monkeypatch.setattr(get_settings(), "job_retry_base_seconds", 5.0)
    runner = JobRunner(db_session_factory, storage, FakeAIProvider([]))

    wait = await runner.run(job_id)

    job = await _load(db_session_factory, job_id)
    assert wait == 5.0
    assert (job.status, job.stage, job.retry_count, job.attempts, job.dead_letter) == ("pending", "retrying", 1, 1, False)
    assert job.error_message is None and job.failure_reason == "transient:ConnectionError"
    assert job.payload == {"text": _TEXT}  # needed again
    assert await storage.get("jobs/j/input") == b"upload"

    assert await runner.run(job_id) is None

    job = await _load(db_session_factory, job_id)
    assert (job.status, job.attempts, job.retry_count, job.result) == ("succeeded", 2, 1, {"documentId": "made"})
    assert job.payload == {}  # the text goes once the job is done


def test_the_backoff_doubles_up_to_a_cap(monkeypatch):
    monkeypatch.setattr(get_settings(), "job_retry_base_seconds", 5.0)
    monkeypatch.setattr(get_settings(), "job_retry_max_seconds", 30.0)

    assert [policy.backoff_seconds(done) for done in range(5)] == [5.0, 10.0, 20.0, 30.0, 30.0]


async def test_the_in_process_queue_waits_and_runs_a_transient_failure_again(db_session_factory, tmp_path, monkeypatch):
    job_id = await _add(db_session_factory, await _user(db_session_factory))
    monkeypatch.setitem(runner_module.KINDS, IMPORT_TEXT, _failing_with(ConnectionError("1"), TimeoutError("2")))

    await EagerQueue(_runner(db_session_factory, tmp_path)).enqueue(job_id)

    job = await _load(db_session_factory, job_id)
    assert (job.status, job.attempts, job.retry_count, job.dead_letter) == ("succeeded", 3, 2, False)


async def test_the_arq_worker_defers_a_retry(db_session_factory, tmp_path, monkeypatch):
    from arq import Retry

    from app import worker

    job_id = await _add(db_session_factory, await _user(db_session_factory))
    monkeypatch.setitem(runner_module.KINDS, IMPORT_TEXT, _failing_with(ConnectionError("down")))
    monkeypatch.setattr(get_settings(), "job_retry_base_seconds", 7.0)

    with pytest.raises(Retry) as retry:
        await worker.run_job({"runner": _runner(db_session_factory, tmp_path)}, job_id)

    assert retry.value.defer_score == 7000  # milliseconds
    await worker.run_job({"runner": _runner(db_session_factory, tmp_path)}, job_id)  # the retry arq makes
    assert (await _load(db_session_factory, job_id)).status == "succeeded"
    assert worker.WorkerSettings.max_tries > get_settings().job_max_attempts  # arq never gives up before the runner


async def test_a_transient_failure_that_never_clears_is_a_dead_letter_after_the_last_attempt(db_session_factory, tmp_path, monkeypatch):
    storage = LocalStorageProvider(tmp_path)
    await storage.put("jobs/d/input", b"upload", "application/octet-stream")
    job_id = await _add(db_session_factory, await _user(db_session_factory), input_key="jobs/d/input")
    monkeypatch.setitem(runner_module.KINDS, IMPORT_TEXT, _failing_with(*[ConnectionError("down")] * 10))
    runner = JobRunner(db_session_factory, storage, FakeAIProvider([]))

    waits = [await runner.run(job_id), await runner.run(job_id), await runner.run(job_id)]

    job = await _load(db_session_factory, job_id)
    assert [wait is not None for wait in waits] == [True, True, False]  # three attempts, then it gives up
    assert (job.status, job.stage, job.attempts, job.retry_count) == ("failed", "failed", 3, 2)
    assert (job.dead_letter, job.failure_reason) == (True, "transient:ConnectionError")
    assert job.error_message == "This couldn't be finished after 3 tries. Please try again in a few minutes."
    assert "ConnectionError" not in job.error_message
    assert job.payload == {} and not (tmp_path / "jobs" / "d" / "input").exists()
    assert await runner.run(job_id) is None  # a dead letter is not run again
    assert (await _load(db_session_factory, job_id)).attempts == 3


_USER_ERRORS = [
    DocxParseError("This is not a valid .docx file."),
    PdfParseError("This PDF can't be read."),
    UnsafeFileError("This isn't a Word (.docx) file."),
    UnsupportedFileTypeError("Unsupported."),
    JobError("The document no longer exists."),
    PlanLimitError("maxDocuments", "Your plan allows 3 documents.", limit=3, used=3),
    ValueError("a bug"),
    KeyError("a bug"),
]


@pytest.mark.parametrize("error", _USER_ERRORS, ids=lambda error: type(error).__name__)
async def test_what_the_user_sent_and_bugs_are_never_retried(db_session_factory, tmp_path, monkeypatch, error):
    job_id = await _add(db_session_factory, await _user(db_session_factory))
    monkeypatch.setitem(runner_module.KINDS, IMPORT_TEXT, _failing_with(error, error, error))

    wait = await _runner(db_session_factory, tmp_path).run(job_id)

    job = await _load(db_session_factory, job_id)
    assert wait is None
    assert (job.status, job.attempts, job.retry_count, job.dead_letter) == ("failed", 1, 0, False)


async def test_an_error_that_is_a_users_and_also_a_network_one_is_not_retried(db_session_factory, tmp_path, monkeypatch):
    class OddlyBoth(DocxParseError, ConnectionError):
        pass

    assert not policy.is_transient(OddlyBoth("x"))
    job_id = await _add(db_session_factory, await _user(db_session_factory))
    monkeypatch.setitem(runner_module.KINDS, IMPORT_TEXT, _failing_with(OddlyBoth("x")))

    assert await _runner(db_session_factory, tmp_path).run(job_id) is None
    assert (await _load(db_session_factory, job_id)).attempts == 1


def test_which_exceptions_are_transient():
    import anthropic
    import httpx
    from botocore.exceptions import ConnectTimeoutError
    from sqlalchemy.exc import OperationalError, ProgrammingError

    request = httpx.Request("POST", "https://api.example.invalid")
    transient = [
        ConnectionResetError(),
        TimeoutError(),
        OperationalError("select 1", {}, Exception("server closed the connection")),
        anthropic.APIConnectionError(request=request),
        anthropic.APITimeoutError(request=request),
        anthropic.RateLimitError("slow down", response=httpx.Response(429, request=request), body=None),
        anthropic.InternalServerError("oops", response=httpx.Response(500, request=request), body=None),
        ConnectTimeoutError(endpoint_url="https://s3.example.invalid"),
    ]
    permanent = [
        ValueError(),
        KeyError(),
        FileNotFoundError(),
        ProgrammingError("select", {}, Exception("no such table")),
        anthropic.BadRequestError("bad", response=httpx.Response(400, request=request), body=None),
        *_USER_ERRORS,
    ]
    assert all(policy.is_transient(error) for error in transient)
    assert not any(policy.is_transient(error) for error in permanent)


def test_an_unreadable_upload_fails_once_with_the_reason_and_is_not_retried(api_db):
    job = client.post("/api/v1/jobs/import-file", files={"file": ("broken.docx", _docx()[:200], _DOCX)})
    assert job.status_code == 400  # refused before a job exists

    # A package with no Word document inside: only reading it finds out.
    import zipfile

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as package:
        package.writestr("[Content_Types].xml", '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>')
    job = client.post("/api/v1/jobs/import-file", files={"file": ("broken.docx", buffer.getvalue(), _DOCX)}).json()

    row = _rows(api_db)[-1]
    assert (job["status"], job["retryCount"], job["deadLetter"]) == ("failed", 0, False)
    assert (row.attempts, row.retry_count, row.dead_letter) == (1, 0, False)


async def test_the_calls_a_failed_attempt_made_are_still_counted(db_session_factory, tmp_path, monkeypatch):
    from app.services.usage_service import AI_OPERATIONS

    async def kind(ctx):
        ctx.usage.append(AI_OPERATIONS)
        raise ConnectionError("down")

    monkeypatch.setitem(runner_module.KINDS, IMPORT_TEXT, kind)
    job_id = await _add(db_session_factory, await _user(db_session_factory))

    await _runner(db_session_factory, tmp_path).run(job_id)

    from app.db.models import UsageRecord

    async with db_session_factory() as session:
        assert await session.scalar(select(func.count()).select_from(UsageRecord).where(UsageRecord.metric == AI_OPERATIONS)) == 1


# -- (3) a timeout per job type ---------------------------------------------------


def test_every_job_type_has_its_own_timeout():
    assert set(policy.JOB_TIMEOUTS) == {job_type.value for job_type in JobType}
    assert policy.timeout_for(JobType.EXPORT.value) != policy.timeout_for(JobType.IMPORT_FILE.value)
    assert all(0 < seconds <= get_settings().ai_seconds_per_job for seconds in policy.JOB_TIMEOUTS.values())


async def test_a_job_past_its_timeout_fails_with_a_message_for_people_and_is_not_retried(db_session_factory, tmp_path, monkeypatch):
    async def slow(ctx):
        await asyncio.sleep(5)
        return {"documentId": "never"}

    monkeypatch.setitem(runner_module.KINDS, IMPORT_TEXT, slow)
    monkeypatch.setitem(policy.JOB_TIMEOUTS, IMPORT_TEXT, 0.05)
    job_id = await _add(db_session_factory, await _user(db_session_factory))

    wait = await _runner(db_session_factory, tmp_path).run(job_id)

    job = await _load(db_session_factory, job_id)
    assert wait is None
    assert (job.status, job.stage, job.attempts, job.dead_letter, job.result) == ("failed", "failed", 1, False, None)
    assert job.failure_reason == "timeout"
    assert job.error_message == "This import took longer than 0 seconds, so it was stopped. Please try again; if it keeps happening, try a smaller document."
    assert "Timeout" not in job.error_message and "asyncio" not in job.error_message


async def test_each_job_type_is_held_to_its_own_timeout(db_session_factory, tmp_path, monkeypatch):
    async def a_while(ctx):
        await asyncio.sleep(0.3)
        return {"documentId": "made"}

    monkeypatch.setitem(runner_module.KINDS, IMPORT_TEXT, a_while)
    monkeypatch.setitem(runner_module.KINDS, JobType.EXPORT.value, a_while)
    monkeypatch.setitem(policy.JOB_TIMEOUTS, IMPORT_TEXT, 0.05)
    monkeypatch.setitem(policy.JOB_TIMEOUTS, JobType.EXPORT.value, 5)
    user = await _user(db_session_factory)
    short, long = await _add(db_session_factory, user), await _add(db_session_factory, user, job_type=JobType.EXPORT.value)

    await _runner(db_session_factory, tmp_path).run(short)
    await _runner(db_session_factory, tmp_path).run(long)

    assert [(await _load(db_session_factory, job_id)).status for job_id in (short, long)] == ["failed", "succeeded"]


def test_the_timeout_message_names_the_minutes_and_the_kind_of_job():
    assert "5 minutes" in policy.timeout_message(JobType.EXPORT.value) and "export" in policy.timeout_message(JobType.EXPORT.value)
    assert "10 minutes" in policy.timeout_message(JobType.FORMAT.value) and "formatting" in policy.timeout_message(JobType.FORMAT.value)


async def test_a_timeout_raised_by_the_work_itself_is_a_transient_error_not_the_jobs_deadline(db_session_factory, tmp_path, monkeypatch):
    monkeypatch.setitem(runner_module.KINDS, IMPORT_TEXT, _failing_with(TimeoutError("the provider timed out")))
    job_id = await _add(db_session_factory, await _user(db_session_factory))

    assert await _runner(db_session_factory, tmp_path).run(job_id) is not None
    assert (await _load(db_session_factory, job_id)).failure_reason == "transient:TimeoutError"


# -- (4) cancellation -------------------------------------------------------------


def _cancel(job_id: str, http: TestClient = client):
    return http.post(f"/api/v1/jobs/{job_id}/cancel")


async def test_a_job_that_has_not_started_is_cancelled_and_never_runs(db_session_factory, tmp_path):
    storage = LocalStorageProvider(tmp_path)
    await storage.put("jobs/c/input", b"upload", "application/octet-stream")
    user = await _user(db_session_factory)
    job_id = await _add(db_session_factory, user, input_key="jobs/c/input")

    async with db_session_factory() as session:
        cancelled = await JobService(session, user_id=user[0], storage=storage).cancel(job_id)

    assert (cancelled.status, cancelled.stage, cancelled.finished_at is not None) == ("cancelled", "cancelled", True)
    assert cancelled.payload == {} and not (tmp_path / "jobs" / "c" / "input").exists()
    assert await JobRunner(db_session_factory, storage, FakeAIProvider([])).run(job_id) is None
    job = await _load(db_session_factory, job_id)
    assert (job.status, job.attempts, job.result) == ("cancelled", 0, None)


def test_the_owner_cancels_over_the_api_and_a_second_cancel_changes_nothing(api_db, monkeypatch):
    from app.jobs.queue import get_job_queue

    class Parked:  # a queue that never gets to the job: it stays pending
        async def enqueue(self, job_id, *, priority=False):
            pass

    app.dependency_overrides[get_job_queue] = lambda: Parked()
    try:
        job = _import().json()
    finally:
        app.dependency_overrides.pop(get_job_queue, None)
    assert job["status"] == "pending"

    first, second = _cancel(job["id"]), _cancel(job["id"])

    assert (first.status_code, second.status_code) == (200, 200)
    assert first.json()["status"] == "cancelled" and first.json()["finishedAt"]
    assert second.json() == first.json()
    assert client.get(f"/api/v1/jobs/{job['id']}").json() == first.json()


def test_cancelling_a_finished_job_is_not_an_error_and_changes_nothing(api_db):
    done = _import().json()
    assert done["status"] == "succeeded"

    answer = _cancel(done["id"])

    assert answer.status_code == 200 and answer.json() == done
    with pytest.MonkeyPatch.context() as patch:
        patch.setitem(runner_module.KINDS, IMPORT_TEXT, _failing_with(ValueError("a bug")))
        failed = _import(text="# Fails\n\nThis one fails.").json()
    assert failed["status"] == "failed"
    assert _cancel(failed["id"]).json() == failed


def test_cancelling_someone_elses_job_looks_exactly_like_a_missing_job(api_db):
    mine = _import().json()
    other = TestClient(app, base_url="https://testserver")
    assert other.post("/api/v1/auth/register", json={"email": "not-owner@example.com", "password": _PASSWORD}).status_code == 201

    foreign, missing = _cancel(mine["id"], other), _cancel("does-not-exist", other)

    assert foreign.status_code == missing.status_code == 404
    assert error_body(foreign) == error_body(missing)
    assert client.get(f"/api/v1/jobs/{mine['id']}").json() == mine


def test_cancelling_a_pending_job_somebody_else_owns_leaves_it_pending(api_db):
    from app.jobs.queue import get_job_queue

    class Parked:
        async def enqueue(self, job_id, *, priority=False):
            pass

    app.dependency_overrides[get_job_queue] = lambda: Parked()
    try:
        mine = _import().json()
    finally:
        app.dependency_overrides.pop(get_job_queue, None)
    other = TestClient(app, base_url="https://testserver")
    other.post("/api/v1/auth/register", json={"email": "intruder@example.com", "password": _PASSWORD})

    assert _cancel(mine["id"], other).status_code == 404

    assert _rows(api_db)[0].status == "pending"


async def test_a_running_job_notices_the_cancel_at_its_next_check_and_writes_no_result(db_session_factory, tmp_path, monkeypatch):
    storage = LocalStorageProvider(tmp_path)
    await storage.put("jobs/r/input", b"upload", "application/octet-stream")
    user = await _user(db_session_factory)
    job_id = await _add(db_session_factory, user, input_key="jobs/r/input")
    reached = []

    async def kind(ctx):
        await ctx.report("analyzing", 15)
        reached.append("first step")
        async with db_session_factory() as other_session:  # the owner cancels, in another request
            await JobService(other_session, user_id=user[0], storage=storage).cancel(job_id)
        await ctx.report("finalizing", 85)  # the next check
        reached.append("after the cancel")
        return {"documentId": "must not be written"}

    monkeypatch.setitem(runner_module.KINDS, IMPORT_TEXT, kind)

    assert await JobRunner(db_session_factory, storage, FakeAIProvider([])).run(job_id) is None

    job = await _load(db_session_factory, job_id)
    assert reached == ["first step"]
    assert (job.status, job.stage, job.result, job.error_message, job.dead_letter) == ("cancelled", "cancelled", None, None, False)
    assert job.progress == 15  # the step after the cancel was never written
    assert not (tmp_path / "jobs" / "r" / "input").exists()


async def test_a_cancel_after_the_last_check_still_stands_and_the_result_is_dropped(db_session_factory, tmp_path, monkeypatch):
    storage = LocalStorageProvider(tmp_path)
    user = await _user(db_session_factory)
    job_id = await _add(db_session_factory, user, job_type=JobType.EXPORT.value)

    async def kind(ctx):
        await ctx.report("rendering", 10)
        await ctx.storage.put(f"jobs/{ctx.job_id}/output", b"PDF", "application/pdf")
        async with db_session_factory() as other_session:
            await JobService(other_session, user_id=user[0], storage=storage).cancel(job_id)
        return {"key": f"jobs/{ctx.job_id}/output", "filename": "a.pdf", "contentType": "application/pdf", "size": 3}

    monkeypatch.setitem(runner_module.KINDS, JobType.EXPORT.value, kind)

    await JobRunner(db_session_factory, storage, FakeAIProvider([])).run(job_id)

    job = await _load(db_session_factory, job_id)
    assert (job.status, job.result) == ("cancelled", None)
    assert not (tmp_path / "jobs" / job_id / "output").exists()  # no export file nobody can reach is left behind


async def test_a_job_cancelled_while_it_waits_for_a_retry_does_not_run_again(db_session_factory, tmp_path, monkeypatch):
    storage = LocalStorageProvider(tmp_path)
    user = await _user(db_session_factory)
    job_id = await _add(db_session_factory, user)
    monkeypatch.setitem(runner_module.KINDS, IMPORT_TEXT, _failing_with(ConnectionError("down")))
    runner = JobRunner(db_session_factory, storage, FakeAIProvider([]))
    assert await runner.run(job_id) is not None

    async with db_session_factory() as session:
        assert (await JobService(session, user_id=user[0], storage=storage).cancel(job_id)).status == "cancelled"

    assert await runner.run(job_id) is None
    assert (await _load(db_session_factory, job_id)).attempts == 1


async def test_a_cancel_during_a_failing_attempt_is_not_overwritten_by_the_failure(db_session_factory, tmp_path, monkeypatch):
    storage = LocalStorageProvider(tmp_path)
    user = await _user(db_session_factory)
    job_id = await _add(db_session_factory, user)

    async def kind(ctx):
        async with db_session_factory() as other_session:
            await JobService(other_session, user_id=user[0], storage=storage).cancel(job_id)
        raise ConnectionError("down")

    monkeypatch.setitem(runner_module.KINDS, IMPORT_TEXT, kind)

    assert await JobRunner(db_session_factory, storage, FakeAIProvider([])).run(job_id) is None
    assert (await _load(db_session_factory, job_id)).status == "cancelled"


def test_the_work_cannot_swallow_a_cancel():
    assert not issubclass(JobCancelled, Exception)


# -- (5) stuck detection ----------------------------------------------------------


def _ago(seconds: float) -> datetime:
    return datetime.now(timezone.utc) - timedelta(seconds=seconds)


def _requeue(calls: list):
    async def requeue(job_id: str, retry: int) -> None:
        calls.append((job_id, retry))

    return requeue


async def test_a_job_running_past_its_deadline_is_queued_again_while_attempts_remain(db_session_factory, tmp_path):
    user = await _user(db_session_factory)
    stuck = await _add(db_session_factory, user, status="running", attempts=1, started_at=_ago(policy.timeout_for(IMPORT_TEXT) + 120))
    calls: list = []

    assert await recover_stuck_jobs(db_session_factory, LocalStorageProvider(tmp_path), _requeue(calls)) == (1, 0)

    job = await _load(db_session_factory, stuck)
    assert calls == [(stuck, 1)]
    assert (job.status, job.stage, job.retry_count, job.failure_reason, job.dead_letter) == ("pending", "retrying", 1, "stuck", False)
    assert await _runner(db_session_factory, tmp_path).run(stuck) is None  # and it runs
    job = await _load(db_session_factory, stuck)
    assert (job.status, job.attempts, job.retry_count) == ("succeeded", 2, 1)


async def test_a_stuck_job_on_its_last_attempt_is_a_dead_letter_and_its_upload_goes(db_session_factory, tmp_path):
    storage = LocalStorageProvider(tmp_path)
    await storage.put("jobs/s/input", b"upload", "application/octet-stream")
    user = await _user(db_session_factory)
    stuck = await _add(
        db_session_factory, user, status="running", attempts=3, input_key="jobs/s/input", started_at=_ago(policy.timeout_for(IMPORT_TEXT) + 120)
    )
    calls: list = []

    assert await recover_stuck_jobs(db_session_factory, storage, _requeue(calls)) == (0, 1)

    job = await _load(db_session_factory, stuck)
    assert calls == []
    assert (job.status, job.stage, job.dead_letter, job.failure_reason) == ("failed", "failed", True, "stuck")
    assert job.error_message == "This couldn't be finished after 3 tries. Please try again in a few minutes."
    assert job.payload == {} and job.finished_at is not None and not (tmp_path / "jobs" / "s" / "input").exists()


async def test_a_job_still_inside_its_deadline_is_left_alone(db_session_factory, tmp_path):
    user = await _user(db_session_factory)
    # Past its timeout but inside the grace: its own timeout may be about to fire.
    just_late = await _add(db_session_factory, user, status="running", attempts=1, started_at=_ago(policy.timeout_for(IMPORT_TEXT) + 10))
    fresh = await _add(db_session_factory, user, status="running", attempts=1, started_at=_ago(5))
    waiting = await _add(db_session_factory, user)
    calls: list = []

    assert await recover_stuck_jobs(db_session_factory, LocalStorageProvider(tmp_path), _requeue(calls)) == (0, 0)

    assert [(await _load(db_session_factory, job_id)).status for job_id in (just_late, fresh, waiting)] == ["running", "running", "pending"]
    assert calls == []


async def test_the_deadline_is_the_job_types_own(db_session_factory, tmp_path):
    user = await _user(db_session_factory)
    started = _ago(400)  # an export is allowed 300 s (+60), an import of a file 600 s (+60)
    export = await _add(db_session_factory, user, job_type=JobType.EXPORT.value, status="running", attempts=1, started_at=started)
    import_file = await _add(db_session_factory, user, job_type=JobType.IMPORT_FILE.value, status="running", attempts=1, started_at=started)

    assert await recover_stuck_jobs(db_session_factory, LocalStorageProvider(tmp_path), _requeue([])) == (1, 0)

    assert [(await _load(db_session_factory, job_id)).status for job_id in (export, import_file)] == ["pending", "running"]


async def test_the_in_process_sweep_takes_a_stuck_job_back_and_runs_it(tmp_path, monkeypatch):
    from app.jobs import queue as queue_module
    from app.jobs import recovery

    # A SQLite file, not the shared in-memory connection of db_session_factory: cancelling
    # the sweep mid-query invalidates its connection, and a new in-memory one is empty.
    engine = create_async_engine(f"sqlite+aiosqlite:///{tmp_path / 'jobs.db'}")
    event.listen(engine.sync_engine, "connect", _enable_sqlite_fk)
    async with engine.begin() as connection:
        await connection.run_sync(Base.metadata.create_all)
    db_session_factory = async_sessionmaker(engine, expire_on_commit=False)
    monkeypatch.setattr(queue_module, "get_ai_provider", lambda: FakeAIProvider([]))
    monkeypatch.setattr(recovery, "RECOVER_INTERVAL_SECONDS", 0.05)
    stuck = await _add(db_session_factory, await _user(db_session_factory), status="running", attempts=1, started_at=_ago(2000))

    sweep = asyncio.create_task(queue_module.recover_in_process(db_session_factory, LocalStorageProvider(tmp_path)))
    try:
        async with asyncio.timeout(30):
            while (await _load(db_session_factory, stuck)).status != "succeeded":  # taken back, queued, run
                await asyncio.sleep(0.05)
    finally:
        sweep.cancel()
        await asyncio.gather(sweep, *queue_module._running, return_exceptions=True)

    job = await _load(db_session_factory, stuck)
    await engine.dispose()
    assert (job.status, job.attempts, job.retry_count) == ("succeeded", 2, 1)


async def test_a_job_that_could_not_be_queued_again_is_put_back_for_the_next_sweep(db_session_factory, tmp_path):
    user = await _user(db_session_factory)
    stuck = await _add(db_session_factory, user, status="running", attempts=1, started_at=_ago(2000))

    async def redis_is_down(job_id: str, retry: int) -> None:
        raise ConnectionError("Redis is down")

    assert await recover_stuck_jobs(db_session_factory, LocalStorageProvider(tmp_path), redis_is_down) == (0, 0)

    job = await _load(db_session_factory, job_id := stuck)
    assert (job.status, job.retry_count) == ("running", 0) and job_id
    calls: list = []
    assert await recover_stuck_jobs(db_session_factory, LocalStorageProvider(tmp_path), _requeue(calls)) == (1, 0)


async def test_a_worker_retrying_a_crashed_job_takes_it_only_when_it_is_past_its_deadline(db_session_factory, tmp_path):
    user = await _user(db_session_factory)
    running_elsewhere = await _add(db_session_factory, user, status="running", attempts=1, started_at=_ago(5))
    crashed = await _add(db_session_factory, user, status="running", attempts=1, started_at=_ago(2000))
    runner = _runner(db_session_factory, tmp_path)

    await runner.run(running_elsewhere)
    await runner.run(crashed)

    other = await _load(db_session_factory, running_elsewhere)
    retried = await _load(db_session_factory, crashed)
    assert (other.status, other.attempts) == ("running", 1)  # not run a second time at once
    assert (retried.status, retried.attempts, retried.retry_count) == ("succeeded", 2, 1)


async def test_of_two_workers_that_see_one_waiting_job_only_one_runs_it(db_session_factory, tmp_path, monkeypatch):
    runs = []

    async def kind(ctx):
        runs.append(1)
        return {"documentId": "made"}

    monkeypatch.setitem(runner_module.KINDS, IMPORT_TEXT, kind)
    job_id = await _add(db_session_factory, await _user(db_session_factory))

    await asyncio.gather(_runner(db_session_factory, tmp_path).run(job_id), _runner(db_session_factory, tmp_path).run(job_id))

    assert len(runs) == 1 and (await _load(db_session_factory, job_id)).attempts == 1


# -- (6) the dead letter ----------------------------------------------------------


async def test_the_sweep_never_deletes_a_dead_letter_but_removes_other_old_jobs(db_session_factory, tmp_path):
    user = await _user(db_session_factory)
    old = _ago(30 * 86400)
    dead = await _add(db_session_factory, user, status="failed", dead_letter=True, failure_reason="transient:ConnectionError", created_at=old, finished_at=old)
    plain = await _add(db_session_factory, user, status="failed", created_at=old, finished_at=old)

    assert await sweep_job_files(db_session_factory, LocalStorageProvider(tmp_path)) == (0, 1)

    assert await _load(db_session_factory, dead) is not None
    assert await _load(db_session_factory, plain) is None


def test_a_dead_letter_is_queryable_through_the_api_and_says_it_gave_up(api_db, monkeypatch):
    monkeypatch.setitem(runner_module.KINDS, IMPORT_TEXT, _failing_with(*[ConnectionError("down")] * 10))
    gave_up = _import().json()
    monkeypatch.setitem(runner_module.KINDS, IMPORT_TEXT, _failing_with(ValueError("a bug")))
    plain_failure = _import(text="# Two\n\nAnother text.").json()
    monkeypatch.setitem(runner_module.KINDS, IMPORT_TEXT, runner_module._import_text)
    fine = _import(text="# Three\n\nA third text.").json()

    assert (gave_up["status"], gave_up["deadLetter"], gave_up["retryCount"]) == ("failed", True, 2)
    assert "3 tries" in gave_up["error"]
    assert (plain_failure["status"], plain_failure["deadLetter"]) == ("failed", False)
    listed = client.get("/api/v1/jobs", params={"dead_letter": "true"}).json()
    assert [job["id"] for job in listed] == [gave_up["id"]]
    assert {job["id"] for job in client.get("/api/v1/jobs", params={"dead_letter": "false"}).json()} == {plain_failure["id"], fine["id"]}
    assert {job["id"] for job in client.get("/api/v1/jobs", params={"status": "failed"}).json()} == {gave_up["id"], plain_failure["id"]}
    row = next(row for row in _rows(api_db) if row.id == gave_up["id"])
    assert (row.dead_letter, row.failure_reason) == (True, "transient:ConnectionError")


async def test_the_job_a_restart_cut_off_says_why(db_session_factory, tmp_path):
    from app.jobs.queue import fail_interrupted_jobs

    job_id = await _add(db_session_factory, await _user(db_session_factory), status="running", started_at=_ago(5))

    assert await fail_interrupted_jobs(db_session_factory, LocalStorageProvider(tmp_path)) == 1

    job = await _load(db_session_factory, job_id)
    assert (job.status, job.failure_reason, job.dead_letter) == ("failed", "restart", False)
