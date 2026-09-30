"""Footnotes and endnotes (tracker DOCX-024): notes stay notes. The editor shows them at
the end of the document, each reference as its label; a Word export writes real
references where the labels are and the notes into their parts again, with their ids,
so a block nobody changed is copied with its reference. Measured in Word before this:
a10's 3 footnotes and 2 endnotes came back as 0 and 0, paragraphs at the end."""

import io
import re
import zipfile
from pathlib import Path

import pytest
from docx import Document as DocxDocument
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from fastapi.testclient import TestClient

from app.export.docx_export import build_docx
from app.export.package_check import package_problems
from app.export.provenance import stamp
from app.fidelity.report import ReportBuilder
from app.formatting.engine import recompute_styles
from app.main import app
from app.models.document import Document, ElementType, InlineRun
from app.parsers.docx import parse_docx

A10 = Path(__file__).parent / "fixtures" / "word" / "a10-notes.docx"
_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
client = TestClient(app, base_url="https://testserver")


def _part(data: bytes, name: str) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as package:
        return package.read(name).decode("utf-8") if name in package.namelist() else ""


def _references(data: bytes) -> list[tuple[str, str]]:
    return re.findall(r'<w:(footnote|endnote)Reference w:id="(\d+)"', _part(data, "word/document.xml"))


def _notes(data: bytes, kind: str) -> dict[str, str]:
    """Each note in its part, by id: its text."""
    xml = _part(data, f"word/{kind}s.xml")
    return {
        note_id: "".join(re.findall(r"<w:t(?: [^>]*)?>([^<]*)</w:t>", body))
        for note_id, body in re.findall(rf'<w:{kind} w:id="(\d+)"[^>]*>(.*?)</w:{kind}>', xml, re.S)
        if note_id not in ("-1", "0")
    }


def _a10() -> tuple[bytes, Document]:
    source = A10.read_bytes()
    document = parse_docx(source, A10.name)
    recompute_styles(document)
    stamp(document)  # as an upload does
    return source, document


_FOOTNOTES = {"1": " Footnote 1 text with bold words", "2": " Footnote 2 text with bold words", "3": " Footnote 3 text with bold words"}
_ENDNOTES = {"1": " Endnote 1 text.", "2": " Endnote 2 text."}


def test_the_import_keeps_where_each_note_is_referred_to():
    _, document = _a10()

    references = [
        fragment["note"]
        for element in document.elements
        for fragment in (element.preservedAttributes or {}).get("ooxml") or []
        if fragment["kind"] == "note"
    ]
    notes = [element.preservedAttributes["note"] for element in document.elements if element.type == ElementType.FOOTNOTE]

    assert references == ["footnote:1", "footnote:2", "footnote:3", "endnote:1", "endnote:2"]
    assert [(note["note"], note["label"]) for note in notes] == [
        ("footnote:1", "1"), ("footnote:2", "2"), ("footnote:3", "3"), ("endnote:1", "i"), ("endnote:2", "ii")
    ]
    assert "Footnotes and endnotes are shown at the end of the document; a Word export puts them back as notes." in document.unsupportedFeatures


@pytest.mark.parametrize("into_the_original", [False, True], ids=["a new file", "into the original"])
def test_notes_go_back_into_word_as_notes(into_the_original):
    source, document = _a10()

    exported = build_docx(document, source=source) if into_the_original else build_docx(document)

    assert _references(exported) == [("footnote", "1"), ("footnote", "2"), ("footnote", "3"), ("endnote", "1"), ("endnote", "2")]
    assert _notes(exported, "footnote") == _FOOTNOTES and _notes(exported, "endnote") == _ENDNOTES
    assert "<w:b/>" in _part(exported, "word/footnotes.xml")  # their bold words bold
    assert "Footnote 1 text" not in _part(exported, "word/document.xml")  # not at the end of the body any more
    assert package_problems(exported) == []


def test_an_unchanged_block_is_copied_with_its_reference():
    source, document = _a10()
    report = ReportBuilder()

    exported = build_docx(document, source=source, report=report)

    features = {item.feature for item in report.items()}
    assert "export.docx.original_blocks" in features and "export.docx.rewritten_blocks" not in features
    assert '<w:footnoteReference w:id="1"/>' in _part(exported, "word/document.xml")


def test_a_note_changed_here_is_written_as_it_is_now():
    source, document = _a10()
    second = next(element for element in document.elements if element.content.startswith("2 "))
    second.inline = [*(second.inline or []), InlineRun(text=" Revised.")]
    second.content = f"{second.content} Revised."

    exported = build_docx(document, source=source)

    assert _notes(exported, "footnote")["2"] == " Footnote 2 text with bold words Revised."
    assert package_problems(exported) == []


def test_a_note_deleted_here_leaves_its_label_as_text():
    source, document = _a10()
    document.elements = [element for element in document.elements if (element.preservedAttributes or {}).get("note", {}).get("note") != "footnote:2"]

    exported = build_docx(document, source=source)

    assert _references(exported) == [("footnote", "1"), ("footnote", "3"), ("endnote", "1"), ("endnote", "2")]
    assert list(_notes(exported, "footnote")) == ["1", "3"]
    assert "Sentence 2 with a footnote.2" in [p.text for p in DocxDocument(io.BytesIO(exported)).paragraphs]
    assert package_problems(exported) == []  # no reference to a note that isn't there


def test_a_note_referred_to_from_a_table_is_written_at_the_end_and_said_so():
    source, document = _a10()
    word = DocxDocument(io.BytesIO(source))
    cell = word.add_table(rows=1, cols=1).cell(0, 0)
    reference = word.element.body.find(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}footnoteReference").getparent()
    cell.paragraphs[0]._p.append(reference)  # footnote 1, referred to from a cell
    buffer = io.BytesIO()
    word.save(buffer)
    moved = buffer.getvalue()
    document = parse_docx(moved, "moved.docx")
    report = ReportBuilder()

    exported = build_docx(document, report=report)

    assert "1" not in _notes(exported, "footnote")
    # Labelled as the import numbers them, in the order they are referred to: the cell's comes last.
    assert "3  Footnote 1 text with bold words" in [p.text for p in DocxDocument(io.BytesIO(exported)).paragraphs]
    assert "export.docx.notes_at_end" in {item.feature for item in report.items()}
    assert package_problems(exported) == []


def test_a_style_is_never_defined_twice():
    """Found here: a Word export asked for "Footnote Text", found no style of that name
    beside Word's own "footnote text", and added a second FootnoteText."""
    source, document = _a10()

    exported = build_docx(document, source=source)

    assert _part(exported, "word/styles.xml").count('w:styleId="FootnoteText"') == 1
    assert package_problems(exported) == []


def test_the_package_check_names_a_reference_to_a_note_that_isnt_there():
    word = DocxDocument()
    word.add_paragraph()._p.append(parse_xml(f'<w:r {nsdecls("w")}><w:footnoteReference w:id="9"/></w:r>'))
    buffer = io.BytesIO()
    word.save(buffer)

    assert "word/document.xml: footnote '9' isn't defined" in package_problems(buffer.getvalue())


def test_a_pdf_says_the_notes_are_printed_at_the_end(api_db):
    client.cookies.clear()
    assert client.post("/api/v1/auth/register", json={"email": "notes@example.com", "password": "long enough password"}).status_code == 201
    document = client.post("/api/v1/documents/upload", files={"file": (A10.name, A10.read_bytes(), _DOCX)}).json()

    job = client.post("/api/v1/jobs/export", json={"documentId": document["id"], "format": "pdf"}).json()
    finished = client.get(f"/api/v1/jobs/{job['id']}").json()
    client.cookies.clear()

    assert "export.pdf.notes" in {item["feature"] for item in finished["result"]["fidelity"]["items"]}
