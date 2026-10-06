"""Storage counts what is stored, in bytes (tracker PLAN-005): documents as saved,
their version history as stored (compressed, PERF-004; an older uncompressed row as
its JSON), pictures and the Word files imports were made from -- compared here with
the bytes actually stored, on SQLite and on PostgreSQL (TEST-031). length() of text
counts characters on both databases, so Cyrillic text is what tells bytes apart.
A job's upload and an export's file aren't counted: the workspace doesn't keep them."""

import io

import pytest
import pytest_asyncio
from docx import Document as DocxDocument
from fastapi.testclient import TestClient
from PIL import Image as PILImage
from sqlalchemy import Text, cast, create_engine, event, select
from sqlalchemy.ext.asyncio import async_sessionmaker
from sqlalchemy.pool import NullPool

from app.db.base import Base
from app.db.models import Document as DocumentRow
from app.db.models import DocumentAsset, DocumentVersion
from app.db.session import make_engine
from app.main import app
from app.services.auth_service import AuthService
from app.services.document_service import DocumentService
from app.services.usage_service import storage_bytes
from app.storage.local_provider import LocalStorageProvider
from tests.conftest import _enable_sqlite_fk
from tests.fakes import FakeAIProvider

client = TestClient(app, base_url="https://testserver")
_CYRILLIC = "Годишен доклад за качеството на водите в региона"


@pytest.fixture(params=["sqlite", pytest.param("postgres", marks=pytest.mark.postgres)])
def database(request):
    return request.getfixturevalue("postgres_sessions") if request.param == "postgres" else None


@pytest_asyncio.fixture
async def sessions(database, tmp_path):
    if database is not None:
        yield database
        return
    path = tmp_path / "storage.db"
    sync_engine = create_engine(f"sqlite:///{path}")
    Base.metadata.create_all(sync_engine)
    sync_engine.dispose()
    engine = make_engine(f"sqlite+aiosqlite:///{path}", poolclass=NullPool)
    event.listen(engine.sync_engine, "connect", _enable_sqlite_fk)
    yield async_sessionmaker(engine, expire_on_commit=False)
    await engine.dispose()


def _word_file() -> bytes:
    picture = io.BytesIO()
    PILImage.new("RGB", (40, 30), (200, 40, 40)).save(picture, format="PNG")
    picture.seek(0)
    document = DocxDocument()
    document.add_heading(_CYRILLIC, level=1)
    for number in range(5):
        document.add_paragraph(f"{_CYRILLIC}: параграф {number}, с още малко текст на кирилица.")
    document.add_picture(picture)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


async def _stored_bytes(session, workspace_id: str, root) -> tuple[int, int]:
    """The bytes stored for the workspace, worked out apart from storage_bytes: each
    JSON text as the database holds it, encoded here; each compressed version's bytes;
    each asset's file as it lies in storage. And the same documents in characters."""
    texts = (await session.execute(select(cast(DocumentRow.data, Text)).where(DocumentRow.workspace_id == workspace_id))).scalars().all()
    versions = (
        await session.execute(
            select(DocumentVersion.compressed_data, cast(DocumentVersion.legacy_data, Text))
            .join(DocumentRow, DocumentRow.id == DocumentVersion.document_id)
            .where(DocumentRow.workspace_id == workspace_id)
        )
    ).all()
    keys = (await session.execute(select(DocumentAsset.storage_key).where(DocumentAsset.workspace_id == workspace_id))).scalars().all()
    files = sum((root / key).stat().st_size for key in keys)
    in_bytes = sum(len(text.encode("utf-8")) for text in texts)
    in_bytes += sum(len(packed) if packed is not None else len(legacy.encode("utf-8")) for packed, legacy in versions)
    return in_bytes + files, sum(len(text) for text in texts)


async def test_storage_is_the_bytes_actually_stored(sessions, tmp_path):
    root = tmp_path / "files"
    async with sessions() as session:
        user = await AuthService(session).register("stored@example.com", "long enough password", None)
        workspace_id = await AuthService(session).default_workspace_id(user.id)
        await session.commit()
    async with sessions() as session:
        service = DocumentService(session, user_id=user.id, storage=LocalStorageProvider(root))
        created = await service.create_from_bytes(_word_file(), "доклад.docx", None, FakeAIProvider([]))
        await service.rename(created.id, title="Доклад, преименуван")  # a compressed undo step
    async with sessions() as session:
        # A version written before PERF-004: uncompressed JSON with Cyrillic in it.
        session.add(
            DocumentVersion(document_id=created.id, revision_number=99, kind="change", legacy_data={"title": _CYRILLIC}, created_by=user.id)
        )
        await session.commit()

    async with sessions() as session:
        counted = await storage_bytes(session, workspace_id)
        stored, characters = await _stored_bytes(session, workspace_id, root)
        kinds = (await session.execute(select(DocumentAsset.content_type).where(DocumentAsset.workspace_id == workspace_id))).scalars().all()

    assert counted == stored
    assert sorted(kinds) == ["application/vnd.openxmlformats-officedocument.wordprocessingml.document", "image/png"]  # kept original, picture
    assert stored - characters > len(_CYRILLIC)  # bytes, not characters: the Cyrillic text counts twice


@pytest.fixture
def signed_in(api_db):
    client.cookies.clear()
    assert client.post("/api/v1/auth/register", json={"email": "kept@example.com", "password": "long enough password"}).status_code == 201
    yield
    client.cookies.clear()


def test_a_jobs_upload_and_an_exports_file_aren_t_storage(signed_in):
    imported = client.post(
        "/api/v1/jobs/import-file",
        files={"file": ("доклад.docx", _word_file(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    ).json()
    document_id = imported["result"]["documentId"]
    before = client.get("/api/v1/usage").json()["storageBytes"]

    exported = client.post("/api/v1/jobs/export", json={"documentId": document_id, "format": "pdf"}).json()

    assert exported["status"] == "succeeded" and exported["result"]["size"] > 0
    assert client.get("/api/v1/usage").json()["storageBytes"] == before
