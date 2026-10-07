"""Tracked changes rejected, and deleted paragraph marks (tracker DOCX-022A). Rejecting them all
is a third choice beside keeping them for Word and accepting them: the document is read again
from its Word file with every change rejected -- deletions back, insertions out, formatting
changes undone -- and that file is kept instead. And a paragraph whose mark was deleted while
tracking runs into the next one in the accepted reading, as Word does on accepting (before, the
import showed the two apart). What Word does was measured in Word on the file built here: the
expectations below are its results."""

import io
import re
import zipfile
from pathlib import Path

import pytest
from docx import Document as DocxDocument
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from fastapi.testclient import TestClient

from app.export.docx_export import build_docx
from app.export.provenance import stamp
from app.fidelity.imports import TRACKED_REJECTED
from app.formatting.engine import recompute_styles
from app.main import app
from app.parsers.docx import parse_docx
from app.parsers.docx_revisions import reject_all
from app.services.ingestion_service import build_document_from_docx

A07 = Path(__file__).parent / "fixtures" / "word" / "a07-review.docx"
_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_NS = nsdecls("w")
client = TestClient(app, base_url="https://testserver")


def _marks(data: bytes) -> int:
    """How many marks of a tracked change the body holds."""
    xml = zipfile.ZipFile(io.BytesIO(data)).read("word/document.xml").decode("utf-8")
    return len(re.findall(r"<w:(?:ins|del|moveFrom|moveTo|rPrChange|pPrChange|cellIns|cellDel)\b", xml))


def _tracked_file(*, unsafe_field: bool = False) -> bytes:
    """A heading whose mark was deleted, then body text; a paragraph inserted whole; deleted and
    inserted words and a bolding; a paragraph deleted whole; a heading; a table with an inserted
    and a deleted row."""
    attributes = 'w:id="{}" w:author="Tester" w:date="2026-10-01T00:00:00Z"'
    doc = DocxDocument()
    body = doc.element.body

    def paragraph(xml: str) -> None:
        body.insert(len(body) - 1, parse_xml(f"<w:p {_NS}>{xml}</w:p>"))

    paragraph(f'<w:pPr><w:pStyle w:val="Heading1"/><w:rPr><w:del {attributes.format(1)}/></w:rPr></w:pPr><w:r><w:t xml:space="preserve">Heading text </w:t></w:r>')
    paragraph("<w:r><w:t>then body text.</w:t></w:r>")
    paragraph(f'<w:pPr><w:rPr><w:ins {attributes.format(2)}/></w:rPr></w:pPr><w:ins {attributes.format(3)}><w:r><w:t>Inserted paragraph.</w:t></w:r></w:ins>')
    paragraph(
        f'<w:r><w:t xml:space="preserve">Keep </w:t></w:r><w:del {attributes.format(4)}><w:r><w:delText xml:space="preserve">old </w:delText></w:r></w:del>'
        f'<w:ins {attributes.format(5)}><w:r><w:t xml:space="preserve">new </w:t></w:r></w:ins>'
        f"<w:r><w:rPr><w:b/><w:rPrChange {attributes.format(6)}><w:rPr/></w:rPrChange></w:rPr><w:t>bolded.</w:t></w:r>"
    )
    paragraph(f'<w:pPr><w:rPr><w:del {attributes.format(7)}/></w:rPr></w:pPr><w:del {attributes.format(8)}><w:r><w:delText>Gone paragraph.</w:delText></w:r></w:del>')
    paragraph('<w:pPr><w:pStyle w:val="Heading2"/></w:pPr><w:r><w:t>Last heading</w:t></w:r>')
    if unsafe_field:
        paragraph('<w:fldSimple w:instr=" INCLUDETEXT secret.docx "><w:r><w:t>Included text.</w:t></w:r></w:fldSimple>')
    table = doc.add_table(rows=3, cols=1)
    for row, text in zip(table.rows, ["Row one", "Row inserted", "Row deleted"]):
        row.cells[0].paragraphs[0].add_run(text)
    table.rows[1]._tr.get_or_add_trPr().append(parse_xml(f"<w:ins {_NS} {attributes.format(9)}/>"))
    table.rows[2]._tr.get_or_add_trPr().append(parse_xml(f"<w:del {_NS} {attributes.format(10)}/>"))
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def test_rejecting_gives_what_word_gives_on_reject_all():
    rejected = DocxDocument(io.BytesIO(reject_all(_tracked_file())))
    # Measured in Word (Revisions.RejectAll) on the same file.
    assert [(paragraph.text, paragraph.style.name) for paragraph in rejected.paragraphs] == [
        ("Heading text ", "Heading 1"),
        ("then body text.", "Normal"),
        ("Keep old bolded.", "Normal"),
        ("Gone paragraph.", "Normal"),
        ("Last heading", "Heading 2"),
    ]
    assert not any(run.bold for run in rejected.paragraphs[2].runs)  # the bolding undone
    assert [row.cells[0].text for row in rejected.tables[0].rows] == ["Row one", "Row deleted"]
    assert _marks(reject_all(_tracked_file())) == 0


def test_the_accepted_reading_joins_a_paragraph_whose_mark_was_deleted_to_the_next():
    document = build_document_from_docx(_tracked_file(), "tracked.docx", None)
    # Measured in Word (Revisions.AcceptAll): the heading's text runs into the body text, which
    # keeps its own style; the paragraph deleted whole is gone, into the heading after it.
    shown = [(element.type.value, element.content, element.sourceBlocks) for element in document.elements]
    assert shown == [
        ("paragraph", "Heading text then body text.", [0, 1]),
        ("paragraph", "Inserted paragraph.", [2]),
        ("paragraph", "Keep new bolded.", [3]),
        ("heading", "Last heading", [4, 5]),
        ("table", "Row one\nRow inserted", [6]),
    ]
    assert document.trackedChanges == "kept" and document.importReport.contentStatus == "verified"


def test_a_joined_paragraph_left_alone_is_copied_with_both_its_paragraphs():
    source = _tracked_file()
    document = parse_docx(source, "tracked.docx")
    recompute_styles(document)
    stamp(document)
    exported = build_docx(document, source=source)
    assert _marks(exported) == _marks(source)  # every change still there for Word, the deleted marks too
    assert parse_docx(exported, "again.docx").elements[0].content == "Heading text then body text."


def test_rejecting_a07_leaves_no_tracked_change_and_reads_its_deletions_not_its_insertions():
    source = A07.read_bytes()
    rejected = reject_all(source)
    assert _marks(source) > 0 and _marks(rejected) == 0  # opened in Word: 0 revisions, its paragraphs as Word's RejectAll
    accepted_words = set(" ".join(element.content for element in parse_docx(source, "a.docx").elements).split())
    rejected_document = parse_docx(rejected, "r.docx")
    rejected_words = set(" ".join(element.content for element in rejected_document.elements).split())
    assert accepted_words - rejected_words and rejected_words - accepted_words  # insertions out, deletions back
    assert rejected_document.trackedChanges is None


@pytest.fixture
def signed_in(api_db):
    client.cookies.clear()
    assert client.post("/api/v1/auth/register", json={"email": "rejects@example.com", "password": "long enough password"}).status_code == 201
    yield
    client.cookies.clear()


def _upload(data: bytes) -> dict:
    response = client.post("/api/v1/documents/upload", files={"file": ("tracked.docx", data, _DOCX)})
    assert response.status_code == 201, response.text
    return response.json()


def _contents(document: dict) -> list[str]:
    return [element["content"] for element in document["elements"]]


def test_rejecting_all_reads_the_document_again_and_undo_brings_the_changes_back(signed_in):
    uploaded = _upload(_tracked_file(unsafe_field=True))
    assert uploaded["trackedChanges"] == "kept"

    response = client.put(f"/api/v1/documents/{uploaded['id']}/tracked-changes", json={"choice": "rejected"})
    assert response.status_code == 200, response.text
    rejected = response.json()
    assert _contents(rejected)[:5] == ["Heading text ", "then body text.", "Keep old bolded.", "Gone paragraph.", "Last heading"]
    assert rejected["trackedChanges"] == "rejected" and rejected["id"] == uploaded["id"]
    features = {item["feature"]: item for item in rejected["importReport"]["items"]}
    assert features["docx.tracked_changes"]["reason"] == TRACKED_REJECTED and features["docx.tracked_changes"]["contentChanged"]
    assert "docx.field.unsafe" in features  # what was made safe in the file is still said
    exported = client.get(f"/api/v1/documents/{uploaded['id']}/export/docx")
    assert exported.status_code == 200 and _marks(exported.content) == 0
    assert "Gone paragraph." in [paragraph.text for paragraph in DocxDocument(io.BytesIO(exported.content)).paragraphs]

    again = client.put(f"/api/v1/documents/{uploaded['id']}/tracked-changes", json={"choice": "kept"})
    assert again.status_code == 409  # nothing left to keep: undo it to choose again

    undone = client.post(f"/api/v1/documents/{uploaded['id']}/undo").json()
    assert undone["trackedChanges"] == "kept" and _contents(undone) == _contents(uploaded)
    exported = client.get(f"/api/v1/documents/{uploaded['id']}/export/docx")
    assert _marks(exported.content) == _marks(_tracked_file(unsafe_field=True))  # the first file, with its changes


def test_rejecting_all_asks_first_when_changes_made_here_would_go(signed_in):
    uploaded = _upload(_tracked_file())
    elements = uploaded["elements"]
    elements[1] = {**elements[1], "content": "Edited here.", "inline": [{"text": "Edited here.", "marks": []}]}
    saved = client.put(f"/api/v1/documents/{uploaded['id']}/content", json={"elements": elements}, headers={"If-Match": str(uploaded["revision"])})
    assert saved.status_code == 200, saved.text

    refused = client.put(f"/api/v1/documents/{uploaded['id']}/tracked-changes", json={"choice": "rejected"})
    assert refused.status_code == 409 and refused.json()["code"] == "edits_would_be_lost"
    assert "Edited here." in _contents(client.get(f"/api/v1/documents/{uploaded['id']}").json())  # nothing written

    confirmed = client.put(f"/api/v1/documents/{uploaded['id']}/tracked-changes", json={"choice": "rejected", "discardEdits": True})
    assert confirmed.status_code == 200 and "Edited here." not in _contents(confirmed.json())
