"""The preservation layer (корекции.docx §11): what the editor can't show yet --
equations, Word fields, bookmarks, links to bookmarks, comments -- stays with its
paragraph and goes back into the file when the document is exported to Word."""

import io
import zipfile

from docx import Document as DocxDocument
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from app.export.docx_export import build_docx
from app.models.document import Document, InlineRun
from app.parsers.docx import parse_docx

_NS = nsdecls("w", "m", "r")
_EQUATION = (
    f'<w:p {_NS}><w:r><w:t xml:space="preserve">Area: </w:t></w:r><m:oMath><m:r><m:t>A=π</m:t></m:r>'
    "<m:sSup><m:e><m:r><m:t>r</m:t></m:r></m:e><m:sup><m:r><m:t>2</m:t></m:r></m:sup></m:sSup></m:oMath>"
    '<w:r><w:t xml:space="preserve"> in m².</w:t></w:r></w:p>'
)
_DATE_FIELD = (
    f'<w:p {_NS}><w:r><w:t xml:space="preserve">Printed on </w:t></w:r>'
    '<w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText xml:space="preserve"> DATE \\@ "d.M.yyyy" </w:instrText></w:r>'
    '<w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>25.9.2026</w:t></w:r><w:r><w:fldChar w:fldCharType="end"/></w:r>'
    '<w:r><w:t xml:space="preserve">, final.</w:t></w:r></w:p>'
)


def _docx(*paragraphs: str, build=None) -> bytes:
    doc = DocxDocument()
    body = doc.element.body
    for xml in paragraphs:
        body.insert(len(body) - 1, parse_xml(xml))  # before the final sectPr
    if build is not None:
        build(doc)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _document_xml(docx_bytes: bytes) -> str:
    return zipfile.ZipFile(io.BytesIO(docx_bytes)).read("word/document.xml").decode("utf-8")


def _round_trip(source: bytes) -> tuple[Document, bytes, Document]:
    first = parse_docx(source, "in.docx")
    exported = build_docx(first)
    return first, exported, parse_docx(exported, "out.docx")


def _retype(document: Document, text: str) -> None:
    """What typing in the editor does: new content, the element's other fields kept."""
    element = document.elements[0]
    element.inline = [InlineRun(text=text)]
    element.content = text


def test_an_equation_goes_back_into_word_as_an_equation():
    first, exported, back = _round_trip(_docx(_EQUATION))

    xml = _document_xml(exported)
    assert "<m:oMath" in xml and "<m:sSup>" in xml
    assert "<w:t>A=π" not in xml  # the linear text isn't written next to it
    assert first.elements[0].content == back.elements[0].content == "Area: A=πr2 in m²."


def test_an_equation_whose_text_was_changed_is_exported_as_that_text():
    first = parse_docx(_docx(_EQUATION), "in.docx")
    _retype(first, "Area: A=πd²/4 in m².")

    xml = _document_xml(build_docx(first))

    assert "<m:oMath" not in xml
    assert "A=πd²/4" in xml


def test_a_field_goes_back_with_its_last_result():
    first, exported, back = _round_trip(_docx(_DATE_FIELD))

    [field] = first.elements[0].preservedAttributes["ooxml"]
    assert (field["kind"], field["instr"].strip(), field["text"]) == ("field", 'DATE \\@ "d.M.yyyy"', "25.9.2026")
    xml = _document_xml(exported)
    assert xml.count('w:fldCharType="begin"') == 1 and 'DATE \\@ "d.M.yyyy"' in xml
    assert back.elements[0].content == "Printed on 25.9.2026, final."
    assert back.elements[0].preservedAttributes["ooxml"][0]["instr"] == field["instr"]


def test_a_kept_field_is_found_again_after_text_before_it_changes():
    first = parse_docx(_docx(_DATE_FIELD), "in.docx")
    _retype(first, "Report printed on 25.9.2026, final.")

    back = parse_docx(build_docx(first), "out.docx")

    [field] = back.elements[0].preservedAttributes["ooxml"]
    assert (field["start"], field["text"]) == (18, "25.9.2026")


def test_a_simple_field_is_kept_too():
    source = _docx(
        f'<w:p {_NS}><w:r><w:t xml:space="preserve">File: </w:t></w:r>'
        '<w:fldSimple w:instr=" FILENAME "><w:r><w:t>report.docx</w:t></w:r></w:fldSimple></w:p>'
    )

    first, exported, back = _round_trip(source)

    assert "FILENAME" in _document_xml(exported)
    assert back.elements[0].preservedAttributes["ooxml"][0]["text"] == "report.docx"


def test_bookmarks_and_links_to_them_go_back_into_word():
    source = _docx(
        f'<w:p {_NS}><w:bookmarkStart w:id="3" w:name="Results"/><w:r><w:t>Results</w:t></w:r><w:bookmarkEnd w:id="3"/></w:p>',
        f'<w:p {_NS}><w:r><w:t xml:space="preserve">As shown, </w:t></w:r>'
        '<w:hyperlink w:anchor="Results"><w:r><w:t>see the results</w:t></w:r></w:hyperlink></w:p>',
    )

    _, exported, back = _round_trip(source)

    xml = _document_xml(exported)
    assert 'w:name="Results"' in xml and "w:bookmarkEnd" in xml
    assert 'w:anchor="Results"' in xml
    assert back.elements[0].preservedAttributes["ooxml"][0]["name"] == "Results"
    assert back.elements[1].preservedAttributes["ooxml"][0] == {
        "kind": "link",
        "anchor": "Results",
        "start": 10,
        "end": 25,
        "text": "see the results",
    }


def test_comments_go_back_into_word_with_their_author_and_text():
    def comment(doc: DocxDocument) -> None:
        paragraph = doc.add_paragraph("Before ")
        disputed = paragraph.add_run("the disputed claim")
        paragraph.add_run(" after.")
        doc.add_comment(disputed, text="Source?", author="Reviewer", initials="RV")

    first, exported, back = _round_trip(_docx(build=comment))

    [kept] = first.elements[0].preservedAttributes["ooxml"]
    assert (kept["kind"], kept["text"], kept["comment"], kept["author"], kept["initials"]) == (
        "comment",
        "the disputed claim",
        "Source?",
        "Reviewer",
        "RV",
    )
    assert [(c.author, c.text) for c in DocxDocument(io.BytesIO(exported)).comments] == [("Reviewer", "Source?")]
    assert back.elements[0].preservedAttributes["ooxml"][0]["text"] == "the disputed claim"
    assert any(note.startswith("Comments aren't shown in the editor yet") for note in first.unsupportedFeatures)


def test_what_sits_inside_a_table_is_kept_only_as_text_and_said_so():
    def table_with_equation(doc: DocxDocument) -> None:
        cell = doc.add_table(rows=1, cols=1).cell(0, 0)
        cell.paragraphs[0]._p.append(parse_xml(f"<m:oMath {_NS}><m:r><m:t>x+1</m:t></m:r></m:oMath>"))

    document = parse_docx(_docx(build=table_with_equation), "in.docx")

    assert document.elements[0].table.rows[0].cells[0].inline[0].text == "x+1"
    assert "Equations inside lists, tables, footnotes or code were kept only as their text." in document.unsupportedFeatures


def test_malformed_kept_data_is_left_out_of_the_export():
    """preservedAttributes comes back from the browser; nothing in it is trusted."""
    document = parse_docx(_docx(_EQUATION), "in.docx")
    document.elements[0].preservedAttributes = {
        "ooxml": [
            {"kind": "equation", "start": 6, "end": 11, "text": "A=πr2", "xml": f"<w:p {_NS}/>"},  # not an equation
            {"kind": "equation", "start": 6, "end": 11, "text": "A=πr2", "xml": "<not xml"},
            {"kind": "bookmark", "start": 0, "end": 0, "text": "", "name": "bad name!"},
            {"kind": "field", "start": 0, "end": 4, "text": "Area", "instr": "  "},
            {"kind": "link", "start": -1, "end": 4, "text": "Area", "anchor": "x"},
            {"kind": "script", "start": 0, "end": 0},
            "garbage",
        ]
    }

    exported = build_docx(document)

    xml = _document_xml(exported)
    assert "<m:oMath" not in xml and "bad name" not in xml and "fldChar" not in xml and "w:anchor" not in xml
    assert parse_docx(exported, "out.docx").elements[0].content == "Area: A=πr2 in m²."  # the text itself stays
