"""What a Word export keeps of what the app doesn't hold, and says so (tracker
DOCX-013 part 2, FID-007). With the original file kept, a block nobody changed
is copied as it is (DOCX-028): the import report names what lives in such
blocks as kept while they are unchanged, and an export that writes a block anew
names what that block lost. The language text is in is held by the model, so
it is never lost; a link the app doesn't allow is never copied back."""

import io
import zipfile

from docx import Document as DocxDocument
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from fastapi.testclient import TestClient

from app.export.docx_export import build_docx
from app.fidelity.docx_detect import detect_docx_features
from app.main import app
from app.parsers.docx import parse_docx

client = TestClient(app, base_url="https://testserver")
_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_W = nsdecls("w")
_W14 = 'xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml"'


def _run(paragraph, text: str, properties: str) -> None:
    paragraph._p.append(parse_xml(f'<w:r {_W} {_W14}><w:rPr>{properties}</w:rPr><w:t xml:space="preserve">{text}</w:t></w:r>'))


def _word_file() -> bytes:
    word = DocxDocument()  # en-US by default
    greeting = word.add_paragraph()
    _run(greeting, "Добър ден", '<w:lang w:val="bg-BG"/>')
    _run(greeting, " and good day.", '<w:lang w:val="en-US"/>')
    _run(word.add_paragraph(), "Stretched wide", '<w:w w:val="150"/>')
    _run(word.add_paragraph(), "In the shade", "<w:shadow/>")
    _run(word.add_paragraph(), "Glowing", '<w14:glow w14:rad="63500"><w14:srgbClr w14:val="FF0000"/></w14:glow>')
    word.element.body.insert(
        len(word.element.body) - 1,
        parse_xml(
            f'<w:sdt {_W}><w:sdtPr><w:alias w:val="Status"/></w:sdtPr>'
            "<w:sdtContent><w:p><w:r><w:t>Draft for review</w:t></w:r></w:p></w:sdtContent></w:sdt>"
        ),
    )
    word.add_paragraph("An ordinary last paragraph.")
    buffer = io.BytesIO()
    word.save(buffer)
    return buffer.getvalue()


def _upload(email: str, data: bytes) -> dict:
    client.cookies.clear()
    assert client.post("/api/v1/auth/register", json={"email": email, "password": "long enough password"}).status_code == 201
    response = client.post("/api/v1/documents/upload", files={"file": ("copy.docx", data, _DOCX)})
    assert response.status_code == 201, response.text
    return response.json()


def _export(document_id: str) -> tuple[dict, str, str]:
    job = client.post("/api/v1/jobs/export", json={"documentId": document_id, "format": "docx"}).json()
    finished = client.get(f"/api/v1/jobs/{job['id']}").json()
    content = client.get(f"/api/v1/jobs/{job['id']}/file").content
    with zipfile.ZipFile(io.BytesIO(content)) as package:
        body = package.read("word/document.xml").decode("utf-8")
        rels = package.read("word/_rels/document.xml.rels").decode("utf-8")
    return {item["feature"]: item for item in finished["result"]["fidelity"]["items"]}, body, rels


def _edit(document: dict, old: str, new: str) -> None:
    elements = document["elements"]
    element = next(element for element in elements if element["content"] == old)
    element["content"], element["inline"] = new, [{"text": new, "marks": []}]
    assert client.put(f"/api/v1/documents/{document['id']}/content", json={"elements": elements}).status_code == 200


def test_the_language_text_is_in_is_kept_where_it_isnt_the_documents_own():
    document = parse_docx(_word_file(), "copy.docx")

    runs = [(run.text, [mark.lang for mark in run.marks]) for run in document.elements[0].inline]
    assert runs == [("Добър ден", ["bg-BG"]), (" and good day.", [])]  # en-US is the document's own
    exported = build_docx(document)
    with zipfile.ZipFile(io.BytesIO(exported)) as package:
        assert '<w:lang w:val="bg-BG"/>' in package.read("word/document.xml").decode("utf-8")
    assert parse_docx(exported, "again.docx").elements[0].inline == document.elements[0].inline


def test_stretched_text_and_text_effects_are_named():
    items = {item.feature: item for item in detect_docx_features(_word_file())}

    assert items["docx.character_scale"].sourceState == "e.g. “Stretched wide”"
    assert items["docx.text_effects"].count == 2  # the shadow and the glow


def test_with_the_file_kept_they_are_kept_while_unchanged(api_db):
    document = _upload("kept@example.com", _word_file())
    client.cookies.clear()

    items = {item["feature"]: item for item in document["importReport"]["items"]}
    for feature in ("docx.content_control", "docx.character_scale", "docx.text_effects"):
        assert items[feature]["policy"] == "detected_not_editable", feature
        assert "while the paragraph that holds it isn't changed or restyled here" in items[feature]["reason"]


def test_a_block_written_anew_names_what_it_lost(api_db):
    document = _upload("rewritten@example.com", _word_file())
    items, body, _ = _export(document["id"])
    assert "export.docx.rewritten_blocks" not in items and '<w:w w:val="150"/>' in body  # unchanged: copied

    _edit(document, "Stretched wide", "Stretched wide, edited")
    _edit(document, "Draft for review", "Final")
    items, body, _ = _export(document["id"])
    client.cookies.clear()

    rewritten = items["export.docx.rewritten_blocks"]
    assert rewritten["policy"] == "lossy" and rewritten["count"] == 2
    assert "content controls" in rewritten["reason"] and "stretched text" in rewritten["reason"]
    assert '<w:w w:val="150"/>' not in body and "<w:shadow/>" in body  # the edited one anew, the others as they were


def test_a_link_the_app_doesnt_allow_is_never_copied_back(api_db):
    word = DocxDocument()
    paragraph = word.add_paragraph("Click ")
    rel = word.part.relate_to("javascript:alert(1)", RT.HYPERLINK, is_external=True)
    paragraph._p.append(parse_xml(f'<w:hyperlink {nsdecls("w", "r")} r:id="{rel}"><w:r><w:t>here</w:t></w:r></w:hyperlink>'))
    word.add_paragraph("A safe paragraph.")
    buffer = io.BytesIO()
    word.save(buffer)
    document = _upload("unsafe@example.com", buffer.getvalue())

    items, body, rels = _export(document["id"])
    client.cookies.clear()

    assert "javascript:" not in rels and "Click here" in body.replace("</w:t></w:r><w:r><w:t>", "")
    assert {item["feature"] for item in document["importReport"]["items"]} >= {"docx.link.unsafe"}


def test_a_pdf_says_they_are_in_a_word_export_only(api_db):
    document = _upload("pdfonly@example.com", _word_file())

    job = client.post("/api/v1/jobs/export", json={"documentId": document["id"], "format": "pdf"}).json()
    finished = client.get(f"/api/v1/jobs/{job['id']}").json()
    client.cookies.clear()

    assert "export.pdf.word_only" in {item["feature"] for item in finished["result"]["fidelity"]["items"]}
