"""SVG, which can carry script (tracker SEC-017): refused everywhere a picture comes in --
pasted, from a Word file, from Markdown -- never opened as a picture, never drawn in an
export, and a stored asset is served so a browser can't run anything in it."""

import base64
import io
import re
import zipfile

import pytest
from docx import Document as DocxDocument
from fastapi.testclient import TestClient
from PIL import Image as PILImage
from sqlalchemy import func, select
from sqlalchemy.orm import Session as OrmSession

from app.db.models import DocumentAsset
from app.export.docx_export import build_docx
from app.export.pdf_export import build_pdf
from app.fidelity.report import ReportBuilder
from app.formatting.engine import recompute_styles
from app.main import app
from app.models.document import Document, DocumentMetadata, Element, ElementType, ImageContent
from app.parsers.docx import parse_docx
from app.parsers.markdown import parse_markdown
from app.security.files import picture_problem

pytestmark = pytest.mark.security

client = TestClient(app, base_url="https://testserver")
SVG = b'<svg xmlns="http://www.w3.org/2000/svg" onload="alert(document.cookie)" width="10" height="10"><script>alert(1)</script></svg>'
_SVG_URI = "data:image/svg+xml;base64," + base64.b64encode(SVG).decode()


@pytest.fixture
def signed_in(api_db):
    client.cookies.clear()
    assert client.post("/api/v1/auth/register", json={"email": "svg@example.com", "password": "long enough password"}).status_code == 201
    yield api_db
    client.cookies.clear()


def _png() -> bytes:
    buffer = io.BytesIO()
    PILImage.new("RGB", (4, 3), (40, 160, 40)).save(buffer, format="PNG")
    return buffer.getvalue()


def test_svg_is_never_opened_as_a_picture():
    for content_type in ("image/svg+xml", "image/png", None):
        assert picture_problem(SVG, content_type).kind == "unreadable"


def test_a_pasted_svg_is_removed_and_nothing_is_stored(signed_in):
    document = client.post("/api/v1/documents", json={"text": "# Pictures\n\nSome text."}).json()
    as_svg = {"type": "image", "content": "", "order": 5, "image": {"src": _SVG_URI}}
    as_png = {"type": "image", "content": "", "order": 6, "image": {"src": "data:image/png;base64," + base64.b64encode(SVG).decode()}}

    saved = client.put(f"/api/v1/documents/{document['id']}/content", json={"elements": document["elements"] + [as_svg, as_png]}).json()

    assert not [element for element in saved["elements"] if element["type"] == "image"]
    assert any("isn't a valid PNG, JPEG, GIF, WebP or BMP" in note for note in saved["unsupportedFeatures"])
    with OrmSession(signed_in) as session:
        assert session.scalar(select(func.count(DocumentAsset.id))) == 0


def _word_file(*, fallback: bool) -> bytes:
    """A Word picture that is SVG: with Word's own PNG fallback beside it (svgBlip), or SVG only."""
    doc = DocxDocument()
    doc.add_picture(io.BytesIO(_png()))
    buffer = io.BytesIO()
    doc.save(buffer)
    with zipfile.ZipFile(io.BytesIO(buffer.getvalue())) as package:
        parts = {name: package.read(name) for name in package.namelist()}
    parts["word/media/drawing.svg"] = SVG
    parts["[Content_Types].xml"] = parts["[Content_Types].xml"].replace(
        b"</Types>", b'<Default Extension="svg" ContentType="image/svg+xml"/></Types>'
    )
    parts["word/_rels/document.xml.rels"] = parts["word/_rels/document.xml.rels"].replace(
        b"</Relationships>",
        b'<Relationship Id="rIdSvg" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="media/drawing.svg"/></Relationships>',
    )
    body = parts["word/document.xml"]
    if fallback:
        svg_blip = (
            b'<a:extLst><a:ext uri="{96DAC541-7B7A-43D3-8B79-37D633B846F1}">'
            b'<asvg:svgBlip xmlns:asvg="http://schemas.microsoft.com/office/drawing/2016/SVG/main" r:embed="rIdSvg"/></a:ext></a:extLst>'
        )
        body = re.sub(rb"(<a:blip [^>]*?)/>", lambda match: match.group(1) + b">" + svg_blip + b"</a:blip>", body, count=1)
    else:
        body = re.sub(rb'(<a:blip [^>]*r:embed=")[^"]+(")', rb"\1rIdSvg\2", body, count=1)
    parts["word/document.xml"] = body
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as package:
        for name, content in parts.items():
            package.writestr(name, content)
    return out.getvalue()


def test_a_word_svg_picture_is_read_as_its_png_fallback_or_not_at_all():
    with_fallback = parse_docx(_word_file(fallback=True), "fallback.docx")
    [picture] = [element for element in with_fallback.elements if element.type == ElementType.IMAGE]
    assert picture.image.mime == "image/png" and "svg" not in picture.image.src

    svg_only = parse_docx(_word_file(fallback=False), "svg.docx")
    assert not [element for element in svg_only.elements if element.type == ElementType.IMAGE]
    assert any("unsupported format (image/svg+xml)" in note for note in svg_only.unsupportedFeatures)


def test_markdown_svg_is_never_read_as_a_picture():
    # markdown-it takes no data: address but a web picture's, so this stays the text it is.
    document = parse_markdown(f"# Title\n\n![logo]({_SVG_URI})\n")
    assert not [element for element in document.elements if element.type == ElementType.IMAGE]
    assert "<svg" not in document.model_dump_json() and _SVG_URI in document.elements[1].content


def test_no_export_draws_an_svg_it_is_handed():
    """As a document stored before any of this could hold one."""
    document = Document(
        metadata=DocumentMetadata(title="SVG"),
        elements=[Element(type=ElementType.IMAGE, content="", order=0, image=ImageContent(src=_SVG_URI))],
    )
    recompute_styles(document)
    for build, feature in ((build_docx, "export.docx.image_format"), (build_pdf, "export.pdf.image_unreadable")):
        report = ReportBuilder()
        output = build(document, report=report)
        assert feature in {item.feature for item in report.items()}, build.__name__
        assert b"<script" not in output and b"onload" not in output


def test_a_stored_asset_is_served_so_nothing_in_it_can_run(signed_in):
    document = client.post("/api/v1/documents", json={"text": "# Pictures\n\nSome text."}).json()
    picture = {"type": "image", "content": "", "order": 5, "image": {"src": "data:image/png;base64," + base64.b64encode(_png()).decode()}}
    saved = client.put(f"/api/v1/documents/{document['id']}/content", json={"elements": document["elements"] + [picture]}).json()
    asset_id = next(element["image"]["assetId"] for element in saved["elements"] if element["type"] == "image")

    served = client.get(f"/api/v1/assets/{asset_id}")

    assert served.headers["content-type"] == "image/png"
    assert served.headers["x-content-type-options"] == "nosniff"
    assert "sandbox" in served.headers["content-security-policy"] and "default-src 'none'" in served.headers["content-security-policy"]
