"""Header fields beyond page numbers (tracker DOCX-020A): a STYLEREF, DATE or document-property
field in a header or footer is kept as {FIELD <instruction>|<last result>} -- shown here and in a
PDF as its result, written back to Word as the field -- so a header edited here keeps it; before,
only {PAGE} and {NUMPAGES} were fields and every other one became its last result, as text."""

import io

from docx import Document as DocxDocument
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from pypdf import PdfReader

from app.export.docx_export import build_docx
from app.export.pdf_export import build_pdf
from app.formatting import header_fields
from app.formatting.engine import recompute_styles
from app.models.document import Document, Element, ElementType, FormattingProperty, FormattingRule, InlineRun
from app.parsers.docx import parse_docx

_NS = nsdecls("w")


def _complex(instr: str, result: str) -> str:
    return (
        f'<w:r {_NS}><w:fldChar w:fldCharType="begin"/></w:r>'
        f'<w:r {_NS}><w:instrText xml:space="preserve"> {instr} </w:instrText></w:r>'
        f'<w:r {_NS}><w:fldChar w:fldCharType="separate"/></w:r>'
        f"<w:r {_NS}><w:t>{result}</w:t></w:r>"
        f'<w:r {_NS}><w:fldChar w:fldCharType="end"/></w:r>'
    )


def _word_file() -> bytes:
    """A header with a STYLEREF field (complex), a DATE field (simple) and a refused INCLUDETEXT
    one; a footer with a page number and a document property."""
    doc = DocxDocument()
    doc.add_paragraph("Body text.")
    header = doc.sections[0].header.paragraphs[0]
    header.add_run("Chapter: ")
    for run in parse_xml(f"<w:p {_NS}>{_complex('STYLEREF Heading1', 'Introduction')}</w:p>"):
        header._p.append(run)
    header.add_run(" on ")
    header._p.append(parse_xml(f'<w:fldSimple {_NS} w:instr=" DATE "><w:r><w:t>1.10.2026</w:t></w:r></w:fldSimple>'))
    header._p.append(parse_xml(f'<w:fldSimple {_NS} w:instr=" INCLUDETEXT x.docx "><w:r><w:t> (outside)</w:t></w:r></w:fldSimple>'))
    footer = doc.sections[0].footer.paragraphs[0]
    footer._p.append(parse_xml(f'<w:fldSimple {_NS} w:instr=" PAGE "><w:r><w:t>1</w:t></w:r></w:fldSimple>'))
    footer.add_run(" / ")
    for run in parse_xml(f"<w:p {_NS}>{_complex('DOCPROPERTY Company', 'Example Ltd')}</w:p>"):
        footer._p.append(run)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _fields(data: bytes, part: str) -> list[tuple[str, str]]:
    """Each complex field in a header or footer part: its instruction and the result Word shows."""
    doc = DocxDocument(io.BytesIO(data))
    root = (doc.sections[0].header if part == "header" else doc.sections[0].footer)._element
    found: list[tuple[str, str]] = []
    instr, result, state = "", "", None
    for node in root.iter(qn("w:fldChar"), qn("w:instrText"), qn("w:t")):
        if node.tag == qn("w:fldChar"):
            kind = node.get(qn("w:fldCharType"))
            if kind == "begin":
                instr, result, state = "", "", "instr"
            elif kind == "separate":
                state = "result"
            elif kind == "end" and state is not None:
                found.append((" ".join(instr.split()), result))
                state = None
        elif node.tag == qn("w:instrText") and state == "instr":
            instr += node.text or ""
        elif node.tag == qn("w:t") and state == "result":
            result += node.text or ""
    return found


def test_a_header_s_fields_are_read_as_placeholders_holding_their_last_result():
    document = parse_docx(_word_file(), "fields.docx")
    assert document.settings.header == "Chapter: {FIELD STYLEREF Heading1|Introduction} on {FIELD DATE|1.10.2026} (outside)"
    assert document.settings.footer == "{PAGE} / {FIELD DOCPROPERTY Company|Example Ltd}"


def _document(header: str, footer: str | None = None) -> Document:
    """A document whose header and footer were set here (the page settings' rules)."""
    rules = [FormattingRule(target="Document", property=FormattingProperty.HEADER, value=header, source="user")]
    if footer is not None:
        rules.append(FormattingRule(target="Document", property=FormattingProperty.FOOTER, value=footer, source="user"))
    document = Document(
        elements=[Element(type=ElementType.PARAGRAPH, content="Body text.", inline=[InlineRun(text="Body text.")], order=0)],
        formattingRules=rules,
    )
    recompute_styles(document)
    assert document.settings.header == header
    return document


def test_a_header_edited_here_keeps_its_fields_in_a_word_export():
    imported = parse_docx(_word_file(), "fields.docx")
    edited = _document("Now: " + imported.settings.header, imported.settings.footer)  # its text changed, its fields kept
    exported = build_docx(edited)
    assert _fields(exported, "header") == [("STYLEREF Heading1", "Introduction"), ("DATE", "1.10.2026")]
    assert ("DOCPROPERTY Company", "Example Ltd") in _fields(exported, "footer")
    again = parse_docx(exported, "again.docx")
    assert again.settings.header == "Now: Chapter: {FIELD STYLEREF Heading1|Introduction} on {FIELD DATE|1.10.2026} (outside)"
    assert again.settings.footer == imported.settings.footer


def test_a_field_the_policy_refuses_typed_by_hand_is_written_as_its_result():
    exported = build_docx(_document("Before {FIELD INCLUDETEXT secret.docx|kept as text} after"))
    xml = DocxDocument(io.BytesIO(exported)).sections[0].header._element.xml
    assert "INCLUDETEXT" not in xml and "instrText" not in xml
    assert parse_docx(exported, "x.docx").settings.header == "Before kept as text after"


def test_a_pdf_and_the_pages_show_each_field_s_last_result():
    text = PdfReader(io.BytesIO(build_pdf(_document("Chapter {FIELD STYLEREF Heading1|Introduction}", "{FIELD DATE|1.10.2026} p. {PAGE}")))).pages[0].extract_text()
    assert "Chapter Introduction" in text and "1.10.2026 p. 1" in text and "FIELD" not in text


def test_a_placeholder_holds_only_what_it_can_and_text_too_long_keeps_its_results():
    assert header_fields.field_token(" STYLEREF   Heading1 ", "Intro") == "{FIELD STYLEREF Heading1|Intro}"
    assert header_fields.field_token("INCLUDETEXT x", "Intro") is None  # refused by the field policy
    assert header_fields.field_token("STYLEREF Heading1", "a {b}") is None  # a result the placeholder can't hold
    assert header_fields.field_token('DATE \\@ "d|M"', "1|10") is None
    assert header_fields.parts("p. {PAGE} {FIELD DATE|1.10.2026}") == [("text", "p. ", ""), ("page", "", ""), ("text", " ", ""), ("field", "DATE", "1.10.2026")]
    long = " ".join(["{FIELD STYLEREF Heading1|Intro}"] * 20)  # 639 characters as placeholders
    assert header_fields.fitted(long) == " ".join(["Intro"] * 20)
    assert header_fields.fitted("x" * 600) == "x" * 500


def test_into_the_original_a_header_left_alone_stays_as_it_was_and_one_edited_keeps_its_fields():
    source = _word_file()
    imported = parse_docx(source, "fields.docx")
    original = DocxDocument(io.BytesIO(source)).sections[0].header._element.xml
    untouched = build_docx(imported, source=source)
    assert DocxDocument(io.BytesIO(untouched)).sections[0].header._element.xml == original  # the same text: not rewritten
    edited = imported.model_copy(deep=True)
    edited.formattingRules.append(FormattingRule(target="Document", property=FormattingProperty.HEADER, value="Edited: " + imported.settings.header, source="user"))
    recompute_styles(edited)
    exported = build_docx(edited, source=source)
    assert _fields(exported, "header") == [("STYLEREF Heading1", "Introduction"), ("DATE", "1.10.2026")]
    assert parse_docx(exported, "again.docx").settings.header == "Edited: " + imported.settings.header
