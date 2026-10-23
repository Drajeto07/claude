"""B-035A: a document file too large for one request comes in parts (app/jobs/uploads.py, api/jobs.py).

On the Preview a 4.91 MB and an 8.81 MB Word file were refused with 413 (Vercel's 4.5 MB request limit, owner's
test 2026-10-23). Here: such files come in 3 MiB parts and are imported exactly as a file sent whole is -- the
same bytes (the imported document's own record of its source, SHA-256), the same checks, one job; and every
way it can go wrong -- parts missing, repeated, out of order, damaged, a stranger, an unsafe file, the plan, a
restart, an abandoned upload -- ends in a stated refusal or the same single job, never in a half-imported file.
"""

import asyncio
import hashlib
import io
import random
import zipfile
from datetime import datetime, timedelta, timezone

import httpx
import pytest
from docx import Document as DocxDocument
from docx.shared import Cm
from fastapi.testclient import TestClient
from PIL import Image
from sqlalchemy import func, select
from sqlalchemy.orm import Session as OrmSession

from app.db.models import ProcessingJob, UsageRecord
from app.jobs import uploads
from app.main import app
from app.storage.factory import get_storage_provider
from app.storage.local_provider import LocalStorageProvider

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PASSWORD = "long enough password"
CHUNK = uploads.CHUNK_SIZE


def _client(email: str) -> TestClient:
    client = TestClient(app, base_url="https://testserver")
    assert client.post("/api/v1/auth/register", json={"email": email, "password": PASSWORD}).status_code == 201
    return client


@pytest.fixture
def alice(api_db) -> TestClient:
    return _client("alice-chunks@example.com")


@pytest.fixture
def bob(api_db) -> TestClient:
    return _client("bob-chunks@example.com")


def _noise_png(size: int, seed: int) -> bytes:
    """A PNG of about `size` bytes that doesn't compress (random pixels, stored uncompressed), the same every time."""
    side = int((size / 3) ** 0.5)
    pixels = random.Random(seed).randbytes(side * side * 3)
    buffer = io.BytesIO()
    Image.frombytes("RGB", (side, side), pixels).save(buffer, "PNG", compress_level=0)
    return buffer.getvalue()


def _docx(picture_bytes: int = 0, *, seed: int = 1, text: str = "Quarterly report: revenue 1,250 EUR on 2026-10-23.") -> tuple[bytes, bytes | None]:
    document = DocxDocument()
    document.add_paragraph(text)
    picture = _noise_png(picture_bytes, seed) if picture_bytes else None
    if picture:
        document.add_picture(io.BytesIO(picture), width=Cm(12))
        document.add_paragraph("Text after the picture.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue(), picture


def _start(client: TestClient, data: bytes, filename: str = "report.docx", **extra) -> httpx.Response:
    declared = DOCX if filename.endswith(".docx") else None  # as a browser declares it
    body = {"filename": filename, "size": len(data), "sha256": hashlib.sha256(data).hexdigest(), "contentType": declared, **extra}
    return client.post("/api/v1/jobs/import-file/uploads", json=body)


def _parts(data: bytes, chunk: int = CHUNK) -> list[bytes]:
    return [data[offset : offset + chunk] for offset in range(0, len(data), chunk)]


def _put(client: TestClient, upload_id: str, index: int, part: bytes, sha256: str | None = None) -> httpx.Response:
    query = f"?sha256={sha256}" if sha256 else ""
    return client.put(f"/api/v1/jobs/import-file/uploads/{upload_id}/parts/{index}{query}", content=part, headers={"Content-Type": "application/octet-stream"})


def _complete(client: TestClient, upload_id: str) -> httpx.Response:
    return client.post(f"/api/v1/jobs/import-file/uploads/{upload_id}/complete")


def _upload(client: TestClient, data: bytes, *, order: list[int] | None = None, **start) -> httpx.Response:
    started = _start(client, data, **start)
    assert started.status_code == 201, started.text
    upload_id = started.json()["uploadId"]
    parts = _parts(data)
    for index in order if order is not None else range(len(parts)):
        assert _put(client, upload_id, index, parts[index]).status_code == 200
    return _complete(client, upload_id)


def _storage() -> LocalStorageProvider:
    return app.dependency_overrides[get_storage_provider]()


def _keys(prefix: str = uploads.PREFIX) -> list[str]:
    return sorted(key for key, _ in asyncio.run(_storage().list(prefix)))


def _count(api_db, model, *where) -> int:
    with OrmSession(api_db) as session:
        return session.scalar(select(func.count()).select_from(model).where(*where))


# -- Large files, imported as they are ------------------------------------------------------------------------


@pytest.mark.parametrize("picture_mb", [4.9, 8.8], ids=["5MB", "9MB"])
def test_a_large_word_file_comes_in_parts_and_is_imported_byte_for_byte(alice, picture_mb):
    data, picture = _docx(int(picture_mb * 1024 * 1024))
    assert len(data) > 4.5 * 1024 * 1024
    started = _start(alice, data, title="Big report")
    assert started.status_code == 201
    assert started.json()["chunkSize"] == CHUNK
    assert started.json()["chunks"] == -(-len(data) // CHUNK)
    upload_id = started.json()["uploadId"]
    parts = _parts(data)
    assert max(len(part) for part in parts) <= CHUNK < 4.5 * 1024 * 1024
    for index in reversed(range(len(parts))):  # any order
        received = _put(alice, upload_id, index, parts[index], sha256=hashlib.sha256(parts[index]).hexdigest())
        assert received.json() == {"index": index, "size": len(parts[index])}

    job = _complete(alice, upload_id)
    assert job.status_code == 202, job.text
    assert job.json()["status"] == "succeeded", job.json()
    document = alice.get(f"/api/v1/documents/{job.json()['result']['documentId']}").json()
    # The file the import kept as its source is the very file sent: same bytes, same SHA-256.
    assert document["sourcePackage"]["sha256"] == hashlib.sha256(data).hexdigest()
    assert document["sourcePackage"]["size"] == len(data)
    assert document["metadata"]["title"] == "Big report"
    assert "revenue 1,250 EUR on 2026-10-23" in " ".join(element["content"] for element in document["elements"])
    # Its picture came through whole.
    image = next(element for element in document["elements"] if element["type"] == "image")
    assert alice.get(f"/api/v1/assets/{image['image']['assetId']}").content == picture
    # Its parts are gone; only the record of which job it became stays, until the sweep.
    left = _keys()
    assert len(left) == 1 and left[0].endswith(f"/{upload_id}/session.json")


def test_a_small_file_sent_whole_is_imported_as_before_and_as_one_sent_in_parts(alice):
    data, _ = _docx()
    whole = alice.post("/api/v1/jobs/import-file", files={"file": ("small.docx", data, DOCX)})
    assert whole.status_code == 202 and whole.json()["status"] == "succeeded"
    parted = _upload(alice, data, filename="small.docx")
    assert parted.status_code == 202 and parted.json()["status"] == "succeeded"
    first = alice.get(f"/api/v1/documents/{whole.json()['result']['documentId']}").json()
    second = alice.get(f"/api/v1/documents/{parted.json()['result']['documentId']}").json()
    assert [element["content"] for element in first["elements"]] == [element["content"] for element in second["elements"]]
    assert first["sourcePackage"]["sha256"] == second["sourcePackage"]["sha256"] == hashlib.sha256(data).hexdigest()


def test_the_import_options_travel_with_the_upload(alice):
    data = b"Plain text with www.example.com in it.\n"
    job = _upload(alice, data, filename="notes.txt", title="Notes", autolink=True)
    assert job.json()["status"] == "succeeded"
    stored = alice.get(f"/api/v1/documents/{job.json()['result']['documentId']}").json()
    assert stored["metadata"]["title"] == "Notes"


# -- Sizes ------------------------------------------------------------------------------------------------------


def test_the_server_cap_is_taken_exactly_and_a_byte_more_refused(alice):
    cap = 10 * 1024 * 1024
    body = {"filename": "big.docx", "sha256": "0" * 64}
    at_cap = alice.post("/api/v1/jobs/import-file/uploads", json={**body, "size": cap})
    assert at_cap.status_code == 201 and at_cap.json()["chunks"] == 4
    over = alice.post("/api/v1/jobs/import-file/uploads", json={**body, "size": cap + 1})
    assert over.status_code == 413
    assert over.json()["code"] == "too_large" and over.json()["details"]["maxMb"] == 10


def test_a_plan_smaller_than_the_file_refuses_it_before_any_part(alice, monkeypatch):
    from app.services import entitlements_service

    real = entitlements_service.effective
    monkeypatch.setattr(entitlements_service, "effective", lambda found: real(found).model_copy(update={"maxDocumentSizeMb": 5}))
    refused = alice.post("/api/v1/jobs/import-file/uploads", json={"filename": "big.docx", "size": 6 * 1024 * 1024, "sha256": "0" * 64})
    assert refused.status_code == 402 and refused.json()["code"] == "plan_limit"
    assert _keys() == []


def test_a_part_of_the_wrong_size_or_number_is_refused(alice):
    data = random.Random(3).randbytes(CHUNK + 100)
    upload_id = _start(alice, data, filename="x.docx").json()["uploadId"]
    assert _put(alice, upload_id, 0, data[: CHUNK - 1]).json()["code"] == "chunk_size"
    assert _put(alice, upload_id, 1, data[CHUNK:] + b"!").status_code in (400, 413)
    assert _put(alice, upload_id, 2, b"x").json()["code"] == "chunk_out_of_range"
    assert _put(alice, upload_id, -1, b"x").json()["code"] == "chunk_out_of_range"
    damaged = _put(alice, upload_id, 1, data[CHUNK:], sha256="0" * 64)
    assert damaged.status_code == 400 and damaged.json()["code"] == "chunk_corrupt"


# -- Parts missing, repeated, out of order, damaged ------------------------------------------------------------


def test_a_missing_part_is_named_and_completing_after_sending_it_works(alice, api_db):
    data, _ = _docx(int(4.9 * 1024 * 1024))
    upload_id = _start(alice, data).json()["uploadId"]
    parts = _parts(data)
    _put(alice, upload_id, 0, parts[0])
    incomplete = _complete(alice, upload_id)
    assert incomplete.status_code == 409
    assert incomplete.json()["code"] == "upload_incomplete" and incomplete.json()["details"]["missing"] == [1]
    # Nothing was processed or made while it was incomplete.
    assert _count(api_db, ProcessingJob) == 0
    _put(alice, upload_id, 1, parts[1])
    assert _complete(alice, upload_id).json()["status"] == "succeeded"


def test_a_part_sent_twice_and_parts_out_of_order_make_the_same_file(alice):
    data, _ = _docx(int(8.8 * 1024 * 1024), seed=5)
    parts = _parts(data)
    order = [2, 0, 1, 0, 2] + list(range(len(parts)))  # repeats and any order
    job = _upload(alice, data, order=order)
    document = alice.get(f"/api/v1/documents/{job.json()['result']['documentId']}").json()
    assert document["sourcePackage"]["sha256"] == hashlib.sha256(data).hexdigest()


def test_a_part_replaced_by_other_bytes_is_caught_by_the_files_sha256_and_nothing_is_imported(alice, api_db):
    data, _ = _docx(int(4.9 * 1024 * 1024))
    upload_id = _start(alice, data).json()["uploadId"]
    parts = _parts(data)
    _put(alice, upload_id, 0, parts[0])
    _put(alice, upload_id, 1, bytes(reversed(parts[1])))  # the right size, the wrong bytes
    mismatch = _complete(alice, upload_id)
    assert mismatch.status_code == 422 and mismatch.json()["code"] == "upload_mismatch"
    assert _count(api_db, ProcessingJob) == 0
    assert _keys() == []  # discarded: it can't be repaired, only sent again
    assert _complete(alice, upload_id).status_code == 404


def test_a_part_that_failed_to_store_can_be_sent_again(alice, monkeypatch):
    data, _ = _docx(int(4.9 * 1024 * 1024))
    upload_id = _start(alice, data).json()["uploadId"]
    parts = _parts(data)
    real_put = LocalStorageProvider.put
    failures = {"left": 1}

    async def flaky(self, key, value, content_type):
        if key.endswith("part-00001") and failures["left"]:
            failures["left"] -= 1
            raise ConnectionError("storage hiccup")
        return await real_put(self, key, value, content_type)

    monkeypatch.setattr(LocalStorageProvider, "put", flaky)
    safe = TestClient(app, base_url="https://testserver", raise_server_exceptions=False, cookies=alice.cookies)
    _put(safe, upload_id, 0, parts[0])
    assert _put(safe, upload_id, 1, parts[1]).status_code == 500  # the transient failure
    assert _put(safe, upload_id, 1, parts[1]).status_code == 200  # the retry
    assert _complete(safe, upload_id).json()["status"] == "succeeded"


# -- One job, however often it is completed --------------------------------------------------------------------


def test_completing_again_answers_the_same_job_and_charges_once(alice, api_db):
    data, _ = _docx(int(4.9 * 1024 * 1024))
    upload_id = _start(alice, data).json()["uploadId"]
    for index, part in enumerate(_parts(data)):
        _put(alice, upload_id, index, part)
    first = _complete(alice, upload_id).json()
    second = _complete(alice, upload_id).json()
    assert first["id"] == second["id"] and first["status"] == "succeeded"
    assert _count(api_db, ProcessingJob) == 1
    assert _count(api_db, UsageRecord, UsageRecord.metric == "processing_jobs") == 1
    assert alice.get("/api/v1/documents").json()["total"] == 1
    assert _put(alice, upload_id, 0, _parts(data)[0]).json()["code"] == "upload_finished"


def test_two_completions_at_once_make_one_job(alice, api_db):
    data, _ = _docx(int(4.9 * 1024 * 1024))
    upload_id = _start(alice, data).json()["uploadId"]
    for index, part in enumerate(_parts(data)):
        _put(alice, upload_id, index, part)

    async def both() -> list[httpx.Response]:
        transport = httpx.ASGITransport(app=app)
        async with httpx.AsyncClient(transport=transport, base_url="https://testserver", cookies=dict(alice.cookies)) as client:
            return await asyncio.gather(*(client.post(f"/api/v1/jobs/import-file/uploads/{upload_id}/complete") for _ in range(2)))

    answers = asyncio.run(both())
    assert [answer.status_code for answer in answers] == [202, 202], [answer.text for answer in answers]
    assert answers[0].json()["id"] == answers[1].json()["id"]
    assert _count(api_db, ProcessingJob) == 1


# -- Who may, and what may be uploaded ------------------------------------------------------------------------


def test_every_upload_route_needs_a_signed_in_user(api_db):
    anonymous = TestClient(app, base_url="https://testserver")
    some_id = "00000000-0000-4000-8000-000000000000"
    assert anonymous.post("/api/v1/jobs/import-file/uploads", json={"filename": "a.docx", "size": 1, "sha256": "0" * 64}).status_code == 401
    assert anonymous.put(f"/api/v1/jobs/import-file/uploads/{some_id}/parts/0", content=b"x").status_code == 401
    assert anonymous.post(f"/api/v1/jobs/import-file/uploads/{some_id}/complete").status_code == 401


def test_another_user_is_told_the_upload_does_not_exist(alice, bob, api_db):
    data, _ = _docx(int(4.9 * 1024 * 1024))
    upload_id = _start(alice, data).json()["uploadId"]
    parts = _parts(data)
    for index, part in enumerate(parts):
        stranger = _put(bob, upload_id, index, part)
        assert stranger.status_code == 404 and stranger.json()["code"] == "upload_not_found"
    assert _complete(bob, upload_id).status_code == 404
    assert _complete(bob, "not-an-id").status_code == 404
    assert _put(bob, "../../jobs", 0, b"x").status_code == 404
    # Alice's upload is untouched by any of it.
    for index, part in enumerate(parts):
        _put(alice, upload_id, index, part)
    assert _complete(alice, upload_id).json()["status"] == "succeeded"
    assert _count(api_db, ProcessingJob) == 1


def test_a_file_of_the_wrong_type_by_name_is_refused_at_the_start(alice):
    refused = alice.post("/api/v1/jobs/import-file/uploads", json={"filename": "run.exe", "size": 10, "sha256": "0" * 64})
    assert refused.status_code == 400
    assert _keys() == []


@pytest.mark.parametrize(
    ("filename", "data"),
    [
        ("disguised.docx", b"%PDF-1.7\n" + random.Random(1).randbytes(CHUNK)),
        ("random.docx", random.Random(2).randbytes(CHUNK + 10)),
        ("old-word.docx", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + bytes(CHUNK)),
        ("not-a.pdf", b"PK\x03\x04" + bytes(CHUNK)),
    ],
    ids=["pdf-as-docx", "random-bytes", "ole-doc", "zip-as-pdf"],
)
def test_content_that_isnt_what_its_name_says_is_refused_and_nothing_kept(alice, api_db, filename, data):
    refused = _upload(alice, data, filename=filename)
    assert refused.status_code == 400 and refused.json()["code"] == "invalid_file", refused.text
    assert _count(api_db, ProcessingJob) == 0
    assert _keys() == []


def test_a_zip_bomb_is_refused_by_the_same_checks_and_nothing_kept(alice, api_db):
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as package:
        package.writestr("[Content_Types].xml", "<Types/>")
        package.writestr("word/media/padding.bin", bytes(300 * 1024 * 1024))
    bomb = buffer.getvalue()
    refused = _upload(alice, bomb, filename="bomb.docx")
    assert refused.status_code == 400 and refused.json()["code"] == "invalid_file"
    assert _count(api_db, ProcessingJob) == 0
    assert _keys() == []


# -- Restarts, abandoned uploads, expiry -----------------------------------------------------------------------


def test_an_upload_goes_on_after_a_restart_its_state_is_in_the_store(alice, api_db, tmp_path):
    data, _ = _docx(int(4.9 * 1024 * 1024))
    upload_id = _start(alice, data).json()["uploadId"]
    parts = _parts(data)
    _put(alice, upload_id, 0, parts[0])
    # A new instance: a new storage provider object on the same store, nothing kept in memory.
    fresh = LocalStorageProvider(tmp_path / "assets")
    app.dependency_overrides[get_storage_provider] = lambda: fresh
    _put(alice, upload_id, 1, parts[1])
    assert _complete(alice, upload_id).json()["status"] == "succeeded"


def test_an_abandoned_upload_is_swept_and_a_delete_that_fails_is_tried_again(alice, monkeypatch):
    data, _ = _docx(int(4.9 * 1024 * 1024))
    upload_id = _start(alice, data).json()["uploadId"]
    _put(alice, upload_id, 0, _parts(data)[0])
    assert len(_keys()) == 2
    later = datetime.now(timezone.utc) + uploads.SESSION_TTL + timedelta(minutes=1)
    assert asyncio.run(uploads.sweep(_storage(), now=datetime.now(timezone.utc))) == 0  # not old yet
    real_delete = LocalStorageProvider.delete
    failures = {"left": 1}

    async def flaky(self, key):
        if failures["left"]:
            failures["left"] -= 1
            raise ConnectionError("storage hiccup")
        return await real_delete(self, key)

    monkeypatch.setattr(LocalStorageProvider, "delete", flaky)
    asyncio.run(uploads.sweep(_storage(), now=later))
    assert len(_keys()) == 1  # one delete failed: still listed
    asyncio.run(uploads.sweep(_storage(), now=later))
    assert _keys() == []


def test_the_hourly_job_sweep_takes_abandoned_uploads_too(alice, api_db, monkeypatch):
    from app.jobs import files

    data, _ = _docx(int(4.9 * 1024 * 1024))
    upload_id = _start(alice, data).json()["uploadId"]
    _put(alice, upload_id, 0, _parts(data)[0])
    monkeypatch.setattr(uploads, "_now", lambda: datetime.now(timezone.utc) + uploads.SESSION_TTL + timedelta(minutes=1))
    from app.jobs.queue import get_job_session_factory  # the test database

    factory = app.dependency_overrides[get_job_session_factory]()
    asyncio.run(files.sweep_job_files(factory, _storage()))
    assert _keys() == []


def test_starting_an_upload_sweeps_the_workspaces_old_ones(alice, monkeypatch):
    data, _ = _docx(int(4.9 * 1024 * 1024))
    old = _start(alice, data).json()["uploadId"]
    _put(alice, old, 0, _parts(data)[0])
    monkeypatch.setattr(uploads, "_now", lambda: datetime.now(timezone.utc) + uploads.SESSION_TTL + timedelta(minutes=1))
    new = _start(alice, data).json()["uploadId"]
    assert all(old not in key for key in _keys())
    assert any(new in key for key in _keys())


def test_an_expired_upload_is_refused_and_discarded(alice, api_db, monkeypatch):
    data, _ = _docx(int(4.9 * 1024 * 1024))
    upload_id = _start(alice, data).json()["uploadId"]
    for index, part in enumerate(_parts(data)):
        _put(alice, upload_id, index, part)
    monkeypatch.setattr(uploads.UploadSession, "expired", lambda self, now: True)
    expired = _complete(alice, upload_id)
    assert expired.status_code == 410 and expired.json()["code"] == "upload_expired"
    assert _count(api_db, ProcessingJob) == 0
    assert _keys() == []


def test_nothing_is_processed_or_made_before_completion(alice, api_db):
    data, _ = _docx(int(8.8 * 1024 * 1024))
    upload_id = _start(alice, data).json()["uploadId"]
    for index, part in enumerate(_parts(data)):
        _put(alice, upload_id, index, part)
    assert _count(api_db, ProcessingJob) == 0
    assert alice.get("/api/v1/documents").json()["total"] == 0
    assert _complete(alice, upload_id).json()["status"] == "succeeded"
    assert _count(api_db, ProcessingJob) == 1


# -- Hardening: bounded leftovers, races, interrupted completion, S3 ------------------------------------------------


def _filled(client: TestClient, data: bytes, parts: int | None = None, **start) -> str:
    upload_id = _start(client, data, **start).json()["uploadId"]
    for index, part in enumerate(_parts(data)[:parts]):
        assert _put(client, upload_id, index, part).status_code == 200
    return upload_id


def _user_id(api_db, email: str = "alice-chunks@example.com") -> str:
    from app.db.models import User

    with OrmSession(api_db) as session:
        return session.scalar(select(User.id).where(User.email == email))


def _later() -> datetime:
    return datetime.now(timezone.utc) + uploads.SESSION_TTL + timedelta(minutes=1)


def test_a_workspace_holds_at_most_three_unfinished_uploads_so_leftovers_stay_bounded(alice):
    data, _ = _docx(int(4.9 * 1024 * 1024))
    for _ in range(uploads.MAX_OPEN_UPLOADS):
        _filled(alice, data, parts=1)
    _start(alice, data)  # started without parts: holds nothing, doesn't count
    refused = _start(alice, data)
    assert refused.status_code == 429 and refused.json()["code"] == "too_many_uploads"
    assert refused.json()["details"]["limit"] == uploads.MAX_OPEN_UPLOADS


def test_finishing_an_upload_frees_its_place(alice):
    data, _ = _docx(int(4.9 * 1024 * 1024))
    first = _filled(alice, data)
    for _ in range(uploads.MAX_OPEN_UPLOADS - 1):
        _filled(alice, data, parts=1)
    assert _start(alice, data).status_code == 429
    assert _complete(alice, first).json()["status"] == "succeeded"
    assert _start(alice, data).status_code == 201


def test_expired_uploads_neither_count_nor_stay(alice, monkeypatch):
    data, _ = _docx(int(4.9 * 1024 * 1024))
    for _ in range(uploads.MAX_OPEN_UPLOADS):
        _filled(alice, data, parts=1)
    monkeypatch.setattr(uploads, "_now", _later)
    assert _start(alice, data).status_code == 201
    assert len({key.split("/")[2] for key in _keys()}) == 1  # only the new one is left


def test_anyones_start_sweeps_other_workspaces_old_uploads_a_batch_at_a_time(alice, bob, monkeypatch):
    data, _ = _docx(int(4.9 * 1024 * 1024))
    bobs = _filled(bob, data, parts=1)
    monkeypatch.setattr(uploads, "_now", _later)
    monkeypatch.setattr(uploads, "SWEEP_BATCH", 1)
    _start(alice, data)
    assert all(bobs not in key for key in _keys())


def test_sweeps_run_twice_or_at_once_are_safe(alice):
    data, _ = _docx(int(4.9 * 1024 * 1024))
    _filled(alice, data, parts=2)
    later = _later()

    async def together():
        return await asyncio.gather(uploads.sweep(_storage(), now=later), uploads.sweep(_storage(), now=later))

    assert sum(asyncio.run(together())) >= 1
    assert asyncio.run(uploads.sweep(_storage(), now=later)) == 0
    assert _keys() == []


def test_a_late_part_after_completion_is_refused_or_ignored_and_the_import_is_untouched(alice, api_db):
    data, _ = _docx(int(4.9 * 1024 * 1024))
    upload_id = _filled(alice, data)
    workspace = _keys()[0].split("/")[1]
    # A part request that read the session just before completion ...
    before = asyncio.run(uploads.find(_storage(), workspace_id=workspace, user_id=_user_id(api_db), upload_id=upload_id))
    job = _complete(alice, upload_id).json()
    # ... and one that arrives after it is refused.
    assert _put(alice, upload_id, 0, _parts(data)[0]).json()["code"] == "upload_finished"
    # The first writes a part nobody reads again: the job already has its own copy of the file.
    asyncio.run(uploads.put_part(_storage(), before, 0, _parts(data)[0]))
    again = _complete(alice, upload_id).json()
    assert again["id"] == job["id"]
    document = alice.get(f"/api/v1/documents/{job['result']['documentId']}").json()
    assert document["sourcePackage"]["sha256"] == hashlib.sha256(data).hexdigest()
    assert _count(api_db, ProcessingJob) == 1
    asyncio.run(uploads.sweep(_storage(), now=_later()))
    assert _keys() == []


def test_completion_racing_the_sweep_never_imports_part_of_a_file(alice, api_db):
    data, _ = _docx(int(4.9 * 1024 * 1024))
    upload_id = _filled(alice, data)
    # The sweep took a part first: completion finds it missing, and nothing is made.
    asyncio.run(_storage().delete(next(key for key in _keys() if key.endswith("part-00001"))))
    assert _complete(alice, upload_id).json()["code"] == "upload_incomplete"
    assert _count(api_db, ProcessingJob) == 0
    # It took everything: the upload is gone.
    asyncio.run(uploads.sweep(_storage(), now=_later()))
    assert _complete(alice, upload_id).status_code == 404
    assert _count(api_db, ProcessingJob) == 0


def test_a_completion_cut_off_after_the_job_was_made_is_retried_into_the_same_job(alice, api_db, monkeypatch):
    data, _ = _docx(int(4.9 * 1024 * 1024))
    upload_id = _filled(alice, data)
    real = uploads.finished

    async def cut_off(storage, session, job_id):  # the process died before it recorded the job or removed the parts
        return None

    monkeypatch.setattr(uploads, "finished", cut_off)
    first = _complete(alice, upload_id).json()
    assert len(_keys()) == 3  # session.json and both parts are still there
    monkeypatch.setattr(uploads, "finished", real)
    second = _complete(alice, upload_id).json()
    assert second["id"] == first["id"]
    assert _count(api_db, ProcessingJob) == 1
    assert _count(api_db, UsageRecord, UsageRecord.metric == "processing_jobs") == 1
    assert alice.get("/api/v1/documents").json()["total"] == 1
    assert len(_keys()) == 1


def test_an_idempotency_key_reused_for_another_file_is_refused(alice, api_db):
    first, _ = _docx(int(4.9 * 1024 * 1024), seed=1)
    other, _ = _docx(int(4.9 * 1024 * 1024), seed=2)
    a, b = _filled(alice, first), _filled(alice, other)
    key = {"Idempotency-Key": "my-key-1"}
    made = alice.post(f"/api/v1/jobs/import-file/uploads/{a}/complete", headers=key)
    assert made.status_code == 202
    reused = alice.post(f"/api/v1/jobs/import-file/uploads/{b}/complete", headers=key)
    assert reused.status_code == 422 and reused.json()["code"] == "idempotency_key_reused"
    again = alice.post(f"/api/v1/jobs/import-file/uploads/{a}/complete", headers=key)
    assert again.json()["id"] == made.json()["id"]
    assert _count(api_db, ProcessingJob) == 1


# S3-compatible storage (moto: AWS S3's behaviour, not Supabase's -- that is checked on the Preview).


@pytest.fixture
def s3_storage(api_db, monkeypatch):
    import boto3
    from moto import mock_aws

    from app.storage.s3_provider import S3StorageProvider

    monkeypatch.setenv("AWS_ACCESS_KEY_ID", "testing")
    monkeypatch.setenv("AWS_SECRET_ACCESS_KEY", "testing")
    with mock_aws():
        boto3.client("s3", region_name="us-east-1").create_bucket(Bucket="assets")
        provider = S3StorageProvider("assets", region="us-east-1", access_key_id="testing", secret_access_key="testing")
        app.dependency_overrides[get_storage_provider] = lambda: provider
        yield provider


def test_on_s3_a_large_file_comes_in_parts_and_is_imported_byte_for_byte(s3_storage, alice):
    data, picture = _docx(int(4.9 * 1024 * 1024), seed=9)
    job = _upload(alice, data, order=[1, 0, 1])
    assert job.json()["status"] == "succeeded", job.text
    document = alice.get(f"/api/v1/documents/{job.json()['result']['documentId']}").json()
    assert document["sourcePackage"]["sha256"] == hashlib.sha256(data).hexdigest()
    image = next(element for element in document["elements"] if element["type"] == "image")
    assert alice.get(f"/api/v1/assets/{image['image']['assetId']}").content == picture
    left = asyncio.run(s3_storage.list(uploads.PREFIX))
    assert [key.rsplit("/", 1)[-1] for key, _ in left] == ["session.json"]


def test_on_s3_a_prefix_lists_only_its_own_workspace_and_a_failed_delete_is_retried(s3_storage, monkeypatch):
    async def scenario():
        for key in ("uploads/w1/a/part-00000", "uploads/w1/a/session.json", "uploads/w10/b/part-00000"):
            await s3_storage.put(key, b"x", "application/octet-stream")
        assert sorted(key for key, _ in await s3_storage.list("uploads/w1/")) == ["uploads/w1/a/part-00000", "uploads/w1/a/session.json"]
        assert await s3_storage.list("uploads/nobody/") == []
        real = type(s3_storage).delete
        failures = {"left": 1}

        async def flaky(self, key):
            if failures["left"]:
                failures["left"] -= 1
                raise ConnectionError("S3 hiccup")
            return await real(self, key)

        monkeypatch.setattr(type(s3_storage), "delete", flaky)
        later = _later()
        await uploads.sweep(s3_storage, prefix="uploads/w1/", now=later)
        assert len(await s3_storage.list("uploads/w1/")) == 1  # the one that failed
        await uploads.sweep(s3_storage, prefix="uploads/w1/", now=later)
        assert await s3_storage.list("uploads/w1/") == []
        assert len(await s3_storage.list("uploads/w10/")) == 1  # another workspace's, untouched

    asyncio.run(scenario())
