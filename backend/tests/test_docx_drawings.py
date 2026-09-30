"""Drawings the model doesn't hold survive a Word export that writes their paragraph
anew (tracker DOCX-019): charts, SmartArt, shapes and embedded objects go back from
the kept original file -- the server's own copy, nothing the browser sent -- into
the paragraph written for them, where they were in its text, and a paragraph of
nothing but drawings goes back after the block it came with. Text boxes stay the
document's own paragraphs (their text is imported)."""

import io
import zipfile
from pathlib import Path

from docx import Document as DocxDocument
from docx.oxml import parse_xml
from lxml import etree

from app.export.docx_export import build_docx
from app.export.package_check import package_problems
from app.export.provenance import stamp
from app.fidelity.imports import with_source_kept
from app.fidelity.report import ReportBuilder
from app.formatting.engine import apply_formatting
from app.formatting.templates import BUILTIN_TEMPLATES
from app.services.ingestion_service import build_document_from_docx

WORD = Path(__file__).parent / "fixtures" / "word"
_NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "dgm": "http://schemas.openxmlformats.org/drawingml/2006/diagram",
    "o": "urn:schemas-microsoft-com:office:office",
    "wps": "http://schemas.microsoft.com/office/word/2010/wordprocessingShape",
    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
}
_SHAPE = (
    '<w:r xmlns:w="{w}" xmlns:wp="{wp}" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:wps="{wps}">'
    '<w:drawing><wp:inline distT="0" distB="0" distL="0" distR="0"><wp:extent cx="914400" cy="457200"/>'
    '<wp:docPr id="{id}" name="Shape {id}"/><a:graphic><a:graphicData uri="{wps}"><wps:wsp><wps:cNvSpPr/>'
    '<wps:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="914400" cy="457200"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/>'
    "</a:prstGeom></wps:spPr><wps:bodyPr/></wps:wsp></a:graphicData></a:graphic></wp:inline></w:drawing></w:r>"
)


def _shape(shape_id: int):
    return parse_xml(_SHAPE.format(id=shape_id, **_NS))


def _through_a_template(data: bytes, name: str):
    """Imported as an upload imports it, then restyled by a template -- every block
    written anew -- and exported into the original."""
    document = build_document_from_docx(data, name, None)
    document.importReport = with_source_kept(document.importReport, data, document)
    stamp(document)
    template = BUILTIN_TEMPLATES["academic-default"]
    apply_formatting(document, template_id=template.id, template_rules=template.rules, instruction_rules=[])
    noted = ReportBuilder()
    return document, build_docx(document, source=data, report=noted), noted


def _body(data: bytes):
    with zipfile.ZipFile(io.BytesIO(data)) as package:
        return etree.fromstring(package.read("word/document.xml"))


def _count(body, path: str) -> int:
    return len(body.xpath(path, namespaces=_NS))


def test_charts_smartart_and_objects_come_back_into_a_file_written_anew():
    data = (WORD / "a09-objects.docx").read_bytes()
    before = _body(data)

    document, exported, noted = _through_a_template(data, "a09-objects.docx")
    after = _body(exported)

    for path in ("//c:chart", "//dgm:relIds", "//o:OLEObject", "//wps:wsp[not(wps:txbx)]"):
        assert _count(after, path) == _count(before, path), path
    assert _count(after, "//c:chart") == 1 and _count(after, "//dgm:relIds") == 1 and _count(after, "//o:OLEObject") == 1
    assert package_problems(exported) == []
    lost = next((item for item in noted.items() if item.feature == "export.docx.rewritten_blocks"), None)
    assert lost is None or ("charts" not in lost.reason and "embedded objects" not in lost.reason), lost and lost.reason
    kept = {item.feature: item for item in document.importReport.items}
    assert kept["docx.chart"].policy == "detected_not_editable" and not kept["docx.chart"].contentChanged


def test_a_shape_goes_back_where_it_was_in_its_text_and_on_its_own():
    word = DocxDocument()
    paragraph = word.add_paragraph("Before ")
    paragraph._p.append(_shape(50))
    paragraph.add_run("after.")
    word.add_paragraph()._p.append(_shape(51))  # a paragraph of nothing but a shape
    word.add_paragraph("The end.")
    buffer = io.BytesIO()
    word.save(buffer)

    _, exported, _ = _through_a_template(buffer.getvalue(), "shapes.docx")
    paragraphs = _body(exported).xpath("//w:body/w:p", namespaces=_NS)

    def pieces(p):
        return ["shape" if node.xpath(".//wps:wsp", namespaces=_NS) else "".join(node.xpath(".//w:t/text()", namespaces=_NS)) for node in p.xpath("w:r", namespaces=_NS)]

    assert [piece for piece in pieces(paragraphs[0]) if piece] == ["Before ", "shape", "after."]
    assert pieces(paragraphs[1]) == ["shape"] and "".join(pieces(paragraphs[2])) == "The end."
    ids = [node.get("id") for node in _body(exported).iter(f"{{{_NS['wp']}}}docPr")]
    assert len(ids) == len(set(ids)) == 2
    assert package_problems(exported) == []


def test_a_drawing_used_twice_is_named_by_the_package_check():
    word = DocxDocument()
    word.add_paragraph()._p.append(_shape(7))
    word.add_paragraph()._p.append(_shape(7))
    buffer = io.BytesIO()
    word.save(buffer)

    assert package_problems(buffer.getvalue()) == ["word/document.xml: drawing id '7' is used 2 times"]
