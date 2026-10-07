"""Content controls (tracker DOCX-023, brief §31): every kind -- plain and rich text,
checkbox, drop-down, combo box, date, picture, repeating section -- goes back into
Word with its properties, not only while its block is unchanged. The import keeps a
control in a paragraph around its text (preservedAttributes["ooxml"], kind
"control"), one around blocks as where it starts and ends
(preservedAttributes["controls"]), a picture control with its picture
(preservedAttributes["control"]); a legacy form field keeps its settings (ffData).
Measured in Word before this: a08 written anew had none of its 8 controls, and Word
couldn't read its 2 form fields as form fields; now all of them, as in the file."""

import io
import zipfile
from collections import Counter
from pathlib import Path

from docx import Document as DocxDocument
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from lxml import etree

from app.export.docx_export import build_docx
from app.export.package_check import package_problems
from app.export.provenance import stamp
from app.fidelity.docx_detect import detect_docx_features
from app.fidelity.report import ReportBuilder
from app.formatting.engine import recompute_styles
from app.models.document import Document, InlineRun
from app.parsers.docx import parse_docx

A08 = Path(__file__).parent / "fixtures" / "word" / "a08-content-controls.docx"
GOLDEN_KEPT = Path(__file__).parent / "fixtures" / "documents" / "13-kept-blocks.docx"
_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_W14 = "{http://schemas.microsoft.com/office/word/2010/wordml}"
_W15 = "{http://schemas.microsoft.com/office/word/2012/wordml}"
_KINDS = {
    f"{_W}text": "plain text",
    f"{_W14}checkbox": "checkbox",
    f"{_W}dropDownList": "drop-down",
    f"{_W}comboBox": "combo box",
    f"{_W}date": "date",
    f"{_W}picture": "picture",
    f"{_W15}repeatingSection": "repeating section",
    f"{_W15}repeatingSectionItem": "repeating item",
    f"{_W}docPartObj": "building block",
}


def _body(data: bytes) -> etree._Element:
    return etree.fromstring(zipfile.ZipFile(io.BytesIO(data)).read("word/document.xml")).find(f"{_W}body")


def _kind(sdt: etree._Element) -> str:
    properties = sdt.find(f"{_W}sdtPr")
    return next((_KINDS[child.tag] for child in properties if child.tag in _KINDS), "rich text")


def _title(sdt: etree._Element) -> str | None:
    alias = sdt.find(f"{_W}sdtPr/{_W}alias")
    return alias.get(f"{_W}val") if alias is not None else None


def _controls(data: bytes) -> Counter:
    """The file's content controls: each kind, with its title."""
    return Counter((_kind(sdt), _title(sdt)) for sdt in _body(data).iter(f"{_W}sdt"))


def _a08() -> tuple[bytes, Document]:
    source = A08.read_bytes()
    document = parse_docx(source, A08.name)
    recompute_styles(document)
    stamp(document)  # as an upload does
    return source, document


def _retype(document: Document, starting: str, old: str, new: str) -> None:
    element = next(element for element in document.elements if element.content.startswith(starting))
    element.inline = [InlineRun(text=run.text.replace(old, new), marks=run.marks) for run in element.inline or []]
    element.content = element.content.replace(old, new)


_A08 = Counter(
    {
        ("plain text", "Plain text"): 1,
        ("rich text", "Rich text"): 1,
        ("checkbox", "Checkbox"): 1,
        ("drop-down", "Dropdown"): 1,
        ("combo box", "Combo box"): 1,
        ("date", "Date picker"): 1,
        ("picture", "Picture"): 1,
        ("repeating section", "Repeating section"): 1,
        ("repeating item", None): 1,
    }
)


def test_the_import_keeps_each_kind_of_control_with_its_text():
    _, document = _a08()
    kept = [
        (parse_xml(fragment["properties"]).find(qn("w:alias")).get(qn("w:val")), fragment["text"])
        for element in document.elements
        for fragment in (element.preservedAttributes or {}).get("ooxml") or []
        if fragment["kind"] == "control"
    ]
    picture = next(element for element in document.elements if element.image)
    item = next(element for element in document.elements if element.content == "Repeating item text")

    assert kept == [
        ("Plain text", "Plain text value"),
        ("Rich text", "Rich text value"),
        ("Checkbox", "☒"),
        ("Dropdown", "Option B"),
        ("Combo box", "Option B"),
        ("Date picker", "26.09.2026"),
    ]
    assert parse_xml(picture.preservedAttributes["control"]["properties"]).find(qn("w:picture")) is not None
    assert [marker["edge"] for marker in item.preservedAttributes["controls"]] == ["open", "open", "close", "close"]


def test_a_file_written_anew_has_every_control_back_with_its_properties():
    source, document = _a08()
    assert _controls(source) == _A08

    exported = build_docx(document)  # a new file: nothing copied

    assert _controls(exported) == _A08
    body = _body(exported)
    drop_down = next(sdt for sdt in body.iter(f"{_W}sdt") if _kind(sdt) == "drop-down")
    assert [item.get(f"{_W}displayText") for item in drop_down.iter(f"{_W}listItem")] == ["Option A", "Option B"]
    date = next(sdt for sdt in body.iter(f"{_W}sdt") if _kind(sdt) == "date")
    assert date.find(f"{_W}sdtPr/{_W}date/{_W}dateFormat").get(f"{_W}val") == "dd.MM.yyyy"
    assert package_problems(exported) == []


def test_controls_in_blocks_changed_here_go_back_too():
    source, document = _a08()
    _retype(document, "Dropdown", "Option B", "Option A")  # chosen again, by typing
    _retype(document, "Repeating item", "text", "text, edited")
    report = ReportBuilder()

    exported = build_docx(document, source=source, report=report)

    assert _controls(exported) == _A08
    body = _body(exported)
    drop_down = next(sdt for sdt in body.iter(f"{_W}sdt") if _kind(sdt) == "drop-down")
    assert "".join(t.text for t in drop_down.iter(f"{_W}t")) == "Option A"
    section = next(sdt for sdt in body.iter(f"{_W}sdt") if _kind(sdt) == "repeating section")
    assert "".join(t.text for t in section.iter(f"{_W}t")) == "Repeating item text, edited"  # around its paragraph, the item in it
    assert not any("content controls" in item.reason for item in report.items() if item.feature == "export.docx.rewritten_blocks")
    assert package_problems(exported) == []


def test_a_checkbox_follows_the_symbol_it_shows():
    _, document = _a08()
    _retype(document, "Checkbox", "☒", "☐")

    exported = build_docx(document)

    box = next(sdt for sdt in _body(exported).iter(f"{_W}sdt") if _kind(sdt) == "checkbox")
    assert box.find(f"{_W}sdtPr/{_W14}checkbox/{_W14}checked").get(f"{_W14}val") == "0"
    assert "".join(t.text for t in box.iter(f"{_W}t")) == "☐"


def test_a_control_around_a_paragraph_goes_back_when_the_paragraph_changes():
    source = GOLDEN_KEPT.read_bytes()
    document = parse_docx(source, GOLDEN_KEPT.name)
    recompute_styles(document)
    stamp(document)
    _retype(document, "Draft for review", "Draft", "Final draft")

    exported = build_docx(document, source=source)

    [status] = [sdt for sdt in _body(exported).iter(f"{_W}sdt") if sdt.find(f"{_W}sdtPr/{_W}alias") is not None]
    assert status.find(f"{_W}sdtPr/{_W}alias").get(f"{_W}val") == "Status"
    assert "".join(t.text for t in status.iter(f"{_W}t")) == "Final draft for review"


def test_legacy_form_fields_keep_their_settings():
    _, document = _a08()

    exported = build_docx(document)  # written anew

    forms = [char.find(f"{_W}ffData") for char in _body(exported).iter(f"{_W}fldChar") if char.get(f"{_W}fldCharType") == "begin"]
    names = [form.find(f"{_W}name").get(f"{_W}val") for form in forms if form is not None]
    assert names == ["Text1", "Check1"]
    assert any(form is not None and form.find(f"{_W}checkBox") is not None for form in forms)


def test_each_control_has_an_id_of_its_own():
    _, document = _a08()
    plain = next(element for element in document.elements if element.content.startswith("Plain text"))
    twin = plain.model_copy(deep=True, update={"id": "twin", "sourceBlocks": None, "sourceHash": None})
    document.elements.insert(document.elements.index(plain) + 1, twin)  # pasted twice, say

    exported = build_docx(document)

    ids = [node.get(f"{_W}val") for node in _body(exported).iter(f"{_W}id") if node.getparent().tag == f"{_W}sdtPr"]
    assert len(ids) == len(set(ids)) == 10


def test_malformed_control_data_is_left_out():
    """preservedAttributes comes back from the browser; nothing in it is trusted."""
    _, document = _a08()
    for element in document.elements:
        for fragment in (element.preservedAttributes or {}).get("ooxml") or []:
            if fragment["kind"] == "control":
                fragment["properties"] = f'<w:sdtPr {nsdecls("w", "r")}><w:alias w:val="x" r:id="rId1"/></w:sdtPr>'  # a relationship
    exported = build_docx(document)

    kinds = Counter(kind for kind, _ in _controls(exported).elements())
    assert kinds["plain text"] == kinds["drop-down"] == 0  # left out, their text kept
    assert "Plain text: Plain text value" in [p.text for p in DocxDocument(io.BytesIO(exported)).paragraphs]
    assert package_problems(exported) == []


def test_a_control_that_lost_its_last_block_is_left_out_and_the_export_says_so():
    _, document = _a08()
    item = next(element for element in document.elements if element.content == "Repeating item text")
    item.preservedAttributes = {**item.preservedAttributes, "controls": [m for m in item.preservedAttributes["controls"] if m["edge"] == "open"]}
    report = ReportBuilder()

    exported = build_docx(document, report=report)

    assert ("repeating section", "Repeating section") not in _controls(exported)
    assert "export.docx.control_region" in {entry.feature for entry in report.items()}


def _cell_file(inner: str) -> bytes:
    word = DocxDocument()
    cell = word.add_table(rows=1, cols=1).cell(0, 0)
    cell._tc.append(parse_xml(inner))
    cell._tc.remove(cell.paragraphs[0]._p)
    buffer = io.BytesIO()
    word.save(buffer)
    return buffer.getvalue()


def _edited_export(source: bytes) -> tuple[bytes, ReportBuilder]:
    document = parse_docx(source, "cell.docx")
    recompute_styles(document)
    stamp(document)
    table = document.elements[0]
    cell = table.table.rows[0].cells[0]
    if cell.blocks:
        cell.blocks[0].inline = [InlineRun(text="In a cell, edited")]
        cell.blocks[0].content = "In a cell, edited"
    else:
        cell.inline = [InlineRun(text="In a cell, edited")]
    table.content = "In a cell, edited"
    report = ReportBuilder()
    return build_docx(document, source=source, report=report), report


def test_a_control_in_a_table_cell_goes_back_when_its_table_is_written_anew():
    """DOCX-023A: before, a control in a cell was kept only while its table was unchanged."""
    source = _cell_file(
        f'<w:p {nsdecls("w")}><w:sdt><w:sdtPr><w:alias w:val="Cell field"/><w:text/></w:sdtPr>'
        "<w:sdtContent><w:r><w:t>In a cell</w:t></w:r></w:sdtContent></w:sdt></w:p>"
    )
    features = {item.feature: item.policy.value for item in detect_docx_features(source)}
    exported, report = _edited_export(source)

    assert "docx.content_control.nested" not in features and "docx.content_control" in features
    assert not [item for item in report.items() if item.feature == "export.docx.rewritten_blocks" and "content controls" in item.reason]
    [sdt] = DocxDocument(io.BytesIO(exported)).tables[0]._tbl.iter(qn("w:sdt"))
    assert sdt.find(f"{qn('w:sdtPr')}/{qn('w:alias')}").get(qn("w:val")) == "Cell field"
    assert "".join(t.text for t in sdt.iter(qn("w:t"))) == "In a cell"  # around its own text, still there
    cell = DocxDocument(io.BytesIO(exported)).tables[0]._tbl.find(f"{qn('w:tr')}/{qn('w:tc')}")
    assert "".join(t.text for t in cell.iter(qn("w:t"))) == "In a cell, edited"


def test_a_control_around_paragraphs_in_a_cell_is_kept_only_while_its_table_is_unchanged():
    source = _cell_file(
        f'<w:sdt {nsdecls("w")}><w:sdtPr><w:alias w:val="Cell section"/></w:sdtPr><w:sdtContent>'
        "<w:p><w:r><w:t>In a cell</w:t></w:r></w:p><w:p><w:r><w:t>Second line</w:t></w:r></w:p></w:sdtContent></w:sdt>"
    )
    features = {item.feature: item.policy.value for item in detect_docx_features(source)}
    _, report = _edited_export(source)

    assert features.get("docx.content_control.nested") == "lossy"
    [rewritten] = [item for item in report.items() if item.feature == "export.docx.rewritten_blocks"]
    assert "content controls" in rewritten.reason


def test_a_control_around_a_field_holds_all_of_it():
    """Found by the Phase 3 gate: a06's citation control holds its CITATION field; written
    anew, the field began outside the control and ended in it, and Word called the file
    corrupted. Over the same text the control is outermost now."""
    source = (Path(__file__).parent / "fixtures" / "word" / "a06-fields.docx").read_bytes()
    document = parse_docx(source, "a06.docx")

    exported = build_docx(document)

    citation = next(sdt for sdt in _body(exported).iter(f"{_W}sdt") if sdt.find(f"{_W}sdtPr/{_W}citation") is not None)
    kinds = [char.get(f"{_W}fldCharType") for char in citation.iter(f"{_W}fldChar")]
    assert kinds == ["begin", "separate", "end"]  # the whole field, in the control
    assert package_problems(exported) == []


def test_the_package_check_names_a_field_across_a_control():
    word = DocxDocument()
    word.add_paragraph()._p.append(
        parse_xml(
            f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="begin"/><w:instrText>AUTHOR</w:instrText><w:fldChar w:fldCharType="separate"/></w:r>'
        )
    )
    word.paragraphs[-1]._p.append(
        parse_xml(
            f'<w:sdt {nsdecls("w")}><w:sdtPr/><w:sdtContent><w:r><w:t>Ana</w:t></w:r><w:r><w:fldChar w:fldCharType="end"/></w:r>'
            "</w:sdtContent></w:sdt>"
        )
    )
    buffer = io.BytesIO()
    word.save(buffer)

    assert any("a field starts outside a content control or link" in problem for problem in package_problems(buffer.getvalue()))

