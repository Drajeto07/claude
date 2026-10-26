"""B-037: pictures with no size from Word (pasted or picked in the editor) are exported inside the page.

A phone screenshot (946 x 2048 px) pasted into a document went into the DOCX at its file's own size -- 33 x 72 cm, far
past an A4 page -- while the PDF export drew it inside the page (95% of the text area's height, its own proportions)
and the editor ran it across a page edge. Now the DOCX draws such a picture as the PDF does: no wider than the column,
no taller than 95% of the text area, in proportion. The picture's bytes are untouched; a picture with its own size from
Word keeps that size exactly (it is what the author set in Word). Synthetic pictures only.
"""

import base64
import io
import zipfile

import pytest
from docx import Document as DocxDocument
from docx.shared import Cm
from fastapi.testclient import TestClient
from PIL import Image
from pypdf import PdfReader

from app.main import app

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
EMU_PER_CM = 360_000


def _grid(width: int, height: int) -> bytes:
    image = Image.new("RGB", (width, height), (240, 240, 240))
    for y in range(0, height, 128):
        image.paste((40, 90, 200), (0, y, width, min(y + 64, height)))
    buffer = io.BytesIO()
    image.save(buffer, "PNG")
    return buffer.getvalue()


@pytest.fixture
def client(api_db) -> TestClient:
    client = TestClient(app, base_url="https://testserver")
    assert client.post("/api/v1/auth/register", json={"email": "tall@example.com", "password": "long enough password"}).status_code == 201
    return client


def _with_pasted(client: TestClient, *pictures: bytes) -> str:
    document = client.post("/api/v1/documents", json={"text": "# Report\n\nA paragraph before the pictures."}).json()
    pasted = [{"type": "image", "content": "", "order": 90 + n, "image": {"src": "data:image/png;base64," + base64.b64encode(p).decode()}} for n, p in enumerate(pictures)]
    after = {"type": "paragraph", "content": "A paragraph after the pictures.", "order": 99}
    saved = client.put(f"/api/v1/documents/{document['id']}/content", json={"elements": document["elements"] + pasted + [after]})
    assert saved.status_code == 200, saved.text
    return document["id"]


def _docx_layout(data: bytes) -> tuple[list[tuple[float, float]], float, float, list[bytes], str]:
    package = zipfile.ZipFile(io.BytesIO(data))
    assert package.testzip() is None
    word = DocxDocument(io.BytesIO(data))
    section = word.sections[-1]
    text_width = (section.page_width - section.left_margin - section.right_margin) / EMU_PER_CM
    text_height = (section.page_height - section.top_margin - section.bottom_margin) / EMU_PER_CM
    sizes = [(shape.width / EMU_PER_CM, shape.height / EMU_PER_CM) for shape in word.inline_shapes]
    media = [package.read(name) for name in package.namelist() if name.startswith("word/media/")]
    return sizes, text_width, text_height, media, package.read("word/document.xml").decode()


def _pixels(data: bytes) -> bytes:
    return Image.open(io.BytesIO(data)).convert("RGB").tobytes()


def test_a_pasted_screenshot_is_exported_inside_the_page_in_proportion_with_its_own_pixels(client):
    tall = _grid(946, 2048)
    document_id = _with_pasted(client, tall)
    sizes, text_width, text_height, media, xml = _docx_layout(client.get(f"/api/v1/documents/{document_id}/export/docx").content)
    [(width, height)] = sizes
    assert width <= text_width + 0.01 and height <= 0.95 * text_height + 0.01  # was 33.4 x 72.2 cm
    assert height == pytest.approx(0.95 * text_height, abs=0.05)  # as large as the page allows
    assert width / height == pytest.approx(946 / 2048, rel=0.005)
    assert [_pixels(one) for one in media] == [_pixels(tall)]
    assert "A paragraph before the pictures." in xml and "A paragraph after the pictures." in xml

    pdf = PdfReader(io.BytesIO(client.get(f"/api/v1/documents/{document_id}/export/pdf").content))
    drawn = [(image.image.width, image.image.height) for page in pdf.pages for image in page.images]
    assert drawn == [(946, 2048)]  # the PDF holds the very picture, drawn inside a page (pdf_export.py)


def test_several_tall_and_a_wide_picture_each_fit_and_a_small_one_keeps_its_own_size(client):
    pictures = [_grid(946, 2048), _grid(946, 2100), _grid(3000, 1000), _grid(120, 80)]
    document_id = _with_pasted(client, *pictures)
    sizes, text_width, text_height, media, _ = _docx_layout(client.get(f"/api/v1/documents/{document_id}/export/docx").content)
    assert len(sizes) == 4 and len(media) == 4
    for (width, height), picture in zip(sizes, pictures):
        source = Image.open(io.BytesIO(picture))
        assert width <= text_width + 0.01 and height <= 0.95 * text_height + 0.01
        assert width / height == pytest.approx(source.width / source.height, rel=0.005)
    wide = sizes[2]
    assert wide[0] == pytest.approx(text_width, abs=0.01)  # the wide one fills the column, and its height follows
    small = sizes[3]
    assert small[0] < text_width / 2  # a small picture isn't blown up to the column


def test_a_picture_with_its_own_size_from_word_keeps_it_exactly(client):
    word = DocxDocument()
    word.add_paragraph("Before.")
    word.add_picture(io.BytesIO(_grid(946, 2048)), width=Cm(12))  # the author's size in Word: taller than the text area
    buffer = io.BytesIO()
    word.save(buffer)
    document = client.post("/api/v1/documents/upload", files={"file": ("own.docx", buffer.getvalue(), DOCX)}).json()
    sizes, *_ = _docx_layout(client.get(f"/api/v1/documents/{document['id']}/export/docx").content)
    [(width, height)] = sizes
    assert width == pytest.approx(12, abs=0.01) and height == pytest.approx(12 * 2048 / 946, abs=0.02)  # not refitted
