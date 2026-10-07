"""Word tables' geometry and look (tracker DOCX-017, brief §26): column widths, the
table's width, alignment and indent, row heights, borders and cell margins -- a table
style's where the table has none of its own -- cells' vertical and own alignment,
header rows only where the file says so, and the style's name and look. Imported,
written back into Word, drawn in a PDF; a table made here keeps its grid."""

import io
import zipfile

from docx import Document as DocxDocument
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

from app.export.docx_export import build_docx
from app.export.package_check import package_problems
from app.export.pdf_export import build_pdf
from app.models.document import Document, Element, ElementType, InlineRun, TableCell, TableContent, TableRow
from app.parsers.docx import parse_docx

_W = nsdecls("w")


def _save(word) -> bytes:
    buffer = io.BytesIO()
    word.save(buffer)
    return buffer.getvalue()


def _table_of(document):
    [table] = [element.table for element in document.elements if element.type == ElementType.TABLE]
    return table


def _word_table(style: str = "Table Grid") -> bytes:
    word = DocxDocument()
    table = word.add_table(rows=3, cols=2)
    table.style = word.styles[style]
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for column, twips in zip(table._tbl.tblGrid.findall(qn("w:gridCol")), (1134, 4536)):  # 2 cm and 8 cm
        column.set(qn("w:w"), str(twips))
    table._tbl.tblPr.append(parse_xml(f'<w:tblCellMar {_W}><w:left w:w="170" w:type="dxa"/><w:right w:w="170" w:type="dxa"/></w:tblCellMar>'))
    table.rows[1]._tr.get_or_add_trPr().append(parse_xml(f'<w:trHeight {_W} w:val="851" w:hRule="exact"/>'))
    for (row, column), text in {(0, 0): "Item", (0, 1): "Price", (1, 0): "Paper", (1, 1): "12.50", (2, 0): "Ink", (2, 1): "3.20"}.items():
        table.cell(row, column).text = text
    tc_pr = table.cell(1, 1)._tc.get_or_add_tcPr()
    tc_pr.append(parse_xml(f'<w:tcBorders {_W}><w:bottom w:val="double" w:sz="12" w:color="C00000"/></w:tcBorders>'))
    tc_pr.append(parse_xml(f'<w:tcMar {_W}><w:top w:w="57" w:type="dxa"/></w:tcMar>'))
    tc_pr.append(parse_xml(f'<w:vAlign {_W} w:val="center"/>'))
    return _save(word)


def test_a_tables_grid_borders_margins_and_style_are_kept():
    table = _table_of(parse_docx(_word_table(), "prices.docx"))

    assert (table.style, table.columnWidthsCm, table.align) == ("Table Grid", [2.0, 8.0], "center")
    assert table.borders.top == table.borders.insideV == "solid 0.5pt #000000"  # the style's, resolved
    assert (table.cellMargins.leftCm, table.cellMargins.rightCm) == (0.3, 0.3)  # the table's own over the style's
    assert (table.rows[1].heightCm, table.rows[1].heightRule) == (1.5, "exact")
    cell = table.rows[1].cells[1]
    assert (cell.verticalAlign, cell.borders.bottom, cell.margins.topCm) == ("center", "double 1.5pt #C00000", 0.1)
    assert table.headerBold is False and table.look.firstRow


def test_a_header_row_needs_evidence():
    plain = _table_of(parse_docx(_word_table(), "plain.docx"))
    assert not plain.hasHeaderRow and not any(cell.header for row in plain.rows for cell in row.cells)

    word = DocxDocument(io.BytesIO(_word_table()))
    word.tables[0].rows[0]._tr.get_or_add_trPr().append(parse_xml(f"<w:tblHeader {_W}/>"))
    repeated = _table_of(parse_docx(_save(word), "repeated.docx"))
    assert repeated.rows[0].repeatHeader and all(cell.header for cell in repeated.rows[0].cells)

    styled = _table_of(parse_docx(_word_table("Light List Accent 1"), "styled.docx"))  # its first row is bold and filled
    first = styled.rows[0].cells
    assert styled.hasHeaderRow and all(cell.header for cell in first) and not styled.rows[0].repeatHeader
    # Its first row as Word draws it: filled, bold, the text white (DOCX-017A resolves the colour too).
    assert first[0].background is not None and {mark.type for run in first[0].inline for mark in run.marks} == {"bold", "textStyle"}
    assert {mark.color for run in first[0].inline for mark in run.marks if mark.type == "textStyle"} == {"#FFFFFF"}


def test_a_word_export_writes_the_tables_geometry_back():
    document = parse_docx(_word_table("Light List Accent 1"), "styled.docx")

    exported = build_docx(document)
    again = _table_of(parse_docx(exported, "again.docx"))

    assert package_problems(exported) == []
    original = _table_of(document)
    assert again.model_dump(exclude={"rows"}) == original.model_dump(exclude={"rows"})
    assert [row.model_dump(exclude={"id", "cells"}) for row in again.rows] == [row.model_dump(exclude={"id", "cells"}) for row in original.rows]
    looks = lambda table: [(cell.header, cell.verticalAlign, cell.borders, cell.margins, cell.background) for row in table.rows for cell in row.cells]
    assert looks(again) == looks(original)
    with zipfile.ZipFile(io.BytesIO(exported)) as package:
        body = package.read("word/document.xml").decode("utf-8")
    assert 'w:tblStyle w:val="LightList-Accent1"' in body and "TableGrid" not in body  # its own style, not a grid forced on it


def test_a_pdf_draws_the_tables_widths_and_borders():
    from pypdf import PdfReader

    placed: dict[str, float] = {}

    def visit(text, cm_matrix, tm_matrix, font_dict, font_size):
        if text.strip() in ("Paper", "12.50"):
            placed[text.strip()] = cm_matrix[4] + tm_matrix[4]

    PdfReader(io.BytesIO(build_pdf(parse_docx(_word_table(), "prices.docx")))).pages[0].extract_text(visitor_text=visit)

    assert abs((placed["12.50"] - placed["Paper"]) - 2 * 72 / 2.54) < 3  # the second column starts 2 cm on


def test_a_table_made_here_keeps_its_grid_and_bold_header():
    rows = [
        TableRow(cells=[TableCell(inline=[InlineRun(text="Name")], header=True), TableCell(inline=[InlineRun(text="Price")], header=True)]),
        TableRow(cells=[TableCell(inline=[InlineRun(text="Widget")]), TableCell(inline=[InlineRun(text="9.99")])]),
    ]
    element = Element(type=ElementType.TABLE, content="", order=0, table=TableContent(rows=rows, hasHeaderRow=True))

    exported = build_docx(Document(elements=[element]))

    word = DocxDocument(io.BytesIO(exported))
    table = word.tables[0]
    assert table.style.name == "Table Grid" and table.alignment == WD_TABLE_ALIGNMENT.CENTER
    assert table.cell(0, 0).paragraphs[0].runs[0].bold and not table.cell(1, 0).paragraphs[0].runs[0].bold


def test_a_cell_aligned_unlike_its_column_keeps_its_own_alignment():
    word = DocxDocument(io.BytesIO(_word_table()))
    table = word.tables[0]
    table.cell(0, 0).paragraphs[0].alignment = 1  # centred
    table.cell(1, 0).paragraphs[0].alignment = 2  # right

    imported = _table_of(parse_docx(_save(word), "aligned.docx"))
    again = _table_of(parse_docx(build_docx(parse_docx(_save(word), "aligned.docx")), "again.docx"))

    assert imported.alignments is None or imported.alignments[0] is None
    assert [row.cells[0].align for row in imported.rows] == ["center", "right", None]
    assert [row.cells[0].align for row in again.rows] == ["center", "right", None]


def _busy_cell() -> bytes:
    word = DocxDocument()
    table = word.add_table(rows=1, cols=2)
    table.style = word.styles["Table Grid"]
    cell = table.cell(0, 0)
    cell.paragraphs[0].text = "First paragraph."
    cell.add_paragraph("Second paragraph.")
    for text in ("apples", "pears"):
        cell.add_paragraph(text, style="List Bullet")
    inner = cell.add_table(rows=1, cols=2)
    inner.cell(0, 0).text, inner.cell(0, 1).text = "inner a", "inner b"
    cell.add_paragraph("After the table.")
    table.cell(0, 1).text = "Plain."
    return _save(word)


def test_a_cells_paragraphs_lists_and_tables_are_its_blocks():
    table = _table_of(parse_docx(_busy_cell(), "cells.docx"))

    busy, plain = table.rows[0].cells
    # python-docx, as Word, puts an empty paragraph after a table in a cell: it is the cell's too.
    assert [block.type for block in busy.blocks] == ["paragraph", "paragraph", "list", "table", "paragraph", "paragraph"]
    assert [block.content for block in busy.blocks[4:]] == ["", "After the table."]
    assert [block.content for block in busy.blocks[:2]] == ["First paragraph.", "Second paragraph."]
    bullets = busy.blocks[2]
    assert not bullets.ordered and [item.inline[0].text for item in bullets.listItems] == ["apples", "pears"]
    assert [cell.inline[0].text for cell in busy.blocks[3].table.rows[0].cells] == ["inner a", "inner b"]
    assert plain.blocks is None and plain.inline[0].text == "Plain."  # one plain paragraph stays the cell's text


def test_a_cells_blocks_come_back_from_a_word_export():
    from pypdf import PdfReader

    document = parse_docx(_busy_cell(), "cells.docx")

    exported = build_docx(document)
    again = _table_of(parse_docx(exported, "again.docx"))
    pdf_text = "\n".join(page.extract_text() for page in PdfReader(io.BytesIO(build_pdf(document))).pages)

    assert package_problems(exported) == []
    shape = lambda cell: [(block.type, block.content) for block in cell.blocks or []]
    assert shape(again.rows[0].cells[0]) == shape(_table_of(document).rows[0].cells[0])
    assert "inner a" in pdf_text and "pears" in pdf_text and "After the table." in pdf_text


def test_a_floating_table_and_a_row_kept_whole_come_back_from_a_word_export():
    word = DocxDocument(io.BytesIO(_word_table()))
    table = word.tables[0]
    table._tbl.tblPr.insert(1, parse_xml(
        f'<w:tblpPr {_W} w:leftFromText="170" w:rightFromText="170" w:vertAnchor="text" w:horzAnchor="margin" w:tblpX="567" w:tblpYSpec="top"/>'
    ))
    table.rows[2]._tr.get_or_add_trPr().append(parse_xml(f"<w:cantSplit {_W}/>"))
    document = parse_docx(_save(word), "floating.docx")

    imported = _table_of(document)
    again = _table_of(parse_docx(build_docx(document), "again.docx"))

    assert (imported.floating.horizontalAnchor, imported.floating.xCm, imported.floating.yAlign, imported.floating.leftFromTextCm) == ("margin", 1.0, "top", 0.3)
    assert imported.rows[2].cantSplit and not imported.rows[1].cantSplit
    assert again.floating == imported.floating and again.rows[2].cantSplit
    assert "docx.table.floating" in {item.feature for item in document.importReport.items}


_STYLE = f"""<w:style {_W} w:type="table" w:styleId="Banded"><w:name w:val="Banded"/>
  <w:tblPr><w:tblStyleRowBandSize w:val="2"/></w:tblPr>
  <w:tblStylePr w:type="firstRow"><w:rPr><w:b/></w:rPr><w:tcPr><w:shd w:val="clear" w:color="auto" w:fill="1F3864"/></w:tcPr></w:tblStylePr>
  <w:tblStylePr w:type="lastRow"><w:rPr><w:b/><w:color w:val="C00000"/></w:rPr></w:tblStylePr>
  <w:tblStylePr w:type="firstCol"><w:rPr><w:i/></w:rPr><w:tcPr><w:shd w:val="clear" w:color="auto" w:fill="E2EFDA"/></w:tcPr></w:tblStylePr>
  <w:tblStylePr w:type="band1Horz"><w:tcPr><w:shd w:val="clear" w:color="auto" w:fill="DDEBF7"/></w:tcPr></w:tblStylePr>
  <w:tblStylePr w:type="swCell"><w:tcPr><w:shd w:val="clear" w:color="auto" w:fill="FFF2CC"/></w:tcPr></w:tblStylePr>
</w:style>"""


def _banded_table(look: str = 'w:firstRow="1" w:lastRow="1" w:firstColumn="1" w:lastColumn="0" w:noHBand="0" w:noVBand="1"') -> bytes:
    """Seven rows by two in a style that fills banded rows two at a time, bolds the first and last
    rows, italicises the first column and fills the bottom-left corner; one cell is shaded itself."""
    word = DocxDocument()
    word.styles.element.append(parse_xml(_STYLE))
    table = word.add_table(rows=7, cols=2)
    table._tbl.tblPr.append(parse_xml(f'<w:tblStyle {_W} w:val="Banded"/>'))
    for default in table._tbl.tblPr.findall(qn("w:tblLook")):  # python-docx gives every table one
        table._tbl.tblPr.remove(default)
    table._tbl.tblPr.append(parse_xml(f"<w:tblLook {_W} {look}/>"))
    for row in range(7):
        for column in range(2):
            table.cell(row, column).text = f"r{row}c{column}"
    table.cell(1, 1)._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {_W} w:val="clear" w:color="auto" w:fill="00FF00"/>'))
    return _save(word)


def _looks(table) -> list[list[tuple]]:
    def marks(cell):
        found = {mark.type.value + (f":{mark.color}" if mark.color else "") for run in cell.inline for mark in run.marks}
        return tuple(sorted(found))

    return [[(cell.background, marks(cell)) for cell in row.cells] for row in table.rows]


def test_a_table_styles_banded_rows_columns_and_corners_are_drawn_as_word_draws_them():
    # DOCX-017A: only an export written into the original kept these.
    looks = _looks(_table_of(parse_docx(_banded_table(), "banded.docx")))
    header, band = "#1F3864", "#DDEBF7"
    assert looks[0] == [(header, ("bold", "italic")), (header, ("bold",))]  # the first row (and column)
    # Banded rows leave the first row out, two rows a band: rows 1-2 filled, 3-4 not, 5 filled.
    assert [row[1][0] for row in looks[1:6]] == ["#00FF00", band, None, None, band]  # row 1's own shading over its band
    assert looks[0][0][0] == header and looks[6][0][0] == "#FFF2CC"  # the first row and the corner over the first column
    assert looks[1][0] == ("#E2EFDA", ("italic",))  # the first column over the band, in Word's order
    assert looks[6] == [("#FFF2CC", ("bold", "italic", "textStyle:#C00000")), (None, ("bold", "textStyle:#C00000"))]  # the last row, its corner


def test_only_the_parts_a_table_shows_are_drawn():
    plain = 'w:firstRow="0" w:lastRow="0" w:firstColumn="0" w:lastColumn="0" w:noHBand="1" w:noVBand="1"'
    looks = _looks(_table_of(parse_docx(_banded_table(plain), "plain.docx")))
    assert all(cell == (None, ()) for row in looks for cell in row if cell[0] != "#00FF00")


def test_the_drawn_look_reaches_a_pdf_and_a_new_word_file():
    document = parse_docx(_banded_table(), "banded.docx")
    again = _table_of(parse_docx(build_docx(document), "again.docx"))  # written anew: the look is the cells' own now
    assert _looks(again) == _looks(_table_of(document))
    assert build_pdf(document)[:4] == b"%PDF"
