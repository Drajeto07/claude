"""The Word export looks a style up once, not once per paragraph (PERF-008). python-docx's
get_style_id scans every style of the file (about 2 ms with a Word file's hundred and more),
which for headings, lists, quotes and code, in the body and in cells, was the biggest cost of
a long document. The answers are remembered per python-docx document, and forgotten when the
styles change; what comes out is the same bytes (the fixtures are compared part by part in
the report of this task; here the lookups are held to python-docx's own)."""

import time
from collections import Counter

import pytest
from docx import Document as DocxDocument
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml.styles import CT_Styles

from app.export import docx_export
from app.export.docx_export import _add_paragraph_at_end, _named_style, _new_paragraph, _style_id, _style_lookups, _table_style, _word_style, build_docx
from app.formatting.engine import recompute_styles
from app.models.document import Document, Element, ElementType, InlineRun, ListItem, TableCell, TableContent, TableRow

_PARAGRAPH, _CHARACTER, _TABLE = WD_STYLE_TYPE.PARAGRAPH, WD_STYLE_TYPE.CHARACTER, WD_STYLE_TYPE.TABLE


def _text_element(kind: ElementType, index: int) -> Element:
    text = f"Block {index} with some words in it."
    if kind == ElementType.LIST:
        return Element(type=kind, content=text, order=index, listItems=[ListItem(inline=[InlineRun(text="one")]), ListItem(inline=[InlineRun(text="two")])])
    if kind == ElementType.CODE_BLOCK:
        return Element(type=kind, content=text, order=index)
    return Element(type=kind, content=text, order=index, level=1 + index % 3 if kind == ElementType.HEADING else None, inline=[InlineRun(text=text)])


def _blocks(count: int, kinds: tuple[ElementType, ...]) -> Document:
    document = Document(elements=[_text_element(kinds[index % len(kinds)], index) for index in range(count)])
    recompute_styles(document)
    return document


_STYLED = (ElementType.HEADING, ElementType.QUOTE, ElementType.CODE_BLOCK, ElementType.LIST)
_PLAIN_AND_STYLED = (ElementType.HEADING, ElementType.QUOTE, ElementType.CODE_BLOCK, ElementType.PARAGRAPH)


# -- what is looked up is what python-docx gives --------------------------------------------


def test_a_remembered_style_id_is_the_one_python_docx_gives():
    document = DocxDocument()
    for name, kind in [("Normal", _PARAGRAPH), ("Heading 1", _PARAGRAPH), ("Quote", _PARAGRAPH), ("List Bullet", _PARAGRAPH), ("Table Grid", _TABLE), ("Hyperlink", _CHARACTER)]:
        try:
            expected = document.part.get_style_id(name, kind)
        except (KeyError, ValueError) as error:
            expected = type(error)
        for _ in range(2):  # the second answer comes from what was remembered
            try:
                found = _style_id(document, name, kind)
            except (KeyError, ValueError) as error:
                found = type(error)
            assert found == expected, (name, kind)
        style = document.styles[name] if name in document.styles else None
        if style is not None and style.type == kind:
            assert _style_id(document, style, kind) == document.part.get_style_id(style, kind)


def test_the_default_style_has_no_id_and_wrong_kind_or_name_fail_as_in_python_docx_and_are_not_remembered():
    document = DocxDocument()
    assert _style_id(document, "Normal") is None  # the default paragraph style: python-docx writes no id for it
    with pytest.raises(ValueError):
        _style_id(document, "Table Grid", _PARAGRAPH)  # a table style, asked for as a paragraph's
    with pytest.raises(KeyError):
        _style_id(document, "No Such Style")
    document.styles.add_style("No Such Style", _PARAGRAPH)  # ... and now there is one
    assert _style_id(document, "No Such Style") == document.part.get_style_id("No Such Style", _PARAGRAPH) == "NoSuchStyle"


def test_a_new_paragraph_is_the_one_add_paragraph_makes():
    plain, ours = DocxDocument(), DocxDocument()
    for name in ("Heading 2", "Quote", "List Bullet", "Normal", None):
        plain.add_paragraph(style=name)
        _new_paragraph(ours, name)
    assert [p._p.xml for p in plain.paragraphs] == [p._p.xml for p in ours.paragraphs]


# -- and forgotten when the styles change ----------------------------------------------------


def test_a_style_added_after_a_lookup_is_found():
    document = DocxDocument()
    assert _named_style(document, "Made Later") is None
    assert _table_style(document, "Made Later Table") is None
    assert _style_id(document, "Normal") is None
    document.styles.add_style("Made Later", _PARAGRAPH)
    document.styles.add_style("Made Later Table", _TABLE)
    assert _named_style(document, "Made Later") is not None
    assert _table_style(document, "Made Later Table") is not None


def test_a_style_the_export_adds_midway_is_found_by_the_paragraphs_after_it():
    document = DocxDocument()
    styles = document.styles.element
    for style in [s for s in styles.findall(qn("w:style")) if s.find(qn("w:name")).get(qn("w:val")) in ("Quote", "heading 3")]:
        styles.remove(style)
    assert _named_style(document, "Quote") is None
    first = _new_paragraph(document, _word_style(document, "Quote"))  # the style is made now
    second = _new_paragraph(document, "Quote")
    assert first._p.style == second._p.style == "Quote"
    assert _named_style(document, "Quote") is not None
    heading = _new_paragraph(document, _word_style(document, "Heading 3"))
    assert heading._p.style == document.part.get_style_id("Heading 3", _PARAGRAPH)


def test_every_document_has_its_own_lookups():
    first, second = DocxDocument(), DocxDocument()
    _style_id(first, "Heading 1")
    assert _style_lookups(first) is _style_lookups(first)
    assert _style_lookups(first) is not _style_lookups(second)
    assert not _style_lookups(second).ids


def test_an_export_that_adds_styles_as_it_goes_gives_the_same_file_as_one_made_with_python_docx_lookups(monkeypatch):
    """Styles a file lacks are made when their kind of block first comes up. The export
    with the remembered answers and one with every lookup done afresh (python-docx's own,
    forced by forgetting after each) are the same bytes."""
    document = _blocks(60, _STYLED)
    document.elements.append(Element(type=ElementType.TABLE, content="", order=99, table=TableContent(rows=[TableRow(cells=[TableCell(inline=[], blocks=[_text_element(ElementType.HEADING, 1), _text_element(ElementType.QUOTE, 2)])])])))
    kept = build_docx(document)

    real = docx_export._style_lookups

    def forgetting(docx_document):
        docx_document.part._style_lookups = None
        return real(docx_document)

    monkeypatch.setattr(docx_export, "_style_lookups", forgetting)
    afresh = build_docx(document)

    import io
    import re
    import zipfile

    def parts(blob):
        archive = zipfile.ZipFile(io.BytesIO(blob))
        return {name: re.sub(rb"<dcterms:(modified|created)[^>]*>[^<]*</dcterms:(modified|created)>", b"", archive.read(name)) for name in archive.namelist()}

    assert parts(kept) == parts(afresh)


# -- a long document does not pay for them ---------------------------------------------------


def test_the_styles_are_not_scanned_once_per_paragraph(monkeypatch):
    """Deterministic: how often python-docx's scanning lookups run in an export doesn't
    grow with the document."""
    scans: Counter = Counter()
    for owner, name in ((CT_Styles, "default_for"), (CT_Styles, "get_by_name"), (CT_Styles, "get_by_id")):
        real = getattr(owner, name)

        def counted(self, *args, _real=real, _name=name):
            scans[_name] += 1
            return _real(self, *args)

        monkeypatch.setattr(owner, name, counted)

    def scans_for(count: int) -> int:
        scans.clear()
        build_docx(_blocks(count, _STYLED))
        return sum(scans.values())

    small, large = scans_for(40), scans_for(400)
    assert large <= small + 10, (small, large)  # ten times the blocks, no more lookups to speak of


def test_a_big_document_s_export_costs_less_than_a_style_scan_a_paragraph():
    """10,000 blocks, most of them of the kinds that take a style (headings, quotes, code):
    the time for each is compared with what python-docx's own lookup of a style takes on
    this machine (what each of them cost before, at least, and about all it cost with
    them) and with the same document a tenth the size (linear, not more: python-docx also
    searched the whole body to place each paragraph). Both are ratios, so a slow or busy
    machine passes. Lists are left out: each one looks through the document's numbering,
    which is its own quadratic, not this task's."""
    probe = DocxDocument()
    started = time.perf_counter()
    for _ in range(30):
        probe.part.get_style_id("Heading 1", _PARAGRAPH)
    scan = (time.perf_counter() - started) / 30

    def per_block(count: int) -> float:
        document = _blocks(count, _PLAIN_AND_STYLED)
        best = float("inf")
        for _ in range(2):
            started = time.perf_counter()
            build_docx(document)
            best = min(best, (time.perf_counter() - started) / count)
        return best

    small, large = per_block(1_000), per_block(10_000)
    assert large < 0.5 * scan, f"{large * 1000:.2f} ms a block against {scan * 1000:.2f} ms for one python-docx style lookup"
    assert large < 1.5 * small, f"a block takes {small * 1000:.2f} ms in 1,000 and {large * 1000:.2f} ms in 10,000"


def test_a_paragraph_is_placed_where_python_docx_places_it():
    """In the body in front of the section properties, in a cell, after what was there."""
    plain, ours = DocxDocument(), DocxDocument()
    for document, add in ((plain, lambda d: d.add_paragraph()), (ours, _add_paragraph_at_end)):
        document.add_paragraph("first")
        add(document).add_run("second")
        table = document.add_table(rows=1, cols=1)
        add(table.cell(0, 0)).add_run("in a cell")
        add(document).add_run("after the table")
    assert plain.element.xml == ours.element.xml
    assert ours.element.body[-1].tag == qn("w:sectPr")
    assert [p.text for p in ours.paragraphs] == ["first", "second", "after the table"]
