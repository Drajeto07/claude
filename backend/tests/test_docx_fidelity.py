"""What a Word document looks like after import: lists and page breaks that come
through styles, direct and style formatting, character formatting, header and
footer fields, pictures, captions, checklists, table spans and shading, links,
tracked changes, equations, footnotes, text boxes. Each document is built here
with python-docx (plus raw XML where python-docx has no API)."""

import io

from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Mm, Pt, RGBColor
from PIL import Image as PILImage

from app.formatting.engine import apply_formatting
from app.formatting.priorities import Priority
from app.formatting.templates import BUILTIN_TEMPLATES
from app.models.document import Element, ElementType, MarkType
from app.parsers.docx import parse_docx

_NS = nsdecls("w", "m", "r")
_V = 'xmlns:v="urn:schemas-microsoft-com:vml"'


def _parse(doc: DocxDocument):
    buffer = io.BytesIO()
    doc.save(buffer)
    return parse_docx(buffer.getvalue(), "test.docx")


def _append_xml(doc: DocxDocument, xml: str) -> None:
    body = doc.element.body
    body.insert(len(body) - 1, parse_xml(xml))  # before the final sectPr


def _png() -> bytes:
    buffer = io.BytesIO()
    PILImage.new("RGB", (40, 20), "blue").save(buffer, format="PNG")
    return buffer.getvalue()


def _elements(document, kind: ElementType) -> list[Element]:
    return [element for element in document.elements if element.type == kind]


def _own_css(document, element: Element) -> dict[str, str]:
    return document.resolvedStyles.get(element.styleRef or "", {})


def _marks(run) -> dict[MarkType, object]:
    return {mark.type: mark for mark in run.marks}


# -- structure -------------------------------------------------------------------


def test_lists_numbered_through_their_style_keep_their_kind_and_level():
    doc = DocxDocument()
    for text in ("Apple", "Pear"):
        doc.add_paragraph(text, style="List Bullet")
    for text in ("One", "Two"):
        doc.add_paragraph(text, style="List Number")
    doc.add_paragraph("Deeper", style="List Bullet 2")

    document = _parse(doc)

    lists = _elements(document, ElementType.LIST)
    assert [(lst.ordered, [item.inline[0].text for item in lst.listItems]) for lst in lists] == [
        (False, ["Apple", "Pear"]),
        (True, ["One", "Two"]),
        (False, ["Deeper"]),
    ]
    assert _own_css(document, lists[2])["margin-left"] == "0.63cm"  # a second-level list, indented as one


def test_page_breaks_become_page_break_elements():
    doc = DocxDocument()
    doc.add_paragraph("Before")
    doc.add_page_break()
    doc.add_paragraph("After")
    doc.add_paragraph("Starts its own page").paragraph_format.page_break_before = True

    document = _parse(doc)

    assert [element.type for element in document.elements] == [
        ElementType.PARAGRAPH,
        ElementType.PAGE_BREAK,
        ElementType.PARAGRAPH,
        ElementType.PAGE_BREAK,
        ElementType.PARAGRAPH,
    ]


def test_an_empty_paragraph_with_a_border_is_a_horizontal_rule():
    doc = DocxDocument()
    _append_xml(doc, f'<w:p {_NS}><w:pPr><w:pBdr><w:bottom w:val="single" w:sz="6" w:space="1" w:color="auto"/></w:pBdr></w:pPr></w:p>')

    assert [element.type for element in _parse(doc).elements] == [ElementType.HORIZONTAL_RULE]


def test_a_paragraph_set_entirely_in_a_monospace_font_is_a_code_block():
    doc = DocxDocument()
    run = doc.add_paragraph().add_run("def total(x):\n    return x * 2")
    run.font.name = "Courier New"

    document = _parse(doc)

    [code] = _elements(document, ElementType.CODE_BLOCK)
    assert code.content == "def total(x):\n    return x * 2"
    assert _own_css(document, code)["font-family"] == "Courier New"


def test_captions_are_recognised_next_to_pictures_and_tables():
    doc = DocxDocument()
    doc.add_picture(io.BytesIO(_png()))
    doc.add_paragraph("Figure 1: The blue box.")
    doc.add_paragraph("Table 2: Prices")
    doc.add_table(rows=1, cols=1)
    doc.add_paragraph("Figure 3 is discussed later, but this is ordinary text.")

    document = _parse(doc)

    assert [element.type for element in document.elements] == [
        ElementType.IMAGE,
        ElementType.CAPTION,
        ElementType.CAPTION,
        ElementType.TABLE,
        ElementType.PARAGRAPH,
    ]


def test_checkbox_list_items_become_a_checklist():
    doc = DocxDocument()
    for text in ("☐ Write the report", "☑ Book the room"):
        doc.add_paragraph(text, style="List Bullet")
    symbol_item = doc.add_paragraph(style="List Bullet")
    symbol_item._p.append(parse_xml(f'<w:r {_NS}><w:sym w:font="Wingdings" w:char="F0FE"/></w:r>'))
    symbol_item.add_run(" Send invites")

    [checklist] = _elements(_parse(doc), ElementType.LIST)

    assert checklist.ordered is False
    assert [(item.checked, item.inline[0].text) for item in checklist.listItems] == [
        (False, "Write the report"),
        (True, "Book the room"),
        (True, "Send invites"),
    ]


def test_headings_numbered_by_word_show_their_numbers():
    doc = DocxDocument()
    for text in ("Introduction", "Method"):
        heading = doc.add_heading(text, level=1)
        heading._p.get_or_add_pPr().append(parse_xml(f'<w:numPr {_NS}><w:ilvl w:val="0"/><w:numId w:val="5"/></w:numPr>'))

    headings = _elements(_parse(doc), ElementType.HEADING)

    assert [heading.content for heading in headings] == ["1. Introduction", "2. Method"]


def test_text_boxes_are_imported_as_paragraphs():
    doc = DocxDocument()
    _append_xml(
        doc,
        f'<w:p {_NS} {_V}><w:r><w:t>Anchor text</w:t></w:r><w:r><w:pict><v:shape><v:textbox><w:txbxContent>'
        f"<w:p><w:r><w:t>Inside the box</w:t></w:r></w:p></w:txbxContent></v:textbox></v:shape></w:pict></w:r></w:p>",
    )

    document = _parse(doc)

    assert [element.content for element in document.elements] == ["Anchor text", "Inside the box"]
    assert "Text boxes were imported as ordinary paragraphs." in document.unsupportedFeatures


# -- how it looks ----------------------------------------------------------------


def test_the_documents_word_styles_become_its_own_look():
    doc = DocxDocument()
    doc.styles["Normal"].font.name = "Georgia"
    doc.styles["Normal"].font.size = Pt(12)
    doc.styles["Heading 1"].font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    doc.add_heading("Title", level=1)
    doc.add_paragraph("Body text.")

    document = _parse(doc)

    assert document.resolvedStyles["Paragraph"]["font-family"] == "Georgia"
    assert document.resolvedStyles["Paragraph"]["font-size"] == "12pt"
    assert document.resolvedStyles["Heading 1"]["color"] == "#C00000"
    source = {rule.priority for rule in document.formattingRules if rule.source == "source_document"}
    assert source == {Priority.SOURCE_DOCUMENT}


def test_direct_paragraph_formatting_applies_to_that_paragraph_only():
    doc = DocxDocument()
    centered = doc.add_paragraph("Centered")
    centered.alignment = WD_ALIGN_PARAGRAPH.CENTER
    indented = doc.add_paragraph("Indented")
    indented.paragraph_format.first_line_indent = Cm(1.25)
    indented.paragraph_format.line_spacing = 1.5
    indented.paragraph_format.space_before = Pt(6)
    doc.add_paragraph("Plain")

    document = _parse(doc)

    centered_el, indented_el, plain_el = document.elements
    assert _own_css(document, centered_el)["text-align"] == "center"
    indented_css = _own_css(document, indented_el)
    assert (indented_css["text-indent"], indented_css["line-height"], indented_css["margin-top"]) == ("1.25cm", "1.5", "6pt")
    assert plain_el.styleRef == "Paragraph"


def test_character_formatting_shared_by_the_whole_paragraph_moves_to_the_paragraph():
    doc = DocxDocument()
    paragraph = doc.add_paragraph()
    for text in ("Big ", "title"):
        run = paragraph.add_run(text)
        run.font.size = Pt(28)
        run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
        run.bold = True

    document = _parse(doc)

    [title] = document.elements
    css = _own_css(document, title)
    assert (css["font-size"], css["color"]) == ("28pt", "#1F4E79")
    assert [set(_marks(run)) for run in title.inline] == [{MarkType.BOLD}]


def test_character_formatting_that_varies_stays_on_the_text():
    doc = DocxDocument()
    paragraph = doc.add_paragraph()
    paragraph.add_run("Normal ")
    red = paragraph.add_run("red and big")
    red.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)
    red.font.size = Pt(20)
    paragraph.add_run(" and ").font.highlight_color = WD_COLOR_INDEX.YELLOW
    paragraph.add_run("x")
    paragraph.add_run("2").font.superscript = True
    paragraph.add_run(" H")
    paragraph.add_run("2").font.subscript = True
    paragraph.add_run("O ")
    paragraph.add_run("print()").font.name = "Consolas"
    paragraph.add_run(" in ")
    paragraph.add_run("Georgia").font.name = "Georgia"

    [element] = _parse(doc).elements
    by_text = {run.text: _marks(run) for run in element.inline}

    red_style = by_text["red and big"][MarkType.TEXT_STYLE]
    assert (red_style.color, red_style.fontSizePt) == ("#FF0000", 20)
    assert by_text[" and "][MarkType.TEXT_STYLE].backgroundColor == "#FFFF00"
    assert MarkType.SUPERSCRIPT in by_text["2"] or MarkType.SUBSCRIPT in by_text["2"]
    assert [set(_marks(run)) for run in element.inline if run.text == "2"] == [{MarkType.SUPERSCRIPT}, {MarkType.SUBSCRIPT}]
    assert set(by_text["print()"]) == {MarkType.CODE}  # an inline monospace run is code
    assert by_text["Georgia"][MarkType.TEXT_STYLE].fontFamily == "Georgia"


def test_a_template_applied_later_still_restyles_an_imported_document():
    doc = DocxDocument()
    doc.styles["Normal"].font.name = "Georgia"
    doc.styles["Heading 1"].font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    doc.add_heading("Title", level=1)
    centered = doc.add_paragraph("Centered in the original.")
    centered.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document = _parse(doc)
    template = BUILTIN_TEMPLATES["academic-default"]
    apply_formatting(document, template_id=template.id, template_rules=template.rules, instruction_rules=[])

    assert document.resolvedStyles["Paragraph"]["font-family"] == "Times New Roman"
    assert _own_css(document, document.elements[1])["text-align"] == "justify"  # the template's, not the original's
    assert document.resolvedStyles["Heading 1"]["color"] == "#C00000"  # the template sets no colour, so the original's stays


# -- page ------------------------------------------------------------------------


def test_page_size_orientation_and_margins_come_from_the_section():
    doc = DocxDocument()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.left_margin, section.right_margin = Cm(3), Cm(2)
    section.top_margin, section.bottom_margin = Cm(2.5), Cm(2.5)
    doc.add_paragraph("Text")

    settings = _parse(doc).settings

    assert (settings.pageSize, settings.orientation) == ("A4", "portrait")
    assert (settings.marginLeftCm, settings.marginRightCm, settings.marginTopCm) == (3.0, 2.0, 2.5)


def test_header_and_footer_keep_their_page_number_fields():
    doc = DocxDocument()
    section = doc.sections[0]
    section.header.paragraphs[0].text = "Quarterly report"
    footer = section.footer.paragraphs[0]
    footer.add_run("Page ")
    footer._p.append(parse_xml(f'<w:fldSimple {_NS} w:instr=" PAGE "><w:r><w:t>1</w:t></w:r></w:fldSimple>'))
    footer.add_run(" of ")
    for xml in (
        f'<w:r {_NS}><w:fldChar w:fldCharType="begin"/></w:r>',
        f'<w:r {_NS}><w:instrText xml:space="preserve"> NUMPAGES </w:instrText></w:r>',
        f'<w:r {_NS}><w:fldChar w:fldCharType="separate"/></w:r>',
        f"<w:r {_NS}><w:t>4</w:t></w:r>",
        f'<w:r {_NS}><w:fldChar w:fldCharType="end"/></w:r>',
    ):
        footer._p.append(parse_xml(xml))
    doc.add_paragraph("Body")

    settings = _parse(doc).settings

    assert settings.header == "Quarterly report"
    assert settings.footer == "Page {PAGE} of {NUMPAGES}"


def test_a_multi_column_layout_is_reported():
    doc = DocxDocument()
    doc.sections[0]._sectPr.find(qn("w:cols")).set(qn("w:num"), "2")
    doc.add_paragraph("Text")

    assert "The document is laid out in 2 columns; the app shows it in one." in _parse(doc).unsupportedFeatures


# -- pictures and tables ----------------------------------------------------------


def test_pictures_keep_their_size_and_alignment():
    doc = DocxDocument()
    section = doc.sections[0]
    section.page_width, section.left_margin, section.right_margin = Mm(210), Cm(2), Cm(2)  # 17 cm of text width
    doc.add_picture(io.BytesIO(_png()), width=Cm(8))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    document = _parse(doc)

    [image] = _elements(document, ElementType.IMAGE)
    css = _own_css(document, image)
    assert css["width"] == "47.1%"
    assert (css["margin-left"], css["margin-right"]) == ("auto", "auto")


def test_table_cell_shading_and_column_alignment_are_kept():
    doc = DocxDocument()
    table = doc.add_table(rows=2, cols=2)
    for cell in table.rows[0].cells:
        cell._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {_NS} w:val="clear" w:color="auto" w:fill="D9EAF7"/>'))
    for row in table.rows:
        row.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    [table_element] = _elements(_parse(doc), ElementType.TABLE)

    assert [cell.background for cell in table_element.table.rows[0].cells] == ["#D9EAF7", "#D9EAF7"]
    assert table_element.table.rows[1].cells[0].background is None
    assert table_element.table.alignments == ["center", None]


# -- text ------------------------------------------------------------------------


def test_links_are_kept_for_safe_addresses_and_plain_addresses_become_links():
    doc = DocxDocument()
    safe = doc.part.relate_to("https://example.com/docs", RT.HYPERLINK, is_external=True)
    unsafe = doc.part.relate_to("javascript:alert(1)", RT.HYPERLINK, is_external=True)
    _append_xml(
        doc,
        f'<w:p {_NS}><w:hyperlink r:id="{safe}"><w:r><w:t>the docs</w:t></w:r></w:hyperlink>'
        f'<w:r><w:t xml:space="preserve"> and </w:t></w:r>'
        f'<w:hyperlink r:id="{unsafe}"><w:r><w:t>a trap</w:t></w:r></w:hyperlink></w:p>',
    )
    doc.add_paragraph("Write to team@example.org or visit www.example.net.")

    first, second = _parse(doc).elements
    links = {run.text: _marks(run).get(MarkType.LINK) for element in (first, second) for run in element.inline}

    assert links["the docs"].href == "https://example.com/docs"
    assert not any(link for text, link in links.items() if "a trap" in text)  # javascript: -> plain text
    assert links["team@example.org"].href == "mailto:team@example.org"
    assert links["www.example.net"].href == "https://www.example.net"


def test_bookmarks_and_links_to_them_are_reported():
    doc = DocxDocument()
    _append_xml(
        doc,
        f'<w:p {_NS}><w:bookmarkStart w:id="0" w:name="Results"/><w:r><w:t>Results</w:t></w:r><w:bookmarkEnd w:id="0"/></w:p>',
    )
    _append_xml(doc, f'<w:p {_NS}><w:hyperlink w:anchor="Results"><w:r><w:t>see the results</w:t></w:r></w:hyperlink></w:p>')

    document = _parse(doc)

    assert [element.content for element in document.elements] == ["Results", "see the results"]
    assert not any(MarkType.LINK in _marks(run) for element in document.elements for run in element.inline)
    assert document.unsupportedFeatures == ["Bookmarks aren't kept, so links to places inside the document became plain text."]


def test_words_own_hidden_bookmarks_are_not_reported():
    doc = DocxDocument()
    _append_xml(doc, f'<w:p {_NS}><w:bookmarkStart w:id="0" w:name="_GoBack"/><w:r><w:t>Text</w:t></w:r><w:bookmarkEnd w:id="0"/></w:p>')

    assert _parse(doc).unsupportedFeatures == []


def test_tracked_changes_come_in_accepted():
    doc = DocxDocument()
    _append_xml(
        doc,
        f'<w:p {_NS}><w:r><w:t xml:space="preserve">Keep </w:t></w:r>'
        f'<w:ins w:id="1" w:author="A" w:date="2026-01-01T00:00:00Z"><w:r><w:t>added</w:t></w:r></w:ins>'
        f'<w:del w:id="2" w:author="A" w:date="2026-01-01T00:00:00Z"><w:r><w:delText>removed</w:delText></w:r></w:del></w:p>',
    )

    document = _parse(doc)

    assert document.elements[0].content == "Keep added"
    assert any("Tracked changes" in note for note in document.unsupportedFeatures)


def test_equations_become_linear_text_with_superscripts():
    doc = DocxDocument()
    _append_xml(
        doc,
        f'<w:p {_NS}><w:r><w:t xml:space="preserve">Area: </w:t></w:r><m:oMath><m:r><m:t>A=π</m:t></m:r>'
        f"<m:sSup><m:e><m:r><m:t>r</m:t></m:r></m:e><m:sup><m:r><m:t>2</m:t></m:r></m:sup></m:sSup></m:oMath></w:p>",
    )

    document = _parse(doc)

    [element] = document.elements
    assert element.content == "Area: A=πr2"
    assert set(_marks(element.inline[-1])) == {MarkType.ITALIC, MarkType.SUPERSCRIPT}
    assert "Equations were imported as plain text." in document.unsupportedFeatures


def test_footnotes_are_numbered_in_the_text_and_moved_to_the_end():
    doc = DocxDocument()
    notes_xml = (
        f'<w:footnotes {_NS}><w:footnote w:type="separator" w:id="-1"><w:p/></w:footnote>'
        f'<w:footnote w:id="7"><w:p><w:r><w:footnoteRef/></w:r><w:r><w:t xml:space="preserve"> The note text.</w:t></w:r></w:p></w:footnote></w:footnotes>'
    )
    part = Part(
        PackURI("/word/footnotes.xml"),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml",
        notes_xml.encode(),
        doc.part.package,
    )
    doc.part.relate_to(part, RT.FOOTNOTES)
    paragraph = doc.add_paragraph("A claim")
    paragraph._p.append(parse_xml(f'<w:r {_NS}><w:footnoteReference w:id="7"/></w:r>'))
    doc.add_paragraph("More text.")

    document = _parse(doc)

    claim = document.elements[0]
    assert claim.content == "A claim1"
    assert set(_marks(claim.inline[-1])) == {MarkType.SUPERSCRIPT}
    note = document.elements[-1]
    assert (note.type, note.content) == (ElementType.FOOTNOTE, "1  The note text.")
