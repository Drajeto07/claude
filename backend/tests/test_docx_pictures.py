"""Word pictures as the model keeps them (tracker DOCX-018, brief §27): their type and
name, alt text and title apart, the size they're drawn at, what is cropped away, how
they're turned and flipped, and where a floating one sits -- imported, written back
into Word, drawn cropped and turned in a PDF."""

import io
import zipfile
from pathlib import Path

from docx import Document as DocxDocument
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm
from PIL import Image as PILImage

from app.export.docx_export import build_docx
from app.export.package_check import package_problems
from app.export.pdf_export import _build_image, _build_table, build_pdf
from app.formatting.engine import recompute_styles
from app.models.document import ElementType, plain_text_from_inline
from app.parsers.docx import parse_docx


def _png(width: int = 200, height: int = 100) -> bytes:
    picture = PILImage.new("RGB", (width, height), "red")
    picture.paste(PILImage.new("RGB", (width // 2, height), "blue"), (width // 2, 0))
    buffer = io.BytesIO()
    picture.save(buffer, format="PNG")
    return buffer.getvalue()


def _save(word) -> bytes:
    buffer = io.BytesIO()
    word.save(buffer)
    return buffer.getvalue()


def _picture_file(*, floating: bool = False) -> bytes:
    word = DocxDocument()
    word.add_paragraph("Before the picture.")
    shape = word.add_paragraph().add_run().add_picture(io.BytesIO(_png()), width=Cm(4))
    inline = shape._inline
    inline.docPr.set("descr", "A red and blue flag")
    inline.docPr.set("title", "Flag")
    inline.docPr.set("name", "flag.png")
    pic = inline.graphic.graphicData.pic
    pic.blipFill.find(qn("a:blip")).addnext(parse_xml(f'<a:srcRect {nsdecls("a")} l="25000"/>'))
    xfrm = pic.spPr.find(qn("a:xfrm"))
    xfrm.set("rot", "5400000")  # 90 degrees
    xfrm.set("flipH", "1")
    if floating:
        anchor = parse_xml(
            f'<wp:anchor {nsdecls("wp")} distT="0" distB="0" distL="114300" distR="114300" simplePos="0" relativeHeight="251658240" '
            'behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1"><wp:simplePos x="0" y="0"/>'
            '<wp:positionH relativeFrom="margin"><wp:posOffset>720000</wp:posOffset></wp:positionH>'
            '<wp:positionV relativeFrom="paragraph"><wp:align>top</wp:align></wp:positionV></wp:anchor>'
        )
        anchor.append(inline.find(qn("wp:extent")))
        anchor.append(parse_xml(f'<wp:wrapSquare {nsdecls("wp")} wrapText="bothSides"/>'))
        for tag in ("wp:docPr", "wp:cNvGraphicFramePr", "a:graphic"):
            part = inline.find(qn(tag))
            if part is not None:
                anchor.append(part)
        inline.getparent().replace(inline, anchor)
    return _save(word)


GOLDEN_PICTURES = Path(__file__).parent / "fixtures" / "documents" / "17-pictures.docx"


def _picture_of(document):
    [picture] = [element.image for element in document.elements if element.type == ElementType.IMAGE]
    return picture


def test_a_pictures_size_crop_turn_and_name_are_kept():
    picture = _picture_of(parse_docx(_picture_file(), "flag.docx"))

    assert (picture.mime, picture.name, picture.alt, picture.title) == ("image/png", "flag.png", "A red and blue flag", "Flag")
    assert (picture.widthCm, picture.heightCm) == (4.0, 2.0)
    assert (picture.crop.left, picture.crop.right, picture.rotation, picture.flipHorizontal) == (0.25, 0, 90, True)
    assert picture.placement is None  # in line with the text


def test_a_floating_picture_keeps_where_it_floats_in_a_word_export():
    document = parse_docx(_picture_file(floating=True), "floating.docx")
    placement = _picture_of(document).placement

    exported = build_docx(document)
    again = _picture_of(parse_docx(exported, "again.docx"))

    assert (placement.wrap, placement.horizontalFrom, placement.horizontalCm, placement.verticalAlign, placement.distanceLeftCm) == (
        "square",
        "margin",
        2.0,
        "top",
        0.32,
    )
    assert package_problems(exported) == []
    assert again.placement == placement
    # Square wrapping: it floats at its side here and in a PDF (DOCX-018A), 2 cm in from the margin: the left.
    assert placement.side == "left" and "docx.image.floating_wrapped" in {item.feature for item in document.importReport.items}


def test_a_word_export_writes_crop_turn_and_flips_back():
    document = parse_docx(_picture_file(), "flag.docx")

    exported = build_docx(document)
    again = _picture_of(parse_docx(exported, "again.docx"))

    assert package_problems(exported) == []
    kept = ("mime", "name", "alt", "title", "widthCm", "heightCm", "crop", "rotation", "flipHorizontal", "flipVertical", "placement")
    assert again.model_dump(include=set(kept)) == _picture_of(document).model_dump(include=set(kept))


def test_an_unchanged_picture_keeps_its_exact_size_and_a_new_width_rule_still_wins():
    document = parse_docx(_picture_file(), "flag.docx")
    recompute_styles(document)  # as an upload does: its width is a rule now, a share of the text width
    [element] = [element for element in document.elements if element.type == ElementType.IMAGE]
    assert document.resolvedStyles[element.styleRef]["width"] == "26.2%"  # 4 cm of 15.24

    again = _picture_of(parse_docx(build_docx(document), "again.docx"))
    document.resolvedStyles[element.styleRef]["width"] = "50%"
    resized = _picture_of(parse_docx(build_docx(document), "resized.docx"))

    assert (again.widthCm, again.heightCm) == (4.0, 2.0)  # not 26.2% of the text width: 3.99 cm
    assert (resized.widthCm, resized.heightCm) == (7.62, 3.81)  # half the text width, in proportion


def test_a_turned_picture_takes_the_room_of_its_turned_outline():
    document = parse_docx(_picture_file(), "flag.docx")  # 4 by 2 cm, turned a quarter
    [element] = [element for element in document.elements if element.type == ElementType.IMAGE]

    with zipfile.ZipFile(io.BytesIO(build_docx(document))) as package:
        body = package.read("word/document.xml").decode("utf-8")
    drawn = _build_image(element, document, {}, width=1000)

    # Word: its size stays 4 by 2 cm; each side comes in 1 cm and its top and bottom go out 1 cm.
    assert '<wp:extent cx="1440000" cy="720000"/>' in body
    assert '<wp:effectExtent l="-360000" t="360000" r="-360000" b="360000"/>' in body
    # The PDF draws it 2 cm wide and 4 cm high (in points).
    assert (round(drawn.drawWidth / 28.3465, 2), round(drawn.drawHeight / 28.3465, 2)) == (2.0, 4.0)


def test_a_pdf_draws_the_picture_cropped_and_turned():
    from pypdf import PdfReader

    pdf = build_pdf(parse_docx(_picture_file(), "flag.docx"))

    [drawn] = PdfReader(io.BytesIO(pdf)).pages[0].images
    assert drawn.image.size == (100, 150)  # 200x100, a quarter cut off the left, then turned a quarter


def test_in_a_pdf_the_text_by_a_picture_in_a_cell_keeps_the_tables_look():
    document = parse_docx(GOLDEN_PICTURES.read_bytes(), GOLDEN_PICTURES.name)
    recompute_styles(document)
    [element] = [element for element in document.elements if element.type == ElementType.TABLE]

    table = _build_table(element, document, {}, width=450)
    (text, picture), beside = table._cellvalues[0]

    # The cell's own paragraph (it holds a picture too) is drawn as the one-line cell beside
    # it is -- in the table's font and line height -- so the picture below doesn't cover it.
    assert (text.style.fontName, text.style.fontSize, text.style.leading) == (beside.style.fontName, beside.style.fontSize, beside.style.leading)
    assert text.style.leading > text.style.fontSize
    assert (round(picture.drawWidth / 28.3465, 2), round(picture.drawHeight / 28.3465, 2)) == (3.0, 2.0)


# -- pictures in list items (DOCX-027) ------------------------------------------------------


def _list_with_pictures(word) -> None:
    """A numbered list: an item with a picture after its text, one that is only a
    picture, and one with only text."""
    for text, colour in (("Open the box", "purple"), ("", "teal"), ("Close it", None)):
        item = word.add_paragraph(text, style="List Number")
        if colour:
            picture = PILImage.new("RGB", (90, 45), colour)
            buffer = io.BytesIO()
            picture.save(buffer, format="PNG")
            item.add_run().add_picture(io.BytesIO(buffer.getvalue()), width=Cm(1.5))


def _items(document) -> list[tuple]:
    [listing] = [element for element in document.elements if element.type == ElementType.LIST]
    return [
        (plain_text_from_inline(item.inline), [(block.type, block.image.widthCm, block.image.heightCm) for block in item.blocks or []])
        for item in listing.listItems
    ]


def test_a_list_items_pictures_are_what_it_holds_and_come_back_in_its_paragraph():
    word = DocxDocument()
    _list_with_pictures(word)
    document = parse_docx(_save(word), "steps.docx")

    exported = build_docx(document)
    paragraphs = DocxDocument(io.BytesIO(exported)).paragraphs
    numbered = [paragraph._p for paragraph in paragraphs if paragraph._p.pPr is not None and paragraph._p.pPr.numPr is not None]

    assert _items(document) == [
        ("Open the box", [(ElementType.IMAGE, 1.5, 0.75)]),
        ("", [(ElementType.IMAGE, 1.5, 0.75)]),  # an item that is only a picture, no longer left out
        ("Close it", []),
    ]
    assert document.unsupportedFeatures == []
    assert package_problems(exported) == []
    assert [len(paragraph.xpath(".//w:drawing")) for paragraph in numbered] == [1, 1, 0]  # in the items' own paragraphs
    assert len(paragraphs) == 3  # and nowhere else
    assert _items(parse_docx(exported, "again.docx")) == _items(document)


def test_a_list_items_picture_in_a_table_cell_is_what_the_item_holds():
    word = DocxDocument()
    cell = word.add_table(rows=1, cols=1).cell(0, 0)
    cell.paragraphs[0].text = "In the cell:"
    item = cell.add_paragraph("A step", style="List Number")
    item.add_run().add_picture(io.BytesIO(_png(90, 45)), width=Cm(1.5))
    cell.add_paragraph("Another step", style="List Number")

    [table] = [element for element in parse_docx(_save(word), "cell.docx").elements if element.type == ElementType.TABLE]
    blocks = table.table.rows[0].cells[0].blocks

    assert [block.type for block in blocks] == [ElementType.PARAGRAPH, ElementType.LIST]  # one list, not split by the picture
    first, second = blocks[1].listItems
    assert [block.image.widthCm for block in first.blocks] == [1.5] and not second.blocks


def test_a_pdf_draws_a_list_items_pictures_under_its_text():
    from pypdf import PdfReader

    word = DocxDocument()
    _list_with_pictures(word)

    pdf = build_pdf(parse_docx(_save(word), "steps.docx"))

    assert len(PdfReader(io.BytesIO(pdf)).pages[0].images) == 2


import pytest  # noqa: E402

from app.models.document import Document, Element, ImageContent, ImagePlacement, InlineRun  # noqa: E402
from app.parsers.docx_pictures import float_side  # noqa: E402


@pytest.mark.parametrize(
    "placement, width, side",
    [
        ({"wrap": "square", "horizontalAlign": "right"}, 5, "right"),
        ({"wrap": "tight", "horizontalAlign": "outside"}, 5, "right"),
        ({"wrap": "square", "horizontalAlign": "left"}, 5, "left"),
        ({"wrap": "through", "horizontalFrom": "rightMargin", "horizontalCm": 0}, 3, "right"),
        ({"wrap": "square", "horizontalAlign": "center"}, 5, None),  # text on both sides: drawn in line
        ({"wrap": "square", "horizontalFrom": "column", "horizontalCm": 11}, 5, "right"),  # its middle at 13.5 of 17
        ({"wrap": "square", "horizontalFrom": "column", "horizontalCm": 3}, 5, "left"),
        ({"wrap": "square", "horizontalFrom": "page", "horizontalCm": 6}, 5, "left"),  # 4 into the column, its middle at 6.5 of 17
        ({"wrap": "square", "horizontalFrom": "page", "horizontalCm": 12}, 5, "right"),
        ({"wrap": "square", "horizontalFrom": "column", "horizontalCm": 6}, 7, "right"),  # a wide one: its middle at 9.5
        ({"wrap": "square", "horizontalFrom": "page", "horizontalCm": 9}, 0, "left"),  # 7 into the column: the margin counts
        ({"wrap": "topAndBottom", "horizontalAlign": "right"}, 5, None),
        ({"wrap": "behind", "horizontalAlign": "right"}, 5, None),
        ({"wrap": "inFront", "horizontalAlign": "left"}, 5, None),
    ],
)
def test_the_side_a_floating_picture_floats_to(placement, width, side):
    # DOCX-018A: worked out at import from its wrap and position, in a 17 cm column 2 cm in from the page's edge.
    assert float_side(placement, width, 2.0, 17.0) == side


def _floating_document(side: str) -> Document:
    text = "Text that wraps around the picture beside it, line after line, as Word lays it out. " * 6
    picture = Element(
        type=ElementType.IMAGE, content="", order=0,
        image=ImageContent(src=f"data:image/png;base64,{__import__('base64').b64encode(_png(400, 400)).decode()}", widthCm=5, heightCm=5,
                           placement=ImagePlacement(wrap="square", side=side)),
    )
    paragraph = Element(type=ElementType.PARAGRAPH, content=text, inline=[InlineRun(text=text)], order=1)
    document = Document(elements=[picture, paragraph])
    recompute_styles(document)
    return document


def _first_line_x(pdf: bytes) -> float:
    from pdfminer.high_level import extract_pages
    from pdfminer.layout import LTTextLine, LTTextBox

    lines = [line for page in extract_pages(io.BytesIO(pdf)) for box in page if isinstance(box, LTTextBox) for line in box if isinstance(line, LTTextLine)]
    return min(lines, key=lambda line: -line.y1).x0


def test_a_pdf_wraps_the_text_around_a_floating_picture_at_its_side():
    left_margin_pt = 2 * 72 / 2.54
    beside = _first_line_x(build_pdf(_floating_document("left")))
    assert beside > left_margin_pt + 5 * 72 / 2.54  # the text starts right of the 5 cm picture
    alone = _first_line_x(build_pdf(_floating_document("right")))
    assert alone - left_margin_pt < 10  # at the margin (reportlab's padding aside), the picture on the right
