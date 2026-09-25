"""Managing documents (корекции.docx §31, §32, §39, §64): the list with search,
sorting and pages; the kept versions, looked at, restored and compared (the
original against now is "before/after"); and deleting for good."""

import io

import pytest
from docx import Document as DocxDocument
from fastapi.testclient import TestClient

from app.main import app

client = TestClient(app, base_url="https://testserver")
_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_TEXT = "# Report\n\nA first paragraph with enough words to be real text.\n\n## Details\n\nMore text here."


@pytest.fixture(autouse=True)
def signed_in(api_db):
    client.cookies.clear()
    assert client.post("/api/v1/auth/register", json={"email": "manager@example.com", "password": "long enough password"}).status_code == 201
    yield
    client.cookies.clear()


def _create(title: str, text: str = _TEXT) -> dict:
    return client.post("/api/v1/documents", json={"text": text, "title": title}).json()


def _docx() -> bytes:
    doc = DocxDocument()
    doc.add_heading("Uploaded", level=1)
    doc.add_paragraph("A paragraph from a Word file.")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def test_the_list_shows_the_newest_first_with_what_each_document_is():
    first = _create("First")
    uploaded = client.post("/api/v1/documents/upload", files={"file": ("Plan.docx", _docx(), _DOCX)}, data={"title": "Plan"}).json()
    client.post(f"/api/v1/documents/{first['id']}/format", data={"templateId": "modern-report"})

    listing = client.get("/api/v1/documents").json()

    assert listing["total"] == 2
    assert [item["title"] for item in listing["items"]] == ["First", "Plan"]  # formatting made "First" the latest
    formatted, plan = listing["items"]
    assert (formatted["status"], formatted["templateId"], formatted["templateName"]) == ("formatted", "modern-report", "Модерен доклад")
    assert formatted["formattedAt"] is not None
    assert (plan["status"], plan["sourceType"], plan["originalFilename"], plan["templateName"]) == ("draft", "uploaded_docx", "Plan.docx", None)
    assert plan["id"] == uploaded["id"]


def test_the_list_searches_titles_sorts_and_pages():
    for title in ("Budget 2026", "budget notes", "Minutes", "100% done_ok"):
        _create(title)

    assert {item["title"] for item in client.get("/api/v1/documents", params={"q": "BUDGET"}).json()["items"]} == {"Budget 2026", "budget notes"}
    # % and _ are searched for as themselves, not as wildcards.
    assert [item["title"] for item in client.get("/api/v1/documents", params={"q": "0% d"}).json()["items"]] == ["100% done_ok"]
    assert client.get("/api/v1/documents", params={"q": "_"}).json()["total"] == 1

    by_title = client.get("/api/v1/documents", params={"sort": "title"}).json()["items"]
    assert [item["title"] for item in by_title] == ["100% done_ok", "Budget 2026", "budget notes", "Minutes"]

    page = client.get("/api/v1/documents", params={"sort": "title", "limit": 2, "offset": 2}).json()
    assert (page["total"], [item["title"] for item in page["items"]]) == (4, ["budget notes", "Minutes"])
    assert client.get("/api/v1/documents", params={"limit": 0}).status_code == 422


def test_nobody_sees_anyone_elses_documents_in_the_list():
    _create("Mine")
    client.cookies.clear()
    client.post("/api/v1/auth/register", json={"email": "someone@example.com", "password": "long enough password"})

    assert client.get("/api/v1/documents").json() == {"items": [], "total": 0}


def test_versions_say_what_each_change_did_and_which_is_current():
    document = _create("Versions")
    client.patch(f"/api/v1/documents/{document['id']}", json={"title": "Renamed"})
    client.post(f"/api/v1/documents/{document['id']}/format", data={"templateId": "academic-default"})

    versions = client.get(f"/api/v1/documents/{document['id']}/versions").json()

    assert [(v["number"], v["kind"], v["current"]) for v in versions] == [(3, "change", True), (2, "change", False), (1, "created", False)]
    assert [v["description"] for v in versions] == [
        "Applied formatting (template: academic-default)",
        "Renamed to “Renamed”",
        "Created from pasted text",
    ]
    assert versions[0]["author"] == "manager@example.com"

    client.post(f"/api/v1/documents/{document['id']}/undo")
    assert [v["current"] for v in client.get(f"/api/v1/documents/{document['id']}/versions").json()] == [False, True, False]


def test_an_old_version_can_be_looked_at_and_restored():
    document = _create("Original title")
    client.patch(f"/api/v1/documents/{document['id']}", json={"title": "Second title"})

    original = client.get(f"/api/v1/documents/{document['id']}/versions/1").json()
    assert original["metadata"]["title"] == "Original title"
    assert original["resolvedStyles"]  # worked out, ready to show
    assert client.get(f"/api/v1/documents/{document['id']}/versions/9").status_code == 404

    current = client.get(f"/api/v1/documents/{document['id']}").json()
    stale = client.post(f"/api/v1/documents/{document['id']}/versions/1/restore", headers={"If-Match": str(current["revision"] - 1)})
    assert stale.status_code == 412

    restored = client.post(f"/api/v1/documents/{document['id']}/versions/1/restore", headers={"If-Match": str(current["revision"])})
    assert restored.status_code == 200 and restored.json()["metadata"]["title"] == "Original title"
    top = client.get(f"/api/v1/documents/{document['id']}/versions").json()[0]
    assert (top["number"], top["description"], top["current"]) == (3, "Restored version 1", True)
    # A restore is a change like any other: it can be undone.
    assert client.post(f"/api/v1/documents/{document['id']}/undo").json()["metadata"]["title"] == "Second title"


def test_before_and_after_compares_the_original_with_now():
    document = _create("Compare")
    client.post(f"/api/v1/documents/{document['id']}/format", data={"templateId": "academic-default"})
    now = client.get(f"/api/v1/documents/{document['id']}").json()
    elements = now["elements"]
    elements[1]["content"] = "A rewritten first paragraph."
    elements[1]["inline"] = [{"text": "A rewritten first paragraph.", "marks": []}]
    client.put(f"/api/v1/documents/{document['id']}/content", json={"elements": elements})
    client.patch(f"/api/v1/documents/{document['id']}/settings", json={"property": "marginTop", "value": "3", "unit": "cm"})

    comparison = client.get(f"/api/v1/documents/{document['id']}/compare").json()

    assert comparison["fromVersion"] == 1 and comparison["toVersion"] == 4
    assert [(c["change"], c["afterText"]) for c in comparison["structure"]] == [("edited", "A rewritten first paragraph.")]
    body = {c["property"]: (c["before"], c["after"]) for c in comparison["styles"] if c["target"] == "Paragraph"}
    assert body["font-family"][1] == "Times New Roman"
    assert all(c["label"] == "Body text" for c in comparison["styles"] if c["target"] == "Paragraph")
    assert {"property": "marginTopCm", "before": "2.0", "after": "3.0"} in comparison["settings"]

    between = client.get(f"/api/v1/documents/{document['id']}/compare", params={"from": 2, "to": 3}).json()
    assert between["structure"] and not between["settings"]
    assert client.get(f"/api/v1/documents/{document['id']}/compare", params={"from": 7}).status_code == 404


def test_deleting_is_for_good_and_takes_the_history_along():
    document = _create("Doomed")
    client.patch(f"/api/v1/documents/{document['id']}", json={"title": "Still doomed"})

    assert client.delete(f"/api/v1/documents/{document['id']}").status_code == 204

    assert client.get(f"/api/v1/documents/{document['id']}").status_code == 404
    assert client.get(f"/api/v1/documents/{document['id']}/versions").status_code == 404
    assert client.get("/api/v1/documents").json()["total"] == 0
    assert client.delete(f"/api/v1/documents/{document['id']}").status_code == 404
