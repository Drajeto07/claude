"""How much a file may hold before it is refused (tracker SEC-013): what costs time to read
is bounded, so no upload ties a worker up for minutes. Each limit is tested on an ordinary
small file with the limit lowered for the test -- the refusal and its message are what
matter, not the size."""

import io

import pytest
from docx import Document as DocxDocument
from fastapi.testclient import TestClient
from reportlab.pdfgen import canvas

from app.main import app
from app.parsers import docx as docx_module
from app.parsers import pdf as pdf_module
from app.parsers.docx import DocxParseError, import_docx
from app.parsers.pdf import TOO_DENSE, TOO_SLOW, PdfParseError, read_pdf

pytestmark = pytest.mark.security

client = TestClient(app, base_url="https://testserver")
_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@pytest.fixture
def signed_in(api_db):
    client.cookies.clear()
    assert client.post("/api/v1/auth/register", json={"email": "limits@example.com", "password": "long enough password"}).status_code == 201
    yield
    client.cookies.clear()


def _word(paragraphs: int = 0, table: tuple[int, int] | None = None) -> bytes:
    document = DocxDocument()
    for number in range(paragraphs):
        document.add_paragraph(f"Paragraph {number + 1}.")
    if table:
        document.add_table(rows=table[0], cols=table[1])
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _pdf(pages: int = 1) -> bytes:
    buffer = io.BytesIO()
    document = canvas.Canvas(buffer, invariant=1)
    for number in range(pages):
        document.drawString(72, 720, f"Page {number + 1} of a short PDF.")
        document.showPage()
    document.save()
    return buffer.getvalue()


def test_a_word_file_with_more_paragraphs_than_can_be_read_in_seconds_is_refused(monkeypatch):
    monkeypatch.setattr(docx_module, "MAX_PARAGRAPHS", 3)
    assert len(import_docx(_word(3), "three.docx").document.elements) == 3
    with pytest.raises(DocxParseError, match="more than 3 paragraphs"):
        import_docx(_word(4), "four.docx")


def test_a_word_file_with_more_table_cells_than_can_be_read_in_seconds_is_refused(monkeypatch):
    monkeypatch.setattr(docx_module, "MAX_TABLE_CELLS", 4)
    import_docx(_word(table=(2, 2)), "four.docx")
    with pytest.raises(DocxParseError, match="more than 4 cells"):
        import_docx(_word(table=(3, 2)), "six.docx")


def test_a_pdf_page_with_more_on_it_than_can_be_read_is_refused(monkeypatch):
    assert "Page 1" in read_pdf(_pdf()).text
    monkeypatch.setattr(pdf_module, "MAX_PDF_PAGE_CONTENT", 10)
    with pytest.raises(PdfParseError) as refused:
        read_pdf(_pdf())
    assert str(refused.value) == TOO_DENSE


def test_a_pdf_is_read_for_so_long_and_no_longer(monkeypatch):
    monkeypatch.setattr(pdf_module, "MAX_PDF_SECONDS", -1.0)  # the time already up at the first page
    with pytest.raises(PdfParseError) as refused:
        read_pdf(_pdf(2))
    assert str(refused.value) == TOO_SLOW


def test_over_the_api_a_file_past_a_limit_is_a_400_with_the_reason(signed_in, monkeypatch):
    monkeypatch.setattr(docx_module, "MAX_PARAGRAPHS", 3)
    monkeypatch.setattr(pdf_module, "MAX_PDF_PAGE_CONTENT", 10)

    word = client.post("/api/v1/documents/upload", files={"file": ("long.docx", _word(4), _DOCX)})
    pdf = client.post("/api/v1/documents/upload", files={"file": ("dense.pdf", _pdf(), "application/pdf")})

    assert (word.status_code, word.json()["code"]) == (400, "invalid_file") and "more than 3 paragraphs" in word.json()["message"]
    assert (pdf.status_code, pdf.json()["code"], pdf.json()["message"]) == (400, "invalid_file", TOO_DENSE)
