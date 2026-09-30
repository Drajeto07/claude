"""Word's paragraph formatting beyond spacing and indents (tracker DOCX-014): the
right indent, a background colour (shading), keep with next, keep lines
together, widow and orphan control, contextual spacing and the writing
direction are read from the paragraph and its style, kept as formatting rules,
drawn with CSS's own properties, written back into Word in the schema's order,
and followed by the PDF."""

import io
import zipfile

import pytest
from docx import Document as DocxDocument
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from reportlab.lib.enums import TA_RIGHT

from app.export.docx_export import _P_PR_ORDER, build_docx
from app.export.package_check import package_problems
from app.export.pdf_export import _paragraph_style, build_pdf
from app.formatting.engine import recompute_styles
from app.formatting.values import InvalidRuleValue, clean_rule_value
from app.models.document import FormattingProperty
from app.parsers.docx import parse_docx

_W = nsdecls("w")
_PROPERTIES = (
    '<w:keepNext/><w:keepLines/><w:widowControl w:val="0"/>'
    '<w:shd w:val="clear" w:color="auto" w:fill="FFF2CC"/><w:bidi/>'
    '<w:ind w:right="567"/><w:contextualSpacing/>'
)


def _word_file() -> bytes:
    word = DocxDocument()
    word.styles["Normal"].paragraph_format.widow_control = True  # the paragraph below turns it off
    word.styles["Heading 1"].paragraph_format.keep_with_next = True
    word.add_heading("A heading kept with what follows", level=1)
    formatted = word.add_paragraph("Shaded, indented on the right, right to left, kept together.")
    formatted._p.get_or_add_pPr().extend(parse_xml(f"<w:pPr {_W}>{_PROPERTIES}</w:pPr>"))
    word.add_paragraph("An ordinary paragraph.")
    buffer = io.BytesIO()
    word.save(buffer)
    return buffer.getvalue()


def _document():
    document = parse_docx(_word_file(), "paragraphs.docx")
    recompute_styles(document)
    return document


def test_paragraph_formatting_is_read_from_the_paragraph_and_its_style():
    document = _document()
    heading, formatted, ordinary = document.elements

    assert document.resolvedStyles[heading.styleRef]["break-after"] == "avoid"  # from Heading 1
    css = document.resolvedStyles[formatted.styleRef]
    assert css["margin-right"] == "1cm"
    assert css["background-color"] == "#FFF2CC"
    assert css["break-after"] == "avoid" and css["break-inside"] == "avoid"
    assert (css["widows"], css["orphans"]) == ("1", "1")
    assert css["--contextual-spacing"] == "true" and css["direction"] == "rtl"
    assert "background-color" not in document.resolvedStyles[ordinary.styleRef]


def test_a_word_export_writes_them_back_in_the_schemas_order():
    document = _document()

    exported = build_docx(document)

    assert package_problems(exported) == []
    with zipfile.ZipFile(io.BytesIO(exported)) as package:
        body = parse_xml(package.read("word/document.xml"))
    properties = [p.find(qn("w:pPr")) for p in body.iter(qn("w:p")) if "kept together" in "".join(t.text or "" for t in p.iter(qn("w:t")))][0]
    tags = [child.tag.split("}")[1] for child in properties]
    assert {"keepNext", "keepLines", "widowControl", "shd", "bidi", "ind", "contextualSpacing"} <= set(tags)
    assert tags == sorted(tags, key=_P_PR_ORDER.index)
    again = parse_docx(exported, "again.docx")
    recompute_styles(again)
    before = {key: value for key, value in document.resolvedStyles[document.elements[1].styleRef].items() if key != "font-family"}
    after = {key: value for key, value in again.resolvedStyles[again.elements[1].styleRef].items() if key != "font-family"}
    assert {key: after.get(key) for key in ("margin-right", "background-color", "break-after", "break-inside", "widows", "direction", "--contextual-spacing")} == {
        key: before.get(key) for key in ("margin-right", "background-color", "break-after", "break-inside", "widows", "direction", "--contextual-spacing")
    }


def test_a_pdf_follows_them():
    style = _paragraph_style(
        "p",
        {"margin-right": "1cm", "background-color": "#FFF2CC", "break-after": "avoid", "widows": "1", "direction": "rtl"},
    )

    assert round(style.rightIndent) == 28 and style.keepWithNext == 1
    assert (style.allowWidows, style.allowOrphans) == (1, 1)
    assert style.backColor is not None and style.alignment == TA_RIGHT
    assert build_pdf(_document()).startswith(b"%PDF")


@pytest.mark.parametrize(
    ("prop", "value", "unit"),
    [
        (FormattingProperty.KEEP_WITH_NEXT, "maybe", None),
        (FormattingProperty.DIRECTION, "upwards", None),
        (FormattingProperty.SHADING, "url(x)", None),
        (FormattingProperty.INDENT_RIGHT, "30", "cm"),
    ],
)
def test_a_value_that_cant_be_drawn_is_refused(prop, value, unit):
    with pytest.raises(InvalidRuleValue):
        clean_rule_value(prop, value, unit)


def test_values_are_written_one_way():
    assert clean_rule_value(FormattingProperty.KEEP_LINES_TOGETHER, "Yes", None) == ("true", None)
    assert clean_rule_value(FormattingProperty.DIRECTION, "RTL", None) == ("rtl", None)
    assert clean_rule_value(FormattingProperty.INDENT_RIGHT, "1.50", "cm") == ("1.5", "cm")


def test_contextual_spacing_closes_up_paragraphs_of_the_same_kind_in_a_pdf():
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import Paragraph

    from app.export.pdf_export import _close_up

    first = Paragraph("One", ParagraphStyle("one", spaceBefore=6, spaceAfter=12))
    second = Paragraph("Two", ParagraphStyle("two", spaceBefore=6, spaceAfter=12))

    _close_up([first], [second])

    assert (first.style.spaceBefore, first.style.spaceAfter) == (6, 0)
    assert (second.style.spaceBefore, second.style.spaceAfter) == (0, 12)


# -- borders and tab stops (DOCX-014 part 2) ----------------------------------------------

_BOX = '<w:pBdr><w:top w:val="single" w:sz="4" w:space="1" w:color="auto"/><w:left w:val="single" w:sz="4" w:space="4" w:color="auto"/><w:bottom w:val="single" w:sz="4" w:space="1" w:color="auto"/><w:right w:val="single" w:sz="4" w:space="4" w:color="auto"/></w:pBdr>'
_TABS = '<w:tabs><w:tab w:val="right" w:leader="dot" w:pos="9072"/><w:tab w:val="left" w:pos="1134"/><w:tab w:val="clear" w:pos="720"/></w:tabs>'


def _bordered_file() -> bytes:
    word = DocxDocument()
    heading = word.styles["Heading 1"].element.get_or_add_pPr()
    heading.append(parse_xml(f'<w:pBdr {_W}><w:bottom w:val="double" w:sz="12" w:space="1" w:color="0000FF"/></w:pBdr>'))
    word.add_heading("Ruled under", level=1)
    plain = word.add_heading("Not ruled", level=1)
    plain._p.get_or_add_pPr().append(parse_xml(f'<w:pBdr {_W}><w:bottom w:val="nil"/></w:pBdr>'))
    boxed = word.add_paragraph("Boxed")
    boxed._p.get_or_add_pPr().append(parse_xml(_BOX.replace("<w:pBdr>", f"<w:pBdr {_W}>")))
    tabbed = word.add_paragraph("Chapter one\t7")
    tabbed._p.get_or_add_pPr().append(parse_xml(_TABS.replace("<w:tabs>", f"<w:tabs {_W}>")))
    left = word.add_paragraph("A bar on the left")
    left._p.get_or_add_pPr().append(parse_xml(f'<w:pBdr {_W}><w:left w:val="single" w:sz="24" w:space="4" w:color="C00000"/></w:pBdr>'))
    buffer = io.BytesIO()
    word.save(buffer)
    return buffer.getvalue()


def _bordered():
    document = parse_docx(_bordered_file(), "borders.docx")
    recompute_styles(document)
    return document


def test_borders_and_tab_stops_are_read_from_the_paragraph_and_its_style():
    document = _bordered()
    ruled, plain, boxed, tabbed, left = (document.resolvedStyles[element.styleRef] for element in document.elements)

    assert ruled["border-bottom"] == "double 1.5pt #0000FF"  # from Heading 1
    assert plain["border-bottom"] == "none"  # the paragraph turns its style's off
    assert {boxed[f"border-{side}"] for side in ("top", "bottom", "left", "right")} == {"solid 0.5pt #000000"}
    assert tabbed["--tab-stops"] == "right 16cm dot; left 2cm"  # "clear" isn't a stop
    assert left["border-left"] == "solid 3pt #C00000"


def test_a_word_export_writes_borders_and_tab_stops_back():
    document = _bordered()

    exported = build_docx(document)

    assert package_problems(exported) == []
    again = parse_docx(exported, "again.docx")
    recompute_styles(again)
    for before, after in zip(document.elements, again.elements):
        keys = ("border-top", "border-bottom", "border-left", "border-right", "--tab-stops")
        assert {key: again.resolvedStyles[after.styleRef].get(key) for key in keys} == {
            key: document.resolvedStyles[before.styleRef].get(key) for key in keys
        }, before.content
    with zipfile.ZipFile(io.BytesIO(exported)) as package:
        body = parse_xml(package.read("word/document.xml"))
    for properties in body.iter(qn("w:pPr")):
        tags = [child.tag.split("}")[1] for child in properties]
        assert tags == sorted(tags, key=_P_PR_ORDER.index)


def test_a_pdf_draws_a_box_or_a_line_and_names_the_rest():
    from reportlab.platypus import HRFlowable

    from app.export.pdf_export import _border_lines
    from app.fidelity.report import ReportBuilder

    boxed = _paragraph_style("box", {f"border-{side}": "solid 0.5pt #000000" for side in ("top", "bottom", "left", "right")})
    above, below = _border_lines({"border-bottom": "double 1.5pt #0000FF"})
    report = ReportBuilder()
    build_pdf(_bordered(), report=report)

    assert boxed.borderWidth == 0.5 and boxed.borderColor is not None
    assert above == [] and isinstance(below[0], HRFlowable)
    assert {"export.pdf.tab_stops", "export.pdf.paragraph_borders"} <= {item.feature for item in report.items()}


def test_tabs_in_the_text_are_named_as_not_shown():
    from app.fidelity.docx_detect import detect_docx_features

    item = next(item for item in detect_docx_features(_bordered_file()) if item.feature == "docx.tab_stops")

    assert item.policy == "detected_not_editable" and "Chapter one" in (item.sourceState or "")


@pytest.mark.parametrize(
    ("prop", "value"),
    [
        (FormattingProperty.BORDER_TOP, "wavy 1pt red"),
        (FormattingProperty.BORDER_TOP, "solid 20pt red"),
        (FormattingProperty.BORDER_LEFT, "solid 1pt url(x)"),
        (FormattingProperty.TAB_STOPS, "diagonal 2cm"),
        (FormattingProperty.TAB_STOPS, "; ".join(["left 1cm"] * 31)),
    ],
)
def test_a_border_or_tab_stop_that_cant_be_drawn_is_refused(prop, value):
    with pytest.raises(InvalidRuleValue):
        clean_rule_value(prop, value, None)


def test_templates_check_borders_and_tab_stops_the_same_way():
    from pydantic import ValidationError

    from app.formatting.style_system import TextStyle

    assert TextStyle(borderBottom="solid  1pt #000", tabStops="right 16.00cm dot").model_dump(exclude_none=True) == {
        "borderBottom": "solid 1pt #000",
        "tabStops": "right 16cm dot",
    }
    with pytest.raises(ValidationError):
        TextStyle(borderTop="rainbow")
