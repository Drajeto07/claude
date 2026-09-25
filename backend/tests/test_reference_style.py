"""Format by Example (корекции.docx §17): the style read from a reference Word
document -- its Word styles, the look most of its text really has, headings
recognised by their look or by the AI -- and what is deliberately left out."""

import asyncio
import io
import re
from typing import TypeVar

from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Mm, Pt, RGBColor
from PIL import Image as PILImage
from pydantic import BaseModel

from app.ai.base import AIProvider, AIStructuredOutputError
from app.ai.semantic_labeling import AIParagraphLabel, AIParagraphLabels
from app.formatting.reference_style import ReferenceStyle
from app.services.reference_service import extract_from_docx, suggested_name
from tests.fakes import FakeAIProvider

T = TypeVar("T", bound=BaseModel)

_BODY = "Това е обикновен абзац с достатъчно текст, за да прилича на истинско съдържание на документа. " * 3


def _extract(doc: DocxDocument, provider: AIProvider | None = None) -> ReferenceStyle:
    buffer = io.BytesIO()
    doc.save(buffer)
    return asyncio.run(extract_from_docx(buffer.getvalue(), "reference.docx", provider))


def _set_font(style, name: str) -> None:
    """A named font; python-docx's template gives headings a theme font, which Word would use instead."""
    style.font.name = name
    fonts = style.element.rPr.rFonts
    for attribute in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        fonts.attrib.pop(qn(attribute), None)


def _line(doc: DocxDocument, text: str, *, size: float | None = None, bold: bool = False):
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = bold or None
    if size:
        run.font.size = Pt(size)
    return paragraph


def _messy_report() -> DocxDocument:
    """Headings made by hand: bold, larger Normal paragraphs, no heading styles."""
    doc = DocxDocument()
    doc.styles["Normal"].font.size = Pt(11)
    _line(doc, "1. Въведение", size=16, bold=True)
    doc.add_paragraph(_BODY)
    _line(doc, "1.1 Цел на работата", size=13, bold=True)
    doc.add_paragraph(_BODY)
    _line(doc, "2. Резултати", size=16, bold=True)
    doc.add_paragraph(_BODY)
    doc.add_paragraph(_BODY)
    return doc


# -- Word styles -------------------------------------------------------------------


def test_a_reference_using_word_styles_gives_those_styles():
    doc = DocxDocument()
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Times New Roman", Pt(12)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.5
    heading = doc.styles["Heading 1"]
    _set_font(heading, "Georgia")
    heading.font.size, heading.font.bold, heading.font.color.rgb = Pt(18), True, RGBColor(0x1F, 0x3A, 0x5F)
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = Cm(2.5)
    doc.add_heading("Въведение", level=1)
    doc.add_paragraph(_BODY)
    doc.add_paragraph(_BODY)

    reference = _extract(doc)

    style = reference.style_system
    body = style.paragraph
    assert (body.fontFamily, body.fontSizePt, body.alignment, body.lineSpacing) == ("Times New Roman", 12, "justify", 1.5)
    h1 = style.headings.h1
    assert (h1.fontFamily, h1.fontSizePt, h1.bold, h1.color) == ("Georgia", 18, True, "#1F3A5F")
    assert (style.page.size, style.page.orientation, style.page.marginLeftCm, style.page.marginTopCm) == ("A4", "portrait", 2.5, 2.5)
    assert reference.headings_from == "styles"
    assert reference.heading_counts == {1: 1}
    assert reference.counts["paragraphs"] == 2


def test_the_look_most_of_the_text_has_wins_over_the_style():
    """Formatting applied by hand counts: the body is Times New Roman 12 even
    though its style says Calibri 11, and one odd paragraph doesn't change that."""
    doc = DocxDocument()
    doc.styles["Normal"].font.name, doc.styles["Normal"].font.size = "Calibri", Pt(11)
    for _ in range(3):
        run = doc.add_paragraph().add_run(_BODY)
        run.font.name, run.font.size = "Times New Roman", Pt(12)
    odd = doc.add_paragraph().add_run("A short odd one.")
    odd.font.name = "Arial"

    body = _extract(doc).style_system.paragraph

    assert (body.fontFamily, body.fontSizePt) == ("Times New Roman", 12)


# -- headings without heading styles ------------------------------------------------


def test_headings_are_recognised_by_their_look_when_no_heading_styles_are_used():
    reference = _extract(_messy_report())

    headings = reference.style_system.headings
    assert (headings.h1.fontSizePt, headings.h1.bold) == (16, True)
    assert (headings.h2.fontSizePt, headings.h2.bold) == (13, True)
    # Deeper levels the reference doesn't use look like its deepest heading.
    assert headings.h3 == headings.h2 and headings.h6 == headings.h2
    assert reference.style_system.paragraph.fontSizePt == 11
    assert reference.headings_from == "look"
    assert reference.heading_counts == {1: 2, 2: 1}
    assert any("headings were recognised by their look" in note for note in reference.notes)


def test_short_lines_that_look_like_the_body_text_are_not_headings():
    doc = DocxDocument()
    doc.add_paragraph("Кратък ред")
    doc.add_paragraph(_BODY)
    doc.add_paragraph("Още един кратък ред")
    doc.add_paragraph(_BODY)

    reference = _extract(doc)

    assert reference.headings_from == "none"
    assert reference.heading_counts == {}


class _LabellingAI(AIProvider):
    """Answers like the real model would: one label per paragraph listed in the
    prompt, chosen by its text."""

    def __init__(self, choose) -> None:
        self._choose = choose
        self.calls = 0

    def provider_name(self) -> str:
        return "labelling-fake"

    async def complete(self, prompt: str, *, max_tokens: int = 256) -> str:
        return ""

    async def complete_structured(self, prompt: str, *, response_model: type[T], max_tokens: int = 8192) -> T:
        self.calls += 1
        labels = [
            AIParagraphLabel(id=element_id, **self._choose(text))
            for element_id, text in re.findall(r'id=(\S+) \|[^|]*\| text: "([^"]*)"', prompt)
        ]
        return AIParagraphLabels(labels=labels)


def test_the_ai_can_say_which_paragraphs_are_headings():
    """The AI decides that "1.1 Цел на работата" is body text after all; the look
    of the headings it did pick still comes from the document."""
    ai = _LabellingAI(lambda text: {"role": "heading", "level": 1} if text[0].isdigit() and "." in text[:2] and "Цел" not in text else {"role": "body"})

    reference = _extract(_messy_report(), ai)

    assert ai.calls == 1
    assert reference.headings_from == "ai"
    assert reference.heading_counts == {1: 2}
    assert (reference.style_system.headings.h1.fontSizePt, reference.style_system.headings.h1.bold) == (16, True)
    assert any("the AI picked out its headings" in note for note in reference.notes)


def test_an_ai_answer_naming_unknown_paragraphs_is_not_trusted():
    ai = FakeAIProvider([AIParagraphLabels(labels=[AIParagraphLabel(id="not-a-paragraph", role="heading", level=1)])])

    reference = _extract(_messy_report(), ai)

    assert reference.headings_from == "look"
    assert reference.heading_counts == {1: 2, 2: 1}


def test_an_unavailable_ai_falls_back_to_the_look():
    reference = _extract(_messy_report(), FakeAIProvider([AIStructuredOutputError("no output")]))

    assert reference.headings_from == "look"


def test_the_ai_is_not_asked_when_the_reference_uses_heading_styles():
    doc = DocxDocument()
    doc.add_heading("Title", level=1)
    doc.add_paragraph(_BODY)
    ai = FakeAIProvider([])  # any call would fail the test

    assert _extract(doc, ai).headings_from == "styles"
    assert ai.calls == 0


# -- page, header and footer, pictures ---------------------------------------------


def test_header_text_is_left_out_but_a_page_number_footer_is_kept():
    doc = DocxDocument()
    section = doc.sections[0]
    section.header.paragraphs[0].text = "ACME Corp confidential"
    footer = section.footer.paragraphs[0]
    footer.add_run("Page ")
    footer._p.append(parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr=" PAGE "><w:r><w:t>1</w:t></w:r></w:fldSimple>'))
    doc.add_paragraph(_BODY)

    reference = _extract(doc)

    assert reference.style_system.header.text is None
    assert reference.style_system.footer.text == "Page {PAGE}"
    assert any("ACME Corp confidential" in note and "wasn't copied" in note for note in reference.notes)


def _png() -> bytes:
    buffer = io.BytesIO()
    PILImage.new("RGB", (40, 20), "blue").save(buffer, format="PNG")
    return buffer.getvalue()


def test_the_picture_placement_most_pictures_share_becomes_the_image_rule():
    doc = DocxDocument()
    section = doc.sections[0]
    section.page_width, section.left_margin, section.right_margin = Mm(210), Cm(2), Cm(2)  # 17 cm of text width
    for width, alignment in ((Cm(8), WD_ALIGN_PARAGRAPH.CENTER), (Cm(8), WD_ALIGN_PARAGRAPH.CENTER), (Cm(4), WD_ALIGN_PARAGRAPH.RIGHT)):
        doc.add_picture(io.BytesIO(_png()), width=width)
        doc.paragraphs[-1].alignment = alignment
        doc.add_paragraph(_BODY)

    images = _extract(doc).style_system.images

    assert (images.alignment, images.widthPercent) == ("center", 47.1)


def test_pictures_of_every_size_leave_the_width_to_each_picture():
    doc = DocxDocument()
    for width in (Cm(4), Cm(8), Cm(12)):
        doc.add_picture(io.BytesIO(_png()), width=width)

    assert _extract(doc).style_system.images.widthPercent is None


def test_the_suggested_template_name_comes_from_the_file_name():
    assert suggested_name("Дипломна  работа.docx") == "Дипломна работа style"
    assert suggested_name("   .docx") == "Reference style"
    taken = {"Report style", "Report style 2"}
    assert suggested_name("Report.docx", taken) == "Report style 3"
