"""The Word export of a big table grows with its cells (tracker PERF-001). python-docx's
table.cell() and cell.merge() each walk the whole table, which made a 500 x 8 table take
minutes: the time is checked here as a ratio between sizes (never an absolute number, so
a slow machine passes), and the table that comes out is checked against one made cell by
cell with python-docx itself."""

import copy
import io
import random
import time

from docx import Document as DocxDocument
from docx.oxml.ns import qn
from docx.oxml.table import CT_Tc
from docx.table import Table, _Cell

from app.export.docx_export import _merge_cells, build_docx
from app.export.package_check import package_problems
from app.models.document import Document, Element, ElementType, InlineRun, TableCell, TableContent, TableRow
from app.parsers.docx import parse_docx
from scripts.benchmark_table_export import table_file


def _best_export_time(rows: int, runs: int = 2) -> float:
    data = table_file(rows, 8)
    document = parse_docx(data, "table.docx")
    best = float("inf")
    for _ in range(runs):
        started = time.perf_counter()
        build_docx(document)
        best = min(best, time.perf_counter() - started)
    return best


def test_exporting_a_table_grows_with_its_cells_not_their_square():
    small, large = _best_export_time(50), _best_export_time(200)
    # 4 times the cells: 4 times the time when linear, 16 when it grows with the square. 9 leaves room for a noisy machine.
    assert large / small < 9, f"50 rows took {small:.2f} s, 200 rows {large:.2f} s"


def _merged_table(rows: int, columns: int) -> Document:
    """Every row's first two cells joined, and a block spanning two rows and two columns in every fourth row."""
    table_rows = []
    for row in range(rows):
        cells, column = [], 0
        while column < columns:
            wide = column == 0 or (row % 4 == 0 and column == 3)
            tall = row % 4 == 0 and column == 3 and row + 1 < rows
            if row % 4 == 1 and column == 3 and row > 0:
                column += 2  # under the block above
                continue
            cells.append(TableCell(inline=[InlineRun(text=f"r{row}c{column}")], colspan=2 if wide else 1, rowspan=2 if tall else 1))
            column += 2 if wide else 1
        table_rows.append(TableRow(cells=cells))
    return Document(title="merged", elements=[Element(type=ElementType.TABLE, content="", order=0, table=TableContent(rows=table_rows))])


def test_a_big_table_with_merged_cells_grows_with_its_cells_too():
    started = time.perf_counter()
    build_docx(_merged_table(100, 8))
    small = time.perf_counter() - started
    started = time.perf_counter()
    build_docx(_merged_table(800, 8))
    large = time.perf_counter() - started
    assert large / small < 11, f"100 rows took {small:.2f} s, 800 rows {large:.2f} s"


def test_the_export_never_uses_the_lookups_that_walk_the_whole_table(monkeypatch):
    """Whatever the clock says on this machine: table.cell(), cell.merge() and the row and cell-below lookups behind
    them are the ones whose cost grows with the table, so none of them is used."""
    def refuse(*args, **kwargs):
        raise AssertionError("a lookup that walks the whole table was used")

    for owner, name in ((Table, "cell"), (_Cell, "merge"), (CT_Tc, "merge")):
        monkeypatch.setattr(owner, name, refuse)
    for owner, name in ((Table, "_cells"), (CT_Tc, "_tr_idx"), (CT_Tc, "_tc_below")):  # properties in python-docx
        monkeypatch.setattr(owner, name, property(refuse))
    build_docx(_merged_table(30, 8))
    build_docx(parse_docx(table_file(30, 8), "table.docx"))


def test_a_big_tables_package_is_sound():
    data = table_file(250, 8)
    exported = build_docx(parse_docx(data, "table.docx"), source=data)
    assert package_problems(exported) == []
    assert package_problems(build_docx(_merged_table(250, 8))) == []


def _skeleton(tbl) -> list[list[tuple[str, str | None, str]]]:
    """Each row's cells as (columns spanned, vertical merge, text)."""
    rows = []
    for tr in tbl.findall(qn("w:tr")):
        cells = []
        for tc in tr.findall(qn("w:tc")):
            span = tc.find(f"{qn('w:tcPr')}/{qn('w:gridSpan')}")
            merge = tc.find(f"{qn('w:tcPr')}/{qn('w:vMerge')}")
            cells.append((
                span.get(qn("w:val")) if span is not None else "1",
                None if merge is None else merge.get(qn("w:val")) or "continue",
                "".join(t.text or "" for t in tc.iter(qn("w:t"))),
            ))
        rows.append(cells)
    return rows


def _by_python_docx(document: Document, rows: int, columns: int):
    """The same table made cell by cell as python-docx does it: table.cell() and cell.merge()."""
    word = DocxDocument()
    table = word.add_table(rows=rows, cols=columns)
    occupied: set[tuple[int, int]] = set()
    for row_index, row in enumerate(document.elements[0].table.rows):
        column = 0
        for cell in row.cells:
            while (row_index, column) in occupied:
                column += 1
            for dr in range(cell.rowspan):
                for dc in range(cell.colspan):
                    occupied.add((row_index + dr, column + dc))
            target = table.cell(row_index, column)
            if cell.rowspan > 1 or cell.colspan > 1:
                target = target.merge(table.cell(row_index + cell.rowspan - 1, column + cell.colspan - 1))
            target.text = cell.inline[0].text
            column += cell.colspan
    return table._tbl


def test_a_small_table_comes_out_as_one_made_cell_by_cell():
    document = _merged_table(9, 8)
    [exported_tbl] = DocxDocument(io.BytesIO(build_docx(document))).element.body.findall(qn("w:tbl"))
    assert _skeleton(exported_tbl) == _skeleton(_by_python_docx(document, 9, 8))


def _shape(element) -> tuple:
    """An element's tags, attributes and text, whatever prefixes the namespaces have."""
    return (element.tag, sorted(element.attrib.items()), element.text, [_shape(child) for child in element])


def test_merging_cells_at_hand_is_what_python_docx_merge_does():
    generator = random.Random(28)
    for _ in range(40):
        rows, columns = generator.randint(1, 7), generator.randint(1, 7)
        table = DocxDocument().add_table(rows=rows, cols=columns)
        by_hand = copy.deepcopy(table._tbl)
        grid = [tr.tc_lst for tr in by_hand.tr_lst]  # as the export has it: taken before anything is merged
        free = {(row, column) for row in range(rows) for column in range(columns)}
        for _ in range(6):  # blocks that don't touch each other
            if not free:
                break
            row, column = generator.choice(sorted(free))
            height, width = generator.randint(1, rows - row), generator.randint(1, columns - column)
            block = {(r, c) for r in range(row, row + height) for c in range(column, column + width)}
            if not block <= free:
                continue
            free -= block
            table.cell(row, column).merge(table.cell(row + height - 1, column + width - 1))
            _merge_cells(grid, row, column, row + height - 1, column + width - 1)
        assert _shape(table._tbl) == _shape(by_hand)
