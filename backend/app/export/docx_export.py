import io
import re
from collections.abc import Mapping

from docx import Document as DocxDocument
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.run import Run

from app.export.images import resolve_image_bytes
from app.formatting.colors import NAMED_COLORS
from app.formatting.render_spec import page_size_mm
from app.models.document import (
    Document,
    DocumentSettings,
    Element,
    ElementType,
    InlineRun,
    Mark,
    MarkType,
    TableContent,
    target_for_element,
)

_ALIGNMENT_MAP = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}

_STYLE_FOR_TYPE = {
    ElementType.QUOTE: "Quote",
    ElementType.CAPTION: "Caption",
    ElementType.FOOTNOTE: "Footnote Text",
}

# The Word styles each kind of block is written with (корекции.docx §24: native
# Word semantics, not inline CSS on every run). Each is set explicitly from that
# kind's resolved style, so changing "Heading 1" in Word restyles every heading,
# and nothing of python-docx's own template (blue Calibri Light headings, 10 pt
# after every paragraph) leaks into the file. Styles the template lacks are added.
_WORD_STYLES: dict[str, tuple[str, ...]] = {
    "Paragraph": ("Normal",),
    **{f"Heading {level}": (f"Heading {level}",) for level in range(1, 7)},
    "Quote": ("Quote",),
    "Caption": ("Caption",),
    "Footnote": ("Footnote Text",),
    "List": ("List Bullet", "List Number", "List Paragraph"),
    "Table": ("Table Text",),
    "CodeBlock": ("Code",),
}
_LIST_STYLES = ("List Bullet", "List Number", "List Paragraph")
# The children of w:pPr in schema order, for inserting ones python-docx has no API for.
_P_PR_ORDER = (
    "pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr suppressLineNumbers pBdr shd tabs "
    "suppressAutoHyphens kinsoku wordWrap overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi adjustRightInd "
    "snapToGrid spacing ind contextualSpacing mirrorIndents suppressOverlap jc textDirection textAlignment "
    "textboxTightWrap outlineLvl divId cnfStyle rPr sectPr pPrChange"
).split()
_AFTER_SHADING = tuple(qn(f"w:{name}") for name in _P_PR_ORDER[_P_PR_ORDER.index("shd") + 1 :])
_AFTER_CONTEXTUAL_SPACING = tuple(qn(f"w:{name}") for name in _P_PR_ORDER[_P_PR_ORDER.index("contextualSpacing") + 1 :])

# Background colours Word can show as a real highlight; any other becomes run shading.
_WORD_HIGHLIGHTS = {
    "FFFF00": "yellow",
    "00FF00": "green",
    "00FFFF": "cyan",
    "FF00FF": "magenta",
    "0000FF": "blue",
    "FF0000": "red",
    "000080": "darkBlue",
    "008080": "darkCyan",
    "008000": "darkGreen",
    "800080": "darkMagenta",
    "800000": "darkRed",
    "808000": "darkYellow",
    "808080": "darkGray",
    "C0C0C0": "lightGray",
    "000000": "black",
}

_PAGE_FIELD = re.compile(r"(\{PAGE\}|\{NUMPAGES\})")


def build_docx(
    document: Document,
    *,
    assets: Mapping[str, bytes] | None = None,
    include_headers: bool = True,
    include_page_numbers: bool = True,
    include_page_breaks: bool = True,
) -> bytes:
    """Real, editable .docx (spec §7.17) built from the same resolved styles
    the editor already renders -- no separate style computation. Independent
    of pdf_export.py's reportlab-based builder (LibreOffice isn't available
    on this machine to convert one into the other -- see the Phase 6 plan);
    both read the same document.resolvedStyles, so there's nothing to keep
    in sync beyond that shared source of truth.

    The three include_* flags are export-time-only overrides (spec's export
    options screen) -- they never touch the persisted document.settings, so
    exporting once without page numbers doesn't turn them off for next time.
    All default True, matching this function's behavior before these flags
    existed. With page numbers left out, header/footer text built around a
    page-number field ({PAGE}, {NUMPAGES}) is left out too."""
    docx_document = DocxDocument()
    zoom = docx_document.settings.element.find(qn("w:zoom"))
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")  # required by the schema; python-docx's template leaves it out
    _define_styles(docx_document, document)
    _apply_page_setup(docx_document, document, include_headers=include_headers, include_page_numbers=include_page_numbers)
    for element in document.elements:
        if element.type == ElementType.PAGE_BREAK and not include_page_breaks:
            continue
        _add_element(docx_document, element, document, assets or {})

    buffer = io.BytesIO()
    docx_document.save(buffer)
    return buffer.getvalue()


def _page_dimensions_mm(settings: DocumentSettings) -> tuple[float, float]:
    return page_size_mm(settings.pageSize, settings.orientation)


def _word_style(docx_document: DocxDocument, name: str, kind: WD_STYLE_TYPE = WD_STYLE_TYPE.PARAGRAPH):
    try:
        return docx_document.styles[name]
    except KeyError:
        style = docx_document.styles.add_style(name, kind)
        if kind == WD_STYLE_TYPE.PARAGRAPH:
            style.base_style = docx_document.styles["Normal"]
            style.quick_style = True
        return style


def _set_style(style, css: dict[str, str]) -> None:
    """Everything the renderers show for this kind of block, set on its Word
    style; whatever the resolved style leaves out gets the neutral value the
    editor and the PDF use (left, regular, single spacing, no space)."""
    font = style.font
    family = (css.get("font-family") or "").strip('"')
    if family:
        font.name = family
        r_fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        r_fonts.set(qn("w:cs"), family)
        for theme in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            r_fonts.attrib.pop(qn(theme), None)  # a theme font wins over the name in Word
    if css.get("font-size", "").endswith("pt"):
        font.size = Pt(_parse_pt(css["font-size"]))
    font.bold = css.get("font-weight") == "bold"
    font.italic = css.get("font-style") == "italic"
    font.underline = css.get("text-decoration") == "underline"
    color = _parse_color(css.get("color") or "")
    r_pr = style.element.get_or_add_rPr()
    for old in r_pr.findall(qn("w:color")):
        r_pr.remove(old)  # the template's theme colours included
    if color is not None:
        font.color.rgb = color

    paragraph_format = style.paragraph_format
    paragraph_format.alignment = _ALIGNMENT_MAP.get(css.get("text-align", ""), WD_ALIGN_PARAGRAPH.LEFT)
    line_height = css.get("--line-spacing") or css.get("line-height", "")  # Word's own value
    try:
        paragraph_format.line_spacing = Pt(_parse_pt(line_height)) if line_height.endswith("pt") else float(line_height or 1)
    except ValueError:
        paragraph_format.line_spacing = 1.0
    paragraph_format.space_before = Pt(_parse_pt(css.get("margin-top", "0pt")))
    paragraph_format.space_after = Pt(_parse_pt(css.get("margin-bottom", "0pt")))
    margin_left, text_indent = css.get("margin-left", ""), css.get("text-indent", "")
    paragraph_format.left_indent = Cm(_parse_cm(margin_left)) if margin_left.endswith("cm") else Cm(0)
    paragraph_format.first_line_indent = Cm(_parse_cm(text_indent)) if text_indent.endswith("cm") else Cm(0)


def _contextual_spacing(style) -> None:
    """No space between paragraphs of this style, only after the last one --
    how list items sit together in the editor and the PDF."""
    p_pr = style.element.get_or_add_pPr()
    if p_pr.find(qn("w:contextualSpacing")) is not None:
        return
    element = OxmlElement("w:contextualSpacing")
    successor = next((child for child in p_pr if child.tag in _AFTER_CONTEXTUAL_SPACING), None)
    if successor is not None:
        successor.addprevious(element)
    else:
        p_pr.append(element)


def _define_styles(docx_document: DocxDocument, document: Document) -> None:
    for target, names in _WORD_STYLES.items():
        css = document.resolvedStyles.get(target, {})
        if target == "Table":  # cell text: the space after belongs to the table, not to each cell
            css = {key: value for key, value in css.items() if key not in ("margin-top", "margin-bottom")}
        for name in names:
            style = _word_style(docx_document, name)
            _set_style(style, css)
            if name in _LIST_STYLES:
                _contextual_spacing(style)
    code = docx_document.styles["Code"]
    p_pr = code.element.get_or_add_pPr()
    if p_pr.find(qn("w:shd")) is None:
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="F0F0F0"/>')
        successor = next((child for child in p_pr if child.tag in _AFTER_SHADING), None)
        if successor is not None:
            successor.addprevious(shading)
        else:
            p_pr.append(shading)
    hyperlink = _word_style(docx_document, "Hyperlink", WD_STYLE_TYPE.CHARACTER)
    hyperlink.font.color.rgb = RGBColor(0x05, 0x63, 0xC1)
    hyperlink.font.underline = True


def _own_css(element: Element, document: Document) -> dict[str, str]:
    """What this element sets differently from its kind's style: the only
    formatting it carries itself, the rest comes from its Word style."""
    if not element.styleRef or element.styleRef == target_for_element(element):
        return {}
    kind = document.resolvedStyles.get(target_for_element(element), {})
    return {key: value for key, value in _resolved_css(element, document).items() if kind.get(key) != value}


def _apply_page_setup(docx_document: DocxDocument, document: Document, *, include_headers: bool, include_page_numbers: bool) -> None:
    settings = document.settings
    section = docx_document.sections[0]
    width_mm, height_mm = _page_dimensions_mm(settings)
    section.orientation = WD_ORIENT.LANDSCAPE if settings.orientation == "landscape" else WD_ORIENT.PORTRAIT
    section.page_width = Cm(width_mm / 10)
    section.page_height = Cm(height_mm / 10)
    section.top_margin = Cm(settings.marginTopCm)
    section.bottom_margin = Cm(settings.marginBottomCm)
    section.left_margin = Cm(settings.marginLeftCm)
    section.right_margin = Cm(settings.marginRightCm)

    header = _page_text(settings.header, include_headers, include_page_numbers)
    footer = _page_text(settings.footer, include_headers, include_page_numbers)
    if header:
        _write_page_text(section.header.paragraphs[0], header)
    if footer:
        _write_page_text(section.footer.paragraphs[0], footer)
    if include_page_numbers and settings.showPageNumbers:
        page_number_paragraph = section.footer.add_paragraph() if footer else section.footer.paragraphs[0]
        page_number_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _append_field(page_number_paragraph, "PAGE")


def _page_text(text: str | None, include_headers: bool, include_page_numbers: bool) -> str | None:
    if not include_headers or not text:
        return None
    if not include_page_numbers and _PAGE_FIELD.search(text):
        return None
    return text


def _write_page_text(paragraph, text: str) -> None:
    """Header/footer text, with {PAGE}/{NUMPAGES} written as the Word fields."""
    for part in _PAGE_FIELD.split(text):
        if part == "{PAGE}":
            _append_field(paragraph, "PAGE")
        elif part == "{NUMPAGES}":
            _append_field(paragraph, "NUMPAGES")
        elif part:
            paragraph.add_run(part)


def _append_field(paragraph, instruction: str) -> None:
    """python-docx has no high-level API for field codes -- raw XML, as for
    hyperlinks below."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def _content_width_cm(document: Document) -> float:
    width_mm, _ = _page_dimensions_mm(document.settings)
    return width_mm / 10 - document.settings.marginLeftCm - document.settings.marginRightCm


def _resolved_css(element: Element, document: Document) -> dict[str, str]:
    if not element.styleRef:
        return {}
    return document.resolvedStyles.get(element.styleRef, {})


def _parse_pt(value: str) -> float:
    try:
        return float(value.replace("pt", "").strip())
    except ValueError:
        return 0.0


def _parse_cm(value: str) -> float:
    try:
        return float(value.replace("cm", "").strip())
    except ValueError:
        return 0.0


def _hex6(value: str | None) -> str | None:
    """#rgb/#rrggbb/named colour as RRGGBB, or None."""
    if not value:
        return None
    value = value.strip().lower()
    if value.startswith("#"):
        hex_value = value.lstrip("#")
        if len(hex_value) == 3:
            hex_value = "".join(ch * 2 for ch in hex_value)
        return hex_value.upper() if re.fullmatch(r"[0-9a-f]{6}", hex_value) else None
    return NAMED_COLORS.get(value)


def _parse_color(value: str) -> RGBColor | None:
    hex_value = _hex6(value)
    return RGBColor.from_string(hex_value) if hex_value else None


def _apply_run_css(run, css: dict[str, str]) -> None:
    font_family = css.get("font-family")
    if font_family:
        run.font.name = font_family.strip('"')
    font_size = css.get("font-size")
    if font_size:
        run.font.size = Pt(_parse_pt(font_size))
    color = css.get("color")
    if color:
        rgb = _parse_color(color)
        if rgb is not None:
            run.font.color.rgb = rgb
    # An element's own "normal"/"none" undoes what its style sets.
    if css.get("font-weight"):
        run.font.bold = css["font-weight"] == "bold"
    if css.get("font-style"):
        run.font.italic = css["font-style"] == "italic"
    if css.get("text-decoration"):
        run.font.underline = css["text-decoration"] == "underline"


def _apply_text_style(run, mark: Mark) -> None:
    """A textStyle mark on this run: its values win over the element's CSS."""
    if mark.fontFamily:
        run.font.name = mark.fontFamily
    if mark.fontSizePt:
        run.font.size = Pt(mark.fontSizePt)
    if (rgb := _parse_color(mark.color or "")) is not None:
        run.font.color.rgb = rgb
    background = _hex6(mark.backgroundColor)
    if background:
        r_pr = run._r.get_or_add_rPr()
        if background in _WORD_HIGHLIGHTS:
            highlight = OxmlElement("w:highlight")
            highlight.set(qn("w:val"), _WORD_HIGHLIGHTS[background])
            r_pr.append(highlight)
        else:
            shading = OxmlElement("w:shd")
            shading.set(qn("w:val"), "clear")
            shading.set(qn("w:color"), "auto")
            shading.set(qn("w:fill"), background)
            r_pr.append(shading)


def _apply_paragraph_css(paragraph, css: dict[str, str]) -> None:
    alignment = _ALIGNMENT_MAP.get(css.get("text-align", ""))
    if alignment is not None:
        paragraph.alignment = alignment
    line_height = css.get("--line-spacing") or css.get("line-height")  # Word's own value
    if line_height:
        try:
            # "12pt" is an exact line height; a plain number is a multiple.
            paragraph.paragraph_format.line_spacing = Pt(_parse_pt(line_height)) if line_height.endswith("pt") else float(line_height)
        except ValueError:
            pass
    margin_top = css.get("margin-top")
    if margin_top:
        paragraph.paragraph_format.space_before = Pt(_parse_pt(margin_top))
    margin_bottom = css.get("margin-bottom")
    if margin_bottom:
        paragraph.paragraph_format.space_after = Pt(_parse_pt(margin_bottom))
    margin_left = css.get("margin-left")
    if margin_left and margin_left.endswith("cm"):
        paragraph.paragraph_format.left_indent = Cm(_parse_cm(margin_left))
    text_indent = css.get("text-indent")
    if text_indent:
        paragraph.paragraph_format.first_line_indent = Cm(_parse_cm(text_indent))


def _add_hyperlink_run(paragraph, text: str, url: str) -> Run:
    """python-docx has no high-level hyperlink API -- same raw-XML pattern
    already used elsewhere in this file for numbering and the page-number
    field. Returns a real Run wrapper around the new <w:r> inside the
    hyperlink, so the caller applies bold/italic/underline/etc. through the
    normal .font API exactly as for any other run."""
    r_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    run_element = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    style = OxmlElement("w:rStyle")
    style.set(qn("w:val"), "Hyperlink")
    run_properties.append(style)
    run_element.append(run_properties)
    hyperlink.append(run_element)
    paragraph._p.append(hyperlink)

    run = Run(run_element, paragraph)
    run.text = text
    return run


def _add_inline_run(paragraph, inline_run: InlineRun, css: dict[str, str]) -> Run:
    marks = {mark.type: mark for mark in inline_run.marks}
    link = marks.get(MarkType.LINK) if marks.get(MarkType.LINK) and marks[MarkType.LINK].href else None
    # run.text turns "\n" into a line break and "\t" into a tab.
    run = _add_hyperlink_run(paragraph, inline_run.text, link.href) if link else paragraph.add_run(inline_run.text)
    _apply_run_css(run, css)
    if MarkType.BOLD in marks:
        run.font.bold = True
    if MarkType.ITALIC in marks:
        run.font.italic = True
    if MarkType.UNDERLINE in marks:
        run.font.underline = True
    if MarkType.STRIKE in marks:
        run.font.strike = True
    if MarkType.SUPERSCRIPT in marks:
        run.font.superscript = True
    elif MarkType.SUBSCRIPT in marks:
        run.font.subscript = True
    if MarkType.CODE in marks:
        run.font.name = "Courier New"
    if MarkType.TEXT_STYLE in marks:
        _apply_text_style(run, marks[MarkType.TEXT_STYLE])
    return run


def _add_inline_runs(paragraph, inline_runs: list[InlineRun], css: dict[str, str]) -> None:
    for inline_run in inline_runs:
        _add_inline_run(paragraph, inline_run, css)


# -- what the import kept for export (корекции.docx §11) -----------------------------

_KEPT_KINDS = ("equation", "field", "bookmark", "link", "comment")
_BOOKMARK_NAME = re.compile(r"[\w.\-]{1,40}")  # Word's own limit is 40 characters
_MATH_ROOTS = (qn("m:oMath"), qn("m:oMathPara"))
_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _short_text(value, limit: int) -> bool:
    return isinstance(value, str) and len(value) <= limit and not _CONTROL.search(value)


def _valid_fragment(fragment) -> bool:
    """preservedAttributes travels through the browser on every autosave, so
    nothing in it is trusted: anything malformed is left out of the export."""
    if not isinstance(fragment, dict) or fragment.get("kind") not in _KEPT_KINDS:
        return False
    start, end = fragment.get("start"), fragment.get("end")
    if not isinstance(start, int) or not isinstance(end, int) or not 0 <= start <= end:
        return False
    if not _short_text(fragment.get("text", ""), 10_000):
        return False
    kind = fragment["kind"]
    if kind == "equation":
        xml = fragment.get("xml")
        if not isinstance(xml, str) or len(xml) > 200_000:
            return False
        try:
            return parse_xml(xml).tag in _MATH_ROOTS
        except Exception:  # noqa: BLE001 -- not XML at all
            return False
    if kind == "field":
        return _short_text(fragment.get("instr"), 2_000) and bool(fragment["instr"].strip())
    if kind == "bookmark":
        return isinstance(fragment.get("name"), str) and bool(_BOOKMARK_NAME.fullmatch(fragment["name"]))
    if kind == "link":
        return isinstance(fragment.get("anchor"), str) and bool(_BOOKMARK_NAME.fullmatch(fragment["anchor"]))
    return all(_short_text(fragment.get(name, ""), limit) for name, limit in (("author", 255), ("initials", 16), ("comment", 20_000)))


def _find_near(text: str, wanted: str, near: int) -> int | None:
    """Where `wanted` is in `text`: at `near` if it's still there, else the closest occurrence."""
    if text.startswith(wanted, near):
        return near
    best, found = None, text.find(wanted)
    while found != -1:
        if best is None or abs(found - near) < abs(best - near):
            best = found
        found = text.find(wanted, found + 1)
    return best


def _place_fragments(text: str, fragments: list) -> list[tuple[int, int, dict]]:
    """Each kept fragment's span in the element's current text. A fragment whose
    text was edited away is dropped (the text itself stays, as edited); so is
    anything that would cut into an equation or cross another internal link."""
    placed: list[tuple[int, int, dict]] = []
    for fragment in fragments:
        if not _valid_fragment(fragment):
            continue
        wanted = fragment.get("text") or ""
        if wanted:
            start = _find_near(text, wanted, fragment["start"])
            if start is not None:
                placed.append((start, start + len(wanted), fragment))
        elif fragment["kind"] != "equation" and fragment["kind"] != "link":
            position = min(fragment["start"], len(text))
            placed.append((position, position, fragment))
    kept: list[tuple[int, int, dict]] = []
    for start, end, fragment in sorted(placed, key=lambda item: (item[0], -item[1])):
        exclusive = [(s, e) for s, e, f in kept if f["kind"] == "equation" or (f["kind"] == "link" and fragment["kind"] == "link")]
        if any(s < start < e or s < end < e or (start < s < end and fragment["kind"] == "equation") for s, e in exclusive):
            continue
        kept.append((start, end, fragment))
    return kept


def _cut(inline_runs: list[InlineRun], cuts: list[int]) -> list[tuple[int, int, InlineRun]]:
    """The runs split at every cut position, each piece with where it starts and ends."""
    pieces: list[tuple[int, int, InlineRun]] = []
    position = 0
    for run in inline_runs:
        start, end = position, position + len(run.text)
        edges = [start, *(cut for cut in cuts if start < cut < end), end]
        pieces.extend(
            (a, b, InlineRun(text=run.text[a - start : b - start], marks=run.marks)) for a, b in zip(edges, edges[1:]) if b > a
        )
        position = end
    return pieces


def _field_char(kind: str) -> OxmlElement:
    element = OxmlElement("w:fldChar")
    element.set(qn("w:fldCharType"), kind)
    return element


def _add_inline_runs_keeping(paragraph, inline_runs: list[InlineRun], css: dict[str, str], fragments: list) -> None:
    """The runs, with the equations, fields, bookmarks, internal links and
    comments the import kept put back around (or, for an equation, in place of)
    their text."""
    text = "".join(run.text for run in inline_runs)
    placed = _place_fragments(text, fragments)
    if not placed:
        _add_inline_runs(paragraph, inline_runs, css)
        return
    cuts = sorted({0, len(text), *(start for start, _, _ in placed), *(end for _, end, _ in placed)})
    pieces = {start: (start, end, run) for start, end, run in _cut(inline_runs, cuts)}
    body = paragraph.part.element.body
    state = {"container": paragraph._p, "skip_until": -1}
    bookmark_ids: dict[int, str] = {}
    made: list[tuple[int, int, Run]] = []

    def into_container(run: Run) -> None:
        if state["container"] is not paragraph._p and run._r.getparent() is paragraph._p:
            state["container"].append(run._r)

    def open_(index: int, start: int, end: int, fragment: dict) -> None:
        kind = fragment["kind"]
        if kind == "equation":
            state["container"].append(parse_xml(fragment["xml"]))
            state["skip_until"] = end
        elif kind == "field":
            run = paragraph.add_run()
            into_container(run)
            instr = OxmlElement("w:instrText")
            instr.set(qn("xml:space"), "preserve")
            instr.text = fragment["instr"]
            run._r.append(_field_char("begin"))
            run._r.append(instr)
            run._r.append(_field_char("separate"))
        elif kind == "bookmark":
            taken = [int(mark.get(qn("w:id"))) for mark in body.iter(qn("w:bookmarkStart")) if (mark.get(qn("w:id")) or "").isdigit()]
            bookmark_ids[index] = str(max(taken, default=-1) + 1)
            mark = OxmlElement("w:bookmarkStart")
            mark.set(qn("w:id"), bookmark_ids[index])
            mark.set(qn("w:name"), fragment["name"])
            state["container"].append(mark)
        elif kind == "link":
            hyperlink = OxmlElement("w:hyperlink")
            hyperlink.set(qn("w:anchor"), fragment["anchor"])
            paragraph._p.append(hyperlink)
            state["container"] = hyperlink

    def close(index: int, fragment: dict) -> None:
        kind = fragment["kind"]
        if kind == "field":
            run = paragraph.add_run()
            into_container(run)
            run._r.append(_field_char("end"))
        elif kind == "bookmark":
            mark = OxmlElement("w:bookmarkEnd")
            mark.set(qn("w:id"), bookmark_ids[index])
            state["container"].append(mark)
        elif kind == "link":
            state["container"] = paragraph._p

    indexed = list(enumerate(placed))
    for position in cuts:
        for index, (start, end, fragment) in sorted(indexed, key=lambda item: -item[1][0]):
            if end == position and start < position:  # the one opened last closes first
                close(index, fragment)
        for index, (start, end, fragment) in indexed:
            if start == end == position:
                open_(index, start, end, fragment)
                close(index, fragment)
        for index, (start, end, fragment) in sorted(indexed, key=lambda item: -item[1][1]):
            if start == position and end > position:  # the one closing last opens first
                open_(index, start, end, fragment)
        piece = pieces.get(position)
        if piece is not None and not piece[0] < state["skip_until"]:
            run = _add_inline_run(paragraph, piece[2], css)
            into_container(run)
            made.append((piece[0], piece[1], run))

    for start, end, fragment in placed:
        if fragment["kind"] != "comment":
            continue
        runs = [run for piece_start, piece_end, run in made if start <= piece_start and piece_end <= end]
        runs = runs or [run for piece_start, piece_end, run in made if piece_start <= start < piece_end] or [run for *_, run in made[-1:]]
        if not runs:
            continue
        comment = paragraph.part.document.add_comment(
            runs=[runs[0], runs[-1]],
            text=fragment.get("comment", ""),
            author=fragment.get("author", ""),
            initials=fragment.get("initials", ""),
        )
        if isinstance(fragment.get("date"), str) and re.fullmatch(r"\d{4}-\d\d-\d\dT[\d:.]+(?:Z|[+-]\d\d:\d\d)?", fragment["date"]):
            comment._comment_elm.set(qn("w:date"), fragment["date"])


def _add_runs(paragraph, element: Element, document: Document) -> None:
    css = _own_css(element, document)
    inline_runs = element.inline or ([InlineRun(text=element.content)] if element.content else [])
    kept = (element.preservedAttributes or {}).get("ooxml")
    if isinstance(kept, list) and kept:
        _add_inline_runs_keeping(paragraph, inline_runs, css, kept)
    else:
        _add_inline_runs(paragraph, inline_runs, css)
    _apply_paragraph_css(paragraph, css)


def _add_heading(docx_document: DocxDocument, element: Element, document: Document) -> None:
    heading = docx_document.add_heading(level=min(max(element.level or 1, 1), 9))
    _add_runs(heading, element, document)


def _add_paragraph(docx_document: DocxDocument, element: Element, document: Document) -> None:
    paragraph = docx_document.add_paragraph(style=_STYLE_FOR_TYPE.get(element.type))
    _add_runs(paragraph, element, document)


def _add_checkbox(paragraph, checked: bool) -> None:
    """A real Word checkbox (a content control): clickable in Word 2010 and
    later, a ☒/☐ character anywhere else -- and imported back as a checklist."""
    glyph = "☒" if checked else "☐"
    paragraph._p.append(
        parse_xml(
            f"<w:sdt {nsdecls('w', 'w14')}><w:sdtPr><w14:checkbox>"
            f'<w14:checked w14:val="{1 if checked else 0}"/>'
            '<w14:checkedState w14:val="2612" w14:font="MS Gothic"/>'
            '<w14:uncheckedState w14:val="2610" w14:font="MS Gothic"/>'
            "</w14:checkbox></w:sdtPr><w:sdtContent><w:r><w:rPr>"
            '<w:rFonts w:ascii="MS Gothic" w:eastAsia="MS Gothic" w:hAnsi="MS Gothic" w:hint="eastAsia"/>'
            f"</w:rPr><w:t>{glyph}</w:t></w:r></w:sdtContent></w:sdt>"
        )
    )
    paragraph.add_run(" ")


_LIST_LEVELS = 9  # Word's maximum
_LEVEL_INDENT_TWIPS = 357  # 0.63 cm per level
_BULLETS = ("•", "◦", "▪")
_NUMBER_FORMATS = ("decimal", "lowerLetter", "lowerRoman")


def _abstract_numbering(numbering, kind: str) -> str:
    """The id of this document's multi-level list definition for `kind`
    ("bullet", "number", or "none" for checklists), added the first time it's
    needed. python-docx's template only has single-level lists."""
    name = f"SmartDoc {kind}"
    for abstract in numbering.findall(qn("w:abstractNum")):
        name_element = abstract.find(qn("w:name"))
        if name_element is not None and name_element.get(qn("w:val")) == name:
            return abstract.get(qn("w:abstractNumId"))
    abstract_id = str(1 + max((int(a.get(qn("w:abstractNumId"))) for a in numbering.findall(qn("w:abstractNum"))), default=-1))
    levels = []
    for ilvl in range(_LIST_LEVELS):
        left = _LEVEL_INDENT_TWIPS * (ilvl + 1)
        if kind == "none":  # just the indent; the checkbox leads the text
            number = '<w:numFmt w:val="none"/><w:suff w:val="nothing"/><w:lvlText w:val=""/>'
            indent = f'<w:ind w:left="{left}" w:hanging="0"/>'
        else:
            fmt, text = ("bullet", _BULLETS[ilvl % 3]) if kind == "bullet" else (_NUMBER_FORMATS[ilvl % 3], f"%{ilvl + 1}.")
            number = f'<w:numFmt w:val="{fmt}"/><w:lvlText w:val="{text}"/>'
            indent = f'<w:ind w:left="{left}" w:hanging="{_LEVEL_INDENT_TWIPS}"/>'
        levels.append(
            f'<w:lvl w:ilvl="{ilvl}"><w:start w:val="1"/>{number}<w:lvlJc w:val="left"/><w:pPr>{indent}</w:pPr></w:lvl>'
        )
    abstract = parse_xml(
        f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{abstract_id}">'
        f'<w:multiLevelType w:val="hybridMultilevel"/><w:name w:val="{name}"/>{"".join(levels)}</w:abstractNum>'
    )
    first_num = numbering.find(qn("w:num"))  # the schema puts every abstractNum before the nums
    if first_num is not None:
        first_num.addprevious(abstract)
    else:
        numbering.append(abstract)
    return abstract_id


def _new_list_numbering(docx_document: DocxDocument, kind: str) -> int:
    """A numbering instance of its own for one list: levels nest for real (Tab
    and Shift+Tab work in Word, and a re-import keeps them), and a numbered
    list starts again at 1 instead of continuing the previous one."""
    numbering = docx_document.part.numbering_part.element
    abstract_id = _abstract_numbering(numbering, kind)
    num_id = 1 + max((int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))), default=0)
    restarts = "".join(
        f'<w:lvlOverride w:ilvl="{ilvl}"><w:startOverride w:val="1"/></w:lvlOverride>' for ilvl in range(_LIST_LEVELS)
    )
    numbering.append(
        parse_xml(f'<w:num {nsdecls("w")} w:numId="{num_id}"><w:abstractNumId w:val="{abstract_id}"/>{restarts}</w:num>')
    )
    return num_id


def _set_numbering(paragraph, num_id: int, level: int) -> None:
    num_pr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
    num_pr.get_or_add_ilvl().val = min(max(level, 0), _LIST_LEVELS - 1)
    num_pr.get_or_add_numId().val = num_id


def _add_list(docx_document: DocxDocument, element: Element, document: Document) -> None:
    full_css = _resolved_css(element, document)
    own = {key: value for key, value in _own_css(element, document).items() if key != "margin-left"}
    checklist = any(item.checked is not None for item in element.listItems or [])
    style_name = "List Paragraph" if checklist else "List Number" if element.ordered else "List Bullet"
    num_id = _new_list_numbering(docx_document, "none" if checklist else "number" if element.ordered else "bullet")
    # The numbering's own indent beats a style's, so a list indent goes on each item.
    base_indent = _parse_cm(full_css.get("margin-left", "")) if full_css.get("margin-left", "").endswith("cm") else 0.0
    for item in element.listItems or []:
        paragraph = docx_document.add_paragraph(style=style_name)
        _set_numbering(paragraph, num_id, item.level)
        if base_indent:  # otherwise the list level sets the indent
            paragraph.paragraph_format.left_indent = Cm(base_indent + 0.63 * (item.level + 1))
        if item.checked is not None:
            _add_checkbox(paragraph, item.checked)
        _add_inline_runs(paragraph, item.inline, own)
        _apply_paragraph_css(paragraph, own)


def _grid_positions(table_content: TableContent) -> tuple[list[tuple[int, int, object]], int]:
    """(row, column, cell) for every cell, with spans taken into account."""
    occupied: set[tuple[int, int]] = set()
    placed: list[tuple[int, int, object]] = []
    width = 0
    for row_index, row in enumerate(table_content.rows):
        column = 0
        for cell in row.cells:
            while (row_index, column) in occupied:
                column += 1
            placed.append((row_index, column, cell))
            for dr in range(cell.rowspan):
                for dc in range(cell.colspan):
                    occupied.add((row_index + dr, column + dc))
            column += cell.colspan
            width = max(width, column)
    return placed, width


def _add_table(docx_document: DocxDocument, element: Element, document: Document) -> None:
    table_content = element.table
    if table_content is None or not table_content.rows:
        return
    css = {key: value for key, value in _own_css(element, document).items() if not key.startswith("margin")}
    placed, width = _grid_positions(table_content)
    if width == 0:
        return
    height = len(table_content.rows)
    table = docx_document.add_table(rows=height, cols=width)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    alignments = table_content.alignments or []
    for row_index, column, cell in placed:
        last_row = min(row_index + cell.rowspan - 1, height - 1)
        last_column = min(column + cell.colspan - 1, width - 1)
        target = table.cell(row_index, column)
        if (last_row, last_column) != (row_index, column):
            target = target.merge(table.cell(last_row, last_column))
        paragraph = target.paragraphs[0]
        paragraph.style = docx_document.styles["Table Text"]
        _add_inline_runs(paragraph, cell.inline, css)
        if cell.header:
            for run in paragraph.runs:
                run.font.bold = True
        alignment = _ALIGNMENT_MAP.get((alignments[column] if column < len(alignments) else None) or "")
        if alignment is not None:
            paragraph.alignment = alignment
        background = _hex6(cell.background)
        if background:
            shading = OxmlElement("w:shd")
            shading.set(qn("w:val"), "clear")
            shading.set(qn("w:color"), "auto")
            shading.set(qn("w:fill"), background)
            target._tc.get_or_add_tcPr().append(shading)


def _image_alignment(css: dict[str, str]):
    left, right = css.get("margin-left"), css.get("margin-right")
    if left == "auto" and right == "auto":
        return WD_ALIGN_PARAGRAPH.CENTER
    if left == "auto":
        return WD_ALIGN_PARAGRAPH.RIGHT
    return None


def _add_image(
    docx_document: DocxDocument, element: Element, document: Document, assets: Mapping[str, bytes]
) -> None:
    image_bytes = resolve_image_bytes(element.image, assets) if element.image else None
    if image_bytes is None:
        return

    css = _resolved_css(element, document)
    width = None
    image_width_css = css.get("width", "")
    if image_width_css.endswith("%"):
        try:
            percent = float(image_width_css.rstrip("%"))
            width = Cm(_content_width_cm(document) * percent / 100)
        except ValueError:
            width = None

    try:
        if width is not None:
            docx_document.add_picture(io.BytesIO(image_bytes), width=width)
        else:
            docx_document.add_picture(io.BytesIO(image_bytes))
    except Exception:
        return  # best-effort, same philosophy as the parser's own image handling
    alignment = _image_alignment(css)
    if alignment is not None:
        docx_document.paragraphs[-1].alignment = alignment


def _add_code_block(docx_document: DocxDocument, element: Element, document: Document) -> None:
    """The Code style carries the monospace font and the grey shading."""
    own = _own_css(element, document)
    paragraph = docx_document.add_paragraph(style="Code")
    _apply_run_css(paragraph.add_run(element.content), own)
    _apply_paragraph_css(paragraph, own)


def _add_horizontal_rule(docx_document: DocxDocument) -> None:
    paragraph = docx_document.add_paragraph()
    paragraph._p.get_or_add_pPr().append(
        parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="9CA3AF"/></w:pBdr>')
    )


def _add_page_break(docx_document: DocxDocument) -> None:
    # python-docx has a real, first-class page break -- a run-level WD_BREAK,
    # not a styled paragraph standing in for one.
    docx_document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _add_element(
    docx_document: DocxDocument, element: Element, document: Document, assets: Mapping[str, bytes]
) -> None:
    if element.type == ElementType.HEADING:
        _add_heading(docx_document, element, document)
    elif element.type == ElementType.LIST:
        _add_list(docx_document, element, document)
    elif element.type == ElementType.TABLE:
        _add_table(docx_document, element, document)
    elif element.type == ElementType.IMAGE:
        _add_image(docx_document, element, document, assets)
    elif element.type == ElementType.CODE_BLOCK:
        _add_code_block(docx_document, element, document)
    elif element.type == ElementType.PAGE_BREAK:
        _add_page_break(docx_document)
    elif element.type == ElementType.HORIZONTAL_RULE:
        _add_horizontal_rule(docx_document)
    else:
        _add_paragraph(docx_document, element, document)
