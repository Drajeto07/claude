"""Validation Pack v1 (S2-100; strategy docs/SMARTDOC_PRODUCT_VISION_2026-10-22.md sections 29-31): the scenarios
R01-R10 as labelled Word files, measured by scripts/doctor_validation.py.

It builds on the Doctor v0 defect corpus (scripts/make_doctor_fixtures.py, d01-d13), which it neither rewrites nor
copies: d01-d13 are labelled here with the same ground truth, by place, and the scenarios they don't cover get new
files (r01-r10), each with a control that has no such defect (section 31: "every category needs documents without
the error"), and adversarial decoys -- things that look like a defect but must be left alone.

    python -m scripts.make_validation_pack

Writes the r*.docx files and validation_pack.json into tests/fixtures/doctor/. Deterministic: the same bytes every
time. All of it is synthetic (the development set of section 31); real documents are S2-107's holdout set.

Ground truth is written here, by construction -- never read back from what Doctor finds. Each label is a *region*
of the document found by its text (so it can be found again in the repaired file):

    {"text": "...", "nth": 0}  the body block (of word/document.xml) whose text contains this -- the nth such
    {"empty_after": "..."}     the empty paragraphs right after that block
    {"breaks_after": "..."}    the page breaks right after that block (a paragraph set to start on a new page counts)
    {"breaks_at_start": 1}     the page breaks before any content
    {"picture": 0}             the nth block holding a drawing
    {"document": 1}            the document as a whole (headers and footers, refusals)

Regions are blocks of the Word file itself, not of the editor's model, so what the import leaves out is still counted.

with `found` (how many blocks the region has in the original) and `defects` (how many of them are the defect);
`defects: 0` is a decoy. `keep` lists what the repaired file must still have, checked in its XML.
"""

import json

from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from scripts.make_doctor_fixtures import FIXED, LOREM, OUT, _fixed_dates, bookmark, new_document, png, save, text_paragraphs

PACK = OUT / "validation_pack.json"

# Section 30, Document Repair. "mvp": what v1 asks of Doctor -- R03, R04, R07, R08 and R10 may be diagnosis or a safe
# refusal only. "detectors": the Document Health checks whose finding counts as detecting the scenario (R06 and R08
# since S2-140: fake_headings, headers_footers; before, none -- every instance a miss).
CATEGORIES = {
    "R01": {"scenario": "Different heading styles", "expected": "detects and proposes unifying", "mvp": "repair", "detectors": ["heading_sizes", "hierarchy"]},
    "R02": {"scenario": "Chaotic spacing", "expected": "proposes normalising", "mvp": "repair", "detectors": ["spacing", "empty_paragraphs"]},
    "R03": {"scenario": "Broken numbering", "expected": "diagnoses the problem", "mvp": "diagnose", "detectors": ["numbering", "lists"]},
    "R04": {"scenario": "Problem table", "expected": "detects unsuitable properties", "mvp": "diagnose", "detectors": ["tables", "broken_tables", "layout"]},
    "R05": {"scenario": "Random fonts", "expected": "detects obvious differences", "mvp": "detect", "detectors": ["fonts", "direct_formatting"]},
    "R06": {"scenario": "Fake headings", "expected": "proposes real Heading styles", "mvp": "repair", "detectors": ["fake_headings"]},
    "R07": {"scenario": "Unwanted page breaks", "expected": "explains the likely cause", "mvp": "diagnose", "detectors": ["page_breaks"]},
    "R08": {"scenario": "Different headers/footers", "expected": "detects the mismatches", "mvp": "diagnose", "detectors": ["headers_footers"]},
    "R09": {"scenario": "Formatting from copied text", "expected": "detects incompatible styles", "mvp": "detect", "detectors": ["duplicated_formatting", "direct_formatting", "fonts"]},
    "R10": {"scenario": "Complex DOCX with objects", "expected": "keeps the objects or refuses the risky repair", "mvp": "safe", "detectors": ["fonts", "layout", "empty_paragraphs"]},
}


def label(id: str, category: str, where: dict, about: str, found: int = 1, defects: int = 1, **extra) -> dict:
    return {"id": id, "category": category, "where": where, "found": found, "defects": defects, "about": about, **extra}


def header_row(table) -> None:
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))


def filled_table(document, rows: int, cols: int, text) -> object:
    table = document.add_table(rows=rows, cols=cols)
    table.style = "Table Grid"
    for r, row in enumerate(table.rows):
        for c, cell in enumerate(row.cells):
            cell.text = text(r, c)
    return table


def hyperlink(paragraph, url: str, text: str, font: str) -> None:
    rid = paragraph.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), font)
    fonts.set(qn("w:hAnsi"), font)
    properties.append(fonts)
    run.append(properties)
    words = OxmlElement("w:t")
    words.text = text
    run.append(words)
    link.append(run)
    paragraph._p.append(link)


def content_control(document, tag: str, text: str) -> None:
    body = document.element.body
    control = OxmlElement("w:sdt")
    properties = OxmlElement("w:sdtPr")
    alias = OxmlElement("w:alias")
    alias.set(qn("w:val"), "Client name")
    properties.append(alias)
    marker = OxmlElement("w:tag")
    marker.set(qn("w:val"), tag)
    properties.append(marker)
    control.append(properties)
    content = OxmlElement("w:sdtContent")
    paragraph, run, words = OxmlElement("w:p"), OxmlElement("w:r"), OxmlElement("w:t")
    words.text = text
    run.append(words)
    paragraph.append(run)
    content.append(paragraph)
    control.append(content)
    body.insert(len(body) - 1, control)


def comment(document, paragraph, text: str, with_style: bool = True) -> None:
    if with_style:  # python-docx refers to this style but its template doesn't define it
        style = document.styles.add_style("annotation reference", WD_STYLE_TYPE.CHARACTER)
        style.style_id = "CommentReference"
    added = document.add_comment(paragraph.runs, text=text, author="Reviewer", initials="R")
    added._comment_elm.set(qn("w:date"), FIXED.strftime("%Y-%m-%dT%H:%M:%SZ"))


def page_break(document) -> None:
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def numbered_headings(document, titles: list[str]) -> None:
    for index, title in enumerate(titles):
        document.add_heading(title, level=1)
        text_paragraphs(document, 1, index)


# Labels for the existing defect corpus (d01-d13): the same ground truth as its manifest.json, placed.
DOCTOR_CORPUS = {
    "d01-empty-spacing.docx": [
        label("d01-a", "R02", {"empty_after": LOREM[0]}, "2 empty paragraphs used as spacing", found=2, defects=2),
        label("d01-b", "R02", {"empty_after": LOREM[1]}, "3 empty paragraphs used as spacing", found=3, defects=3),
        label("d01-c", "R02", {"empty_after": LOREM[2]}, "1 empty paragraph used as spacing"),
    ],
    "d02-mixed-fonts.docx": [
        label("d02-a", "R05", {"text": LOREM[2]}, "paragraph in Comic Sans MS"),
        label("d02-b", "R05", {"text": LOREM[3]}, "paragraph in Comic Sans MS"),
    ],
    "d03-repeated-formatting.docx": [
        label("d03-a", "R09", {"text": LOREM[0], "nth": 1}, "repeats the style's Arial 11 as direct formatting"),
        label("d03-b", "R09", {"text": LOREM[1], "nth": 1}, "repeats the style's Arial 11 as direct formatting"),
        label("d03-c", "R09", {"text": LOREM[2]}, "repeats the style's Arial 11 as direct formatting"),
    ],
    "d04-wide-table.docx": [label("d04-a", "R04", {"text": "Cell 1"}, "table 20 cm wide on 16.51 cm of text width")],
    "d05-wide-picture.docx": [label("d05-a", "R10", {"picture": 0}, "picture 20 cm wide: repair it only with the drawing kept")],
    "d06-page-breaks.docx": [
        label("d06-a", "R07", {"breaks_after": LOREM[0]}, "two page breaks in a row (an empty page)", found=2, defects=1, explains=["empty page"]),
        label("d06-b", "R07", {"breaks_after": LOREM[1]}, "page break after all the content", explains=["after all of it"]),
    ],
    "d07-bookmark-in-empty.docx": [
        label("d07-a", "R02", {"empty_after": LOREM[0]}, "an empty paragraph, then an empty one holding bookmark keep_me", found=2, defects=1),
        label("d07-b", "R02", {"empty_after": LOREM[1]}, "1 empty paragraph used as spacing"),
    ],
    "d08-field-in-mixed-font.docx": [
        label("d08-a", "R05", {"text": LOREM[3]}, "paragraph in Comic Sans MS"),
        label("d08-b", "R10", {"text": "Page"}, "Comic Sans paragraph holding a PAGE field: repair it only if the field stays"),
    ],
    "d09-clean.docx": [],
    "d10-signed.docx": [label("d10-a", "R10", {"document": 1}, "digitally signed: must be refused", found=0, refusal="signed")],
    "d11-macros.docx": [label("d11-a", "R10", {"document": 1}, "macros: must be refused", found=0, refusal="macros")],
    "d12-tracked-changes.docx": [label("d12-a", "R10", {"document": 1}, "tracked changes waiting: must be refused", found=0, refusal="tracked_changes")],
    "d13-mixed-report.docx": [
        label("d13-a", "R02", {"empty_after": LOREM[0]}, "2 empty paragraphs used as spacing", found=2, defects=2),
        label("d13-b", "R05", {"text": LOREM[1]}, "paragraph in Times New Roman"),
        label("d13-c", "R04", {"text": "Cell 1"}, "table 19 cm wide"),
        label("d13-d", "R02", {"empty_after": "Cell 1"}, "empty paragraph right after the table (the manifest's gap: left on purpose)"),
    ],
}

KEEP = {
    "d05-wide-picture.docx": {"drawings": 1},
    "d07-bookmark-in-empty.docx": {"bookmarks": ["keep_me"]},
    "d08-field-in-mixed-font.docx": {"fields": ["PAGE"]},
}


def build() -> tuple[dict[str, bytes], dict]:
    files: dict[str, bytes] = {}
    entries: dict[str, dict] = {}

    def add(name: str, category: str, kind: str, document_or_bytes, labels: list[dict], about: str, keep: dict | None = None) -> None:
        data = document_or_bytes if isinstance(document_or_bytes, bytes) else save(document_or_bytes)
        files[name] = _fixed_dates(data)
        entries[name] = {"category": category, "kind": kind, "origin": "synthetic", "split": "development", "about": about, "labels": labels, "keep": keep or {}}

    # R01 -- a Heading 1 set larger by hand, and a level skipped. Decoy: the title, larger by design.
    d = new_document()
    d.add_heading("Annual report 2026", level=0)
    d.add_heading("Introduction", level=1)
    text_paragraphs(d, 1)
    d.add_heading("", level=1).add_run("Budget overview").font.size = Pt(20)
    text_paragraphs(d, 1, 1)
    d.add_heading("Open questions", level=1)
    text_paragraphs(d, 1, 2)
    d.add_heading("Detailed figures", level=3)
    text_paragraphs(d, 1, 3)
    add("r01-heading-styles.docx", "R01", "defect", d, [
        label("r01-a", "R01", {"text": "Budget overview"}, "Heading 1 in 20 pt, the other Heading 1s in the style's size"),
        label("r01-b", "R01", {"text": "Detailed figures"}, "Heading 3 right after a Heading 1 (a level skipped)"),
        label("r01-x", "R01", {"text": "Annual report 2026"}, "decoy: the title is meant to be larger than the headings", defects=0),
    ], "a heading 1 sized by hand, a skipped level, and a title")

    d = new_document()
    d.add_heading("Annual report 2026", level=0)
    numbered_headings(d, ["Introduction", "Budget overview"])
    d.add_heading("Detailed figures", level=2)
    text_paragraphs(d, 1, 3)
    add("r01-clean.docx", "R01", "control", d, [], "headings of one level alike, levels in order")

    # R02 -- two paragraphs spaced apart from the rest. Decoy: the empty paragraph that keeps two tables apart.
    d = new_document()
    for index in range(6):
        paragraph = d.add_paragraph(f"{LOREM[index % 5]} Item {index}.")
        if index in (2, 4):
            paragraph.paragraph_format.space_after = Pt(24)
            paragraph.paragraph_format.line_spacing = 1.5
    filled_table(d, 2, 2, lambda r, c: f"Plan {r}{c}: {1200 + 10 * r + c} EUR")
    d.add_paragraph("")
    filled_table(d, 2, 2, lambda r, c: f"Actual {r}{c}: {1150 + 10 * r + c} EUR")
    d.add_paragraph("Closing remark, written on 2026-11-30.")
    add("r02-spacing.docx", "R02", "defect", d, [
        label("r02-a", "R02", {"text": "Item 2."}, "24 pt after and 1.5 lines, the rest of the body single-spaced"),
        label("r02-b", "R02", {"text": "Item 4."}, "24 pt after and 1.5 lines, the rest of the body single-spaced"),
        label("r02-x", "R02", {"empty_after": "Plan 00"}, "decoy: the empty paragraph keeps the two tables apart (removed, Word joins them)", defects=0),
    ], "spacing set apart on two paragraphs; an empty paragraph between two tables")

    d = new_document()
    for index in range(5):
        d.add_paragraph(f"{LOREM[index]} Item {index}.")
    add("r02-clean.docx", "R02", "control", d, [], "paragraphs spaced alike, no empty paragraphs")

    # R03 -- a gap in hand-numbered headings, a numbered list whose items also carry typed numbers.
    # Decoy: a heading that starts with a year.
    d = new_document()
    numbered_headings(d, ["1 Introduction", "2 Scope", "4 Budget"])
    d.add_paragraph("1. Approve the plan", style="List Number")
    d.add_paragraph("2. Send the report", style="List Number")
    d.add_heading("2026 targets", level=2)
    text_paragraphs(d, 1, 4)
    add("r03-numbering.docx", "R03", "defect", d, [
        label("r03-a", "R03", {"text": "4 Budget"}, "heading numbered 4 right after 2"),
        label("r03-b", "R03", {"text": "Approve the plan"}, "numbered list whose items also start with a typed number"),
        label("r03-x", "R03", {"text": "2026 targets"}, "decoy: a heading that starts with a year, not a number", defects=0),
    ], "a numbering gap, doubled list numbers, a heading starting with a year")

    d = new_document()
    numbered_headings(d, ["1 Introduction", "2 Scope", "3 Budget"])
    d.add_paragraph("Approve the plan", style="List Number")
    d.add_paragraph("Send the report", style="List Number")
    add("r03-clean.docx", "R03", "control", d, [], "headings numbered in order, a list numbered once")

    # R04 -- a table without the header row the others have, a table with a short row; two empty paragraphs so a
    # repair runs and the tables' content is checked after it.
    d = new_document()
    text_paragraphs(d, 1)
    header_row(filled_table(d, 3, 3, lambda r, c: f"Q{r} 2026-0{c + 1}-15 {r * 100 + c}"))
    text_paragraphs(d, 1, 1)
    filled_table(d, 3, 3, lambda r, c: f"Cost {r}{c} = {r * 7 + c},50")
    text_paragraphs(d, 1, 2)
    short = filled_table(d, 3, 3, lambda r, c: f"Short {r}{c}")
    header_row(short)
    row = short.rows[2]._tr
    row.remove(row.findall(qn("w:tc"))[-1])
    text_paragraphs(d, 1, 3)
    d.add_paragraph("")
    d.add_paragraph("")
    text_paragraphs(d, 1, 4)
    add("r04-tables.docx", "R04", "defect", d, [
        label("r04-a", "R04", {"text": "Cost 00"}, "the one table of three without a header row"),
        label("r04-b", "R04", {"text": "Short 00"}, "a row one cell short of the table's grid"),
        label("r04-c", "R02", {"empty_after": LOREM[3]}, "2 empty paragraphs used as spacing", found=2, defects=2),
    ], "two problem tables, empty paragraphs to repair")

    d = new_document()
    text_paragraphs(d, 1)
    header_row(filled_table(d, 2, 3, lambda r, c: f"Q{r} {2024 + c}"))
    text_paragraphs(d, 1, 1)
    header_row(filled_table(d, 2, 3, lambda r, c: f"Cost {r}{c}"))
    text_paragraphs(d, 1, 2)
    add("r04-clean.docx", "R04", "control", d, [], "two tables alike, both inside the text width")

    # R05 -- words in Comic Sans inside an Arial paragraph. Decoy: a check mark set in Wingdings (another font
    # shows another character).
    d = new_document()
    text_paragraphs(d, 2)
    paragraph = d.add_paragraph("Totals are final ")
    paragraph.add_run("except the travel budget").font.name = "Comic Sans MS"
    paragraph.add_run(" for 2026.")
    paragraph = d.add_paragraph("Approved ")
    paragraph.add_run("").font.name = "Wingdings"
    paragraph.add_run(" by the board on 2026-10-01.")
    text_paragraphs(d, 2, 2)
    add("r05-fonts.docx", "R05", "defect", d, [
        label("r05-a", "R05", {"text": "except the travel budget"}, "words in Comic Sans MS inside an Arial paragraph"),
        label("r05-x", "R05", {"text": "by the board on"}, "decoy: a Wingdings check mark -- its font is the character", defects=0),
    ], "an off-font run, a symbol-font run", keep={"symbols": [["", "Wingdings"]]})

    d = new_document()
    text_paragraphs(d, 4)
    add("r05-clean.docx", "R05", "control", d, [], "one font throughout")

    # R06 -- paragraphs made to look like headings (bold, 16 pt), in the Normal style.
    d = new_document()
    for index, title in enumerate(["Project background", "Delivery plan", "Risks and costs"]):
        run = d.add_paragraph().add_run(title)
        run.bold = True
        run.font.size = Pt(16)
        text_paragraphs(d, 2, index)
    add("r06-fake-headings.docx", "R06", "defect", d, [
        label("r06-a", "R06", {"text": "Project background"}, "bold 16 pt Normal paragraph used as a heading"),
        label("r06-b", "R06", {"text": "Delivery plan"}, "bold 16 pt Normal paragraph used as a heading"),
        label("r06-c", "R06", {"text": "Risks and costs"}, "bold 16 pt Normal paragraph used as a heading"),
    ], "three fake headings")

    d = new_document()
    for index, title in enumerate(["Project background", "Delivery plan", "Risks and costs"]):
        d.add_heading(title, level=1)
        text_paragraphs(d, 2, index)
    add("r06-clean.docx", "R06", "control", d, [], "the same document with real Heading 1s")

    # R07 -- a page break before any content; a manual break right before a paragraph set to start on a new page.
    d = new_document()
    page_break(d)
    text_paragraphs(d, 1)
    page_break(d)
    d.add_paragraph("Appendix starts here.").paragraph_format.page_break_before = True
    text_paragraphs(d, 1, 1)
    add("r07-page-breaks.docx", "R07", "defect", d, [
        label("r07-a", "R07", {"breaks_at_start": 1}, "page break before any content (a blank first page)", explains=["before any content"]),
        label("r07-b", "R07", {"breaks_after": LOREM[0]}, "manual break, and the next paragraph set to 'page break before': an empty page", found=2, defects=1,
              explains=["page break before", "starts on a new page", "new page"]),
    ], "a break at the start; a manual break doubling a paragraph's own page break")

    d = new_document()
    text_paragraphs(d, 1)
    page_break(d)
    text_paragraphs(d, 1, 1)
    add("r07-clean.docx", "R07", "control", d, [], "one page break between content")

    # R08 -- two sections whose headers disagree (name and year) and whose footers differ; an empty paragraph so a
    # repair runs and the headers are checked after it.
    d = new_document()
    d.sections[0].header.paragraphs[0].text = "Northwind Ltd - Annual report 2026"
    d.sections[0].footer.paragraphs[0].text = "Confidential"
    text_paragraphs(d, 2)
    d.add_paragraph("")
    text_paragraphs(d, 1, 2)
    second = d.add_section(WD_SECTION.NEW_PAGE)
    second.header.is_linked_to_previous = False
    second.header.paragraphs[0].text = "Northwind Limited - Annual report 2025"
    second.footer.is_linked_to_previous = False
    text_paragraphs(d, 2, 3)
    add("r08-headers.docx", "R08", "defect", d, [
        label("r08-a", "R08", {"document": 1}, "section 2's header names the company and the year differently", found=0),
        label("r08-b", "R08", {"document": 1}, "section 2 has an empty footer, section 1 says Confidential", found=0),
        label("r08-c", "R02", {"empty_after": LOREM[1]}, "1 empty paragraph used as spacing"),
    ], "headers and footers that disagree between sections")

    d = new_document()
    d.sections[0].header.paragraphs[0].text = "Northwind Ltd - Annual report 2026"
    d.sections[0].footer.paragraphs[0].text = "Confidential"
    text_paragraphs(d, 2)
    d.add_section(WD_SECTION.NEW_PAGE)
    text_paragraphs(d, 2, 2)
    add("r08-clean.docx", "R08", "control", d, [], "two sections sharing one header and footer")

    # R09 -- text pasted with its own font, size and colour; text repeating the style's font and size.
    # Decoy: a warning set bold and red on purpose.
    d = new_document()
    text_paragraphs(d, 2)
    for words in ("Copied from the supplier portal: delivery on 2026-11-03 for 1,250.00 EUR.", "Copied note: payment terms 30 days."):
        run = d.add_paragraph().add_run(words)
        run.font.name, run.font.size, run.font.color.rgb = "Calibri", Pt(12), RGBColor(0x1F, 0x1F, 0x1F)
    paragraph = d.add_paragraph("Read this first: ")
    run = paragraph.add_run("WARNING")
    run.bold, run.font.color.rgb = True, RGBColor(0xC0, 0x00, 0x00)
    paragraph.add_run(" the prices change on 1 January 2027.")
    run = d.add_paragraph().add_run(LOREM[3])
    run.font.name, run.font.size = "Arial", Pt(11)
    text_paragraphs(d, 1, 4)
    add("r09-pasted.docx", "R09", "defect", d, [
        label("r09-a", "R09", {"text": "supplier portal"}, "pasted in Calibri 12 grey, the body Arial 11"),
        label("r09-b", "R09", {"text": "payment terms"}, "pasted in Calibri 12 grey, the body Arial 11"),
        label("r09-c", "R09", {"text": LOREM[3]}, "repeats the style's Arial 11 as direct formatting"),
        label("r09-x", "R09", {"text": "Read this first"}, "decoy: a warning made bold and red on purpose", defects=0),
    ], "pasted formatting, repeated formatting, a deliberate local style", keep={"colours": [["WARNING", "C00000"]]})

    d = new_document()
    text_paragraphs(d, 4)
    add("r09-clean.docx", "R09", "control", d, [], "all text formatted through its style")

    # R10 -- objects next to defects: an off-font run in a paragraph holding a link; a content control; empty
    # paragraphs before a picture; an off-font run in a paragraph holding a bookmark; a comment on a total.
    d = new_document()
    text_paragraphs(d, 1)
    paragraph = d.add_paragraph("See the price list at ")
    paragraph.add_run("our site ").font.name = "Comic Sans MS"
    hyperlink(paragraph, "https://example.com/prices", "example.com/prices", "Comic Sans MS")
    content_control(d, "client", "Client: Northwind Ltd, contract 2026-417")
    d.add_paragraph("")
    d.add_paragraph("")
    d.add_picture(png(200, 100), width=Cm(8))
    text_paragraphs(d, 1, 1)
    paragraph = d.add_paragraph("Signed copy ")
    bookmark(paragraph, "signed_copy", 3)
    paragraph.add_run("on file").font.name = "Comic Sans MS"
    comment(d, d.add_paragraph("Total due: 4,800.00 EUR by 2026-12-31."), "Check the total")
    text_paragraphs(d, 1, 2)
    add("r10-objects.docx", "R10", "defect", d, [
        label("r10-a", "R10", {"text": "See the price list"}, "off-font run in a paragraph holding a link: refuse, or repair with the link kept"),
        label("r10-b", "R10", {"text": "Signed copy"}, "off-font run in a paragraph holding a bookmark: refuse, or repair with the bookmark kept"),
        label("r10-c", "R02", {"empty_after": "Client: Northwind"}, "2 empty paragraphs before the picture, after the content control", found=2, defects=2),
    ], "a link, a content control, a picture, a bookmark and a comment next to defects",
        keep={"links": ["https://example.com/prices"], "content_controls": ["client"], "drawings": 1, "bookmarks": ["signed_copy"], "comments": 1})

    d = new_document()
    text_paragraphs(d, 1)
    comment(d, d.add_paragraph("Total due: 4,800.00 EUR by 2026-12-31."), "Check the total", with_style=False)
    d.add_paragraph("")
    text_paragraphs(d, 1, 1)
    add("r10-broken-package.docx", "R10", "defect", d, [
        label("r10-d", "R10", {"document": 1}, "a comment whose style the package doesn't define: an already damaged package, refused", found=0, refusal="broken_package"),
    ], "a package that is already damaged")

    for name, labels in DOCTOR_CORPUS.items():
        categories = sorted({entry["category"] for entry in labels})
        entries[name] = {
            "category": categories[0] if len(categories) == 1 else ("control" if not categories else "mixed"),
            "kind": "control" if not labels else "defect",
            "origin": "synthetic",
            "split": "development",
            "about": "Doctor v0 defect corpus (scripts/make_doctor_fixtures.py, manifest.json)",
            "labels": labels,
            "keep": KEEP.get(name, {}),
        }
    return files, {"version": 1, "categories": CATEGORIES, "files": dict(sorted(entries.items()))}


def write() -> dict:
    files, pack = build()
    OUT.mkdir(parents=True, exist_ok=True)
    for name, data in files.items():
        (OUT / name).write_bytes(data)
    PACK.write_text(json.dumps(pack, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    return pack


if __name__ == "__main__":
    for name, entry in write()["files"].items():
        print(f"{name}: {entry['category']} {entry['kind']} -- {len(entry['labels'])} labels")
