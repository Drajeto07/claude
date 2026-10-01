"""Times the Word export of big tables (tracker PERF-001): a synthetic python-docx file
with one table of N rows x 8 columns goes through the importer and the exporter, as an
upload does, and the export's time is printed for each size. Synthetic text only.

    python -m scripts.benchmark_table_export [--rows 250 500 1000] [--columns 8] [--into-original] [--profile [N]]

--profile prints the N (default 15) most expensive functions by cumulative time for
each size; --into-original exports into the imported file instead of a new one."""

import argparse
import cProfile
import io
import pstats
import time

from docx import Document as DocxDocument

from app.export.docx_export import build_docx
from app.export.package_check import package_problems
from app.parsers.docx import parse_docx


def table_file(rows: int, columns: int) -> bytes:
    """A Word file with a heading and a header-row table of made-up words."""
    word = DocxDocument()
    word.add_heading("Table", level=1)
    table = word.add_table(rows=rows, cols=columns)
    table.style = word.styles["Table Grid"]
    for row_index, row in enumerate(table.rows):
        for column_index, cell in enumerate(row.cells):
            cell.text = f"Column {column_index}" if row_index == 0 else f"r{row_index} c{column_index} lorem ipsum"
    buffer = io.BytesIO()
    word.save(buffer)
    return buffer.getvalue()


def export_time(rows: int, columns: int, *, into_original: bool = False, profile: int = 0) -> float:
    """Seconds build_docx takes for one table (the import isn't counted)."""
    data = table_file(rows, columns)
    document = parse_docx(data, "table.docx")
    kwargs = {"source": data} if into_original else {}
    profiler = cProfile.Profile() if profile else None
    started = time.perf_counter()
    if profiler:
        profiler.enable()
    exported = build_docx(document, **kwargs)
    if profiler:
        profiler.disable()
    elapsed = time.perf_counter() - started
    if problems := package_problems(exported):
        raise SystemExit(f"{rows}x{columns}: the package check found {problems[:3]}")
    if profiler:
        pstats.Stats(profiler).sort_stats("cumulative").print_stats(profile)
    return elapsed


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    parser.add_argument("--rows", type=int, nargs="+", default=[250, 500, 1000])
    parser.add_argument("--columns", type=int, default=8)
    parser.add_argument("--into-original", action="store_true")
    parser.add_argument("--profile", type=int, nargs="?", const=15, default=0)
    args = parser.parse_args()
    previous = None
    for rows in args.rows:
        seconds = export_time(rows, args.columns, into_original=args.into_original, profile=args.profile)
        growth = f"  x{seconds / previous:.2f} of the one before" if previous else ""
        print(f"{rows} x {args.columns}: {seconds:.2f} s{growth}")
        previous = seconds


if __name__ == "__main__":
    main()
